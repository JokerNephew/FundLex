"""Cloud verification for the recall-first deployment.

Two independent checks:

1. Tunnel check (from this machine): forward a local port to the server's API through
   SSH and load the workbench page, proving the documented secure access method works.
2. Server-side functional check (executed on the server): create a temporary user,
   log in over real HTTP to the local API, inspect capabilities/guides/settings, create
   a recall-first task with a generated DOCX, upload and seal it, verify the task
   snapshot, then remove the temporary rows. No credential is ever printed.

Usage:
  .venv\\Scripts\\python.exe -X utf8 -m tools.deploy.cloud_verify [--skip-tunnel]
"""
from __future__ import annotations

import argparse
import sys
import threading
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
REMOTE_PORT = 8010

SERVER_SCRIPT = r'''
import io, json, os, random, secrets, string, sys, time
import httpx

sys.path.insert(0, "/opt/fundlex-recall-first")
env = {}
for line in open("/etc/fundlex-recall.env"):
    line = line.strip()
    if "=" in line and not line.startswith("#"):
        key, value = line.split("=", 1)
        env[key] = value
for key, value in env.items():
    os.environ.setdefault(key, value)

from fundlex.store import RunStore
from fundlex.tasks import TaskStore
from fundlex.auth import UserStore, SessionStore

store = RunStore(env["FUNDLEX_DATA_DIR"])
tasks = TaskStore(store)
users = UserStore(store.database)
sessions = SessionStore(store.database)

suffix = "".join(random.choices(string.ascii_lowercase + string.digits, k=8))
username = f"deploy-verify-{suffix}"
password = secrets.token_urlsafe(18)
user = users.create_user(username=username, password=password, display_name="部署验证账号", role="member")
base = "http://127.0.0.1:8010"
results = {}
try:
    with httpx.Client(base_url=base, timeout=60) as client:
        login = client.post("/api/auth/login", json={"username": username, "password": password})
        results["login"] = login.status_code
        token = login.json()["token"]
        headers = {"Authorization": "Bearer " + token}
        caps = client.get("/api/capabilities", headers=headers)
        results["capabilities"] = caps.status_code
        caps_json = caps.json()
        engine = caps_json.get("engine") or {}
        results["engine_id"] = engine.get("id")
        results["guide"] = (engine.get("guide") or {}).get("version_id")
        results["guide_status"] = (engine.get("guide") or {}).get("status")
        guides = client.get("/api/guides", headers=headers)
        results["guides"] = guides.status_code
        results["guide_count"] = len(guides.json().get("items") or [])
        settings = client.get("/api/settings/guide", headers=headers)
        results["settings"] = settings.status_code
        results["current_guide"] = (settings.json().get("current") or {}).get("id")

        from docx import Document
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("某某灵活配置混合型证券投资基金基金合同")
        document.add_paragraph("二、释义")
        document.add_paragraph("基金份额类别：本基金将基金份额分为A类和C类不同的类别。在投资者认购、申购基金时收取认购、申购费用而不计提销售服务费的，称为A类基金份额；在投资者认购、申购基金份额时不收取认购、申购费用而是从本类别基金资产中计提销售服务费的，称为C类基金份额。")
        document.save(buffer)
        payload = buffer.getvalue()

        create = client.post("/api/tasks", headers=headers, json={
            "name": f"部署验证-{suffix}",
            "engine_id": "recall_first",
            "files": [{"client_id": "verify-1", "name": "deploy-verify.docx", "size": len(payload)}],
        })
        results["task_create"] = create.status_code
        task = create.json()
        task_id = task.get("id")
        results["task_engine"] = task.get("engine_id")
        results["task_guide"] = (task.get("guide_snapshot") or {}).get("version_id")
        upload = client.post(f"/api/tasks/{task_id}/uploads", headers=headers,
                             data={"client_ids": json.dumps(["verify-1"])},
                             files=[("files", ("deploy-verify.docx", payload,
                                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))])
        results["upload"] = upload.status_code
        receipts = upload.json().get("receipts") or []
        results["upload_status"] = receipts[0].get("status") if receipts else None
        seal = client.post(f"/api/tasks/{task_id}/seal", headers=headers)
        results["seal"] = seal.status_code
        snapshot = client.get(f"/api/tasks/{task_id}", headers=headers).json()
        results["task_status"] = snapshot.get("status")
        documents = client.get(f"/api/tasks/{task_id}/documents", headers=headers).json()
        results["document_count"] = len(documents.get("items") or [])
        if documents.get("items"):
            item = documents["items"][0]
            results["document_engine_bound"] = bool(item.get("guide_snapshot"))
            results["document_status"] = item.get("job_status")
        results["ok"] = True
finally:
    with store.database.connect() as db:
        if 'task_id' in dir():
            pass
    try:
        with store.database.connect() as db:
            rows = db.execute("SELECT id FROM tasks WHERE owner=?", (username,)).fetchall()
            task_ids = [r["id"] if isinstance(r, dict) else r[0] for r in rows]
            for tid in task_ids:
                batches = db.execute("SELECT batch_id FROM task_batches WHERE task_id=?", (tid,)).fetchall()
                for batch in batches:
                    bid = batch["batch_id"] if isinstance(batch, dict) else batch[0]
                    db.execute("DELETE FROM artifacts WHERE document_id IN (SELECT id FROM documents WHERE batch_id=?)", (bid,))
                    db.execute("DELETE FROM documents WHERE batch_id=?", (bid,))
                    db.execute("DELETE FROM task_batches WHERE batch_id=?", (bid,))
                    db.execute("DELETE FROM batches WHERE id=?", (bid,))
                db.execute("DELETE FROM task_files WHERE task_id=?", (tid,))
                db.execute("DELETE FROM tasks WHERE id=?", (tid,))
            db.execute("DELETE FROM sessions WHERE user_id=?", (user["id"],))
            db.execute("DELETE FROM users WHERE id=?", (user["id"],))
        results["cleanup"] = "done"
    except Exception as exc:
        results["cleanup"] = f"failed: {type(exc).__name__}"
print("VERIFY_JSON:" + json.dumps(results, ensure_ascii=False))
'''


def tunnel_check(host: str, user: str, password: str) -> dict:
    """Open a real local port forward (ssh -L equivalent) and exercise the app through it."""
    import paramiko
    import httpx
    import socket
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(hostname=host, port=22, username=user, password=password, timeout=30)
    transport = client.get_transport()
    results: dict[str, object] = {}
    listener = socket.socket()
    listener.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    listener.bind(("127.0.0.1", 0))
    listener.listen(16)
    local_port = listener.getsockname()[1]
    stopping = threading.Event()

    def serve() -> None:
        while not stopping.is_set():
            try:
                connection, _ = listener.accept()
            except OSError:
                return
            threading.Thread(target=_forward, args=(connection, transport), daemon=True).start()

    threading.Thread(target=serve, daemon=True).start()
    try:
        with httpx.Client(timeout=30, trust_env=False) as http:
            page = http.get(f"http://127.0.0.1:{local_port}/")
            results["index_status"] = page.status_code
            results["index_has_guide_page"] = "指引配置" in page.text
            results["index_bytes"] = len(page.content)
            api = http.get(f"http://127.0.0.1:{local_port}/api/capabilities")
            results["api_status"] = api.status_code
            try:
                results["api_body"] = api.json()
            except ValueError:
                results["api_body"] = api.text[:300]
            app_js = http.get(f"http://127.0.0.1:{local_port}/static/app.js")
            results["app_js_status"] = app_js.status_code
            results["app_js_has_recall"] = "recall_first" in app_js.text
            guide_api = http.get(f"http://127.0.0.1:{local_port}/api/guides")
            results["guides_unauth_status"] = guide_api.status_code
        results["tunnel"] = "ok"
    finally:
        stopping.set()
        try:
            listener.close()
        except OSError:
            pass
        client.close()
    return results


def _forward(connection, transport) -> None:
    try:
        channel = transport.open_channel("direct-tcpip", ("127.0.0.1", REMOTE_PORT),
                                         connection.getsockname())
    except Exception:
        connection.close()
        return
    _pump(connection, channel)


def _pump(a, b) -> None:
    """Copy bytes both ways until either side closes."""
    def copy(source, target):
        try:
            while True:
                data = source.recv(8192)
                if not data:
                    break
                target.sendall(data)
        except Exception:
            pass
        finally:
            try:
                target.shutdown_write()
            except Exception:
                pass

    left = threading.Thread(target=copy, args=(a, b), daemon=True)
    right = threading.Thread(target=copy, args=(b, a), daemon=True)
    left.start()
    right.start()
    left.join()
    right.join()
    for channel in (a, b):
        try:
            channel.close()
        except Exception:
            pass


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Verify the cloud deployment")
    parser.add_argument("--skip-tunnel", action="store_true")
    parser.add_argument("--skip-functional", action="store_true")
    args = parser.parse_args(argv)

    sys.path.insert(0, str(REPO / "tools" / "deploy"))
    from ssh_probe import connect, load_credentials
    host, user, password = load_credentials()

    if not args.skip_tunnel:
        print("== tunnel check ==")
        for key, value in tunnel_check(host, user, password).items():
            print(f"  {key}: {value}")

    if not args.skip_functional:
        print("== server-side functional check ==")
        client = connect(host, user, password, port=22, timeout=30)
        try:
            sftp = client.open_sftp()
            with sftp.open("/tmp/fundlex_cloud_verify.py", "w") as handle:
                handle.write(SERVER_SCRIPT)
            sftp.close()
            stdin, stdout, stderr = client.exec_command(
                "/opt/fundlex-recall-first/.venv/bin/python -X utf8 /tmp/fundlex_cloud_verify.py", timeout=600)
            out = stdout.read().decode("utf-8", "replace")
            err = stderr.read().decode("utf-8", "replace")
            code = stdout.channel.recv_exit_status()
            for line in out.splitlines():
                if line.startswith("VERIFY_JSON:"):
                    import json
                    data = json.loads(line[len("VERIFY_JSON:"):])
                    for key, value in data.items():
                        print(f"  {key}: {value}")
                else:
                    print(" ", line)
            if code != 0:
                print("  stderr:", err.strip()[-1500:])
                return 1
            client.exec_command("rm -f /tmp/fundlex_cloud_verify.py")
        finally:
            client.close()
    print("== verification finished ==")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
