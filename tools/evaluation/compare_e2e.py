#!/usr/bin/env python3
"""Real browser acceptance for the independent contract comparison feature.

The probe starts a temporary local FastAPI instance and the independent
``CompareWorker``. It follows the user path (Compare page -> choose DOCX ->
submit -> worker -> viewer) for several scenarios. It never contacts or
deploys a production service.

Example::

    .venv/bin/python tools/evaluation/compare_e2e.py \
      --output /tmp/compare-e2e.json --screenshots /tmp/compare-e2e-shots

Set ``FUNDLEX_BROWSER_HEADLESS=0`` to inspect the browser while it runs.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import shutil
import socket
import tempfile
import threading
import time
import traceback
import zipfile
from pathlib import Path
from typing import Any

import uvicorn
from docx import Document
from lxml import etree


def make_docx(paragraphs: list[str]) -> bytes:
    stream = io.BytesIO()
    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    document.save(stream)
    return stream.getvalue()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def mutate_public_docx(source: Path, destination: Path) -> dict[str, str]:
    """Copy one public contract and make one deterministic visible edit."""
    document = Document(str(source))
    changed = False
    before = after = ""
    for paragraph in document.paragraphs:
        text = paragraph.text
        if not text.strip():
            continue
        if "应当" in text:
            before, after = "应当", "可以"
            paragraph.text = text.replace(before, after, 1)
        elif "不得" in text:
            before, after = "不得", "可以"
            paragraph.text = text.replace(before, after, 1)
        elif "0.40%" in text:
            before, after = "0.40%", "0.30%"
            paragraph.text = text.replace(before, after, 1)
        else:
            before, after = text, text + "（对比测试修改）"
            paragraph.text = after
        changed = True
        break
    if not changed:
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if paragraph.text.strip():
                            before, after = paragraph.text, paragraph.text + "（对比测试修改）"
                            paragraph.text = after
                            changed = True
                            break
                    if changed:
                        break
                if changed:
                    break
            if changed:
                break
    if not changed:
        raise RuntimeError(f"public contract has no visible paragraph: {source}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    return {"source": str(source), "source_sha256": sha256(source),
            "mutated": str(destination), "mutated_sha256": sha256(destination),
            "before": before, "after": after}


def public_contract(source: Path | None) -> Path:
    """Prefer the frozen corpus base; reject contracts with existing revisions."""
    repo = Path(__file__).resolve().parents[2]
    if source is not None:
        candidates = [source] if source.is_file() else sorted(source.rglob("*.docx"))
    else:
        manifest = Path("/tmp/fundlex-compare-corpus-v4/manifest.json")
        if manifest.is_file():
            candidates = [Path(item["path"]) for item in json.loads(manifest.read_text())["bases"]]
        else:
            candidates = sorted((repo / "funds_word").glob("*.docx"))
    for path in candidates:
        with zipfile.ZipFile(path) as package:
            xml = etree.fromstring(package.read("word/document.xml"))
        if not xml.xpath("//*[local-name()='ins' or local-name()='del' or local-name()='moveFrom' or local-name()='moveTo']"):
            return path
    raise RuntimeError("no public base contract without existing revisions was found")


def wait_until(page, expression: str, *, timeout: int = 20) -> None:
    page.wait_for_function(expression, timeout=timeout * 1000)


def run(output: Path, screenshots: Path, corpus_base: Path | None = None) -> dict[str, Any]:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Playwright is required for compare_e2e.py") from exc
    from fundlex.api import create_app
    from fundlex.compare.worker import CompareWorker

    checks: list[dict[str, Any]] = []
    scenarios: list[dict[str, Any]] = []

    def check(name: str, passed: bool, detail: Any = "") -> None:
        checks.append({"name": name, "passed": bool(passed), "detail": detail})

    screenshots.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="fundlex-compare-e2e-") as data_dir:
        root = Path(data_dir)
        app = create_app(root / "data", auth_db=str(root / "auth.sqlite3"))
        sock = socket.socket()
        sock.bind(("127.0.0.1", 0))
        base = f"http://127.0.0.1:{sock.getsockname()[1]}"
        server = uvicorn.Server(uvicorn.Config(app, log_level="critical"))
        server_thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
        server_thread.start()
        deadline = time.monotonic() + 10
        while not server.started and server_thread.is_alive() and time.monotonic() < deadline:
            time.sleep(0.02)
        if not server.started:
            raise RuntimeError("temporary FastAPI server did not start")

        stop_worker = threading.Event()

        def worker_loop() -> None:
            worker = CompareWorker(app.state.compare_store)
            while not stop_worker.is_set():
                worker.run_once()
                stop_worker.wait(0.04)

        worker_thread = threading.Thread(target=worker_loop, daemon=True)
        worker_thread.start()

        def snapshot(page, label: str) -> str:
            target = screenshots / f"{label}.png"
            page.screenshot(path=str(target), full_page=True)
            return str(target)

        def geometry(page) -> tuple[float | None, dict[str, Any]]:
            value = page.evaluate("""() => {
              const svg=document.querySelector('#compare-lines');
              const path=[...svg.querySelectorAll('path')].find(n=>n.classList.contains('active'));
              const id=path?.dataset.diffId;
              const c=document.querySelector('#compare-columns')?.getBoundingClientRect();
              const l=document.querySelector(`#compare-left-pane [data-diff-id="${CSS.escape(id||'')}"]`)?.getBoundingClientRect();
              const r=document.querySelector(`#compare-right-pane [data-diff-id="${CSS.escape(id||'')}"]`)?.getBoundingClientRect();
              const n=(path?.getAttribute('d')||'').match(/-?\\d+(?:\\.\\d+)?/g)||[];
              if(!path||!c||!l||!r||n.length<4) return {error:'missing active path or marker',id};
              const start=[+n[0],+n[1]], end=[+n.at(-2),+n.at(-1)];
              const expected=[[l.right-c.left,l.top+l.height/2-c.top],[r.left-c.left,r.top+r.height/2-c.top]];
              const error=Math.max(Math.abs(start[0]-expected[0][0]),Math.abs(start[1]-expected[0][1]),Math.abs(end[0]-expected[1][0]),Math.abs(end[1]-expected[1][1]));
              return {id,start,end,expected,error};
            }""")
            return (None if value.get("error") else float(value["error"]), value)

        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(headless=os.environ.get("FUNDLEX_BROWSER_HEADLESS", "1") != "0")
                context = browser.new_context(viewport={"width": 1440, "height": 960})
                context.add_init_script("localStorage.setItem('fundlex-dev-user','compare-e2e-user')")
                page = context.new_page()
                page_errors: list[str] = []
                page.on("pageerror", lambda error: page_errors.append(str(error)))
                network: list[dict[str, Any]] = []

                def record_response(response) -> None:
                    if "/api/compare/" not in response.url:
                        return
                    item = {"url": response.url, "method": response.request.method, "status": response.status}
                    try:
                        payload = response.json()
                        if response.request.method == "POST":
                            item["payload"] = payload
                        elif "/view" in response.url:
                            item["payload"] = {k: payload.get(k) for k in ("pair", "summary", "coverage", "result_sha256")}
                        else:
                            item["payload"] = payload
                    except Exception as exc:
                        item["error"] = str(exc)
                    network.append(item)

                page.on("response", record_response)

                def go_compare() -> None:
                    page.goto(base + "/compare#compare")
                    wait_until(page, "document.querySelector('#compare-new') !== null && !document.querySelector('#compare-new').hidden")

                def open_create(name: str) -> None:
                    go_compare()
                    page.locator("#compare-new").click()
                    page.locator("#compare-name").fill(name)

                def set_pair(index: int, left: Path, right: Path) -> None:
                    page.locator(f'[data-compare-left="{index}"]').set_input_files(str(left))
                    page.locator(f'[data-compare-right="{index}"]').set_input_files(str(right))

                def submit_task(name: str, pairs: list[tuple[Path, Path]]) -> dict[str, Any]:
                    open_create(name)
                    for index, (left, right) in enumerate(pairs):
                        if index:
                            page.locator("#compare-add-pair").click()
                        set_pair(index, left, right)
                    page.locator("#compare-submit").click()
                    wait_until(page, f"document.querySelector('#compare-task-list').textContent.includes({json.dumps(name, ensure_ascii=False)})", timeout=30)
                    page.locator("tr", has_text=name).first.locator("[data-compare-task]").click()
                    wait_until(page, f"document.querySelector('#compare-detail').textContent.includes({json.dumps(str(len(pairs)) + ' 组', ensure_ascii=False)})", timeout=20)
                    wait_until(page, f"document.querySelectorAll('#compare-pair-list .compare-pair-status.is-done').length === {len(pairs)}", timeout=45)
                    page.wait_for_timeout(150)
                    uploads = [x.get("payload", {}).get("receipts", []) for x in network if x["method"] == "POST" and "/uploads" in x["url"]]
                    return {"task": name, "pairs": len(pairs), "detail": page.locator("#compare-detail").inner_text(), "upload_receipts": uploads[-len(pairs):], "screenshot": snapshot(page, name)}

                local = root / "fixtures"
                local.mkdir()
                one_l, one_r = local / "one-left.docx", local / "one-right.docx"
                one_l.write_bytes(make_docx(["基金管理费按0.40%年费率计提。"]))
                one_r.write_bytes(make_docx(["基金管理费按0.30%年费率计提。"]))
                single = submit_task("E2E 单处修改", [(one_l, one_r)])
                summary = page.locator("#compare-summary").inner_text()
                lines = page.locator("#compare-lines path").count()
                check("single_modify_summary_one", summary.startswith("共 1 处差异") and "修改 1" in summary, summary)
                check("single_modify_one_svg_line", lines == 1, lines)
                single.update(summary=summary, line_count=lines, screenshot=snapshot(page, "single-modify"))
                scenarios.append(single)

                zero_l, zero_r = local / "zero-left.docx", local / "zero-right.docx"
                zero_bytes = make_docx(["完全相同的合同段落。", "第二个相同段落。"])
                zero_l.write_bytes(zero_bytes); zero_r.write_bytes(zero_bytes)
                zero = submit_task("E2E 零差异", [(zero_l, zero_r)])
                zero_summary = page.locator("#compare-summary").inner_text()
                zero_counter = page.locator("#compare-diff-counter").inner_text()
                zero_lines = page.locator("#compare-lines path").count()
                check("zero_diff_summary_counter_svg", zero_summary.startswith("共 0 处差异") and zero_counter == "0 / 0" and zero_lines == 0, {"summary": zero_summary, "counter": zero_counter, "lines": zero_lines})
                zero_list = page.locator("tr", has_text="E2E 零差异").first.locator("td").nth(3).inner_text()
                check("task_list_zero_differences_not_pair_count", zero_list == "0 处", zero_list)
                zero.update(summary=zero_summary, counter=zero_counter, line_count=zero_lines, screenshot=snapshot(page, "zero-diff"))
                scenarios.append(zero)

                batch_pairs: list[tuple[Path, Path]] = []
                for index in range(3):
                    left, right = local / f"batch-{index}-l.docx", local / f"batch-{index}-r.docx"
                    # Keep this as a one-modify fixture.  A short Chinese
                    # modal-word change such as 应当 -> 可以 may legitimately
                    # be represented as delete+insert by the high-recall diff.
                    left.write_bytes(make_docx([f"Contract {index}: management fee 0.40% per year."]))
                    right.write_bytes(make_docx([f"Contract {index}: management fee 0.30% per year."]))
                    batch_pairs.append((left, right))
                batch = submit_task("E2E 三组批量", batch_pairs)
                counters: list[str] = []
                summaries: list[str] = []
                for index in range(3):
                    page.locator(f'[data-compare-pair="{index}"]').click()
                    wait_until(page, f"document.querySelector('#compare-pair-counter').textContent.trim() === '{index + 1} / 3'")
                    wait_until(page, "document.querySelector('#compare-summary').textContent.includes('处差异')", timeout=30)
                    counters.append(page.locator("#compare-pair-counter").inner_text())
                    summaries.append(page.locator("#compare-summary").inner_text())
                check("batch_three_pair_navigation", counters == ["1 / 3", "2 / 3", "3 / 3"], counters)
                check("batch_three_pair_results", all(value.startswith("共 1 处差异") and "修改 1" in value for value in summaries), summaries)
                batch.update(counters=counters, summaries=summaries, screenshot=snapshot(page, "batch-three-pairs"))
                scenarios.append(batch)

                # The result reader must state the comparison direction and
                # keep the selected pair when a polling response refreshes the
                # task.  A temporarily-running response makes the existing
                # completed task exercise the same polling branch.
                check("comparison_direction_is_explicit", "原版" in page.locator("#compare-detail").inner_text() and "对比版" in page.locator("#compare-detail").inner_text() and "原版" in page.locator(".compare-columns-head").inner_text(), page.locator("#compare-detail").inner_text())
                selected_before_poll = page.locator("#compare-pair-counter").inner_text()
                task_route_calls = {"count": 0}
                def temporary_running(route) -> None:
                    response = route.fetch()
                    payload = response.json()
                    if "/api/compare/tasks/" in route.request.url and "/pairs" not in route.request.url and task_route_calls["count"] == 0:
                        payload["status"] = "running"
                    task_route_calls["count"] += 1
                    route.fulfill(response=response, json=payload)
                page.route("**/api/compare/tasks/*", temporary_running)
                page.locator('[data-compare-pair="2"]').click()
                wait_until(page, "document.querySelector('#compare-pair-counter').textContent.trim() === '3 / 3'")
                page.locator("tr", has_text="E2E 三组批量").first.locator("[data-compare-task]").click()
                page.wait_for_timeout(2800)
                check("polling_keeps_selected_pair", page.locator("#compare-pair-counter").inner_text() == "3 / 3", {"before": selected_before_poll, "after": page.locator("#compare-pair-counter").inner_text(), "route_calls": task_route_calls["count"]})
                page.unroute("**/api/compare/tasks/*", temporary_running)

                all_diff_count = page.locator(".compare-diff-item").count()
                page.locator("#compare-diff-filter").select_option("modify")
                filtered_count = page.locator(".compare-diff-item").count()
                filtered_types = page.locator(".compare-diff-item strong").all_inner_texts()
                check("diff_type_filter", filtered_count <= all_diff_count and all(value == "修改" for value in filtered_types), {"all": all_diff_count, "filtered": filtered_count, "types": filtered_types})
                page.locator("#compare-diff-filter").select_option("")
                page.locator("#compare-left-pane [data-diff-id]").first.focus()
                page.keyboard.press("Enter")
                check("diff_mark_is_keyboard_actionable", page.locator("#compare-left-pane [data-diff-id]").first.get_attribute("role") == "button" and page.locator("#compare-diff-counter").inner_text().startswith("1 /"), page.locator("#compare-diff-counter").inner_text())

                page.set_viewport_size({"width": 390, "height": 844})
                page.wait_for_timeout(150)
                check("mobile_reader_switch_visible", page.locator(".compare-mobile-switch").is_visible(), page.locator(".compare-mobile-switch").inner_text())
                page.locator('[data-compare-side="right"]').click()
                mobile_right = page.evaluate("""() => ({side: document.querySelector('#compare-viewer')?.dataset.mobileSide, left: getComputedStyle(document.querySelector('#compare-left-pane')).display, right: getComputedStyle(document.querySelector('#compare-right-pane')).display})""")
                check("mobile_reader_can_switch_side", mobile_right == {"side": "right", "left": "none", "right": "block"}, mobile_right)
                page.locator('[data-compare-side="left"]').click()
                page.set_viewport_size({"width": 1440, "height": 960})

                page.locator('[data-compare-pair="0"]').click()
                wait_until(page, "document.querySelector('#compare-pair-counter').textContent.trim() === '1 / 3'", timeout=20)
                page.locator("#compare-next-pair").click()
                wait_until(page, "document.querySelector('#compare-pair-counter').textContent.trim() === '2 / 3'", timeout=20)
                page.locator("#compare-prev-pair").click()
                wait_until(page, "document.querySelector('#compare-pair-counter').textContent.trim() === '1 / 3'", timeout=20)
                check("previous_next_pair_controls", True, page.locator("#compare-pair-counter").inner_text())

                many_l, many_r = local / "many-left.docx", local / "many-right.docx"
                many_l.write_bytes(make_docx([f"第{i:02d}条：本条应当履行义务。" for i in range(25)]))
                many_r.write_bytes(make_docx([f"第{i:02d}条：本条可以履行义务。" for i in range(25)]))
                many = submit_task("E2E 二十处以上差异", [(many_l, many_r)])
                many_summary = page.locator("#compare-summary").inner_text()
                many_count = page.locator(".compare-diff-item").count()
                check("twenty_plus_diff_items", many_count >= 20, {"summary": many_summary, "diff_items": many_count})
                many_list = page.locator("tr", has_text="E2E 二十处以上差异").first.locator("td").nth(3).inner_text()
                check("task_list_actual_difference_total", many_list == f"{many_count} 处", many_list)
                many.update(summary=many_summary, diff_items=many_count, screenshot=snapshot(page, "twenty-plus-diffs"))
                scenarios.append(many)

                empty_l, empty_r = local / "empty-left.docx", local / "empty-right.docx"
                empty_l.write_bytes(make_docx([]))
                empty_r.write_bytes(make_docx(["整段新增合同文字。"]))
                empty = submit_task("E2E 空侧虚拟锚点", [(empty_l, empty_r)])
                empty_summary = page.locator("#compare-summary").inner_text()
                left_anchor = page.locator("#compare-left-pane .compare-diff-anchor").count()
                right_mark = page.locator("#compare-right-pane .compare-diff").count()
                empty_lines = page.locator("#compare-lines path").count()
                check("empty_side_virtual_anchor", left_anchor >= 1 and right_mark >= 1 and empty_lines == 1, {"summary": empty_summary, "left_anchors": left_anchor, "right_marks": right_mark, "lines": empty_lines})
                empty.update(summary=empty_summary, left_anchors=left_anchor, right_marks=right_mark, line_count=empty_lines, screenshot=snapshot(page, "empty-side-anchor"))
                scenarios.append(empty)

                mixed_l, mixed_r = local / "mixed-left.docx", local / "mixed-right.docx"
                same = "This is the unchanged anchor clause which stays in both documents."
                modified = "The management fee rate is 0.40% per year and collected daily."
                attack = "<script>window.__compareAttack=1</script><img src=x onerror=window.__compareAttack=2>"
                mixed_l.write_bytes(make_docx([same, "Deleted only from the left contract.", same, modified, same, attack]))
                mixed_r.write_bytes(make_docx([same, same, modified.replace("0.40%", "0.30%"), same, attack, "Entirely new right-side clause: " + attack]))
                mixed = submit_task("E2E 插入删除修改与脚本文本", [(mixed_l, mixed_r)])
                mixed_summary = page.locator("#compare-summary").inner_text()
                mixed_types = page.locator(".compare-diff-item strong").all_inner_texts()
                scripts = page.locator("#compare-left-pane script,#compare-right-pane script,#compare-left-pane img,#compare-right-pane img").count()
                visible_script_text = page.locator("#compare-left-pane,#compare-right-pane").all_inner_texts()
                check("mixed_insert_delete_modify", set(mixed_types) == {"新增", "删除", "修改"}, {"summary": mixed_summary, "types": mixed_types})
                check("script_text_is_escaped", scripts == 0 and all("<script>" in value for value in visible_script_text) and not page.evaluate("window.__compareAttack || 0"), {"script_or_image_nodes": scripts, "visible_text": visible_script_text})
                mixed.update(summary=mixed_summary, diff_types=mixed_types, script_nodes=scripts, screenshot=snapshot(page, "mixed-script-text"))
                scenarios.append(mixed)

                # An unmatched file appears in the explicit pair list and can
                # be completed with the ordinary file picker.
                left_dir, right_dir = local / "folder-left", local / "folder-right"
                left_dir.mkdir(), right_dir.mkdir()
                a_l, b_l, a_r, b_r = left_dir / "a.docx", left_dir / "b.docx", right_dir / "a.docx", right_dir / "b.docx"
                a_l.write_bytes(make_docx(["folder a old"])); b_l.write_bytes(make_docx(["folder b old"]))
                a_r.write_bytes(make_docx(["folder a new"])); b_r.write_bytes(make_docx(["folder b new"]))
                right_only = local / "folder-right-only"
                right_only.mkdir()
                shutil.copy2(a_r, right_only / "a.docx")
                open_create("E2E 文件夹补配")
                page.locator("#compare-left-folder").set_input_files(str(left_dir))
                page.locator("#compare-right-folder").set_input_files(str(right_only))
                wait_until(page, "document.querySelector('#compare-upload-message').textContent.includes('未配对')")
                unmatched_message = page.locator("#compare-upload-message").inner_text()
                rows = page.locator(".compare-pair-input").count()
                folder_rows = page.locator(".compare-pair-input")
                target_row = next((folder_rows.nth(i) for i in range(folder_rows.count())
                                   if "b.docx" in folder_rows.nth(i).locator(".compare-file-label").inner_text()
                                   and not folder_rows.nth(i).locator('input[data-compare-right]').input_value()), None)
                if target_row is None:
                    raise AssertionError(f"could not locate unmatched b.docx row: {folder_rows.all_inner_texts()}")
                target_row.locator('input[data-compare-right]').set_input_files(str(b_r))
                labels = page.locator(".compare-file-label").all_inner_texts()
                check("folder_unmatched_visible_and_manual_pair", "未配对" in unmatched_message and rows == 2 and any("b.docx" in label for label in labels), {"message": unmatched_message, "rows": rows, "labels": labels})
                page.locator("#compare-submit").click()
                try:
                    wait_until(page, "document.querySelector('#compare-task-list').textContent.includes('E2E 文件夹补配')", timeout=30)
                except Exception:
                    output.parent.mkdir(parents=True, exist_ok=True)
                    output.write_text(json.dumps({"checks": checks, "network": network, "message": page.locator('#compare-upload-message').inner_text(), "screenshot": snapshot(page, 'folder-failure')}, ensure_ascii=False, indent=2), encoding='utf-8')
                    raise
                page.locator("tr", has_text="E2E 文件夹补配").first.locator("[data-compare-task]").click()
                wait_until(page, "document.querySelectorAll('#compare-pair-list .compare-pair-status.is-done').length === 2", timeout=45)
                scenarios.append({"task": "E2E 文件夹补配", "message": unmatched_message, "rows": rows, "labels": labels, "screenshot": snapshot(page, "folder-manual-pair")})

                # Terminal processing includes unsuccessful pairs. Create the
                # missing-upload fixture in this disposable store and refresh
                # explicitly (navigating to the same hash does not reload it).
                failed_owner = page.locator("#dev-user").input_value().strip()
                failed_task = app.state.compare_store.create(failed_owner, "E2E 全失败任务", [{
                    "client_pair_id": "missing", "label": "缺失文件组",
                    "left": {"client_id": "missing-l", "name": "left.docx", "size": 1},
                    "right": {"client_id": "missing-r", "name": "right.docx", "size": 1},
                }])
                app.state.compare_store.seal(failed_owner, failed_task["id"])
                page.locator("#compare-refresh").click()
                failed_row = page.locator("tr", has_text="E2E 全失败任务").first
                failed_row.wait_for()
                failed_status = failed_row.locator("td").nth(2).inner_text()
                failed_diffs = failed_row.locator("td").nth(3).inner_text()
                check("all_failed_task_status_is_explicit", "已处理" in failed_status and "失败" in failed_status and "已完成" not in failed_status, failed_status)
                check("all_failed_task_has_no_false_diff_count", "无可用结果" in failed_diffs and "失败" in failed_diffs, failed_diffs)
                failed_row.locator("[data-compare-task]").click()
                page.locator(".compare-pair-error").wait_for()
                failed_error = page.locator(".compare-pair-error").inner_text()
                check("failed_pair_error_is_visible", "COMPARE_PAIR_INCOMPLETE" in failed_error and page.locator(".compare-pair-error").is_visible(), failed_error)
                page.locator('[data-compare-pair="0"]').click()
                check("failed_pair_hides_stale_viewer", page.locator("#compare-viewer").is_hidden())
                scenarios.append({"task": "E2E 全失败任务", "status": failed_status, "differences": failed_diffs, "error": failed_error, "screenshot": snapshot(page, "all-failed-task")})

                # One real public contract, copied and mutated temporarily,
                # goes through the same page upload and worker path.
                source = public_contract(corpus_base)
                real_l, real_r = root / "real-base.docx", root / "real-mutated.docx"
                shutil.copy2(source, real_l)
                real_meta = mutate_public_docx(source, real_r)
                real = submit_task("E2E 真实公开合同", [(real_l, real_r)])
                real_summary = page.locator("#compare-summary").inner_text()
                real_count = page.locator(".compare-diff-item").count()
                check("real_public_contract_mutation_visible", real_count >= 1, {"summary": real_summary, **real_meta})
                real_detail = page.locator("#compare-detail").inner_text()
                wait_until(page, "document.querySelector('#compare-summary').textContent.includes('处差异')", timeout=20)
                viewer_content = page.locator(".compare-diff-item").count() >= 1 and page.locator("#compare-summary").inner_text().startswith("共 ")
                check("incomplete_terminal_result_is_viewable", "覆盖不完整" in real_detail and viewer_content, {"detail": real_detail, "summary": page.locator("#compare-summary").inner_text(), "diff_items": page.locator(".compare-diff-item").count()})
                real.update(summary=real_summary, diff_items=real_count, source=real_meta, screenshot=snapshot(page, "real-public-contract"))
                scenarios.append(real)

                page.goto(base + "/compare#settings")
                wait_until(page, "document.querySelector('#dev-user') !== null && !document.querySelector('#dev-user').hidden", timeout=20)
                page.locator("#dev-user").fill("compare-e2e-other-user")
                page.goto(base + "/compare#compare")
                wait_until(page, "document.querySelector('#compare-task-list') !== null && !document.querySelector('#compare-task-list').textContent.includes('E2E 单处修改')", timeout=20)
                other_tasks = page.locator("#compare-task-list").inner_text()
                page.evaluate("localStorage.setItem('fundlex-dev-user','compare-e2e-user')")
                page.reload()
                page.goto(base + "/compare#compare")
                wait_until(page, "document.querySelector('#compare-task-list').textContent.includes('E2E 单处修改')", timeout=20)
                check("identity_switch_isolation", "E2E 单处修改" not in other_tasks and "E2E 单处修改" in page.locator("#compare-task-list").inner_text(), {"other_user_tasks": other_tasks})

                # Use the long document to measure line endpoints before and
                # after a viewport resize and synchronized left/right scrolling.
                page.goto(base + "/compare#compare")
                wait_until(page, "document.querySelector('#compare-task-list') !== null")
                page.locator("tr", has_text="E2E 二十处以上差异").first.locator("[data-compare-task]").click()
                wait_until(page, "document.querySelector('#compare-summary').textContent.includes('共')", timeout=30)
                page.wait_for_timeout(150)
                before_error, before_geometry = geometry(page)
                page.set_viewport_size({"width": 1100, "height": 800})
                page.wait_for_timeout(250)
                resize_error, resize_geometry = geometry(page)
                sync_result = page.locator("#compare-left-pane").evaluate("node => { const other=document.querySelector('#compare-right-pane'); const max=Math.max(0,node.scrollHeight-node.clientHeight); node.scrollTop=max*.47; return {leftTop:node.scrollTop,leftMax:max,rightTop:other.scrollTop,rightMax:Math.max(0,other.scrollHeight-other.clientHeight)}; }")
                page.wait_for_timeout(180)
                sync_result = page.evaluate("""() => {
                    const left=document.querySelector('#compare-left-pane'), right=document.querySelector('#compare-right-pane');
                    const max=(node)=>Math.max(0,node.scrollHeight-node.clientHeight);
                    return {leftTop:left.scrollTop,leftMax:max(left),rightTop:right.scrollTop,rightMax:max(right)};
                }""")
                sync_error = abs((sync_result["leftTop"] / sync_result["leftMax"] if sync_result["leftMax"] else 0) - (sync_result["rightTop"] / sync_result["rightMax"] if sync_result["rightMax"] else 0))
                reverse_sync = page.locator("#compare-right-pane").evaluate("node => { const other=document.querySelector('#compare-left-pane'); const max=Math.max(0,node.scrollHeight-node.clientHeight); node.scrollTop=max*.73; return true; }")
                page.wait_for_timeout(180)
                reverse_result = page.evaluate("""() => {
                    const left=document.querySelector('#compare-left-pane'), right=document.querySelector('#compare-right-pane');
                    const max=(node)=>Math.max(0,node.scrollHeight-node.clientHeight);
                    return {leftTop:left.scrollTop,leftMax:max(left),rightTop:right.scrollTop,rightMax:max(right)};
                }""")
                reverse_error = abs((reverse_result["leftTop"] / reverse_result["leftMax"] if reverse_result["leftMax"] else 0) - (reverse_result["rightTop"] / reverse_result["rightMax"] if reverse_result["rightMax"] else 0))
                scroll_error, scroll_geometry = geometry(page)
                check("svg_geometry_before_resize", before_error is not None and before_error <= 2, {"error_px": before_error, "geometry": before_geometry})
                check("svg_geometry_after_resize", resize_error is not None and resize_error <= 2, {"error_px": resize_error, "geometry": resize_geometry})
                check("two_panes_scroll_in_sync", sync_error <= 0.02, {"ratio_error": sync_error, "scroll": sync_result})
                check("reverse_scroll_stays_in_sync", reverse_error <= 0.02, {"ratio_error": reverse_error, "scroll": reverse_result})
                check("svg_geometry_after_synchronized_scroll", scroll_error is not None and scroll_error <= 2, {"error_px": scroll_error, "geometry": scroll_geometry})
                snapshot(page, "geometry-after-scroll")
                check("no_browser_page_errors", not page_errors, page_errors)
                browser.close()
        finally:
            stop_worker.set()
            worker_thread.join(timeout=3)
            server.should_exit = True
            server_thread.join(timeout=5)
            sock.close()

    evidence = {"base_url": base, "checks": checks, "scenarios": scenarios,
                "scenario_count": len(scenarios), "passed": all(item["passed"] for item in checks)}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    return evidence


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("docs/evidence/compare-e2e.json"))
    parser.add_argument("--screenshots", type=Path, default=Path("docs/evidence/compare-e2e-browser"))
    parser.add_argument("--corpus-base", type=Path, default=None)
    args = parser.parse_args()
    evidence = run(args.output, args.screenshots, args.corpus_base)
    failed = [item for item in evidence["checks"] if not item["passed"]]
    print(json.dumps({"output": str(args.output), "screenshots": str(args.screenshots), "scenario_count": evidence["scenario_count"], "passed": evidence["passed"], "failed": failed}, ensure_ascii=False, indent=2))
    return 0 if evidence["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
