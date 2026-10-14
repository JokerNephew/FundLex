"""Deterministic per-action audit used as release evidence (no network calls).

Builds independent DOCX fixtures, runs every catalogue action through the real
v2 engine with fixed profiles (scripted transport for the three generation
actions) and reports per-action positive/negative case counts, refusal reasons
and artifact validation results. Exit code 1 when the preregistered thresholds
(``docs/RELEASE_POLICY_V2_INTERNAL.md``) are not met.

Run::

    .venv\\Scripts\\python.exe -X utf8 -m tools.evaluation.action_fixture_audit --output <dir>
"""
from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from fundlex.docx import generate_artifacts, inspect_docx, validate_edit_plan
from fundlex.rules_v2 import load_ruleset_v2
from fundlex.store import engine_build
from fundlex.v2 import plan_v2

THRESHOLDS = {"positive_cases": 1, "negative_cases": 1, "fixture_pass_rate": 1.0, "unsafe_edits": 0}

R10_OLD = ("赎回费用由赎回基金份额的基金份额持有人承担，在基金份额持有人赎回基金份额时收取。"
           "赎回费用归入基金财产的比例应根据相关规定执行，并在招募说明书中列示。")
R10_GOOD = "赎回费用由赎回基金份额的基金份额持有人承担，在基金份额持有人赎回基金份额时收取。赎回费用全额计入基金财产。"
R02_CLAUSE = ("投资者通过直销机构（含基金管理人及其销售子公司，下同）申购本基金A类基金份额时，"
              "不收取认购/申购费用；通过直销机构申购的C类基金份额，不收取销售服务费。"
              "对于投资者通过直销机构之外的其他销售机构申购的C类基金份额且持续持有期限超过一年的，"
              "不再继续收取销售服务费。")


def profile(classes, subscription, service, *, floating=None, count=None):
    entries = []
    for name in classes:
        entries.append({"name": name,
                        "subscription_fee": True if name in subscription else (False if name in service else None),
                        "sales_service_fee": True if name in service else (False if name in subscription else None),
                        "evidence": [{"field": "subscription_fee" if name in subscription else "sales_service_fee",
                                      "quote": f"{name}类基金份额收取{'申购费用' if name in subscription else '销售服务费'}。"}]})
    return {"share_classes": classes, "classes": entries,
            "fee_groups": {"subscription_group": subscription, "service_group": service},
            "fee_groups_verified": True, "class_universe_verified": True,
            "share_class_count": count or len(classes), "share_class_count_verified": True,
            "share_class_count_min": count or len(classes), "floating_management_fee": floating}


MULTI = profile(["A", "C"], ["A"], ["C"], floating=True)
SINGLE = profile(["A"], ["A"], [], floating=True, count=1)
MULTI_NONFLOATING = profile(["A", "C"], ["A"], ["C"], floating=False)
MULTI_NO_GROUPS = profile(["A", "C"], ["A"], ["C"], floating=True)
MULTI_NO_GROUPS["fee_groups_verified"] = False


class Scripted:
    """Transport double for the three generation actions (wiring evidence only)."""

    def __init__(self, responses):
        self.responses = list(responses)
        self.model = "scripted-audit"
        self.max_context_chars = 12000
        self.calls = 0
        self.events: list[dict] = []

    def _take(self, task_type):
        if not self.responses and task_type in {"locate_target", "resolve_scope"}:
            # Unscripted discovery is deliberately inconclusive. It must never
            # invent a target or approve a write to make a negative case pass.
            self.calls += 1
            value = {"quotes": []} if task_type == "locate_target" else {"status": "uncertain", "ranges": []}
            return SimpleNamespace(value=value, model=self.model, attempts=0)
        expected, value = self.responses.pop(0)
        assert expected == task_type, f"expected {expected}, got {task_type}"
        self.calls += 1
        event = {"task_type": task_type, "prompt_version": "scripted", "actual_model": self.model,
                 "request_id": f"audit-{self.calls}", "latency_ms": 0, "http_status": 200, "error_code": None}
        self.events.append(event)
        return SimpleNamespace(value=value, model=self.model, latency_ms=0, prompt_tokens=0,
                               completion_tokens=0, total_tokens=0, request_id=event["request_id"],
                               attempts=1, estimated_cost_usd=0.0, task_type=task_type)

    def locate_target(self, context, *, instruction, **kwargs):
        return self._take("locate_target")

    def resolve_scope(self, context, *, instruction, **kwargs):
        return self._take("resolve_scope")

    def generate_clause(self, context, *, instruction, **kwargs):
        return self._take("generate_clause")

    def verify_content(self, context, *, instruction, **kwargs):
        return self._take("verify_content")


def old_texts(action_id):
    action = load_ruleset_v2().by_id(action_id)
    return action["locate"].get("old_texts") or action["locate"]["triggers"]


def build_fixtures(root: Path) -> dict[str, Path]:
    root.mkdir(parents=True, exist_ok=True)
    multi = root / "multi.docx"
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("释义", level=1)
    document.add_paragraph("基金份额类别")
    document.add_paragraph("基金份额类别：本基金将基金份额分为A类和C类不同的类别。")
    document.add_heading("基金的基本情况", level=1)
    document.add_paragraph("基金的类别：混合型证券投资基金")
    document.add_paragraph("份额类别设置")
    document.add_paragraph("本基金将基金份额分为 A 类和 C 类两种不同的类别。在投资者申购 A 类基金份额时收取申购费用，"
                           "不计提销售服务费；在投资者申购 C 类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费。"
                           "基金管理人可根据实际情况调整并报中国证监会备案。")
    document.add_paragraph("除非基金管理人在未来条件成熟后另行公告开通相关业务，本基金不同基金份额类别之间不得互相转换。")
    document.add_paragraph("有关基金份额类别的具体设置、费率水平等由基金管理人确定。根据基金销售情况，"
                           "基金管理人在履行适当程序后可以增加新的基金份额类别、或者在基金合同规定的范围内"
                           "变更现有基金份额类别的申购费率、调低赎回费率或变更收费方式。")
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    for text in old_texts("R06-A1") + old_texts("R07-A1") + old_texts("R09-A1")[:1] + old_texts("R13-A1")[:1]:
        document.add_paragraph(text)
    document.add_paragraph("基金管理人可以在基金合同约定的范围内调整费率或收费方式。"
                           "本基金对持续持有期少于 7 日的投资者收取不低于 1.50% 的赎回费，并将上述赎回费全额计入基金财产。")
    document.add_heading("基金费用与税收", level=1)
    document.add_paragraph("基金销售服务费")
    for text in old_texts("R14-A1") + old_texts("R15-A1") + old_texts("R15-A2"):
        document.add_paragraph(text)
    document.add_heading("基金的收益与分配", level=1)
    document.add_paragraph("收益分配")
    document.add_paragraph(old_texts("R16-A1")[0] + "，各基金份额类别对应的可供分配利润将有所不同。")
    document.add_heading("基金合同的变更、终止与基金财产的清算", level=1)
    document.add_paragraph("基金财产清算剩余资产的分配")
    document.add_paragraph(old_texts("R17-A2")[0])
    document.add_heading("基金份额的转换", level=1)
    document.add_paragraph("转换业务：本基金不同基金份额类别之间不得互相转换。")
    document.save(multi)

    single = root / "single.docx"
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金的基本情况", level=1)
    document.add_paragraph("基金的类别：混合型证券投资基金")
    document.add_paragraph("份额类别设置")
    document.add_paragraph("本基金不区分份额类别，仅设一种基金份额。")
    document.add_paragraph("有关基金份额类别的具体设置、费率水平等由基金管理人确定。根据基金销售情况，"
                           "基金管理人在履行适当程序后可以增加新的基金份额类别、或者在基金合同规定的范围内"
                           "变更现有基金份额类别的申购费率、调低赎回费率或变更收费方式。")
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    for text in old_texts("R05-A1") + old_texts("R08-A1") + old_texts("R12-A1")[:1]:
        document.add_paragraph(text)
    document.add_heading("基金合同的变更、终止与基金财产的清算", level=1)
    document.add_paragraph("基金财产清算剩余资产的分配")
    document.add_paragraph(old_texts("R17-A1")[0])
    document.save(single)

    r17a3 = root / "r17a3.docx"
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金的基本情况", level=1)
    document.add_paragraph("基金的类别：混合型证券投资基金")
    document.add_paragraph("份额类别设置")
    document.add_paragraph("本基金将基金份额分为 A 类和 C 类两种不同的类别。在投资者申购 A 类基金份额时收取申购费用，"
                           "不计提销售服务费；在投资者申购 C 类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费。")
    document.add_heading("基金合同的变更、终止与基金财产的清算", level=1)
    document.add_paragraph("基金财产清算剩余资产的分配")
    document.add_paragraph(old_texts("R17-A3")[0])
    document.save(r17a3)

    r10 = root / "r10.docx"
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    document.add_paragraph(R10_OLD)
    document.save(r10)

    r11 = root / "r11.docx"
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    document.add_paragraph("赎回费用由赎回基金份额的基金份额持有人承担。" + "本基金对持续持有期少于 7 日的投资者收取不低于 1.50% 的赎回费，"
                           "并将上述赎回费全额计入基金财产。")
    document.save(r11)

    mismatch = root / "mismatch.docx"
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金费用与税收", level=1)
    document.add_paragraph("基金销售服务费")
    document.add_paragraph(old_texts("R15-A1")[0] + "本段还包含其他必须保留的内容。")
    document.save(mismatch)
    return {"multi": multi, "single": single, "r17a3": r17a3, "r10": r10, "r11": r11, "mismatch": mismatch}


def run_action(action_id, doc, profile_value, *, expected, responses=()):
    inspection = inspect_docx(doc, full=True)
    gateway = Scripted(list(responses)) if responses else Scripted([])
    detail: dict = {}
    try:
        outcome = plan_v2(inspection, profile_value, load_ruleset_v2(), gateway,
                          engine_build_value=engine_build(), actions=[action_id])
        action = next(item for item in outcome["actions"] if item["action_id"] == action_id)
        edits = outcome["edits"]
    except Exception as exc:  # noqa: BLE001 - audit records the failure
        return False, False, {"error": type(exc).__name__, "code": getattr(exc, "code", None)}
    passed = False
    artifacts_ok = False
    detail = {"status": action["status"], "failure_reason": action.get("failure_reason"),
              "edits": len(edits),
              "target_reasons": [t.get("reason") for t in action["targets"] if t.get("reason")]}
    if expected == "applied":
        passed = action["status"] == "applied" and len(edits) == 1
        if passed and responses:
            passed = str(edits[0]["replacement"]) == responses[-2][1]["replacement_text"]
        if passed:
            try:
                plan = validate_edit_plan(doc, edits, {}, content_check=outcome["content_check"])
                generate_artifacts(doc, plan, doc.parent / "r.docx", doc.parent / "f.docx")
                artifacts_ok = True
            except Exception as exc:  # noqa: BLE001 - audit records the failure
                detail["artifact_error"] = getattr(exc, "code", type(exc).__name__)
    else:
        passed = action["status"] in {"not_applicable", "failed", "not_matched"} and not edits
    return passed, artifacts_ok, detail


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    root = Path(args.output).resolve()
    fixtures = build_fixtures(root / "fixtures")

    cases = [
        # (action, fixture, profile, expectation, scripted responses)
        ("R03-A3", "multi", MULTI, "applied", ()),
        ("R04-A1", "multi", MULTI, "applied", ()),
        ("R11-A1", "multi", MULTI, "applied", ()),
        ("R01-A1", "multi", MULTI, "applied",
         (("generate_clause", {"source_quote": "基金份额类别：本基金将基金份额分为A类和C类不同的类别。",
                               "replacement_text": R02_CLAUSE, "facts": [], "changes": []}),
          ("verify_content", {"items": [{"requirement_id": "R01-REQ-5", "result": "satisfied",
                                         "evidence_quote": "不收取认购/申购费用", "evidence_source": "generated"}]}))),
        ("R02-A1", "multi", MULTI, "applied",
         (("generate_clause", {"source_quote": "在投资者申购 C 类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费。",
                               "replacement_text": R02_CLAUSE, "facts": [], "changes": []}),
          ("verify_content", {"items": [{"requirement_id": "R02-REQ-5", "result": "satisfied",
                                         "evidence_quote": "不收取认购/申购费用", "evidence_source": "generated"}]}))),
        ("R10-A1", "r10", MULTI, "applied",
         (("locate_target", {"quotes": [{"quote": R10_OLD, "basis": "费用归属条款"}]}),
          ("generate_clause", {"source_quote": R10_OLD, "replacement_text": R10_GOOD, "facts": [], "changes": []}),
          ("verify_content", {"items": [{"requirement_id": "R10-REQ-3", "result": "satisfied",
                                         "evidence_quote": "赎回费用全额计入基金财产", "evidence_source": "generated"},
                                        {"requirement_id": "R10-REQ-6", "result": "satisfied",
                                         "evidence_quote": "赎回费用由赎回基金份额的基金份额持有人承担", "evidence_source": "contract"}]}))),
        ("R03-A1", "multi", MULTI, "applied", ()),
        ("R03-A2", "single", SINGLE, "applied", ()),
        ("R05-A1", "single", SINGLE, "applied", ()),
        ("R06-A1", "multi", MULTI, "applied", ()),
        ("R07-A1", "multi", MULTI, "applied", ()),
        ("R08-A1", "single", SINGLE, "applied", ()),
        ("R09-A1", "multi", MULTI, "applied", ()),
        ("R12-A1", "single", SINGLE, "applied", ()),
        ("R13-A1", "multi", MULTI, "applied", ()),
        ("R14-A1", "multi", MULTI, "applied", ()),
        ("R15-A1", "multi", MULTI, "applied", ()),
        ("R15-A2", "multi", MULTI, "applied", ()),
        ("R16-A1", "multi", MULTI, "applied", ()),
        ("R17-A1", "single", SINGLE, "applied", ()),
        ("R17-A2", "multi", MULTI, "applied", ()),
        ("R17-A3", "r17a3", MULTI_NONFLOATING, "applied", ()),
        # Negative cases: conditions or scope must refuse without edits.
        ("R03-A1", "single", SINGLE, "refused", ()),
        ("R03-A2", "multi", MULTI, "refused", ()),
        ("R05-A1", "multi", MULTI, "refused", ()),
        ("R06-A1", "single", SINGLE, "refused", ()),
        ("R07-A1", "single", SINGLE, "refused", ()),
        ("R08-A1", "multi", MULTI, "refused", ()),
        ("R09-A1", "single", SINGLE, "refused", ()),
        ("R12-A1", "multi", MULTI, "refused", ()),
        ("R13-A1", "single", SINGLE, "refused", ()),
        ("R14-A1", "single", SINGLE, "refused", ()),
        ("R15-A1", "mismatch", MULTI, "refused", ()),
        ("R16-A1", "single", SINGLE, "refused", ()),
        ("R17-A1", "multi", MULTI, "refused", ()),
        ("R17-A3", "multi", MULTI, "refused", ()),
        ("R06-A1", "multi", MULTI_NO_GROUPS, "refused", ()),
        ("R04-A1", "single", SINGLE, "refused", ()),
        ("R11-A1", "r11", MULTI, "refused", ()),
        ("R10-A1", "multi", MULTI, "refused", ()),
        ("R01-A1", "single", SINGLE, "refused", ()),
        ("R02-A1", "single", SINGLE, "refused", ()),
        ("R03-A3", "single", SINGLE, "refused", ()),
        ("R15-A2", "single", SINGLE, "refused", ()),
        ("R17-A2", "single", SINGLE, "refused", ()),
    ]
    report: dict[str, dict] = {}
    for action_id, fixture_key, profile_value, expectation, responses in cases:
        doc = fixtures[fixture_key]
        passed, artifacts_ok, detail = run_action(action_id, doc, profile_value,
                                                  expected=expectation, responses=responses)
        entry = report.setdefault(action_id, {"positive_cases": 0, "positive_passed": 0,
                                              "negative_cases": 0, "negative_passed": 0,
                                              "unsafe_edits": 0, "artifact_checks": None,
                                              "cases": []})
        kind = "positive" if expectation == "applied" else "negative"
        entry[f"{kind}_cases"] += 1
        if passed:
            entry[f"{kind}_passed"] += 1
        if expectation == "applied":
            entry["artifact_checks"] = 1.0 if artifacts_ok else 0.0
        else:
            entry.setdefault("artifact_checks", 1.0)
        entry["cases"].append({"fixture": fixture_key, "expectation": expectation,
                               "passed": passed, **detail})
    total_positive = total_negative = passed_positive = passed_negative = 0
    failures = []
    for action_id, entry in sorted(report.items()):
        entry["fixture_pass_rate"] = round((entry["positive_passed"] + entry["negative_passed"])
                                           / max(1, entry["positive_cases"] + entry["negative_cases"]), 4)
        total_positive += entry["positive_cases"]
        total_negative += entry["negative_cases"]
        passed_positive += entry["positive_passed"]
        passed_negative += entry["negative_passed"]
        if (entry["positive_cases"] < THRESHOLDS["positive_cases"] or entry["negative_cases"] < THRESHOLDS["negative_cases"]
                or entry["fixture_pass_rate"] < THRESHOLDS["fixture_pass_rate"]
                or (entry["artifact_checks"] or 0) < 1.0):
            failures.append(action_id)
    summary = {"schema": "action-fixture-audit-1",
               "generated_at": datetime.now(timezone.utc).isoformat(),
               "engine_build": engine_build(), "ruleset_sha256": load_ruleset_v2().sha256,
               "actions_audited": len(report), "positive_cases": total_positive,
               "positive_passed": passed_positive, "negative_cases": total_negative,
               "negative_passed": passed_negative, "thresholds": THRESHOLDS,
               "failed_actions": failures, "actions": report}
    root.mkdir(parents=True, exist_ok=True)
    (root / "audit.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in summary.items() if k != "actions"}, ensure_ascii=False))
    if failures:
        raise SystemExit(f"audit failed for: {failures}")


if __name__ == "__main__":
    main()
