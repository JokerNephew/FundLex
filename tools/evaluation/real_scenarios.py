"""Real ModelGateway acceptance through HTTP submission and the actual worker.

No credentials, request bodies, or full responses are written to the summary.
Generated DOCX inputs and diagnostics remain in the ignored evaluation folder.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import shutil
import os
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient
from fundlex.api import create_app
from fundlex.jobs import Worker
from fundlex.model import ModelGateway
from fundlex.pipeline import process
from fundlex.config import local_data_dir

def build_contract(path: Path, existing_revision: bool = False):
    doc=Document()
    for text in ('独立边界示例证券投资基金基金合同','第四部分 基金的基本情况',
                 '二、基金的类别','混合型证券投资基金','三、基金的运作方式','契约型开放式',
                 '四、基金份额类别设置','本基金不区分份额类别。',
                 '第十二部分 基金的投资','本基金投资于债券的比例不低于基金资产的80%。',
                 '第六部分 基金份额的申购与赎回','申购费用','申购费用由投资人承担，不列入基金财产。'):
        doc.add_paragraph(text)
    if existing_revision:
        insertion=OxmlElement('w:ins')
        insertion.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id','33')
        run=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text='未结修订'
        run.append(text); insertion.append(run); doc.paragraphs[-1]._p.append(insertion)
    doc.save(path)

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--scenario',choices=['normal','structure_failure','truncated'],required=True)
    parser.add_argument('--output',default='var/evaluation/scenarios')
    args=parser.parse_args()
    root=Path(args.output)/args.scenario
    root.mkdir(parents=True,exist_ok=True)
    source=root/'input.docx';build_contract(source,args.scenario=='structure_failure')
    source_hash=hashlib.sha256(source.read_bytes()).hexdigest()
    gateway=ModelGateway.from_env()
    if args.scenario=='truncated': gateway.max_completion_tokens=1
    state_root=local_data_dir()/'eval'/hashlib.sha256(str(root.resolve()).encode()).hexdigest()[:32]
    app=create_app(state_root); client=TestClient(app)
    headers={'X-FundLex-User':'acceptance-real'}
    created=client.post('/api/batches',headers=headers,data={'ruleset_version':'v1'},
        files=[('files',('input.docx',source.read_bytes()))])
    created.raise_for_status(); batch_id=created.json()['id']
    diagnostics=root/'diagnostics';diagnostics.mkdir(exist_ok=True)
    def processor(input_path,output_dir,version):
        report=process(input_path,output_dir,version,gateway)
        for item in output_dir.iterdir():
            if item.is_file():shutil.copy2(item,diagnostics/item.name)
        return report
    assert Worker(app.state.store,processor).run_once()
    detail=client.get(f'/api/batches/{batch_id}',headers=headers).json()
    document=detail['documents'][0]
    response=client.get(document['report_url'],headers=headers)
    if response.status_code!=200:
        raise AssertionError(f"Report not ready: {document['job_status']} {document.get('error_code')}")
    report=response.json()
    assert hashlib.sha256(source.read_bytes()).hexdigest()==source_hash
    assert gateway.metrics()['http_attempts']>=1
    assert report['delivery_status']!='complete'
    assert 'final' not in document['downloads'] and 'redline' not in document['downloads']
    if args.scenario=='normal':
        assert report['qualification_status']=='eligible',report.get('error_code')
        assert report['edit_plan']['validated']
        assert report['output_validation']['not_a_final_contract']
        assert (diagnostics/'edits.json').is_file()
        assert (diagnostics/'evaluation_final.docx').is_file()
    else:
        assert report.get('error_code'),report
        assert report['delivery_status']=='not_generated'
        assert not (diagnostics/'evaluation_final.docx').exists()
    summary={'scenario':args.scenario,'batch_id':batch_id,'document_id':document['document_id'],
        'source_sha256':source_hash,'qualification_status':report['qualification_status'],
        'delivery_status':report['delivery_status'],'error_code':report.get('error_code'),
        'model':gateway.metrics(),'budget':gateway.budget.snapshot() if gateway.budget else None,
        'ruleset_sha256':report.get('ruleset_sha256'),'plan':report.get('edit_plan'),
        'output_validation':report.get('output_validation'),'source_unchanged':True,
        'api_report_download_verified':client.get(document['downloads']['report_json'],headers=headers).status_code==200,
        'final_delivery_withheld':True}
    (root/'acceptance.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    gateway.close()
    print(json.dumps(summary,ensure_ascii=False))

if __name__=='__main__':main()
