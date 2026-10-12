"""Persistent fixed-template fixtures, independent of core tests and model output."""
from pathlib import Path
from hashlib import sha256
from dataclasses import asdict
import json, sys, zipfile
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from fundlex.docx import inspect_docx, validate_edit_plan, generate_artifacts, DocxError, Q, NS

def main():
    folder=Path('var/evaluation/word_fixture');folder.mkdir(parents=True,exist_ok=True)
    source=folder/'source.docx'; doc=Document()
    doc.sections[0].header.paragraphs[0].text='HEADER_UNCHANGED'
    doc.sections[0].footer.paragraphs[0].text='FOOTER_UNCHANGED'
    doc.add_paragraph('目录')
    p=doc.add_paragraph(); field=OxmlElement('w:fldSimple');field.set(Q('instr'),' TOC \\o "1-3" \\h ');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='第一部分 正文……1';r.append(t);field.append(r);p._p.append(field)
    doc.add_heading('第一部分 正文',1)
    p=doc.add_paragraph('跨run前缀：');p.add_run('旧跨').bold=True;p.add_run('run条款').italic=True;p.add_run('；尾部保留。')
    doc.add_paragraph('NON_TARGET_原有重复文本与收费条款。')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='普通表格旧条款';table.cell(0,1).text='TABLE_NON_TARGET'
    p=doc.add_paragraph('原生换行前缀');p.add_run().add_break();p.add_run('换行后旧条款')
    doc.add_paragraph('追加锚点。')
    doc.add_heading('内容摘要',1);doc.add_paragraph('摘要旧条款')
    doc.add_heading('附录',1);doc.add_paragraph('APPENDIX_NON_TARGET')
    doc.save(source); inspected=inspect_docx(source,full=True)
    edits=[];templates={}
    cases=[('cross_run','旧跨run条款','固定跨run新条款'),('table','普通表格旧条款','固定表格新条款'),('native_newline','换行后旧条款','固定第一行\n固定第二行'),('summary','摘要旧条款','固定摘要新条款'),('append_zero_width','追加锚点。','固定追加条款。')]
    for name,old,replacement in cases:
        p=next(p for p in inspected.paragraphs if old in p.text);start=p.text.index(old);end=start+len(old);insertion=name=='append_zero_width'
        edit={'edit_id':name,'source_sha256':inspected.source_sha256,'location':p.location,'zone':p.zone,'start':end if insertion else start,'end':end,'expected_old_text':'' if insertion else old,'expected_old_hash':sha256(('' if insertion else old).encode()).hexdigest(),'template_id':name,'replacement':replacement,'verified':True,'fingerprint':p.fingerprint,'mapping_version':p.mapping_version}
        if insertion:edit.update(operation='insert',anchor={'start':start,'end':end,'text':old,'sha256':sha256(old.encode()).hexdigest()})
        edits.append(edit);templates[name]=replacement
    plan=validate_edit_plan(source,edits,templates);result=generate_artifacts(source,plan,folder/'redline.docx',folder/'final.docx',author='FundLexAcceptance')
    # Independent non-target canonical XML comparison, without editor snapshot helpers.
    with zipfile.ZipFile(source) as z: original={n:z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(folder/'final.docx') as z: final={n:z.read(n) for n in z.namelist()}
    a=etree.fromstring(original['word/document.xml']);b=etree.fromstring(final['word/document.xml']);target_texts=[old for _,old,_ in cases]
    non_targets=[(idx,p) for idx,p in enumerate(a.xpath('//w:p',namespaces=NS)) if not any(old in ''.join(p.xpath('.//w:t/text()',namespaces=NS)) for old in target_texts)]
    ps=b.xpath('//w:p',namespaces=NS)
    non_target_ok=all(etree.tostring(p,method='c14n')==etree.tostring(ps[idx],method='c14n') for idx,p in non_targets)
    untouched_parts=all(final[n]==v for n,v in original.items() if n!='word/document.xml')
    toc=next(p for p in inspected.paragraphs if p.zone=='toc')
    probe=dict(edits[0],location=toc.location,zone=toc.zone)
    try:validate_edit_plan(source,[probe],templates);toc_error='NOT_REJECTED'
    except DocxError as exc:toc_error=exc.code
    newline_p=next(p for p in inspected.paragraphs if '原生换行前缀' in p.text)
    crossing=dict(edits[2],start=0,end=len(newline_p.text),expected_old_text=newline_p.text,expected_old_hash=sha256(newline_p.text.encode()).hexdigest())
    try:validate_edit_plan(source,[crossing],templates);newline_error='NOT_REJECTED'
    except DocxError as exc:newline_error=exc.code
    metadata={'cases':[c[0] for c in cases],'source_sha256':inspected.source_sha256,'source_unchanged':sha256(source.read_bytes()).hexdigest()==inspected.source_sha256,'validation':asdict(result),'independent_non_target_paragraphs':len(non_targets),'independent_non_target_equal':non_target_ok,'non_document_parts_equal':untouched_parts,'toc_edit_rejection':toc_error,'native_newline_crossing_rejection':newline_error,'toc_updated':False,'edits':edits,'templates':templates}
    (folder/'fixture.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(metadata,ensure_ascii=False))
if __name__=='__main__':main()
