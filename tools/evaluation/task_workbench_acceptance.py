"""Isolated browser acceptance server; synthetic DOCX and processor, no models.

The production frontend and HTTP/store/worker paths are real. Business reports
are deliberately marked fixtures, never evidence of model quality or approval.
No final/redline contract is generated or exposed.
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path
import sys
import threading
import time


def prepare(root: Path, count: int) -> Path:
    from docx import Document

    folder = root / "fixture-inputs"
    folder.mkdir(parents=True, exist_ok=True)
    names = ["正常资格", "类型排除", "证据失败关闭", "结果不完整"]
    for index in range(count):
        case = index % 4
        document = Document()
        document.add_paragraph(f"WORKBENCH_ACCEPTANCE_FIXTURE:{case}:{index}")
        document.add_heading("工作台交互验收构造合同", level=1)
        document.add_paragraph("仅验证文件上传、队列、报告呈现与下载，不含真实合同资料。")
        document.add_paragraph("这是明确标注的测试证据，用于检查中文长文本换行和证据详情阅读体验。" * 12)
        prefix = "工作台验收_" + names[case]
        if index == 3:
            prefix += "_广发粤港澳大湾区创新产业与多资产配置养老目标日期基金合同长文件名展示"
        document.save(folder / f"{prefix}_{index + 1:03d}.docx")
    return folder


def fixture_processor(source: Path, output: Path, ruleset: str) -> dict:
    from docx import Document

    paragraphs = [p.text for p in Document(source).paragraphs]
    marker = next(text for text in paragraphs if text.startswith("WORKBENCH_ACCEPTANCE_FIXTURE:"))
    case = int(marker.split(":")[1])
    # Allow a browser to observe actual queued/running transitions.
    time.sleep(0.08)
    quote = paragraphs[-1]
    report = {
        "acceptance_fixture": True,
        "fixture_notice": "构造 DOCX / 替身处理器报告；没有模型调用，不是业务准确率验收",
        "qualification_status": "eligible",
        "qualification": {"status": "eligible", "fund_type": "ordinary_fund",
                          "evidence": {"quote": quote, "part_uri": "word/document.xml", "paragraph_id": "p4"}},
        "delivery_status": "not_generated",
        "release_gate": {"status": "closed", "auto_write_actions": []},
        "edit_plan": {"count": 0, "executable": False, "reason": "RELEASE_GATE_CLOSED"},
        "has_toc_field": False,
        "toc_update_status": "not_applicable",
        "profile": {"classes": [], "fee_groups_verified": False},
        "actions": [], "errors": [], "error_code": None,
        "model": {"calls": 0, "configured_model": "acceptance-fixture-no-model"},
        "artifacts": {},
    }
    if case == 1:
        report["qualification_status"] = "excluded_by_fund_type"
        report["qualification"].update(status="excluded_by_fund_type", fund_type="money_market")
    elif case == 2:
        report["qualification_status"] = "failed"
        report["qualification"].update(status="failed", fund_type="unknown")
        report["error_code"] = "FUND_TYPE_EVIDENCE_AMBIGUOUS"
        report["errors"] = [{"code": "FUND_TYPE_EVIDENCE_AMBIGUOUS", "stage": "qualification"}]
    else:
        reason = "PROFILE_EVIDENCE_INCOMPLETE" if case == 3 else "ACTION_RELEASE_GATE_CLOSED"
        if case == 3:
            report["delivery_status"] = "incomplete"
        report["actions"] = [{"action_id": "R03-A1", "status": "failed", "auto_write": False,
                              "failure_reason": reason, "candidates": [{"status": "failed", "reason": reason,
                              "quote": quote, "region": "body", "section_path": "验收夹具 / 长文本证据",
                              "location": "word/document.xml#p4", "char_start": 0, "char_end": len(quote)}]}]
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", required=True, type=Path)
    parser.add_argument("--port", type=int, default=8766)
    parser.add_argument("--count", type=int, default=240)
    parser.add_argument("--prepare-only", action="store_true")
    args = parser.parse_args()
    root = args.data_dir.resolve()
    if args.count < 4 or args.count > 500:
        parser.error("count must be between 4 and 500 synthetic fixtures")
    root.mkdir(parents=True, exist_ok=True)
    folder = prepare(root, args.count)
    print(f"Synthetic fixtures: {folder}", flush=True)
    if args.prepare_only:
        return
    # Set isolated storage BEFORE importing api: its module creates an app.
    os.environ["FUNDLEX_DATA_DIR"] = str(root / "runtime")
    os.environ["FUNDLEX_AUTH_MODE"] = "dev"
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
    from fundlex.api import create_app
    from fundlex.jobs import Worker
    import uvicorn

    app = create_app(root / "runtime")
    worker = Worker(app.state.store, fixture_processor)
    stop = threading.Event()

    def work() -> None:
        while not stop.is_set():
            if not worker.run_once():
                stop.wait(0.2)

    thread = threading.Thread(target=work, daemon=True)
    thread.start()
    print("Fixture processor only. No model calls. No production data.", flush=True)
    try:
        uvicorn.run(app, host="127.0.0.1", port=args.port, log_level="warning")
    finally:
        stop.set()
        thread.join(timeout=5)


if __name__ == "__main__":
    main()
