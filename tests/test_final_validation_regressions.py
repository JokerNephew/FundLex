"""Generic defects reproduced by exact-HEAD live validation and independent review."""
import json
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import httpx
import pytest

from fundlex.model import ModelError, ModelGateway
from fundlex.recall import Target, dedupe_operation_targets, plan_target
from fundlex.generation.claim_replay import replay_claims
from fundlex.obligations import compile_obligations


def request():
    return dict(task_type='recall_inventory', bounded_context='public contract',
                response_schema={'type': 'object', 'required': ['ok']},
                instruction='extract', prompt_version='v1')


def test_record_single_flight_full_request_and_failure_replay(tmp_path):
    calls = []
    def provider(req):
        calls.append(req)
        payload = json.loads(req.content)
        if 'failure' in payload['messages'][1]['content']:
            content, reason = '{"ok":', 'length'
        else:
            content, reason = json.dumps({'ok': len(calls)}), 'stop'
        return httpx.Response(200, json={'model': 'm', 'choices': [
            {'message': {'content': content}, 'finish_reason': reason}]})
    with ModelGateway(base_url='https://example.invalid', model='m', api_key='secret',
            mode='record', cassette_dir=tmp_path, max_calls=30,
            client=httpx.Client(transport=httpx.MockTransport(provider))) as gateway:
        with ThreadPoolExecutor(max_workers=4) as pool:
            values = list(pool.map(lambda _: gateway.call(**request()).value, range(4)))
        assert values == [{'ok': 1}] * 4
        assert len(calls) == 1
        # Same context with a different instruction is a distinct full request.
        assert gateway.call(**{**request(), 'instruction': 'other'}).value == {'ok': 2}
        with pytest.raises(ModelError, match='MODEL_OUTPUT_TRUNCATED'):
            gateway.call(**{**request(), 'bounded_context': 'failure'})
        assert len(calls) == 3
    with ModelGateway(base_url='', model='m', api_key='', mode='replay', cassette_dir=tmp_path) as gateway:
        assert gateway._client is None
        assert gateway.call(**request()).value == {'ok': 1}
        with pytest.raises(ModelError, match='MODEL_OUTPUT_TRUNCATED'):
            gateway.call(**{**request(), 'bounded_context': 'failure'})
        with pytest.raises(ModelError, match='MODEL_REPLAY_MISS'):
            gateway.call(**{**request(), 'bounded_context': 'missing'})
        assert gateway.metrics()['http_attempts'] == 0


def test_role_replay_needs_models_but_no_credentials(tmp_path, monkeypatch):
    from fundlex.models.roles import build_role_gateway
    for name in list(__import__('os').environ):
        if name.startswith(('LLM_', 'FUNDLEX_PRIMARY_', 'FUNDLEX_REVIEW_', 'FUNDLEX_MODEL_')):
            monkeypatch.delenv(name)
    monkeypatch.setenv('FUNDLEX_MODEL_MODE', 'replay')
    monkeypatch.setenv('FUNDLEX_MODEL_CASSETTE_DIR', str(tmp_path))
    monkeypatch.setenv('FUNDLEX_PRIMARY_MODEL', 'm1')
    monkeypatch.setenv('FUNDLEX_REVIEW_MODEL', 'm2')
    with build_role_gateway() as gateway:
        assert gateway.primary._client is gateway.review._client is None
        assert gateway.budget is None
        for task in ('recall_facts', 'recall_audit'):
            with pytest.raises(ModelError, match='MODEL_REPLAY_MISS'):
                gateway.call(**{**request(), 'task_type': task})


class Inspection:
    def __init__(self, text):
        self.paragraphs = [SimpleNamespace(location='p0', text=text)]
    def by_location(self, _):
        return self.paragraphs[0]


def test_dedupe_keeps_disjoint_required_occurrences_and_claim_blocks_omission():
    old, new = 'A类基金份额', 'C类基金份额'
    before = '对于收取申购费的A类基金份额应披露费率。另对于收取申购费的A类基金份额应披露申购费用。'
    rule = {'rule_id': 'G02', 'class_substitution': True, 'title': '份额类别', 'operations': [
        {'op': 'replace', 'old_text': old, 'new_text': new, 'condition': {'kind': 'none'},
         'operation_id': 'G02:0', 'group_kind': 'INDEPENDENT'}]}
    a = before.index(old)
    b = before.index(old, a + 1)
    targets = [Target(f'T{i}', 'G02', 'p0', old, 'needs_edit', 'verified occurrence',
        start=p, end=p+len(old), operations=('G02:0',), occurrence_index=i,
        resolution='resolved', segment_ok=True) for i, p in enumerate((a, b))]
    source = Inspection(before)
    _, removed = dedupe_operation_targets(targets, {'G02': rule}, source)
    assert removed == []
    ledger = compile_obligations({'groups': [rule]})
    for record in ledger.records:
        record.applicability = SimpleNamespace(status='TRUE')
    edit = {'rule_id': 'G02', 'target_id': 'T0', 'location': 'p0', 'start': a,
            'end': a+len(old), 'replacement': new, 'operation_keys': ['G02:0'], 'content_source': 'guide'}
    result = replay_claims([edit], source, Inspection(before[:a]+new+before[a+len(old):]),
        rules_by_id={'G02': rule}, targets_by_id={t.target_id: t.as_dict() for t in targets}, obligations=ledger)
    assert result['verified'] is False
    assert result['failures']


def test_delete_prefix_site_cannot_prove_continuation_absent():
    """A split delete reference must not pass from a prefix-only site quote."""
    from fundlex.generation.effects import evaluate_operation_effect

    old = ('本基金对持续持有期少于7日的投资者收取不低于1.50%的赎回费，'
           '并将上述赎回费全额计入基金财产。')
    rule = {'rule_id': 'G13', 'title': '删除持续持有期条款', 'operations': [
        {'operation_id': 'G13:0', 'op': 'delete', 'old_text': old,
         'condition': {'kind': 'none'}, 'group_kind': 'INDEPENDENT'}]}
    from fundlex.obligations import compile_obligations
    from test_operation_execution_closure import Facts
    ledger = compile_obligations({'groups': [rule]})
    record = ledger.records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    class MultiInspection:
        def __init__(self, texts):
            self.paragraphs = [SimpleNamespace(location=f'p{i}', text=text)
                               for i, text in enumerate(texts)]
        def by_location(self, location):
            return next(item for item in self.paragraphs if item.location == location)
    before = MultiInspection([
        '赎回费用由持有人承担。其中，对持续持有期少于',
        '7 日的投资者收取不低于 1.5%的赎回费，并全额计入基金财产。'])
    after = MultiInspection([
        '赎回费用由持有人承担。赎回费用全额计入基金财产。',
        '7 日的投资者收取不低于 1.5%的赎回费，并全额计入基金财产。'])
    proof = evaluate_operation_effect(record, before, after, [
        {'location': 'p0', 'quote': '其中，对持续持有期少于'}], [], Facts())
    assert proof.state == 'FAIL'
    assert proof.artifact_status == 'RESIDUAL_OLD_TEXT'
    assert any(site['location'] == 'p1' and site['state'] == 'FAIL'
               for site in proof.site_effects)


def test_delete_full_site_quote_can_prove_already_absent():
    from fundlex.generation.effects import evaluate_operation_effect
    from fundlex.obligations import compile_obligations
    from test_operation_execution_closure import Facts
    old = '本基金对持续持有期少于7日的投资者收取不低于1.50%的赎回费。'
    rule = {'rule_id': 'G13', 'title': '删除持续持有期条款', 'operations': [
        {'operation_id': 'G13:0', 'op': 'delete', 'old_text': old,
         'condition': {'kind': 'none'}, 'group_kind': 'INDEPENDENT'}]}
    record = compile_obligations({'groups': [rule]}).records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    before, after = Inspection(old), Inspection('其它条款。')
    proof = evaluate_operation_effect(record, before, after,
        [{'location': 'p0', 'quote': old}], [], Facts())
    assert proof.state == 'PASS'
    assert proof.artifact_status == 'VERIFIED_ABSENT'


def test_delete_verified_wording_variant_is_not_proven_absent():
    from fundlex.generation.effects import evaluate_operation_effect
    from test_operation_execution_closure import Facts
    canonical = '本基金对持续持有期少于7日的投资者收取不低于1.50%的赎回费，并将上述赎回费全额计入基金财产。'
    actual = canonical.replace('不低于', '')
    rule = {'rule_id': 'G13', 'operations': [{'op': 'delete', 'old_text': canonical,
        'condition': {'kind': 'none'}, 'operation_id': 'G13:0', 'group_kind': 'INDEPENDENT'}]}
    record = compile_obligations({'groups': [rule]}).records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    proof = evaluate_operation_effect(record, Inspection(actual), Inspection(actual),
        [{'location': 'p0', 'quote': actual}], [], Facts())
    assert proof.state == 'FAIL'
    assert proof.artifact_status == 'RESIDUAL_OLD_TEXT'


@pytest.mark.parametrize('text', [
    '本基金对持续持有期少于 7 日的投资者收取 1.50%的赎回费，并将上述赎回费全额计入基金财产。',
    '对于持续持有期少于 7 日的投资者收取不低于 1.5%的赎回费，并将上述赎回费全额计入基金财产。',
])
def test_g13_source_owned_wording_variant_is_deleted_without_compose(tmp_path, text):
    """Spacing and the optional ``不低于`` wording do not require model compose."""
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, plan_target
    from test_recall_core import _facts

    source = tmp_path / 'source.docx'
    document = Document()
    document.add_paragraph('前置条款。' + text)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    rule = {'rule_id': 'G13', 'operations': [{
        'operation_id': 'G13:0', 'op': 'delete',
        'old_text': ('本基金对持续持有期少于7日的投资者收取不低于1.50%的赎回费，'
                     '并将上述赎回费全额计入基金财产。'),
        'condition': {'kind': 'none'},
    }]}
    target = Target('T13', 'G13', 'p0', text, 'needs_edit', 'source',
                    start=len('前置条款。'), end=len('前置条款。') + len(text),
                    resolution='resolved', operations=('G13:0',))
    edit = plan_target(rule, target, inspection, _facts(), None, compose_budget=[0])
    assert edit is not None
    assert edit.replacement == ''
    assert edit.content_source == 'short_holding_delete'


def test_g13_split_source_owned_tail_is_deleted_without_borrowing_prefix(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, plan_target
    from test_recall_core import _facts

    first = '赎回费用由持有人承担。其中，对持续持有期少于'
    tail = '7 日的投资者收取不低于 1.5%的赎回费，并全额计入基金财产。 '
    source = tmp_path / 'source.docx'
    document = Document()
    document.add_paragraph(first)
    document.add_paragraph(tail)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    rule = {'rule_id': 'G13', 'operations': [{
        'operation_id': 'G13:0', 'op': 'delete',
        'old_text': ('本基金对持续持有期少于7日的投资者收取不低于1.50%的赎回费，'
                     '并将上述赎回费全额计入基金财产。'),
        'condition': {'kind': 'none'},
    }]}
    target = Target('T13', 'G13', 'p1', tail, 'needs_edit', 'source',
                    start=0, end=len(tail), resolution='resolved',
                    operations=('G13:0',))
    edit = plan_target(rule, target, inspection, _facts(), None, compose_budget=[0])
    assert edit is not None and edit.replacement == ''
    assert edit.end == len(tail.rstrip())
    assert inspection.by_location('p0').text == first


@pytest.mark.parametrize('unit_verdict', ['PASS', 'FAIL'])
def test_merged_delete_requires_own_unit_proof_even_when_plan_matches(unit_verdict):
    """A sibling insertion never upgrades a failed delete authorization (FL-D224)."""
    from fundlex.generation.effects import evaluate_operation_effect
    from fundlex.obligations import compile_obligations
    from test_operation_execution_closure import Facts

    old = ('C类基金份额的销售服务费每日计提，按月支付。由基金管理人向基金托管人发送'
           '基金销售服务费划付指令，经基金托管人复核后于次月首日起10个工作日内支付。')
    new_policy = '销售服务费每日计提。直销机构持有的基金份额在赎回时返还。'
    rule = {'rule_id': 'G15', 'operations': [{
        'operation_id': 'G15:0', 'op': 'delete', 'old_text': old,
        'condition': {'kind': 'none'}, 'group_kind': 'INDEPENDENT'}]}
    record = compile_obligations({'groups': [rule]}).records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    before, after = Inspection(old), Inspection(new_policy)
    edits = [
        {'rule_id': 'G15', 'target_id': 'T15', 'location': 'p0', 'start': 0,
         'end': len(old), 'replacement': '', 'operation_keys': ['G15:0']},
        {'rule_id': 'G16', 'target_id': 'T16', 'location': 'p0', 'start': len(old),
         'end': len(old), 'replacement': new_policy, 'operation_keys': ['G16:0']},
    ]
    units = [{'location': 'p0', 'target_id': 'T15', 'operation_id': 'G15:0',
              'verdict': unit_verdict}]
    proof = evaluate_operation_effect(record, before, after,
        [{'location': 'p0', 'quote': old}], edits, Facts(),
        validated_units=units)
    assert proof.state == unit_verdict
    assert proof.artifact_status == ('VERIFIED_ABSENT' if unit_verdict == 'PASS'
                                     else 'WRONG_EFFECT')
    if unit_verdict == 'FAIL':
        assert proof.failure_code == 'CLAIM_PLANNED_EFFECT_MISSING'


@pytest.mark.parametrize('unit_verdict', ['PASS', 'FAIL'])
def test_merged_append_requires_own_unit_proof_even_with_full_policy(unit_verdict):
    from fundlex.generation.effects import evaluate_operation_effect
    from test_operation_execution_closure import Facts
    old, new = '销售服务费使用范围不包括基金募集期间的上述费用。', '销售服务费每日计提。直销机构返还投资者。'
    rule = {'rule_id': 'G16', 'operations': [dict(op='append', new_text=new,
        condition={'kind': 'none'}, operation_id='G16:0', group_kind='INDEPENDENT')]}
    record = compile_obligations({'groups': [rule]}).records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    edits = [dict(rule_id='G17', target_id='T17', location='p0', start=0,
                  end=len(old), replacement='', operation_keys=['G17:0']),
             dict(rule_id='G16', target_id='T16', location='p0', start=len(old),
                  end=len(old), replacement=new, operation_keys=['G16:0'])]
    proof = evaluate_operation_effect(record, Inspection(old), Inspection(new),
        [dict(location='p0', quote=old)], edits, Facts(), validated_units=[
            dict(location='p0', target_id='T16', operation_id='G16:0', verdict=unit_verdict)])
    assert proof.state == unit_verdict
    if unit_verdict == 'FAIL':
        assert proof.failure_code == 'CLAIM_PLANNED_EFFECT_MISSING'


def test_operation_proof_requires_each_exact_source_duplicate():
    from fundlex.generation.effects import evaluate_operation_effect
    from test_operation_execution_closure import Facts
    old = '基金销售服务费每日计提，逐日累计至每月月末，按月支付。'
    new = '新销售服务费直销返还及非直销一年后返还政策。'
    rule = {'rule_id': 'G16', 'operations': [{'op': 'append', 'new_text': new,
        'condition': {'kind': 'none'}, 'operation_id': 'G16:0', 'group_kind': 'INDEPENDENT'}]}
    record = compile_obligations({'groups': [rule]}).records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    def inspection(values):
        spans = [SimpleNamespace(location=f'p{i}', text=text) for i, text in enumerate(values)]
        return SimpleNamespace(paragraphs=spans,
            by_location=lambda address: next(span for span in spans if span.location == address))
    original = inspection([old, old])
    for omitted in (old, ''):
        artifact = inspection([new, omitted])
        proof = evaluate_operation_effect(record, original, artifact,
            [{'location': 'p0', 'quote': old}], [], Facts())
        assert proof.state == 'FAIL'
        assert set(proof.sites) == {'p0', 'p1'}
        assert proof.site_effects[1]['artifact_status'] == 'MISSING_EFFECT'


def test_delete_allows_only_same_site_certified_new_policy_reuse():
    from fundlex.generation.effects import evaluate_operation_effect
    from test_operation_execution_closure import Facts
    old = '基金销售服务费每日计提，按月支付。'
    new = old + '直销机构计提的销售服务费应返还投资者。'
    rule = {'rule_id': 'G15', 'operations': [{'op': 'delete', 'old_text': old,
        'condition': {'kind': 'none'}, 'operation_id': 'G15:0', 'group_kind': 'INDEPENDENT'}]}
    record = compile_obligations({'groups': [rule]}).records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    before, after = Inspection(old), Inspection(new)
    sites = [{'location': 'p0', 'quote': old}]
    allowed = evaluate_operation_effect(record, before, after, sites, [], Facts(),
        authorized_new_texts={'p0': [new]})
    assert allowed.state == 'PASS'
    elsewhere = evaluate_operation_effect(record, before, after, sites, [], Facts(),
        authorized_new_texts={'p1': [new]})
    assert elsewhere.state == 'FAIL'
    unsupported = evaluate_operation_effect(record, before, after, sites, [], Facts())
    assert unsupported.state == 'FAIL'


def test_delete_no_punctuation_does_not_claim_unrelated_follower():
    from fundlex.generation.effects import evaluate_operation_effect
    from test_operation_execution_closure import Facts
    old = '本基金对持有期少于7日的投资者收取赎回费'
    rule = {'rule_id': 'G13', 'operations': [{'op': 'delete', 'old_text': old,
        'condition': {'kind': 'none'}, 'operation_id': 'G13:0', 'group_kind': 'INDEPENDENT'}]}
    record = compile_obligations({'groups': [rule]}).records[0]
    record.applicability = SimpleNamespace(status='TRUE')
    def inspection(first):
        spans = [SimpleNamespace(location='p0', text=first),
                 SimpleNamespace(location='p1', text='投资者应在开放日办理业务。')]
        return SimpleNamespace(paragraphs=spans,
            by_location=lambda address: next(span for span in spans if span.location == address))
    proof = evaluate_operation_effect(record, inspection(old), inspection('其它条款。'),
        [{'location': 'p0', 'quote': old}], [], Facts())
    assert proof.state == 'PASS'
    assert proof.sites == ('p0',)


@pytest.mark.parametrize('names', [('A', 'C'), ('A类', 'C类'), ('A 类', 'C 类'),
                                  ('A类基金份额', 'C类基金份额')])
def test_complete_class_set_accepts_equivalent_display_names(names):
    from fundlex.fact_evidence import verify_class_set
    quote = '本基金将基金份额分为A类、C类两种不同的类别。'
    classes = [SimpleNamespace(name=name, current=True) for name in names]
    assert verify_class_set(Inspection(quote), classes, quote)['set_complete_verified']
    classes.append(SimpleNamespace(name='E类', current=True))
    assert not verify_class_set(Inspection(quote), classes, quote)['set_complete_verified']


def test_class_set_quote_can_cross_adjacent_docx_paragraphs_but_not_skip_one():
    from fundlex.fact_evidence import verify_class_set

    first = '本基金根据申购费用、销售服务费收取方式的不同，将基金份额分为不同的类别。'
    second = 'A 类份额：投资人申购时收取申购费用，但不从本类别基金资产中计提销售服务费。'
    third = 'C 类份额：从本类别基金资产中计提销售服务费而不收取申购费用。'
    classes = [SimpleNamespace(name='A类份额', current=True),
               SimpleNamespace(name='C类份额', current=True)]
    inspection = SimpleNamespace(paragraphs=[
        SimpleNamespace(location='p0', text=first),
        SimpleNamespace(location='p1', text=second),
        SimpleNamespace(location='p2', text=third),
    ])
    quote = first + second + third
    assert verify_class_set(inspection, classes, quote)['set_complete_verified']

    skipped = first + third
    result = verify_class_set(inspection, classes, skipped)
    assert result['quote_verified'] is False
    assert result['set_complete_verified'] is False


def test_class_set_cross_paragraph_prefix_cannot_hide_extra_class():
    from fundlex.fact_evidence import verify_class_set

    first = '本基金将基金份额分为不同的类别。'
    second = 'A 类份额：收取申购费用。C 类份额：计提销售服务费。D 类份额：收取申购费用。'
    classes = [SimpleNamespace(name='A类份额', current=True),
               SimpleNamespace(name='C类份额', current=True),
               SimpleNamespace(name='D类份额', current=True)]
    inspection = SimpleNamespace(paragraphs=[
        SimpleNamespace(location='p0', text=first),
        SimpleNamespace(location='p1', text=second),
    ])
    prefix = first + 'A 类份额：收取申购费用。C 类份额：计提销售服务费。'
    result = verify_class_set(inspection, classes, prefix)
    assert result['quote_verified'] is True
    assert result['set_complete_verified'] is False


def test_audit_missed_second_occurrence_is_not_hidden_by_same_operation(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _absorb_missed_targets
    from test_operation_execution_closure import Facts
    text = '甲条旧业务。乙条旧业务。'
    path = tmp_path / 'source.docx'
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    source = inspect_docx(path, full=True)
    loc = source.paragraphs[0].location
    rule = {'rule_id': 'G99', 'operations': [dict(op='delete', old_text='旧业务', condition={'kind': 'none'})]}
    first = Target('T', 'G99', loc, '甲条旧业务。', 'needs_edit', '', start=0, end=6,
                   operations=('G99:0',), resolution='resolved')
    state = {'missed': {'G99': [dict(location=loc, quote='乙条旧业务。', operation_id='G99:0')]}}
    synthetic, origins = [], {}
    assert _absorb_missed_targets(state, [first], synthetic, origins, source,
                                  rules_by_id={'G99': rule}, facts=Facts()) == 1
    assert synthetic[0].start == 6
    assert _absorb_missed_targets(state, [first], synthetic, origins, source,
                                  rules_by_id={'G99': rule}, facts=Facts()) == 0


def test_g05_token_splice_keeps_coordination_punctuation():
    from test_operation_execution_closure import Facts, Inspection, rule
    text = '基金管理人可调整基金份额类别，调低赎回费率或变更收费方式。'
    target = Target('T', 'G05', 'p0', text, 'needs_edit', '', start=0, end=len(text),
                    operations=('G05:0',), resolution='resolved')
    edit = plan_target(rule(), target, Inspection([text]), Facts(), None, compose_budget=[0])
    assert edit.replacement == text.replace('赎回费率', '销售服务费率')


def test_append_refuses_incomplete_cross_paragraph_clause(monkeypatch):
    from test_operation_execution_closure import Facts, Inspection
    from fundlex.recall import RecallError
    monkeypatch.setenv('FUNDLEX_GENERATION_STRATEGY', 'legacy')
    text = '基金份额类别：指本基金分设A类和C类两种不同的'
    rule = {'rule_id': 'G01', 'operations': [dict(op='append', new_text='新政策。', condition={'kind': 'none'})]}
    target = Target('T', 'G01', 'p0', text, 'needs_edit', '', start=0, end=len(text),
                    operations=('G01:0',), resolution='resolved')
    with pytest.raises(RecallError, match='RECALL_APPEND_CLAUSE_INCOMPLETE'):
        plan_target(rule, target, Inspection([text, '类别。其他内容。']), Facts(), None, compose_budget=[0])


def test_append_complete_clause_without_final_stop_appends_at_paragraph_end():
    from test_operation_execution_closure import Facts, Inspection
    text = ("基金份额类别：本基金将基金份额分为A类和C类不同的类别。"
            "在投资者认购、申购基金份额时收取认购、申购费用，不计提销售服务费的，"
            "称为A类基金份额；在投资者认购、申购基金份额时不收取认购、申购费用，"
            "而是从本类别基金资产中计提销售服务费的，称为C类基金份额 ")
    rule = {"rule_id": "G01", "operations": [{
        "op": "append", "new_text": "新政策。", "condition": {"kind": "none"}}]}
    target = Target("T", "G01", "p0", text, "needs_edit", "",
                    start=0, end=len(text), operations=("G01:0",),
                    resolution="resolved", segment_ok=True)
    edit = plan_target(rule, target, Inspection([text]), Facts(), None,
                       compose_budget=[0])
    assert edit.operation == "insert"
    assert edit.location == "p0"
    assert edit.start == edit.end == len(text.rstrip())
    assert edit.replacement.startswith("。")


def test_absent_delete_target_at_real_class_setup_proves_verified_absent():
    from test_operation_execution_closure import Facts
    setup = ("有关基金份额类别的具体设置等由基金管理人确定。基金管理人可以增加新的"
             "基金份额类别、变更现有基金份额类别的申购费率、调低赎回费率或者停止现有"
             "基金份额类别的销售等，调整前需及时公告。")
    source = Inspection(setup)
    rule = {"rule_id": "G05", "operations": [{
        "op": "delete", "old_text": "并报中国证监会备案",
        "condition": {"kind": "none"}}]}
    ledger = compile_obligations({"groups": [rule]})
    ledger.records[0].applicability = SimpleNamespace(status="TRUE")
    target = {"target_id": "absent", "rule_id": "G05", "location": "p0",
              "quote": "并报中国证监会备案", "state": "absent",
              "operations": ["G05:0"]}
    result = replay_claims([], source, source, rules_by_id={"G05": rule},
                           facts=Facts(), obligations=ledger,
                           targets_by_id={"absent": target})
    proof = result["obligations"][ledger.records[0].obligation_id]
    assert proof["state"] == "PASS"
    assert proof["artifact_status"] == "VERIFIED_ABSENT"
    assert proof["sites"] == ["p0"]


def test_audit_model_execution_hint_cannot_route_missing_plan_to_writer():
    from fundlex.recall_pipeline import _audit_recovery_contexts
    from fundlex.recovery import FailureClass
    target = Target('T', 'G99', 'p0', '旧业务', 'needs_edit', '', start=0, end=3,
                    operations=('G99:0',), resolution='resolved')
    issue = dict(code='AUDIT_NOT_COMPLETED', rule_id='G99', target_id='T', location='p0',
                 failure_class='EXECUTION', affected_obligation='free text model guess')
    contexts = _audit_recovery_contexts([issue], targets=[target], planned={'edits': []},
        artifact_failures=[dict(code='CLAIM_OPERATION_NOT_PLANNED', operation_id='G99:0', location='p0')])
    assert contexts[0].resolved_class is FailureClass.GENERATION
    assert contexts[0].operation_id == 'G99:0'
    assert contexts[0].obligation_id != issue['affected_obligation']


def test_audit_identity_keeps_sibling_operation_negative_verdict():
    from fundlex.recall_pipeline import _merge_audit_state, _audit_issues
    targets = [Target(str(i), 'G05', 'p0', '同址旧业务句。', 'needs_edit', '', start=0, end=8,
                      operations=(f'G05:{i}',), resolution='resolved') for i in range(2)]
    state = {'results': {}, 'missed': {}}
    results = [dict(target_id=str(i), rule_id='G05', verdict=verdict,
                    evidence_verified=True, reason='verified local quote')
               for i, verdict in enumerate(('wrong', 'completed'))]
    _merge_audit_state(state, dict(rules=['G05'], results=results, missed_targets=[]), targets)
    assert len(state['results']) == 2
    issues = _audit_issues(state, targets)
    assert len(issues) == 1 and issues[0]['target_id'] == '0'


def test_recovery_snapshot_handles_bound_and_unbound_sites():
    from fundlex.recall_pipeline import _recovery_snapshot
    targets = [Target('T1', 'G99', 'p0', '旧业务', 'needs_edit', '', start=0, end=3),
               Target('T2', 'G99', 'p0', '旧业务', 'needs_edit', '')]
    state = dict(failures=[])
    one = _recovery_snapshot(targets, {'edits': []}, Inspection('旧业务'), state)
    assert one == _recovery_snapshot(list(reversed(targets)), {'edits': []}, Inspection('旧业务'), state)
    targets[1].start = 0
    assert one != _recovery_snapshot(targets, {'edits': []}, Inspection('旧业务'), state)


def test_audit_sibling_operations_route_each_own_root_cause():
    from fundlex.recall_pipeline import _audit_recovery_contexts
    from fundlex.recovery import FailureClass
    target = Target('T', 'G99', 'p0', '旧业务', 'needs_edit', '', start=0, end=3,
                    operations=('G99:0', 'G99:1'), resolution='resolved')
    issue = dict(code='AUDIT_NOT_COMPLETED', rule_id='G99', target_id='T', location='p0',
                 failure_class='EXECUTION')
    failures = [dict(code='CLAIM_OPERATION_NOT_PLANNED', operation_id='G99:0', location='p0'),
                dict(code='CLAIM_EXECUTION_EFFECT_MISSING', operation_id='G99:1', location='p0')]
    edit = dict(rule_id='G99', target_id='T', location='p0', operation_keys=['G99:1'])
    for order in (failures, list(reversed(failures))):
        contexts = _audit_recovery_contexts([issue], targets=[target], planned={'edits': [edit]},
                                           artifact_failures=order)
        assert {ctx.operation_id: ctx.resolved_class for ctx in contexts} == {
            'G99:0': FailureClass.GENERATION, 'G99:1': FailureClass.EXECUTION}


def test_scope_mapping_preserves_guide_order_for_frozen_inventory(tmp_path):
    """Scope and inventory prompt order must survive a new interpreter process."""
    from fundlex.recall import scope_map, structure_index, plan_windows
    from fundlex.docx import inspect_docx
    from test_recall_core import build_contract
    path = tmp_path / 'contract.docx'
    build_contract(path)
    original = inspect_docx(path, full=True)
    sections = structure_index(original)
    # All mapped rules share one window; model response order must not affect
    # the order of rules in its prompt, including unmentioned retry entries.
    rules = [{'rule_id': f'G{i:02}', 'title': str(i)} for i in range(22, 0, -1)]
    class Gateway:
        max_context_chars = 1000000
        def call(self, **kwargs):
            return SimpleNamespace(value={'mappings': [
                {'rule_id': r['rule_id'], 'section_ids': [sections[0].id]}
                for r in reversed(rules)]})
    mapping = scope_map(rules, sections, Gateway())
    expected = [r['rule_id'] for r in rules]
    assert list(mapping) == expected
    windows = plan_windows(sections, original, mapping, max_chars=1000000)
    assert windows and windows[0][1] == expected


def test_g06_site_guard_rejects_generic_business_rule_clause():
    """A structurally resolvable G06 quote still needs the conversion anchor."""
    from fundlex.recall import target_site_ok
    span = SimpleNamespace(location="p431", text=(
        "基金管理人、登记机构、基金销售机构在法律法规规定或中国证监会许可的范围内"
        "调整有关认购、申购、赎回、转换、非交易过户、转托管等业务规则；"))
    inspection = SimpleNamespace(paragraphs=[span])
    assert target_site_ok("G06", span, inspection) is False


@pytest.mark.parametrize("valid", [False, True])
def test_generation_checks_composite_business_site_without_losing_split_anchor(tmp_path, valid):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _plan_and_generate
    from test_recall_core import _facts
    parts = (["除非基金管理人在未来条件成熟后另行公告开通相关业务，本基金不同基金",
              "份额类别之间不得互相转换。"] if valid else
             ["基金管理人调整有关认购、申购、赎回、转换、非交易过户等业务规则；",
              "基金管理人应及时公告。"])
    source = tmp_path / "source.docx"
    doc = Document()
    for part in parts:
        doc.add_paragraph(part)
    doc.save(source)
    inspection = inspect_docx(source, full=True)
    target = Target("T", "G06", "p0", parts[0], "needs_edit", "survey reference",
        start=0, end=len(parts[0]), resolution="composite", operations=("G06:0",),
        composite=tuple(f"p{i}:0:{len(text)}" for i, text in enumerate(parts)))
    rule = {"rule_id": "G06", "operations": [{"op": "delete", "old_text": (
        "除非基金管理人在未来条件成熟后另行公告开通相关业务，本基金不同基金份额类别之间不得互相转换。"),
        "condition": {"kind": "none"}}]}
    out = tmp_path / "out"
    out.mkdir()
    planned, edits, _, _, actual = _plan_and_generate(None, {"G06": rule}, [target],
        inspection, _facts(), source, out, max_compose_calls=0)
    if valid:
        assert not target.out_of_scope
        assert len(edits) == 2
        assert all(not span.text for span in actual.paragraphs)
    else:
        assert target.out_of_scope
        assert not planned["edits"] and not edits
        assert [span.text for span in actual.paragraphs] == parts


def test_sweep_drops_g17_payment_clause_as_wrong_business_site(tmp_path):
    """G17 only owns the usage clause, not a sibling G15 payment mechanism."""
    from fundlex.recall_pipeline import _Chunk, _sweep_batch
    from fundlex.docx import inspect_docx
    from types import SimpleNamespace
    from docx import Document
    path = tmp_path / "contract.docx"
    doc = Document()
    text = "基金销售服务费每日计提，按月支付，经基金管理人与基金托管人核对一致后支付。"
    doc.add_paragraph(text)
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    location = inspection.paragraphs[0].location
    rule = {"rule_id": "G17", "operations": [{"op": "delete",
        "old_text": "销售服务费主要用于支付销售机构佣金、基金行销广告费、促销活动费、基金份额持有人服务费等。"}]}
    class Gateway:
        def call(self, **kwargs):
            return SimpleNamespace(value={"findings": [{"rule_id": "G17", "location": location,
                "quote": text, "issue": "unaddressed_requirement", "reason": "wrong site"}]})
    result = _sweep_batch(Gateway(), {"G17": rule}, "G17 requirements",
        [_Chunk("C1", (location,), f"[{location}] {text}")], inspection, round_number=0,
        counter=[0], rule_locations={"G17": {location}}, scope_by_chunk={"C1": {"G17"}})
    assert not result["findings"]
    assert any(item["code"] == "SWEEP_SEMANTIC_SITE_INVALID" for item in result["dropped"])


def test_mapped_sweep_retains_unknown_wording_with_valid_rule_site(tmp_path):
    from fundlex.recall_pipeline import _Chunk, _sweep_batch
    from fundlex.docx import inspect_docx
    from docx import Document
    text = "销售服务费使用范围仅限于向投资者提供咨询服务。"
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G17", "operations": [{"op": "delete",
        "old_text": "销售服务费主要用于支付销售机构佣金及推广费用。"}]}
    class Gateway:
        def call(self, **kwargs):
            return SimpleNamespace(value={"findings": [{"rule_id": "G17", "location": "p0",
                "quote": text, "issue": "residual_old_text", "reason": "unknown usage wording"}]})
    result = _sweep_batch(Gateway(), {"G17": rule}, "G17 requirements",
        [_Chunk("C1", ("p0",), f"[p0] {text}")], inspection, round_number=0,
        counter=[0], rule_locations={"G17": {"p0"}}, scope_by_chunk={"C1": {"G17"}})
    assert len(result["findings"]) == 1
    assert result["findings"][0]["status"] == "pending"


def test_atomic_planning_selects_explicit_one_of_before_composing():
    from fundlex.recall import PlannedEdit, _atomic_paragraph_edits, merge_edits
    from test_operation_execution_closure import Inspection
    old = "本基金A类基金份额的申购费用由投资人承担，不列入基金财产。本基金C类基金份额不收取申购费用。"
    inner = "申购费用由投资人承担，不列入基金财产。"
    fixed = "通过直销机构申购本基金A类基金份额不收取申购费。本基金C类基金份额不收取申购费用。"
    source = Inspection([old])
    operations = [{"op": "replace", "old_text": text, "new_text": fixed,
                   "condition": {"kind": "none"}, "group_kind": "ONE_OF",
                   "one_of_group": "G11:history:wording"} for text in (old, inner)]
    ledger = compile_obligations({"groups": [{"rule_id": "G11", "operations": operations}]})
    edits = [PlannedEdit("outer", "G11", "p0", 0, len(old), fixed, "replace", "tier1_template",
                        evidence={"fixed_text": fixed}, operation_keys=("G11:0",)),
             PlannedEdit("inner", "G11", "p0", old.index(inner), old.index(inner)+len(inner),
                         fixed, "replace", "tier1_template", evidence={"fixed_text": fixed},
                         operation_keys=("G11:1",))]
    candidates = _atomic_paragraph_edits(edits, source, ledger)
    merged, conflicts, covered = merge_edits(candidates, source, ledger)
    assert not conflicts and covered == ["inner"]
    assert len(merged) == 1
    assert merged[0][0].replacement == fixed
    assert merged[0][0].replacement.count("通过直销机构") == 1


def test_variant_replace_rebinds_truncated_target_to_full_source_sentence(tmp_path, monkeypatch):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import plan_target
    from test_recall_core import _facts
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "legacy")
    reference = ("依据基金财产清算的分配方案，将基金财产清算后的全部剩余资产扣除基金财产清算费用、"
                 "交纳所欠税款并清偿基金债务后，按基金份额持有人持有的基金份额比例进行分配。")
    actual = reference.replace("持有的基金份额", "持有的各类基金份额")
    fixed = reference[:-1] + "，并加上返还基金份额持有人的销售服务费（如有）。"
    source = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph(actual + "无关条款保持原样。")
    doc.save(source)
    inspection = inspect_docx(source, full=True)
    rule = {"rule_id": "G22", "operations": [{"op": "replace", "old_text": reference,
        "new_text": fixed, "condition": {"kind": "none"}}]}
    target = Target("T", "G22", "p0", actual[:-2], "needs_edit", "fuzzy window",
        start=0, end=len(actual)-2, resolution="resolved", operations=("G22:0",))
    edit = plan_target(rule, target, inspection, _facts(), None, compose_budget=[0])
    assert edit.content_source == "variant_replace"
    assert edit.end == len(actual)
    after = inspection.by_location("p0").text[:edit.start] + edit.replacement + inspection.by_location("p0").text[edit.end:]
    assert after == fixed + "无关条款保持原样。"
    assert "）。配。" not in after


def test_sweep_unsupported_residual_uses_own_site_proof_without_hiding_other_copy_failure():
    from fundlex.recall_pipeline import _apply_artifact_proof_to_findings
    from test_operation_execution_closure import Inspection
    rule = {"rule_id": "G05", "operations": [
        {"op": "replace", "old_text": "基金管理人可以调低赎回费率", "new_text": "基金管理人可以调低销售服务费率"},
        {"op": "delete", "old_text": "并报中国证监会备案"}]}
    ledger = compile_obligations({"groups": [rule]})
    states = {}
    for record in ledger.records:
        operation = record.spec.operation_id
        effects = [dict(rule_id="G05", operation_id=operation, obligation_id=record.obligation_id,
                        site_id="p0", location="p0", state="PASS", failure_code="",
                        artifact_status="VERIFIED_ABSENT" if record.spec.op == "delete" else "EXPECTED_EFFECT_PRESENT",
                        proof_kind="site_final_plan_reconstruction")]
        if record.spec.op == "replace":
            effects.append({**effects[0], "site_id": "p1", "location": "p1", "state": "FAIL",
                            "artifact_status": "RESIDUAL_OLD_TEXT", "failure_code": "CLAIM_OPERATION_NOT_PLANNED"})
        states[record.obligation_id] = dict(rule_id="G05", operation_id=operation,
            obligation_id=record.obligation_id, state="FAIL" if len(effects) > 1 else "PASS",
            artifact_status=effects[-1]["artifact_status"], failure_code=effects[-1]["failure_code"],
            sites=[effect["site_id"] for effect in effects], evidence=effects, site_effects=effects)
    actual = Inspection(["基金管理人可以调低销售服务费率，调整前需及时公告。", "基金管理人可以调低赎回费率。"])
    findings = [dict(rule_id="G05", location="p0", quote=actual.paragraphs[0].text,
                     issue="residual_old_text", reason="备案句没有删除", status="pending"),
                dict(rule_id="G05", location="p1", quote="基金管理人可以调低赎回费率",
                     issue="residual_old_text", status="pending")]
    _apply_artifact_proof_to_findings(findings, {"obligations": states}, actual, ledger, {"G05": rule})
    assert findings[0]["status"] == "resolved"
    assert findings[0]["proof_kind"] == "REMOVAL_UNCONFIRMED"
    assert findings[1]["status"] == "pending"
    assert findings[1]["proof_kind"] == "RESIDUAL_CONFIRMED"
    assert states[ledger.records[0].obligation_id]["state"] == "FAIL"


def _negative_review_fixture(tmp_path, *, residual=False):
    from docx import Document
    from fundlex.docx import inspect_docx
    from test_recall_core import _facts
    old, new = "旧的业务规则应当删除。", "新的业务规则已经生效。"
    source, output = tmp_path / "source.docx", tmp_path / "after.docx"
    for path, lines in ((source, [old, old]), (output, [old if residual else new, old])):
        doc = Document()
        for line in lines:
            doc.add_paragraph(line)
        doc.save(path)
    original, actual = inspect_docx(source, full=True), inspect_docx(output, full=True)
    rule = {"rule_id": "G99", "title": "测试业务规则", "section_hint": "测试章节",
        "operations": [dict(op="replace", old_text=old,
        new_text=new, condition={"kind": "none"})]}
    ledger = compile_obligations({"groups": [rule]})
    from fundlex.obligations import apply_facts
    apply_facts(ledger, _facts())
    targets = [Target(f"T{i}", "G99", f"p{i}", old, "needs_edit", "verified site",
        start=0, end=len(old), resolution="resolved", operations=("G99:0",)) for i in range(2)]
    edit = dict(rule_id="G99", target_id="T0", location="p0", start=0, end=len(old),
        replacement=new, operation="replace", content_source="guide_text",
        operation_keys=["G99:0"], evidence={"fixed_text": new})
    planned = {"edits": [edit], "ledger": [dict(target_id="T0", status="planned",
        operation="replace", replacement=new), dict(target_id="T1", status="unresolved")]}
    class Gateway:
        max_context_chars = 100000
        def call(self, **kwargs):
            return SimpleNamespace(value={"results": [dict(target_id=f"T{i}", verdict="wrong",
                evidence_quote=new if not residual else old, reason="unsupported negative opinion")
                for i in range(2)], "missed_targets": []})
    return Gateway(), {"G99": rule}, targets, planned, original, actual, _facts(), ledger


def test_negative_audit_new_text_uses_local_proof_without_borrowing_other_site(tmp_path):
    from fundlex.recall_pipeline import _run_audit, _merge_audit_state, _audit_issues
    gateway, rules, targets, planned, original, actual, facts, ledger = _negative_review_fixture(tmp_path)
    result = _run_audit(gateway, rules, targets, planned, original, actual,
                        facts=facts, obligations=ledger)
    first, second = result["results"]
    assert first["review_status"] == "REVIEW_UNPROVEN"
    assert first["review_evidence_status"] == "EXPECTED_EFFECT_QUOTED_AS_ERROR"
    assert first["artifact_proven"] is True
    assert second["artifact_proven"] is False
    assert second["evidence_verified"] is False
    state = {"results": {}, "missed": {}}
    _merge_audit_state(state, result, targets)
    issues = _audit_issues(state, targets)
    assert len(issues) == 1 and issues[0]["target_id"] == "T1"


def test_negative_audit_true_residual_keeps_blocking(tmp_path):
    from fundlex.recall_pipeline import _run_audit, _merge_audit_state, _audit_issues
    gateway, rules, targets, planned, original, actual, facts, ledger = _negative_review_fixture(
        tmp_path, residual=True)
    result = _run_audit(gateway, rules, targets, planned, original, actual,
                        facts=facts, obligations=ledger)
    state = {"results": {}, "missed": {}}
    _merge_audit_state(state, result, targets)
    assert all(item["evidence_verified"] and not item.get("artifact_proven")
               for item in result["results"])
    assert len(_audit_issues(state, targets)) == 2


def test_audit_receives_verified_fee_groups_and_counts_them_in_context_limit(tmp_path):
    from fundlex.recall import RecallError
    from fundlex.recall_pipeline import _run_audit, _audit_context
    import json
    import pytest

    gateway, rules, targets, planned, original, actual, facts, ledger = _negative_review_fixture(tmp_path)
    received = []
    original_call = gateway.call

    def call(**kwargs):
        received.append(kwargs["bounded_context"])
        return original_call(**kwargs)

    gateway.call = call
    _run_audit(gateway, rules, targets, planned, original, actual,
               facts=facts, obligations=ledger)
    context = received[0]
    supplied = json.loads(context.split("[FACTS]\n", 1)[1].split("\n\n[BEFORE]", 1)[0])
    assert supplied == facts.as_ledger()
    assert supplied["service_group"] == ["C类基金份额"]
    assert supplied["subscription_group"] == ["A类基金份额"]
    with pytest.raises(RecallError, match="RECALL_AUDIT_CONTEXT_LIMIT"):
        _audit_context(rules["G99"], targets, {}, original, actual,
                       limit=100, facts=facts)


def test_audit_reviews_satisfied_target_to_catch_false_closure(tmp_path):
    from fundlex.recall_pipeline import _run_audit, _merge_audit_state, _audit_issues
    gateway, rules, targets, planned, original, actual, facts, ledger = _negative_review_fixture(tmp_path)
    targets[0].state = "satisfied"
    planned["edits"] = []
    actual = original
    result = _run_audit(gateway, rules, targets, planned, original, actual,
                        facts=facts, obligations=ledger)
    assert {item["target_id"] for item in result["results"]} == {"T0", "T1"}
    state = {"results": {}, "missed": {}}
    _merge_audit_state(state, result, targets)
    issues = _audit_issues(state, targets)
    assert any(item["target_id"] == "T0" for item in issues)


def test_sweep_wrong_text_requires_authorized_new_wording_and_own_site(tmp_path):
    from fundlex.recall_pipeline import _apply_artifact_proof_to_findings
    _, rules, targets, planned, original, actual, facts, ledger = _negative_review_fixture(tmp_path)
    replay = replay_claims(planned["edits"], original, actual, rules_by_id=rules,
        targets_by_id={target.target_id: target.as_dict() for target in targets},
        obligations=ledger, facts=facts)
    new = planned["edits"][0]["replacement"]
    findings = [dict(rule_id="G99", operation_id="G99:0", location=site,
                     quote=quote, issue="wrong_text", status="pending")
                for site, quote in (("p0", new), ("p1", targets[1].quote), ("p1", new))]
    _apply_artifact_proof_to_findings(findings, replay, actual, ledger, rules,
                                     facts=facts, edits=planned["edits"], original=original)
    assert findings[0]["status"] == "resolved" and findings[0]["proof_kind"] == "REVIEW_UNPROVEN"
    assert findings[1]["status"] == "pending" and findings[1]["proof_kind"] == "RESIDUAL_CONFIRMED"
    assert findings[2]["status"] == "pending"


def test_local_proof_keeps_occurrence_identity_and_rejects_unknown_site():
    from fundlex.recall_pipeline import _local_operation_proofs
    from fundlex.generation.effects import effect_proof_passed
    effect = dict(operation_id="G04:1", obligation_id="G04:B1:O2", location="p0",
        site_id="p0#occurrence:0", state="PASS", artifact_status="ALREADY_SATISFIED",
        failure_code="", proof_kind="source_operation_policy")
    proof = {**effect, "sites": [effect["site_id"]], "evidence": [effect], "site_effects": [effect]}
    assert effect_proof_passed(_local_operation_proofs(proof, "G04:1", "G04:B1:O2", "p0")[0])
    assert _local_operation_proofs(proof, "G04:1", "G04:B1:O2", "p1") == []


def test_verified_class_fact_quote_does_not_support_invented_fee_group():
    from dataclasses import replace
    from fundlex.recall import ClassFact
    from fundlex.recall_pipeline import _quote_confirms_class_facts
    from test_recall_core import _facts
    facts = replace(_facts(), classes=(_facts().classes + (ClassFact(
        name="D类基金份额", subscription_fee=True, service_fee=False, current=True,
        rate=None, quote="D类"),)), value_verified=True, set_complete_verified=True,
                    quotes_verified={"floating": True, "class:A类基金份额": True,
                                     "class:C类基金份额": True, "class:D类基金份额": True})
    enumeration = "本基金将基金份额分为A类、C类和D类三种不同的类别"
    fees = "在投资者申购A类及D类基金份额时收取申购费用，不计提销售服务费；在投资者申购C类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费。"
    assert _quote_confirms_class_facts(enumeration, facts)
    assert _quote_confirms_class_facts(fees, facts)
    assert not _quote_confirms_class_facts(fees.replace("A类及D类", "A类及B类"), facts)
    assert not _quote_confirms_class_facts(fees.replace("不计提销售服务费", "计提销售服务费"), facts)
    assert not _quote_confirms_class_facts(fees, replace(facts, set_complete_verified=False))


def test_expected_effect_quote_does_not_hide_mixed_old_residual():
    from types import SimpleNamespace
    from fundlex.recall_pipeline import _quote_is_expected_effect
    record = SimpleNamespace(spec=SimpleNamespace(op="replace", old_text="旧句。",
        alternatives=(), new_text="新句。", operation_id="G99:0"))
    assert not _quote_is_expected_effect("新句。旧句。", record, None,
        actual_text="新句。旧句。", edits=[], location="p0")


@pytest.mark.parametrize("whole_quote", [False, True])
def test_negative_review_compose_contradiction_cannot_authorize_itself(tmp_path, whole_quote):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _run_audit, _apply_artifact_proof_to_findings
    _, rules, targets, planned, original, _, facts, ledger = _negative_review_fixture(tmp_path)
    policy = rules["G99"]["operations"][0]["new_text"]
    contradiction = "但本规则实际上不产生任何效力。"
    replacement = policy + contradiction
    planned["edits"][0].update(replacement=replacement, content_source="compose")
    planned["edits"][0]["evidence"]["fixed_text"] = replacement
    artifact = tmp_path / "contradiction.docx"
    doc = Document()
    doc.add_paragraph(replacement)
    doc.add_paragraph(targets[1].quote)
    doc.save(artifact)
    actual = inspect_docx(artifact, full=True)
    quote = replacement if whole_quote else contradiction
    class Gateway:
        max_context_chars = 100000
        def call(self, **kwargs):
            return SimpleNamespace(value={"results": [dict(target_id="T0", verdict="wrong",
                evidence_quote=quote, reason="exact newly inserted contradictory assertion")]})
    result = _run_audit(Gateway(), rules, targets, planned, original, actual,
                        facts=facts, obligations=ledger)["results"][0]
    assert result["evidence_verified"] is True
    assert not result.get("artifact_proven")
    assert result.get("review_status") != "REVIEW_UNPROVEN"
    replay = replay_claims(planned["edits"], original, actual, rules_by_id=rules,
        targets_by_id={target.target_id: target.as_dict() for target in targets},
        obligations=ledger, facts=facts)
    finding = dict(rule_id="G99", operation_id="G99:0", location="p0", quote=quote,
                   issue="wrong_text", status="pending")
    _apply_artifact_proof_to_findings([finding], replay, actual, ledger, rules,
        facts=facts, edits=planned["edits"], original=original)
    assert finding["status"] == "pending"


def test_full_paragraph_negative_quote_can_prove_unchanged_source_suffix(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _run_audit, _apply_artifact_proof_to_findings
    _, rules, targets, planned, _, _, facts, ledger = _negative_review_fixture(tmp_path)
    suffix = "原有补充条款保持有效。"
    new = planned["edits"][0]["replacement"]
    inspections = []
    for name, text in (("source-suffix", targets[0].quote + suffix), ("after-suffix", new + suffix)):
        doc = Document()
        doc.add_paragraph(text)
        doc.add_paragraph(targets[1].quote)
        path = tmp_path / (name + ".docx")
        doc.save(path)
        inspections.append(inspect_docx(path, full=True))
    original, actual = inspections
    class Gateway:
        max_context_chars = 100000
        def call(self, **kwargs):
            return SimpleNamespace(value={"results": [dict(target_id="T0", verdict="wrong",
                evidence_quote=new + suffix, reason="unsupported opinion on preserved source suffix")]})
    result = _run_audit(Gateway(), rules, targets, planned, original, actual,
                        facts=facts, obligations=ledger)["results"][0]
    assert result["evidence_verified"] is True
    assert result["artifact_proven"] is True
    assert result["review_status"] == "REVIEW_UNPROVEN"
    replay = replay_claims(planned["edits"], original, actual, rules_by_id=rules,
        targets_by_id={target.target_id: target.as_dict() for target in targets},
        obligations=ledger, facts=facts)
    finding = dict(rule_id="G99", operation_id="G99:0", location="p0", quote=new + suffix,
                   issue="wrong_text", status="pending")
    _apply_artifact_proof_to_findings([finding], replay, actual, ledger, rules,
        facts=facts, edits=planned["edits"], original=original)
    assert finding["status"] == "resolved" and finding["proof_kind"] == "REVIEW_UNPROVEN"


@pytest.mark.parametrize(
    "case_id,rule_id,operation_id,old_text",
    [
        ("D02", "G05", "G05:2", "并报中国证监会备案"),
        ("D02", "G06", "G06:0", "不同基金份额类别之间不得互相转换"),
        ("D02", "G13", "G13:0", "持续持有期少于7日的投资者收取不低于1.50%的赎回费"),
        ("D11", "G06", "G06:0", "不同基金份额类别之间不得互相转换"),
    ],
)
@pytest.mark.parametrize("state", ["absent", "needs_edit", "satisfied"])
def test_absent_inventory_placeholder_never_becomes_verified_operation_site(
        case_id, rule_id, operation_id, old_text, state):
    """Window-level ``absent`` is not an operation/site absence proof (D02/D11 P0)."""
    from fundlex.generation.effects import evaluate_operation_effect
    from test_operation_execution_closure import Facts

    operations = [{"op": "noop", "condition": {"kind": "none"}}
                  for _ in range(int(operation_id.rsplit(":", 1)[1]))]
    operations.append({"op": "delete", "old_text": old_text,
                       "condition": {"kind": "none"},
                       "operation_id": operation_id})
    guide = {"groups": [{"rule_id": rule_id, "operations": operations}]}
    ledger = compile_obligations(guide)
    for record in ledger.records:
        if record.spec.operation_id == operation_id:
            record.applicability = SimpleNamespace(status="TRUE")
            target_record = record
            break
    else:  # pragma: no cover - protects the fixture if guide compilation changes
        raise AssertionError(operation_id)
    before = Inspection("一、基金份额持有人、基金管理人和基金托管人的权利与义务")
    if state == "absent":
        absent_site = {"location": "p0", "quote": before.paragraphs[0].text,
                       "state": "absent", "operation_id": operation_id}
        proof = evaluate_operation_effect(target_record, before, before,
                                          [absent_site], [], Facts())
        # The heading is never a verified site.  The operation can still close
        # only through the independent document-wide literal-absence proof.
        assert not proof.sites
        assert proof.state == "PASS"
        assert proof.artifact_status == "VERIFIED_ABSENT"
        assert proof.evidence[0]["proof_kind"] == "document_literal_absent"

    result = replay_claims(
        [], before, before, rules_by_id={rule_id: guide["groups"][0]}, facts=Facts(),
        targets_by_id={"placeholder": {"target_id": "placeholder", "rule_id": rule_id,
            "location": "p0", "quote": before.paragraphs[0].text,
            "state": state, "operations": [operation_id]}},
        obligations=ledger,
    )
    item = result["obligations"][target_record.obligation_id]
    assert item["sites"] == []
    assert item["state"] == "PASS"
    assert item["artifact_status"] == "VERIFIED_ABSENT"
    assert item["evidence"][0]["proof_kind"] == "document_literal_absent"


def test_satisfied_operation_site_can_close_as_already_satisfied():
    """A real satisfied target remains eligible for operation effect proof."""
    from test_operation_execution_closure import Facts
    guide = {"groups": [{"rule_id": "G02", "operations": [{
        "op": "replace", "old_text": "A类基金份额", "new_text": "A类基金份额",
        "condition": {"kind": "none"}, "operation_id": "G02:0"}]}]}
    ledger = compile_obligations(guide)
    record = ledger.records[0]
    record.applicability = SimpleNamespace(status="TRUE")
    before = Inspection("基金份额类别：收取申购费用且不计提销售服务费的，称为 A类基金份额。")
    target_record = {"target_id": "satisfied", "rule_id": "G02", "location": "p0",
                     "quote": "A类基金份额", "state": "satisfied",
                     "operations": ["G02:0"]}
    result = replay_claims([], before, before, rules_by_id={"G02": guide["groups"][0]},
                           facts=Facts(), targets_by_id={"satisfied": target_record},
                           obligations=ledger)
    assert result["verified"] is True
    proof = result["obligations"][record.obligation_id]
    assert proof["state"] == "PASS"
    assert proof["artifact_status"] == "ALREADY_SATISFIED"
    assert proof["sites"] == ["p0"]


def test_satisfied_delete_with_already_absent_literal_keeps_own_business_site():
    from test_operation_execution_closure import Facts
    guide = {"groups": [{"rule_id": "G05", "operations": [{
        "op": "delete", "old_text": "并报中国证监会备案", "condition": {"kind": "none"}}]}]}
    ledger = compile_obligations(guide)
    ledger.records[0].applicability = SimpleNamespace(status="TRUE")
    text = "基金管理人可以增加新的基金份额类别，调整前应及时公告后实施。"
    source = Inspection(text)
    target = {"target_id": "satisfied", "rule_id": "G05", "location": "p0",
              "quote": text, "state": "satisfied", "operations": ["G05:0"]}
    replay = replay_claims([], source, source, rules_by_id={"G05": guide["groups"][0]},
                           facts=Facts(), obligations=ledger,
                           targets_by_id={"satisfied": target})
    proof = replay["obligations"][ledger.records[0].obligation_id]
    assert replay["verified"] is True
    assert proof["state"] == "PASS"
    assert proof["artifact_status"] == "VERIFIED_ABSENT"
    assert proof["sites"] == ["p0"]
    assert proof["evidence"][0]["before_quote"] == text


def test_unattributed_target_cannot_lend_claimed_operation_to_delete_proof():
    from test_operation_execution_closure import Facts
    guide = {"groups": [{"rule_id": "G06", "operations": [{
        "op": "delete", "old_text": "本基金不同基金份额类别之间不得互相转换。",
        "condition": {"kind": "none"}}]}]}
    ledger = compile_obligations(guide)
    ledger.records[0].applicability = SimpleNamespace(status="TRUE")
    text = "本基金不同基金份额类别之间不得互相转换。"
    source = Inspection(text)
    replay = replay_claims([], source, source, rules_by_id={"G06": guide["groups"][0]},
        facts=Facts(), obligations=ledger, targets_by_id={"bad": {
            "target_id": "bad", "rule_id": "G06", "location": "p0", "quote": text,
            "state": "needs_edit", "operations": [], "claimed_operations": ["G06:0"],
            "attribution_error": "RECALL_TARGET_OPERATION_UNATTRIBUTED"}})
    proof = replay["obligations"][ledger.records[0].obligation_id]
    assert replay["verified"] is False
    assert proof["state"] == "OPEN"
    assert proof["sites"] == []


def test_g06_generic_business_adjustment_clause_is_not_a_delete_site():
    from test_operation_execution_closure import Facts
    guide = {"groups": [{"rule_id": "G06", "operations": [{
        "op": "delete", "old_text": "本基金不同基金份额类别之间不得互相转换。",
        "condition": {"kind": "none"}}]}]}
    ledger = compile_obligations(guide)
    ledger.records[0].applicability = SimpleNamespace(status="TRUE")
    text = "基金管理人可调整有关申购、赎回、转换、非交易过户等业务规则，应及时公告。"
    source = Inspection(text)
    replay = replay_claims([], source, source, rules_by_id={"G06": guide["groups"][0]},
        facts=Facts(), obligations=ledger, targets_by_id={"bad": {
            "target_id": "bad", "rule_id": "G06", "location": "p0", "quote": text,
            "state": "needs_edit", "operations": ["G06:0"]}})
    # The generic paragraph is not a G06 site, but the authorized literal is
    # absent from the whole ORIGINAL document: that independent proof is PASS.
    proof = replay["obligations"][ledger.records[0].obligation_id]
    assert proof["state"] == "PASS"
    assert proof["artifact_status"] == "VERIFIED_ABSENT"
    assert proof["sites"] == []
    assert proof["evidence"][0]["proof_kind"] == "document_literal_absent"


@pytest.mark.parametrize("old_text,new_text", [
    ("调低赎回费率", "调低销售服务费率"),
    ("并报中国证监会备案", ""),
])
@pytest.mark.parametrize("source,valid", [
    # Class-setting authority is a required G05 site, including the independent
    # filing operation whose old phrase is already absent in the first variant.
    ("有关基金份额类别的具体设置、费率水平等由基金管理人确定。基金管理人在履行适当程序后"
     "可以增加新的基金份额类别、或者在法律法规和基金合同规定的范围内变更现有基金份额类别"
     "的申购费率、调低赎回费率或变更收费方式，或者停止现有基金份额类别的销售等。"
     "调整前，基金管理人需履行适当程序并及时公告。", True),
    ("有关基金份额类别的具体设置等由基金管理人确定。基金管理人在履行适当程序后可以增加"
     "新的基金份额类别、或者变更现有基金份额类别的申购费率、调低赎回费率或变更收费方式、"
     "或者停止现有基金份额类别的销售等，调整前基金管理人需及时公告并报中国证监会备案。", True),
    # Body and summary copies of meeting exemptions remain outside G05 even
    # where the literal matches, or there is no filing phrase to delete.
    ("在法律法规和《基金合同》规定的范围内调整本基金全部或者部分份额类别的申购费率、"
     "调低全部或者部分份额类别的赎回费率、调低销售服务费率或变更收费方式；", False),
    ("在法律法规和《基金合同》规定的范围内调整本基金的申购费率、调低赎回费率、"
     "调低销售服务费率或变更收费方式；", False),
])
def test_g05_effect_proof_binds_class_setting_not_meeting_exemption(source, valid,
                                                                 old_text, new_text):
    from fundlex.generation.effects import operation_site_bound
    from test_operation_execution_closure import Facts
    rule = {"rule_id": "G05", "operations": [{
        "op": "replace" if new_text else "delete", "old_text": old_text,
        "new_text": new_text, "condition": {"kind": "none"}}]}
    spec = compile_obligations({"groups": [rule]}).records[0].spec
    inspection = Inspection(source)
    assert operation_site_bound(spec, source, inspection, "p0", Facts()) is valid
    if not valid:
        ledger = compile_obligations({"groups": [rule]})
        ledger.records[0].applicability = SimpleNamespace(status="TRUE")
        replay = replay_claims([], inspection, inspection, facts=Facts(),
            rules_by_id={"G05": rule}, obligations=ledger, targets_by_id={"unrelated": {
                "target_id": "unrelated", "rule_id": "G05", "location": "p0",
                "quote": source, "state": "needs_edit", "operations": ["G05:0"]}})
        proof = replay["obligations"][ledger.records[0].obligation_id]
        assert proof["sites"] == []
        if new_text or old_text in source:
            assert proof["state"] == "OPEN"
        else:
            assert proof["state"] == "PASS"
            assert proof["artifact_status"] == "VERIFIED_ABSENT"
            assert proof["evidence"][0]["proof_kind"] == "document_literal_absent"


def test_atomic_replan_unbound_components_are_recorded_once(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import plan_edits, replan_sites
    from test_recall_core import _facts
    doc = Document()
    doc.add_paragraph("旧业务甲句。旧业务乙句。")
    path = tmp_path / "source.docx"
    doc.save(path)
    source = inspect_docx(path, full=True)
    rule = {"rule_id": "G99", "operations": [
        {"op": "replace", "old_text": "旧业务甲句。", "new_text": "新业务甲句。"},
        {"op": "delete", "old_text": "旧业务乙句。"}]}
    targets = [Target(f"Old{i}", "G99", "p0", source.paragraphs[0].text,
        "needs_edit", "", start=0, end=len(source.paragraphs[0].text),
        operations=(f"G99:{i}",), resolution="resolved", segment_ok=True) for i in range(2)]
    class Gateway:
        def call(self, **kwargs):
            raise AssertionError("deterministic replan called model")
    facts = _facts()
    original = plan_edits({"G99": rule}, targets, source, facts, Gateway())
    assert len(original["edits"]) == 1
    repaired = replan_sites(original, set(), {"G99": rule}, targets[:1], source,
                           facts, Gateway(), max_compose_calls=0)
    assert repaired["edits"] == []
    assert len(repaired["ledger"]) == 2
    assert len({entry["target_id"] for entry in repaired["ledger"]}) == 2
    assert all(entry["reason"] == "RECALL_PLAN_REFERENCE_UNBOUND"
               and entry["status"] == "unresolved" for entry in repaired["ledger"])


def test_g02_already_correct_class_group_accepts_ji_delimiter():
    """D11 p91 uses ``A类及D类``; that is the same class set as the guide's ``A类、D类``."""
    from fundlex.generation.claim_replay import replay_claims
    from test_operation_execution_closure import Facts
    guide = {"groups": [{"rule_id": "G02", "operations": [{
        "op": "replace", "old_text": "A类基金份额", "new_text": "A类、D类基金份额",
        "condition": {"kind": "none"}, "operation_id": "G02:0"}]}]}
    ledger = compile_obligations(guide)
    ledger.records[0].applicability = SimpleNamespace(status="TRUE")
    quote = "在投资者申购A类及D类基金份额时收取申购费用，不计提销售服务费"
    source = Inspection("基金份额类别：" + quote + "。")
    target = {"target_id": "satisfied", "rule_id": "G02", "location": "p0",
              "quote": quote, "state": "satisfied", "operations": ["G02:0"]}
    replay = replay_claims([], source, source, rules_by_id={"G02": guide["groups"][0]},
                           facts=Facts(), obligations=ledger,
                           targets_by_id={"satisfied": target})
    proof = replay["obligations"][ledger.records[0].obligation_id]
    assert proof["state"] == "PASS"
    assert proof["artifact_status"] == "ALREADY_SATISFIED"
    assert proof["sites"] == ["p0"]


def test_g02_split_definition_reclassifies_out_of_scope_target_before_plan():
    """A valid split p91/p92 definition must not stay out_of_scope from an earlier pass."""
    from fundlex.recall import Target, assign_target_operations, plan_edits
    from test_operation_execution_closure import Facts, Inspection as ClosureInspection
    text = ("类别。在投资者申购A类及D类基金份额时收取申购费用，不计提销售服务费；"
            "在投资者申购C类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费 ")
    rule = {"rule_id": "G02", "class_substitution": True, "title": "份额类别",
            "operations": [
                {"op": "replace", "old_text": "A类基金份额", "new_text": "A类、B类基金份额",
                 "condition": {"kind": "none"}, "operation_id": "G02:0"},
                {"op": "replace", "old_text": "C类基金份额", "new_text": "C类、D类基金份额",
                 "condition": {"kind": "none"}, "operation_id": "G02:1"}]}
    facts = Facts()
    facts.subscription_group = ("A类", "D类")
    facts.service_group = ("C类",)
    facts.statuses = {"share_class_structure": "VERIFIED",
                      "subscription_fee_group": "VERIFIED",
                      "service_fee_group": "VERIFIED"}
    target = Target("T", "G02", "p0", text, "needs_edit",
                    "TARGET_SOURCE_MISMATCH", start=0, end=len(text),
                    operations=("G02:0",), out_of_scope=True,
                    resolution="resolved", segment_ok=True)
    inspection = ClosureInspection([text])
    assign_target_operations([target], {"G02": rule}, inspection, facts)
    assert target.out_of_scope is False
    assert set(target.operations) == {"G02:0", "G02:1"}
    planned = plan_edits({"G02": rule}, [target], inspection, facts, gateway=None)
    assert planned["edits"] == []
    assert [entry["status"] for entry in planned["ledger"]] == ["satisfied"]


def test_duplicate_satisfied_class_discoveries_reconcile_without_collapsing_occurrences():
    text = "A类基金份额收取申购费；C类基金份额计提销售服务费。"
    rule = {"rule_id": "G02", "operations": [{"op": "replace", "old_text": "A类基金份额"}]}
    targets = [Target(str(i), "G02", "p0", text, "satisfied", "discovery",
                      start=0, end=len(text), operations=("G02:0",)) for i in range(2)]
    _, removed = dedupe_operation_targets(targets, {"G02": rule}, Inspection(text))
    assert removed == ["1"]
    targets = [Target(str(i), "G02", "p0", text, "satisfied", "occurrence",
                      start=0, end=len(text), operations=("G02:0",), occurrence_index=i)
               for i in range(2)]
    assert dedupe_operation_targets(targets, {"G02": rule}, Inspection(text))[1] == []


def test_late_split_class_targets_pass_replay_after_planning(tmp_path):
    import json
    from pathlib import Path
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import assign_target_operations, rebind_split_class_definition_targets
    from fundlex.recall_pipeline import _plan_and_generate
    from test_recall_core import _facts
    from dataclasses import replace
    from fundlex.recall import ClassFact
    guide = json.loads((Path(__file__).resolve().parents[1] / "guides/builtin_guide_v1.json").read_text())
    rule = next(g for g in guide["groups"] if g["rule_id"] == "G02")
    first = "基金份额类别：本基金将基金份额分为A类、C类和D类三种不同的"
    second = "类别。在投资者申购A类及D类基金份额时收取申购费用，不计提销售服务费；在投资者申购C类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费。"
    source = tmp_path / "source.docx"
    document = Document(); document.add_paragraph(first); document.add_paragraph(second); document.save(source)
    inspection = inspect_docx(source, full=True)
    facts = _facts()
    facts = replace(facts, classes=(*facts.classes, ClassFact("D类基金份额", True, False, True, None, "D类")))
    targets = [Target(str(i), "G02", "p0", "", "needs_edit", "late discovery",
                      operations=(f"G02:{i}",)) for i in range(2)]
    rebind_split_class_definition_targets(targets, inspection)
    assign_target_operations(targets, {"G02": rule}, inspection, facts)
    planned, _, _, _, actual = _plan_and_generate(SimpleNamespace(max_context_chars=60000),
        {"G02": rule}, targets, inspection, facts, source, tmp_path, max_compose_calls=0)
    ledger = compile_obligations({"groups": [rule]})
    for record in ledger.records:
        record.applicability = SimpleNamespace(status="TRUE")
    replay = replay_claims(planned["edits"], inspection, actual, rules_by_id={"G02": rule},
        facts=facts, obligations=ledger, targets_by_id={t.target_id: t.as_dict() for t in targets})
    assert replay["verified"], replay["failures"]
    assert all(row["state"] == "PASS" for row in replay["obligations"].values())


@pytest.mark.parametrize("delay", [
    "",
    "若遇法定节假日、休息日或不可抗力致使无法按时支付等，支付日期顺延。",
    "如遇法定节假日、公休日，支付日期相应顺延。",
])
def test_g16_head_and_tail_discoveries_write_once_per_fee_clause(tmp_path, delay):
    from pathlib import Path
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _plan_and_generate
    from test_recall_core import _facts
    import json
    guide = json.loads((Path(__file__).resolve().parents[1] / "guides/builtin_guide_v1.json").read_text())
    rule = next(g for g in guide["groups"] if g["rule_id"] == "G16")
    source = tmp_path / "source.docx"
    document = Document()
    for title in ("十五、基金费用与税收", "二十、基金合同内容摘要"):
        document.add_paragraph(title)
        document.add_paragraph("销售服务费每日计提，按月支付。")
        document.add_paragraph(delay + "销售服务费主要用于支付销售机构佣金。")
        document.add_paragraph("销售服务费使用范围不包括基金募集期间的上述费用。")
    document.save(source)
    inspection = inspect_docx(source, full=True)
    targets = [Target(str(i), "G16", loc, inspection.by_location(loc).text,
                      "needs_edit", "discovery", start=0,
                      end=len(inspection.by_location(loc).text), resolution="resolved",
                      operations=("G16:0",)) for i, loc in enumerate(("p1", "p3", "p5", "p7"))]
    planned, _, _, _, actual = _plan_and_generate(SimpleNamespace(max_context_chars=60000),
        {"G16": rule}, targets, inspection, _facts(), source, tmp_path,
        max_compose_calls=0)
    assert len(planned["edits"]) == 2
    assert {e["location"] for e in planned["edits"]} == {"p3", "p7"}
    assert sum(p.text.count("（1）直销机构") for p in actual.paragraphs) == 2
    ledger = compile_obligations({"groups": [rule]})
    for record in ledger.records:
        record.applicability = SimpleNamespace(status="TRUE")
    replay = replay_claims(planned["edits"], inspection, actual,
        rules_by_id={"G16": rule}, facts=_facts(), obligations=ledger,
        targets_by_id={t.target_id: t.as_dict() for t in targets})
    assert replay["verified"], replay["failures"]


@pytest.mark.parametrize("text", [
    "在法律法规和《基金合同》规定的范围内调整本基金的申购费率、调低赎回费率、调低销售服务费率或变更收费方式；",
    "在法律法规和《基金合同》规定的范围内调整全部或者部分份额类别的申购费率、调低全部或者部分份额类别的赎回费率、调低销售服务费率；",
])
def test_g05_planner_rejects_meeting_exemption_before_any_edit(tmp_path, text):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import RecallError, target_site_ok
    from test_recall_core import _facts
    source = tmp_path / "meeting.docx"
    document = Document(); document.add_paragraph(text); document.save(source)
    inspection = inspect_docx(source, full=True)
    assert not target_site_ok("G05", inspection.by_location("p0"), inspection)
    rule = {"rule_id": "G05", "operations": [{"op": "replace", "old_text": "调低赎回费率",
            "new_text": "调低销售服务费率", "operation_id": "G05:0", "condition": {"kind": "none"}}]}
    target = Target("T", "G05", "p0", text, "needs_edit", "literal match",
                    start=0, end=len(text), resolution="resolved", operations=("G05:0",))
    with pytest.raises(RecallError, match="TARGET_SOURCE_MISMATCH|RECALL_TARGET_OPERATION_UNATTRIBUTED"):
        plan_target(rule, target, inspection, _facts(), None, compose_budget=[0])
    from fundlex.recall import assign_target_operations
    target.state = "satisfied"
    assign_target_operations([target], {"G05": rule}, inspection, _facts())
    assert target.out_of_scope


def test_replan_conflict_cannot_refer_to_reassigned_target_id():
    from fundlex.recall import replan_sites
    from test_recall_core import _facts
    source = Inspection("旧业务甲句。")
    metadata = dict(rule_id="G99", location="p0", start=0, end=6,
                    operation_keys=["G99:0"], quote=source.paragraphs[0].text)
    original = {"edits": [], "covered": [], "ledger": [dict(metadata,
        target_id="Old", status="unresolved", reason="RECALL_EDIT_OVERLAP")],
        "conflicts": [{"code": "RECALL_EDIT_OVERLAP", "location": "p0", "targets": ["Old"]}]}
    current = Target("New", "G99", "p0", source.paragraphs[0].text, "needs_edit", "",
                     start=0, end=6, operations=("G99:0",))
    repaired = replan_sites(original, set(), {}, [current], source, _facts(), None,
                           max_compose_calls=0)
    assert repaired["ledger"][0]["target_id"] == "New"
    assert repaired["conflicts"][0]["targets"] == ["New"]
    assert repaired["ledger"][0]["reason"] == "RECALL_EDIT_OVERLAP"


def test_claim_site_repair_binds_a_non_exact_summary_copy(tmp_path):
    """D11 G18: Claim proved p1207; an active p800 target must not suppress repair."""
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.targeting import target_for_claim_site

    body = ("由于本基金 A 类基金份额和 D 类基金份额不收取销售服务费，"
            "而 C 类基金份额收取销售服务费，各基金份额类别对应的可分配收益将有所不同。"
            "本基金同一类别的每一基金份额享有同等分配权； ")
    summary = body + "5、法律法规或监管机关另有规定的，从其规定。"
    path = tmp_path / "summary.docx"
    doc = Document()
    doc.add_paragraph(body)
    doc.add_paragraph(summary)
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    target = Target("T019", "G18", "p0",
                    "由于本基金 A 类基金份额和 D 类基金份额不收取销售服务费，"
                    "而 C 类基金份额收取销售服务费",
                    "needs_edit", "verified body site", start=0, end=47,
                    resolution="resolved", operations=("G18:0",))
    record = SimpleNamespace(spec=SimpleNamespace(
        rule_id="G18", operation_id="G18:0", op="replace",
        old_text="由于本基金A类基金份额不收取销售服务费，而C类基金份额收取销售服务费",
        new_text="由于本基金A类基金份额和C类基金份额费用收取不同",
        alternatives=()))
    repair = target_for_claim_site(
        record, "p1", [target],
        {"G18": {"rule_id": "G18", "operations": []}}, SimpleNamespace(
            as_ledger=lambda: {}), inspection)
    assert repair is not None
    assert repair.location == "p1" and repair.operations == ("G18:0",)
    assert repair.reason.startswith("claim_required_site:")
    span = inspection.by_location("p1")
    assert span.text[repair.start:repair.end].startswith("由于本基金 A 类")


def test_delete_audit_opinion_must_quote_its_own_residual(tmp_path):
    """G13 cannot be blocked because a sibling G12 replacement removed old wording."""
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _delete_opinion_scope_unproven

    path = tmp_path / "after.docx"
    doc = Document()
    doc.add_paragraph("赎回费用全额计入基金财产。")
    doc.save(path)
    actual = inspect_docx(path, full=True)
    record = SimpleNamespace(spec=SimpleNamespace(
        op="delete", old_text="对持续持有期少于7日的投资者收取不低于1.5%的赎回费，并全额计入基金财产"))
    assert _delete_opinion_scope_unproven(
        "赎回费用全额计入基金财产。", [record], actual, "p0")
    assert not _delete_opinion_scope_unproven(
        "对持续持有期少于7日的投资者收取不低于1.5%的赎回费，并全额计入基金财产",
        [record], actual, "p0")


def test_post_proof_target_reconciliation_removes_stale_candidate():
    from fundlex.recall_pipeline import _reconcile_candidate_targets
    rule = {"rule_id": "G12", "title": "赎回费用", "operations": [
        {"op": "replace", "old_text": "旧费用条款", "new_text": "新费用条款",
         "condition": {"kind": "none"}, "operation_id": "G12:0",
         "group_kind": "INDEPENDENT"}]}
    ledger = compile_obligations({"groups": [rule]})
    target = Target("T3001", "G12", "p219", "旧费用条款", "needs_edit",
                    "T3_open_operation:G12:0", start=0, end=5,
                    operations=("G12:0",))
    proof = {"state": "PASS", "operation_id": "G12:0",
             "obligation_id": "G12:B1:O1", "artifact_status": "ALREADY_SATISFIED",
             "sites": ["p221"], "failure_code": "",
                 "evidence": [{"site_id": "p221", "proof_kind": "source_operation_policy",
                               "operation_id": "G12:0", "obligation_id": "G12:B1:O1"}],
                     "site_effects": [{"state": "PASS", "site_id": "p221",
                                        "artifact_status": "ALREADY_SATISFIED",
                                        "proof_kind": "source_operation_policy",
                                        "operation_id": "G12:0",
                                        "obligation_id": "G12:B1:O1"}]}
    changed = _reconcile_candidate_targets(
        [target], {"ledger": [{"target_id": "T3001", "status": "not_planned"}]},
        {"obligations": {"G12:B1:O1": proof}}, ledger)
    assert changed == 0 and not target.out_of_scope

    target.composite = ("p221:0:5",)
    changed = _reconcile_candidate_targets(
        [target], {"ledger": [{"target_id": "T3001", "status": "not_planned"}]},
        {"obligations": {"G12:B1:O1": proof}}, ledger)
    assert changed == 1 and target.out_of_scope
    assert "POST_PROOF_SUPERSEDED" in target.reason


def test_claim_site_reuse_plan_avoids_recomposing_the_summary_copy(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import plan_edits

    text = ("由于本基金 A 类基金份额和 D 类基金份额不收取销售服务费，"
            "而 C 类基金份额收取销售服务费，各基金份额类别对应的可分配收益将有所不同。")
    path = tmp_path / "copy.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G18", "title": "收益分配", "operations": [{
        "op": "replace",
        "old_text": "由于本基金A类基金份额不收取销售服务费，而C类基金份额收取销售服务费",
        "new_text": "由于本基金A类基金份额和C类基金份额费用收取不同",
        "condition": {"kind": "none"}}]}
    quote = ("由于本基金 A 类基金份额和 D 类基金份额不收取销售服务费，"
             "而 C 类基金份额收取销售服务费")
    target = Target("R027", "G18", "p0", quote, "needs_edit",
                    "claim_required_site:G18:0", start=0, end=len(quote),
                    resolution="resolved", operations=("G18:0",))
    reuse = {"target_id": "T019", "rule_id": "G18", "location": "p800",
             "start": 0, "end": len(quote),
             "replacement": "由于本基金A类、D类基金份额和C类基金份额费用收取不同",
             "operation": "replace", "content_source": "composed",
             "operation_keys": ["G18:0"], "expected_old_text": quote}

    class Gateway:
        def call(self, **kwargs):
            raise AssertionError("composer must not be called for a verified copy")

    result = plan_edits({"G18": rule}, [target], inspection, SimpleNamespace(),
                        Gateway(), reuse_edits=[reuse])
    assert result["edits"][0]["location"] == "p0"
    assert result["edits"][0]["content_source"] == "claim_site_reuse"
    assert result["edits"][0]["replacement"] == reuse["replacement"]


def test_g02_split_fee_group_paragraph_is_a_real_site(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import target_site_ok

    path = tmp_path / "split.docx"
    doc = Document()
    doc.add_paragraph("基金份额类别：本基金将基金份额分为 A 类、C 类和 D 类三种不同的")
    doc.add_paragraph("类别。在投资者申购 A 类及 D 类基金份额时收取申购费用，不计提销售服务费；"
                      "在投资者申购 C 类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费")
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    assert target_site_ok("G02", inspection.by_location("p1"), inspection)
    assert target_site_ok("G04", inspection.by_location("p1"), inspection)


def test_g18_heading_is_not_a_business_site(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import target_site_ok

    path = tmp_path / "g18.docx"
    doc = Document()
    doc.add_paragraph("基金收益分配原则")
    doc.add_paragraph("由于本基金A类、D类基金份额和C类基金份额费用收取不同，各基金份额类别对应的可分配收益将有所不同。")
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    assert not target_site_ok("G18", inspection.by_location("p0"), inspection)
    assert target_site_ok("G18", inspection.by_location("p1"), inspection)


def test_expected_class_group_quote_accepts_and_separator():
    from types import SimpleNamespace
    from fundlex.recall_pipeline import _quote_is_expected_effect

    record = SimpleNamespace(spec=SimpleNamespace(
        op="replace", old_text="A类基金份额", new_text="A类、D类基金份额",
        alternatives=(), operation_id="G04:0", rule_id="G04"))
    quote = "在投资者申购 A 类及 D 类基金份额时收取申购费用，不计提销售服务费"
    assert _quote_is_expected_effect(quote, record, None, location="p0", actual_text=quote)
    assert not _quote_is_expected_effect(
        "在投资者申购 A 类及 B 类基金份额时收取申购费用", record, None,
        location="p0", actual_text="在投资者申购 A 类及 B 类基金份额时收取申购费用")


def test_g02_g04_group_member_is_not_residual_at_satisfied_site():
    """A complete C/E/F group necessarily still contains the old C-class label."""
    from fundlex.generation.effects import evaluate_operation_effect
    from test_operation_execution_closure import Facts, Inspection

    text = ("本基金 A 类、Y 类基金份额收取申购费用，C 类基金份额、E 类基金份额和 "
            "F 类基金份额不收取申购费用，但从该类别基金资产中计提销售服务费。")
    source = Inspection([text])
    facts = Facts()
    for rule_id in ("G02", "G04"):
        obligation = SimpleNamespace(
            obligation_id=f"{rule_id}:B1:O2",
            applicability=SimpleNamespace(status="TRUE"),
            spec=SimpleNamespace(
                rule_id=rule_id, operation_id=f"{rule_id}:1", op="replace",
                old_text="C类基金份额", new_text="C类、E类和F类基金份额",
                alternatives=()))
        effect = evaluate_operation_effect(
            obligation, source, source,
            [{"location": "p0", "quote": text, "state": "satisfied"}], [], facts)
        assert effect.state == "PASS", (rule_id, effect.as_dict())
        assert effect.artifact_status == "ALREADY_SATISFIED"


def test_g02_sweep_introductory_quote_uses_same_paragraph_fee_proof():
    """A scope quote may name classes while the fee proof follows in the paragraph."""
    from types import SimpleNamespace
    from fundlex.recall_pipeline import _quote_is_expected_effect

    record = SimpleNamespace(spec=SimpleNamespace(
        rule_id="G02", operation_id="G02:0", op="replace",
        old_text="A类基金份额", new_text="A类、Y类基金份额",
        alternatives=()), obligation_id="G02:B1:O1")
    quote = "基金份额类别：本基金将基金份额分为A类、C类、E类、F类和Y类五种不同的类别。"
    paragraph = (quote + "在投资者申购A类、Y类基金份额时收取申购费用，不计提销售服务费；"
                 "在投资者申购C类、E类和F类基金份额时不收取申购费用，而从该类别基金资产中计提销售服务费。")
    assert _quote_is_expected_effect(quote, record, None, location="p98",
                                     actual_text=paragraph)


def test_expected_effect_accepts_reworded_fee_difference_quote():
    from types import SimpleNamespace
    from fundlex.recall_pipeline import _quote_is_expected_effect

    record = SimpleNamespace(spec=SimpleNamespace(
        rule_id="G18", operation_id="G18:0", op="replace",
        old_text="由于本基金A类基金份额不收取销售服务费，而C类基金份额收取销售服务费",
        new_text="由于本基金A类基金份额和C类基金份额费用收取不同",
        alternatives=()))
    paragraph = ("本基金A类、C类基金份额分别设置代码。由于基金费用的不同，本基金A类基金份额和"
                 "C类基金份额将分别计算基金份额净值。")
    assert _quote_is_expected_effect(paragraph, record, None, location="p0",
                                     actual_text=paragraph)


def test_expected_effect_accepts_class_substituted_final_edit():
    from fundlex.recall_pipeline import _quote_is_expected_effect

    quote = "由于本基金A类、D类基金份额和C类基金份额费用收取不同，各基金份额类别对应的可分配收益将有所不同。"
    record = SimpleNamespace(
        obligation_id="G18:B1:O1",
        spec=SimpleNamespace(rule_id="G18", op="replace", operation_id="G18:0",
                             old_text="由于本基金A类基金份额不收取销售服务费，而C类基金份额收取销售服务费",
                             new_text="由于本基金A类基金份额和C类基金份额费用收取不同",
                             alternatives=()))
    proof = {"state": "PASS", "operation_id": "G18:0", "obligation_id": "G18:B1:O1",
             "artifact_status": "EXPECTED_EFFECT_PRESENT", "sites": ["p0"],
             "failure_code": "",
             "evidence": [{"site_id": "p0", "location": "p0",
                           "operation_id": "G18:0", "obligation_id": "G18:B1:O1",
                           "proof_kind": "site_final_plan_reconstruction"}],
             "site_effects": [{"site_id": "p0", "location": "p0",
                               "operation_id": "G18:0", "obligation_id": "G18:B1:O1",
                               "state": "PASS"}]}
    edit = {"location": "p0", "replacement": quote, "operation_keys": ["G18:0"],
            "components": [{"location": "p0", "replacement": quote,
                            "operation_keys": ["G18:0"]}]}
    assert _quote_is_expected_effect(quote, record, None, edits=[edit], location="p0",
                                     actual_text=quote, obligations=SimpleNamespace(),
                                     proofs={"G18:B1:O1": proof})


def test_overlapping_composite_deletes_merge_as_joint_delete(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import PlannedEdit, _resolve_overlaps_with_joint_compose

    path = tmp_path / "tail.docx"
    doc = Document()
    doc.add_paragraph("7 日的投资者收取不低于 1.5%的赎回费，并全额计入基金财产。")
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    length = len(inspection.by_location("p0").text)
    edits = [
        PlannedEdit("T12", "G12", "p0", 0, length, "", "replace",
                    "composite_delete", {}, operation_keys=("G12:1",)),
        PlannedEdit("T13", "G13", "p0", 0, length, "", "replace",
                    "composite_delete", {}, operation_keys=("G13:0",)),
    ]

    class Gateway:
        def call(self, **kwargs):
            raise AssertionError("two deletes must not call Joint Compose")

    resolved, remaining = _resolve_overlaps_with_joint_compose(
        [{"targets": ["T12", "T13"], "location": "p0"}], edits, inspection,
        {"G12": {"rule_id": "G12", "operations": []},
         "G13": {"rule_id": "G13", "operations": []}}, None, Gateway())
    assert remaining == []
    assert len(resolved) == 1
    assert resolved[0].content_source == "joint_delete"
    assert resolved[0].replacement == ""
    assert set(resolved[0].operation_keys) == {"G12:1", "G13:0"}


def test_transitive_overlap_conflicts_are_composed_as_one_component(tmp_path):
    """A-B and B-C conflicts must not consume B and leave C unresolved."""
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import PlannedEdit, _resolve_overlaps_with_joint_compose

    path = tmp_path / "transitive-overlap.docx"
    document = Document()
    document.add_paragraph("甲乙丙丁戊己庚")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    edits = [
        PlannedEdit("A", "G01", "p0", 0, 3, "", "replace", "composite_delete", {},
                    operation_keys=("G01:0",)),
        PlannedEdit("B", "G02", "p0", 2, 5, "", "replace", "composite_delete", {},
                    operation_keys=("G02:0",)),
        PlannedEdit("C", "G03", "p0", 4, 7, "", "replace", "composite_delete", {},
                    operation_keys=("G03:0",)),
    ]

    class Gateway:
        def call(self, **kwargs):
            raise AssertionError("connected deletes must not call Joint Compose")

    resolved, remaining = _resolve_overlaps_with_joint_compose(
        [{"targets": ["A", "B"], "location": "p0"},
         {"targets": ["B", "C"], "location": "p0"}],
        edits, inspection,
        {"G01": {"rule_id": "G01", "operations": []},
         "G02": {"rule_id": "G02", "operations": []},
         "G03": {"rule_id": "G03", "operations": []}},
        None, Gateway())
    assert remaining == []
    assert len(resolved) == 1
    assert resolved[0].content_source == "joint_delete"
    assert set(resolved[0].operation_keys) == {"G01:0", "G02:0", "G03:0"}


def test_planning_clamps_trailing_whitespace_to_a_text_segment(tmp_path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, plan_edits

    path = tmp_path / "segment.docx"
    doc = Document()
    doc.add_paragraph("基金份额类别：本基金将基金份额分为A类、C类。有关具体设置等由基金管理人确定。")
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    span = inspection.by_location("p0")
    quote = span.text
    rule = {"rule_id": "G03", "title": "份额类别设置", "operations": [{
        "op": "append", "new_text": "新增条款。", "condition": {"kind": "none"}}]}
    target = Target("T1", "G03", "p0", quote, "needs_edit", "append site",
                    start=0, end=len(quote) + 1, resolution="resolved",
                    operations=("G03:0",))
    from test_recall_core import _facts
    result = plan_edits({"G03": rule}, [target], inspection, _facts(),
                        SimpleNamespace(call=lambda **kwargs: (_ for _ in ()).throw(
                            AssertionError("append should stay deterministic"))))
    edit = result["edits"][0]
    assert edit["end"] <= max(s.end for s in span.segments)


def test_satisfied_target_with_false_branch_is_not_auditable(tmp_path):
    """G20 on a fixed-fee product must not be audited as an unchanged target.

    A sibling G22 edit legitimately changes the same liquidation paragraph; the audit
    otherwise treats that as a contradictory G20 change (live G20 p1321).
    """
    from dataclasses import replace
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, assign_target_operations
    from test_recall_core import _facts

    path = tmp_path / "g20.docx"
    document = Document()
    document.add_paragraph("依据基金财产清算的分配方案，按基金份额持有人持有的基金份额比例进行分配。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G20", "operations": [{
        "operation_id": "G20:0", "op": "replace", "old_text": "分配。",
        "new_text": "分配，并加上返还基金份额持有人的销售服务费（如有）。",
        "condition": {"kind": "is_floating_fee", "value": True}}]}
    target = Target("T1", "G20", "p0", "分配。", "satisfied", "not applicable",
                    operations=("G20:0",))
    facts = replace(_facts(), quotes_verified={"floating": True},
                    floating_management_fee=False)
    assign_target_operations([target], {"G20": rule}, inspection, facts)
    assert target.out_of_scope is True
    assert "TARGET_OPERATION_NOT_APPLICABLE" in target.reason


def test_g21_noop_target_leaves_sibling_g22_auditable(tmp_path):
    """A fixed-fee multi-class G21 noop exits target audit without hiding G22."""
    from dataclasses import replace
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, assign_target_operations
    from fundlex.recall_pipeline import _audit_issues, _target_identity
    from test_recall_core import _facts

    path = tmp_path / "g21-g22.docx"
    text = "依据基金财产清算的分配方案，按基金份额持有人持有的基金份额比例进行分配。"
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    inspection = inspect_docx(path, full=True)
    g21 = {"rule_id": "G21", "operations": [{"operation_id": "G21:0", "op": "noop",
        "text": text, "condition": {"kind": "compound", "value": [
            {"kind": "is_floating_fee", "value": False},
            {"kind": "class_count", "value": 1}]}}]}
    g22 = {"rule_id": "G22", "operations": [{"operation_id": "G22:0", "op": "replace",
        "old_text": "分配。", "new_text": "分配，并加上返还基金份额持有人的销售服务费（如有）。",
        "condition": {"kind": "compound", "value": [
            {"kind": "is_floating_fee", "value": False},
            {"kind": "has_service_fee", "value": True},
            {"kind": "class_count", "value": 2}]}}]}
    facts = _facts()
    noop = Target("noop", "G21", "p0", text, "satisfied", "certified noop",
                  operations=("G21:0",))
    sibling = Target("sibling", "G22", "p0", "分配。", "needs_edit", "sibling",
                     start=text.index("分配"), end=len(text), resolution="resolved",
                     operations=("G22:0",))
    assign_target_operations([noop, sibling], {"G21": g21, "G22": g22}, inspection, facts)
    assert noop.out_of_scope is True
    assert "TARGET_OPERATION_NOT_APPLICABLE" in noop.reason
    state = {"results": {_target_identity(sibling): {
        "rule_id": "G22", "verdict": "completed", "artifact_proven": True,
        "evidence_verified": True}}, "missed": {}}
    assert _audit_issues(state, [noop, sibling]) == []


def test_g18_reworded_fee_difference_proves_already_satisfied(tmp_path):
    """A satisfied G18 site proves ALREADY_SATISFIED without the exact new text."""
    from types import SimpleNamespace
    from fundlex.generation.effects import (evaluate_operation_effect,
                                            operation_site_bound)
    from fundlex.obligations import compile_obligations
    from test_operation_execution_closure import Facts

    text = ("本基金A类、C类基金份额分别设置代码。由于基金费用的不同，本基金A类基金份额和"
            "C类基金份额将分别计算基金份额净值。")
    rule = {"rule_id": "G18", "title": "基金合同-基金的收益与分配",
            "requirement": "将费用差异表述改为费用收取不同。",
            "operations": [{"operation_id": "G18:0", "op": "replace",
                            "old_text": "由于本基金A类基金份额不收取销售服务费，而C类基金份额收取销售服务费",
                            "new_text": "由于本基金A类基金份额和C类基金份额费用收取不同",
                            "condition": {"kind": "none"}, "group_kind": "INDEPENDENT"}]}
    record = compile_obligations({"groups": [rule]}).records[0]
    record.applicability = SimpleNamespace(status="TRUE")
    before = after = Inspection(text)
    site = {"location": "p0", "quote": text}
    assert operation_site_bound(record.spec, text, before, "p0", Facts()) is True
    proof = evaluate_operation_effect(record, before, after, [site], [], Facts())
    assert proof.state == "PASS"
    assert proof.artifact_status == "ALREADY_SATISFIED"


def test_satisfied_g18_carries_its_operation_identity(tmp_path):
    """A reordered fee-difference sentence is ALREADY_SATISFIED, not site-less.

    The contract writes "由于基金费用的不同，本基金A/C类基金份额将分别计算…" while
    the guide writes "…A/C类基金份额费用收取不同".  The rule-level signal already found
    the policy, but ``satisfied_operation_keys`` required the exact new text, leaving
    G18 without an operation/site (live G18 CLAIM_OPERATION_UNPROVEN).
    """
    import json
    from pathlib import Path
    from fundlex.recall import Facts, satisfied_operation_keys
    from test_recall_core import _facts

    guide = json.loads((Path(__file__).resolve().parents[1]
                        / "guides/builtin_guide_v1.json").read_text(encoding="utf-8"))
    rule = next(item for item in guide["groups"] if item["rule_id"] == "G18")
    body = ("由于基金费用的不同，本基金A类基金份额和C类基金份额将分别计算基金份额净值。"
            "计算公式为：计算日某类基金份额净值＝计算日该类基金份额的基金资产净值/"
            "计算日该类基金份额余额总数。")
    assert satisfied_operation_keys(rule, _facts(), body) == {"G18:0"}


def test_g08_subscription_fee_target_cannot_bind_the_purchase_fee_operation(tmp_path):
    """A 认购费率 quote must not execute G08's 申购费率 replacement (live p156)."""
    import json
    import pytest
    from pathlib import Path
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, _selected_operation_indices
    from test_recall_core import _facts

    path = tmp_path / "g08-flip.docx"
    document = Document()
    document.add_paragraph("本基金的认购费率由基金管理人决定，并在招募说明书中列示，基金认购费用不列入基金财产。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    guide = json.loads((Path(__file__).resolve().parents[1]
                        / "guides/builtin_guide_v1.json").read_text(encoding="utf-8"))
    rule = next(item for item in guide["groups"] if item["rule_id"] == "G08")
    from dataclasses import replace

    facts = replace(_facts(), fact_status={"share_class_structure": "VERIFIED"})
    quote = "本基金的认购费率由基金管理人决定，并在招募说明书中列示"
    target = Target("T1", "G08", "p0", quote, "needs_edit", "test",
                    start=0, end=len(quote), resolution="resolved",
                    operations=("G08:0",))
    with pytest.raises(Exception) as excinfo:
        _selected_operation_indices(rule, target, inspection, facts)
    assert getattr(excinfo.value, "code", "") == "RECALL_TARGET_OPERATION_UNATTRIBUTED"
    assert excinfo.value.detail.get("reason") == "operation_topic_flip"


def test_user_facing_target_counts_are_terminal_and_reconcile():
    from fundlex.api import _recall_target_counts
    from fundlex.recall_pipeline import _target_disposition_counts

    targets = [
        {"target_id": "T1", "status": "planned", "state": "needs_edit"},
        {"target_id": "T2", "status": "satisfied", "state": "satisfied"},
        {"target_id": "T3", "status": "not_applicable", "state": "needs_edit",
         "out_of_scope": True},
        {"target_id": "T4", "status": "unresolved", "state": "needs_edit",
         "out_of_scope": True},
    ]
    counts = _target_disposition_counts(targets)
    assert counts == {"targets": 4, "needs_edit": 1, "satisfied": 1, "absent": 2}
    report = {"targets": targets, "unresolved": []}
    api = _recall_target_counts(report)
    assert api == {"total": 4, "completed": 4, "unresolved": 0}


def test_final_omission_review_is_operation_scoped_and_fail_closed(tmp_path):
    from types import SimpleNamespace
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _final_omission_review

    path = tmp_path / "omission.docx"
    from docx import Document
    document = Document()
    document.add_paragraph("原始收费条款。")
    document.save(path)
    original = inspect_docx(path, full=True)
    actual = original
    record = SimpleNamespace(
        obligation_id="G01:B1:O1",
        applicability=SimpleNamespace(status="TRUE"),
        completion="UNRESOLVED",
        spec=SimpleNamespace(rule_id="G01", operation_id="G01:0", op="replace"),
    )
    replay = {"obligations": {"G01:B1:O1": {
        "state": "FAIL", "operation_id": "G01:0", "obligation_id": "G01:B1:O1",
        "sites": ["p0"], "failure_code": "CLAIM_NO_CHANGE_APPLIED",
    }}}
    obligations = SimpleNamespace(records=[record])
    result = _final_omission_review(
        rules_by_id={"G01": {"rule_id": "G01", "operations": []}},
        rules=[], original=original, actual=actual, planned={"edits": []},
        targets=[], facts=SimpleNamespace(), obligations=obligations,
        replay=replay, findings=[])
    assert result["status"] == "FAIL"
    assert result["checked_operation_count"] == 1
    assert any(item["code"] == "FINAL_OPERATION_UNPROVEN" for item in result["issues"])
