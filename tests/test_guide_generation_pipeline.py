from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from fundlex.guide_generation.human import update_guide_item
from fundlex.guide_generation.pipeline import GuidePipeline, PipelineConfig
from fundlex.guide_generation.store import GuideGenerationStore
from fundlex.guide_generation.worker import GuideGenerationWorker
from fundlex.store import Conflict
from fundlex.store import RunStore

from guide_generation_fakes import FakeGuideRouter


def _contract(path: Path, fund: str):
    doc = Document()
    doc.add_paragraph(f"{fund}基金合同")
    doc.add_paragraph("第一部分  前言")
    doc.add_paragraph("订立本基金合同的目的是保护投资人合法权益。")
    doc.add_paragraph("第二部分  基金费用与税收")
    doc.add_paragraph("一、基金费用的种类")
    doc.add_paragraph("本基金A类基金份额不收取销售服务费，C类基金份额的销售服务费按前一日C类基金份额资产净值的0.40%年费率计提。")
    doc.add_paragraph("第三部分  基金份额的申购与赎回")
    doc.add_paragraph("本基金A类基金份额的申购费率由基金管理人决定，并在招募说明书中列示。")
    doc.add_paragraph("赎回费用由赎回基金份额的基金份额持有人承担，在基金份额持有人赎回基金份额时收取。")
    doc.save(path)


def _regulation(path: Path):
    doc = Document()
    doc.add_paragraph("测试销售费用管理规定")
    doc.add_paragraph("第一条")
    doc.add_paragraph("基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。")
    doc.add_paragraph("第二条")
    doc.add_paragraph("对于投资者持续持有期限少于七日的基金份额，应当收取不低于赎回金额1.5%的赎回费，"
                      "并将上述赎回费全额计入基金财产。")
    doc.add_paragraph("第三条")
    doc.add_paragraph("基金管理人应当在基金合同、基金销售协议等法律文件中与相关主体约定上述费用的收取标准。")
    doc.save(path)


def _prepare(tmp_path: Path):
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    _contract(corpus / "测试基金基金合同_20240101.docx", "测试基金")
    _contract(corpus / "另一基金基金合同_20240101.docx", "另一基金")
    regulation = tmp_path / "regulation.docx"
    _regulation(regulation)
    store = RunStore(tmp_path / "run")
    gg = GuideGenerationStore(store)
    gg.ensure_schema()
    data = regulation.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    run = gg.create_run(owner="u", filename=regulation.name, source_type="docx",
                        source_sha256=digest, title="测试销售费用管理规定",
                        payload={"config": {"corpus_dir": str(corpus)}})
    relpath = f"guide_generation/{run['id']}/source.docx"
    target = store.object_path(relpath)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    gg.update_run(run["id"], payload={**(run.get("payload") or {}), "source_relpath": relpath,
                                      "source_object_key": relpath})
    return store, gg, run["id"], corpus


def test_pipeline_end_to_end_with_fake_models(tmp_path):
    store, gg, run_id, corpus = _prepare(tmp_path)
    config = PipelineConfig(corpus_dir=str(corpus), path_b_latest_only=False,
                            path_b_max_calls=20, path_c_enabled=False,
                            semantic_applicability=False, max_residual_rounds=2,
                            residual_max_calls_per_round=10, soft_max_calls=500)
    result = GuidePipeline(gg, run_id, FakeGuideRouter(), config).run()
    assert result["status"] == "complete", result
    assert result["qualification_status"] == "candidate_ready"
    assert result["payload"]["finalize"]["unhandled_requirements"] == 0
    assert result["payload"]["finalize"]["requirement_gaps"] == 0
    assert result["payload"]["finalize"]["unresolved_impacts"] == 0
    assert result["payload"]["finalize"]["residual_saturated"] is True
    items = gg.get_guide_items(run_id)
    assert items
    assert all(item.requirement_ids for item in items)
    assert all(item.impact_ids for item in items)
    assert all(item.new_text_origin for item in items)
    artifact = gg.get_artifact(run_id, "guide")
    assert artifact and store.object_path(artifact["object_key"]).is_file()
    evidence = gg.get_artifact(run_id, "evidence_json")
    assert evidence and store.object_path(evidence["object_key"]).is_file()
    package = json.loads(store.object_path(evidence["object_key"]).read_text(encoding="utf-8"))
    assert package["traceability_chain"]
    assert package["requirements"]
    assert package["impacts"]
    assert package["coverage"]

    # Human edit round-trip: field diff, version and re-export.
    item = items[0]
    updated = update_guide_item(gg, run_id, item.guide_item_id,
                                {"title": "人工修订标题",
                                 "new_text": "通过直销机构申购本基金不收取申购费（人工确认）。"}, "reviewer")
    assert updated["version_no"] == 1
    assert len(updated["changes"]) == 2
    versions = gg.list_human_versions(run_id)
    assert len(versions) == 1
    changes = gg.list_human_changes(run_id)
    assert len(changes) == 2
    assert changes[0]["ai_value"] != changes[0]["human_value"]
    assert gg.get_artifact(run_id, "guide_ai")
    assert gg.get_artifact(run_id, "guide")["sha256"] != artifact["sha256"]
    package_after = json.loads(store.object_path(gg.get_artifact(run_id, "evidence_json")["object_key"])
                               .read_text(encoding="utf-8"))
    assert package_after["human_versions"]
    assert package_after["human_changes"]
    candidate = gg.get_candidate(run_id, owner="u")
    assert candidate["status"] == "human_draft"
    assert len(candidate["versions"]) == 3
    assert candidate["versions"][1]["status"] == "candidate_ready"
    assert candidate["versions"][2]["status"] == "human_draft"


def test_candidate_versions_are_snapshot_bound_and_stale_writes_fail(tmp_path):
    store, gg, run_id, corpus = _prepare(tmp_path)
    config = PipelineConfig(corpus_dir=str(corpus), path_b_latest_only=False,
                            path_b_max_calls=20, path_c_enabled=False,
                            semantic_applicability=False, max_residual_rounds=1,
                            residual_max_calls_per_round=10, soft_max_calls=500)
    GuidePipeline(gg, run_id, FakeGuideRouter(), config).run()

    candidate = gg.get_candidate(run_id, owner="u")
    first = gg.get_candidate_version(run_id, candidate["current_version_id"], owner="u")
    first_payload = json.loads(json.dumps(first["payload"], ensure_ascii=False))
    second = gg.create_candidate_version(
        run_id, status="human_draft", payload={"guide_items": first_payload["guide_items"],
                                                "marker": "second"},
        content_sha256="second-sha", expected_current_version_id=first["candidate_version_id"])
    assert second["version_no"] == first["version_no"] + 1
    assert second["parent_version_id"] == first["candidate_version_id"]
    assert gg.get_candidate_version(run_id, first["candidate_version_id"])["payload"] == first_payload

    with pytest.raises(Conflict, match="candidate version changed"):
        gg.create_candidate_version(run_id, status="human_draft", payload={"marker": "stale"},
                                    expected_current_version_id=first["candidate_version_id"])
    current = gg.get_candidate(run_id, owner="u")
    assert current["current_version_id"] == second["candidate_version_id"]
    assert [version["version_no"] for version in current["versions"]] == [1, 2, 3]


def test_limited_snapshot_and_findings_require_explicit_human_version(tmp_path):
    from fundlex.guide_generation.corpus_catalog import sync_corpus
    from fundlex.guide_generation.api import public_pipeline_config

    store, gg, run_id, corpus = _prepare(tmp_path)
    # A previously indexed wider corpus must not leak back into this run.
    wide = sync_corpus(gg, corpus).snapshot
    config = PipelineConfig(corpus_dir=str(corpus),
                            corpus_limit=public_pipeline_config('{"corpus_limit":1}')["corpus_limit"],
                            path_b_latest_only=False, path_b_max_calls=20, path_c_enabled=False,
                            semantic_applicability=False, max_residual_rounds=1,
                            required_clean_rounds=2, residual_max_calls_per_round=10, soft_max_calls=500)

    class ObservedRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.contexts = []

        def call(self, **kwargs):
            self.contexts.append(kwargs["bounded_context"])
            return super().call(**kwargs)

    router = ObservedRouter()
    result = GuidePipeline(gg, run_id, router, config).run()
    assert result["qualification_status"] == "candidate_with_findings"
    assert result["status"] == "complete_with_unresolved"
    snapshot = gg.get_snapshot(result["corpus_snapshot_id"])
    assert snapshot.counts["contracts"] == 1
    assert len(snapshot.contract_version_ids) == 1
    excluded_ids = set(wide.contract_version_ids) - set(snapshot.contract_version_ids)
    assert len(excluded_ids) == 1
    assert not any(excluded in context for excluded in excluded_ids for context in router.contexts)
    assert all(impact.contract_version_id in snapshot.contract_version_ids for impact in gg.get_impacts(run_id))

    candidate = gg.get_candidate(run_id)
    version_id = candidate["current_version_id"]
    assert candidate["status"] == "candidate_with_findings"
    with pytest.raises(Conflict, match="GUIDE_CANDIDATE_FINDINGS_REVIEW_REQUIRED"):
        gg.approve_candidate(run_id, actor="u", owner="u", version_id=version_id)
    # Also defend legacy snapshots incorrectly marked ready/approved.
    with gg._connect() as db:
        db.execute("UPDATE gg_candidate_versions SET status='candidate_ready' WHERE candidate_version_id=?", (version_id,))
    with pytest.raises(Conflict, match="GUIDE_CANDIDATE_FINDINGS_REVIEW_REQUIRED"):
        gg.approve_candidate(run_id, actor="u", owner="u", version_id=version_id)
    with gg._connect() as db:
        db.execute("UPDATE gg_candidate_versions SET status='approved' WHERE candidate_version_id=?", (version_id,))
    with pytest.raises(Conflict, match="GUIDE_CANDIDATE_FINDINGS_REVIEW_REQUIRED"):
        gg.publish_candidate(run_id, actor="u", owner="u", version_id=version_id)

    item = gg.get_guide_items(run_id)[0]
    unchanged = update_guide_item(gg, run_id, item.guide_item_id, {"title": item.title}, "u")
    assert not unchanged["changes"]
    assert gg.get_candidate(run_id)["current_version_id"] == version_id
    edited = update_guide_item(gg, run_id, item.guide_item_id, {"title": "人工审阅后修订"}, "u")
    human_version = edited["candidate_version"]
    assert human_version["status"] == "human_draft"
    assert human_version["payload"]["human_review_override"] is True
    assert human_version["payload"]["source_qualification_status"] == "candidate_with_findings"
    gg.approve_candidate(run_id, actor="u", owner="u", version_id=human_version["candidate_version_id"])
    gg.publish_candidate(run_id, actor="u", owner="u", version_id=human_version["candidate_version_id"])
    assert gg.get_run(run_id)["qualification_status"] == "candidate_with_findings"


def test_candidate_artifact_hashes_and_worker_fencing_are_enforced(tmp_path):
    store, gg, run_id, corpus = _prepare(tmp_path)
    config = PipelineConfig(corpus_dir=str(corpus), path_b_latest_only=False,
                            path_b_max_calls=20, path_c_enabled=False,
                            semantic_applicability=False, max_residual_rounds=1,
                            residual_max_calls_per_round=10, soft_max_calls=500)
    GuidePipeline(gg, run_id, FakeGuideRouter(), config).run()
    for artifact in gg.list_artifacts(run_id):
        path = store.object_path(artifact["object_key"])
        data = path.read_bytes()
        assert path.is_file()
        assert artifact["size"] == len(data)
        assert artifact["sha256"] == hashlib.sha256(data).hexdigest()

    job_id = gg.enqueue(run_id)
    first = gg.claim(lease_seconds=60)
    assert first and first["id"] == job_id and first["fencing_token"] == 1
    assert gg.claim(lease_seconds=60) is None
    with pytest.raises(Conflict, match="stale guide generation worker lease"):
        gg.heartbeat(job_id, 0, lease_seconds=60)
    gg.heartbeat(job_id, first["fencing_token"], lease_seconds=60)
    with gg._connect() as db:
        db.execute("UPDATE gg_jobs SET lease_until=? WHERE id=?", (0, job_id))
    second = gg.claim(lease_seconds=60)
    assert second and second["id"] == job_id and second["fencing_token"] == 2
    with pytest.raises(Conflict, match="stale guide generation worker lease"):
        gg.finish_job(job_id, first["fencing_token"], status="completed")
    gg.finish_job(job_id, second["fencing_token"], status="completed")


def test_api_upload_worker_human_edit_and_download(tmp_path, monkeypatch):
    from fundlex.api import create_app
    monkeypatch.setenv("FUNDLEX_AUTH_MODE", "dev")
    root = tmp_path / "api"
    app = create_app(root)
    client = TestClient(app)
    regulation = tmp_path / "regulation.docx"
    _regulation(regulation)
    data = regulation.read_bytes()
    response = client.post(
        "/api/guide-generation/runs", headers={"X-FundLex-User": "alice"},
        files={"file": ("regulation.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"config": json.dumps({"corpus_limit": 2,
                                    "path_b_latest_only": False, "path_b_max_calls": 10,
                                    "path_c_enabled": False, "semantic_applicability": False,
                                    "max_residual_rounds": 1, "residual_max_calls_per_round": 5})})
    assert response.status_code == 201, response.text
    run_id = response.json()["id"]
    # Owner isolation.
    assert client.get(f"/api/guide-generation/runs/{run_id}",
                      headers={"X-FundLex-User": "bob"}).status_code == 404
    worker = GuideGenerationWorker(RunStore(root), router_factory=lambda: FakeGuideRouter())
    assert worker.run_once() is True
    detail = client.get(f"/api/guide-generation/runs/{run_id}", headers={"X-FundLex-User": "alice"}).json()
    assert detail["status"] in {"complete", "complete_with_unresolved"}
    guide = client.get(f"/api/guide-generation/runs/{run_id}/guide",
                       headers={"X-FundLex-User": "alice"}).json()
    assert guide["items"]
    item_id = guide["items"][0]["guide_item_id"]
    edited = client.patch(f"/api/guide-generation/runs/{run_id}/guide/items/{item_id}",
                          headers={"X-FundLex-User": "alice"},
                          json={"instruction": "人工确认后的指引"})
    assert edited.status_code == 200, edited.text
    assert edited.json()["version_no"] == 1
    download = client.get(f"/api/guide-generation/runs/{run_id}/download/guide",
                          headers={"X-FundLex-User": "alice"})
    assert download.status_code == 200
    assert download.content[:2] == b"PK"
    evidence = client.get(f"/api/guide-generation/runs/{run_id}/download/evidence_json",
                          headers={"X-FundLex-User": "alice"})
    assert evidence.status_code == 200
    package = json.loads(evidence.content)
    assert package["human_changes"]
    report = client.get(f"/api/guide-generation/runs/{run_id}/report",
                        headers={"X-FundLex-User": "alice"})
    assert report.status_code == 200
    assert "候选修改指引" in report.text or "Evidence" in report.text

    candidate = client.get(f"/api/guide-generation/runs/{run_id}/candidate",
                           headers={"X-FundLex-User": "alice"})
    assert candidate.status_code == 200
    current = candidate.json()["current_version_id"]
    assert candidate.json()["status"] == "human_draft"
    assert client.get(f"/api/guide-generation/runs/{run_id}/candidate",
                      headers={"X-FundLex-User": "bob"}).status_code == 404
    approved = client.post(f"/api/guide-generation/runs/{run_id}/candidate/approve",
                           headers={"X-FundLex-User": "alice"},
                           json={"version_id": current})
    assert approved.status_code == 200, approved.text
    published = client.post(f"/api/guide-generation/runs/{run_id}/candidate/publish",
                            headers={"X-FundLex-User": "alice"},
                            json={"version_id": current})
    assert published.status_code == 200, published.text
    assert published.json()["runtime_guide_import"] == "not_requested"
    rejected_edit = client.patch(
        f"/api/guide-generation/runs/{run_id}/guide/items/{item_id}",
        headers={"X-FundLex-User": "alice"}, json={"instruction": "再次编辑"})
    assert rejected_edit.status_code == 409
