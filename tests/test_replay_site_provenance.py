"""Operation identity and source satisfaction stay local to each physical site."""
import json
from types import SimpleNamespace

import pytest
from docx import Document

from fundlex.docx import inspect_docx
from fundlex.generation.claim_replay import replay_claims
from fundlex.generation.effects import evaluate_operation_effect
from fundlex.obligations import compile_obligations
from fundlex.recall_pipeline import _run_survey
from fundlex.recall import _compact, resolve_source_quote
from test_recall_core import _facts


def _inspection(tmp_path, name, texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    path = tmp_path / f"{name}.docx"
    doc.save(path)
    return inspect_docx(path, full=True)


def _ledger(rules):
    ledger = compile_obligations({"groups": rules})
    for record in ledger.records:
        record.applicability = SimpleNamespace(status="TRUE")
    return ledger


@pytest.mark.parametrize("merged", [False, True])
def test_superseded_g03_cannot_borrow_g01_edit_or_parent_keys(tmp_path, merged):
    definition = "基金份额类别：本基金将基金份额分为A类和C类不同的类别。"
    policy = "通过直销机构申购A类基金份额时，不收取申购费用。"
    rules = [{"rule_id": rid, "operations": [{"op": "append", "new_text": policy,
              "condition": {"kind": "none"}}]} for rid in ("G01", "G03")]
    ledger = _ledger(rules)
    source = _inspection(tmp_path, "source", [definition, "基金管理人可以增加新的基金份额类别。"])
    actual = _inspection(tmp_path, "actual", [definition + policy, source.paragraphs[1].text])
    edit = dict(target_id="g01", rule_id="G01", location="p0", start=len(definition),
                end=len(definition), replacement=policy, operation_keys=["G01:0"])
    if merged:
        # A parent union does not assign the sibling component to G03.
        edit = {**edit, "operation_keys": ["G01:0", "G03:0"], "components": [dict(edit)]}
    result = replay_claims([edit], source, actual, rules_by_id={r["rule_id"]: r for r in rules},
        facts=_facts(), obligations=ledger, targets_by_id={"old-g03": {
            "target_id": "old-g03", "rule_id": "G03", "location": "p1",
            "quote": source.paragraphs[1].text, "state": "needs_edit",
            "out_of_scope": True, "operations": ["G03:0"]}})
    proof = result["obligations"]["G03:B1:O1"]
    assert not result["verified"]
    assert proof["state"] == "OPEN"
    assert proof["sites"] == []
    assert proof["edit_ids"] == []


OLD = "由于本基金A类基金份额不收取销售服务费，而C类基金份额收取销售服务费"
NEW = "由于本基金A类基金份额和C类基金份额费用收取不同"
SATISFIED = "1、由于基金费用的不同，本基金 A 类基金份额和 C 类基金份额将分别计算基金份额净值。计算公式为："


def _g18():
    return {"rule_id": "G18", "requirement": "将销售服务费表述改为费用收取不同。",
            "operations": [{"op": "replace", "old_text": OLD, "new_text": NEW,
                            "condition": {"kind": "none"}}]}


@pytest.mark.parametrize("changed_satisfied_site", [False, True])
def test_g18_source_satisfaction_is_site_local_and_requires_preserved_artifact(
        tmp_path, changed_satisfied_site):
    rule = _g18()
    record = _ledger([rule]).records[0]
    source = _inspection(tmp_path, "source", [SATISFIED, OLD + "，享有同等分配权。"])
    actual = _inspection(tmp_path, "actual", [
        SATISFIED.replace("A 类", "B 类") if changed_satisfied_site else SATISFIED,
        NEW + "，享有同等分配权。"])
    edit = dict(target_id="edit", rule_id="G18", location="p1", start=0, end=len(OLD),
                expected_old_text=OLD, replacement=NEW, operation_keys=["G18:0"])
    sites = [dict(location=f"p{i}", quote=text, state=state) for i, text, state in
             ((0, SATISFIED, "satisfied"), (1, OLD, "needs_edit"))]
    proof = evaluate_operation_effect(record, source, actual, sites, [edit], _facts(),
                                      validated_units=[{
                                          "location": "p1", "target_id": "edit",
                                          "operation_id": "G18:0", "verdict": "PASS",
                                          "before": OLD, "replacement": NEW,
                                      }])
    local = {s["location"]: s for s in proof.site_effects}
    assert local["p1"]["state"] == "PASS"
    if changed_satisfied_site:
        assert proof.state != "PASS"
        assert local["p0"]["state"] != "PASS"
    else:
        assert proof.state == "PASS"
        assert local["p0"]["artifact_status"] == "ALREADY_SATISFIED"
        assert local["p0"]["proof_kind"] == "source_operation_policy_equivalent"


def test_g18_satisfied_survey_binds_real_source_offsets(tmp_path):
    rule = _g18()
    source = _inspection(tmp_path, "source", [SATISFIED, OLD])
    synthetic, origins, rounds, gaps = [], {}, [], []
    _run_survey(None, [rule], {"G18": rule}, _facts(), source, [], synthetic,
                origins, rounds, gaps, [])
    target = next(t for t in synthetic if t.state == "satisfied")
    assert target.location == "p0"
    assert target.operations == ("G18:0",)
    assert target.resolution == "resolved"
    assert target.segment_ok
    assert target.start is not None and target.end > target.start
    assert _compact(target.quote) == _compact(
        source.by_location("p0").text[target.start:target.end])
    assert not gaps


def test_unbound_satisfied_survey_hit_is_a_gap_not_coverage(tmp_path, monkeypatch):
    from fundlex.recall import SurveyHit
    import fundlex.recall_pipeline as pipeline
    rule = _g18()
    source = _inspection(tmp_path, "source", [SATISFIED])
    monkeypatch.setattr(pipeline, "survey_targets", lambda *a, **kw: [SurveyHit(
        "G18", "satisfied", ("p0",), "不存在的引用", 0, 0, 1.0,
        operation_keys=("G18:0",))])
    synthetic, gaps = [], []
    _run_survey(None, [rule], {"G18": rule}, _facts(), source, [], synthetic,
                {}, [], gaps, [])
    assert synthetic == []
    assert any(gap["code"] == "SURVEY_QUOTE_UNRESOLVED" for gap in gaps)


def test_source_quote_binding_allows_word_spacing_and_terminal_punctuation(tmp_path):
    inspection = _inspection(tmp_path, "spacing", [
        "本基金 A 类基金份额不收取销售服务费。",
    ])
    span = inspection.by_location("p0")
    bound = resolve_source_quote(span, "本基金A类基金份额不收取销售服务费；")
    assert bound is not None
    assert span.text[bound[0]:bound[1]] == span.text


def test_source_quote_binding_rejects_duplicate_occurrences(tmp_path):
    inspection = _inspection(tmp_path, "duplicate", [
        "本基金 A 类基金份额不收取销售服务费。随后再次说明：本基金A类基金份额不收取销售服务费。",
    ])
    span = inspection.by_location("p0")
    assert resolve_source_quote(span, "本基金A类基金份额不收取销售服务费。") is None


def test_source_quote_binding_rejects_sentence_prefix_without_terminal_boundary(tmp_path):
    inspection = _inspection(tmp_path, "prefix", [
        "本基金 A 类基金份额不收取销售服务费，并按合同约定进行管理。",
    ])
    span = inspection.by_location("p0")
    assert resolve_source_quote(span, "本基金A类基金份额不收取销售服务费") is None


@pytest.mark.parametrize("quote", [
    "本基金A类基金份额不收取销售服务费",
    "本基金A类基金份额不收取销售服务费；",
])
def test_source_quote_binding_preserves_offsets_with_zero_width_and_spacing(tmp_path, quote):
    text = "本基金 A\u200b 类基金份额不收取销售服务费。"
    span = _inspection(tmp_path, "zero-width", [text]).by_location("p0")
    assert resolve_source_quote(span, quote) == (0, len(text))


def test_source_quote_binding_rejects_protected_and_cross_segment_ranges(tmp_path):
    from dataclasses import replace
    from fundlex.docx import BARRIER, TextSegment
    text = "本基金A类基金份额不收取销售服务费。"
    span = _inspection(tmp_path, "protected", [text]).by_location("p0")
    assert resolve_source_quote(replace(span, segments=()), text) is None
    assert resolve_source_quote(replace(span, segments=(TextSegment(0, 5),
                                TextSegment(5, len(text)))), text) is None
    with_barrier = text[:5] + BARRIER + text[5:]
    blocked = replace(span, text=with_barrier,
                      segments=(TextSegment(0, len(with_barrier)),))
    assert resolve_source_quote(blocked, text) is None


def _guide_rule(rule_id):
    guide = json.loads(open("guides/builtin_guide_v1.json", encoding="utf-8").read())
    return next(rule for rule in guide["groups"] if rule["rule_id"] == rule_id)


def test_single_class_structure_is_verified_only_without_class_markers(tmp_path):
    from fundlex.recall import Facts, single_class_structure_verified
    plain = _inspection(tmp_path, "single", ["本基金不设基金份额类别，管理费固定。"])
    assert single_class_structure_verified(plain) is True
    labelled = _inspection(tmp_path, "labelled", ["本基金分为 A 类和 C 类基金份额。"])
    assert single_class_structure_verified(labelled) is False
    facts = Facts("open_regular", "", "", (), False, "", "", "", {"fund_type": True}, (),
                  single_class_verified=True)
    assert facts.share_class_structure == "SINGLE"
    assert facts.class_count == 1


def test_unnamed_single_class_normalizes_model_default_and_verifies_no_service_fee(tmp_path):
    from fundlex.recall import facts_from_value
    from fundlex.docx import inspect_docx
    from docx import Document
    path = tmp_path / "unnamed-single.docx"
    doc = Document()
    doc.add_paragraph("申购费用由投资人承担，不列入基金财产。")
    doc.add_paragraph("本基金的管理费按前一日基金资产净值的0.80%年费率计提。")
    doc.save(path)
    facts = facts_from_value(inspect_docx(path, full=True), {
        "fund_type": {"value": "open_regular", "quote": "本基金的管理费"},
        "classes": [{"name": "默认份额类别", "subscription_fee": True,
                     "service_fee": False, "current": True,
                     "quote": "申购费用由投资人承担，不列入基金财产。",
                     "subscription_quote": "申购费用由投资人承担，不列入基金财产。"}],
        "floating_management_fee": {
            "value": False,
            "quote": "本基金的管理费按前一日基金资产净值的0.80%年费率计提。",
        },
    })
    assert facts.share_class_structure == "SINGLE"
    assert facts.class_count == 1
    assert facts.classes[0].name == "__single__"
    assert facts.status_of("share_class_structure") == "VERIFIED"
    assert facts.status_of("subscription_fee_group") == "VERIFIED"
    assert facts.status_of("service_fee_group") == "VERIFIED"
    assert facts.service_group == ()


def test_g16_partial_tail_is_reclassified_and_replaced_with_full_clause(tmp_path):
    from fundlex.recall import (Target, assign_target_operations, plan_target)
    tail = ("对于持续持有期限超过一年的基金份额，持有超过一年后继续计提的销售服务费，"
            "在投资者赎回基金份额或基金合同终止时，随赎回款（或清算款）一并返还给投资者。")
    inspection = _inspection(tmp_path, "g16-tail", [tail])
    rule = _guide_rule("G16")
    target = Target("T", "G16", "p0", tail, "satisfied", "",
                    start=0, end=len(tail), resolution="resolved",
                    operations=("G16:0",))
    assign_target_operations([target], {"G16": rule}, inspection, _facts())
    assert target.state == "needs_edit"
    assert "G16_PARTIAL_CLAUSE" in target.reason
    edit = plan_target(rule, target, inspection, _facts(), None, compose_budget=[0])
    assert edit is not None and edit.content_source == "g16_partial_clause_replace"
    assert edit.start == 0 and edit.end == len(tail)
    assert all(token in edit.replacement for token in
               ("每日计提", "直销机构", "其他销售机构", "超过一年"))


def test_satisfied_class_label_fragment_binds_a_unique_paragraph_site(tmp_path):
    from fundlex.recall import Target
    text = ("基金份额类别：本基金将基金份额分为 A 类和 C 类不同的类别。"
            "在投资者认购、申购基金时收取认购、申购费用，不计提销售服务费的，称为 A 类基金份额；"
            "在投资者认购、申购基金份额时不收取认购、申购费用，而是从本类别基金资产中计提销售服务费的，"
            "称为 C 类基金份额 55、发起式基金：指符合中国证监会有关规定。")
    inspection = _inspection(tmp_path, "class-fragment", [text])
    rule = _guide_rule("G02")
    ledger = _ledger([rule])
    target = Target("T", "G02", "p0", "称为 C 类基金份额", "satisfied", "",
                    start=120, end=130, resolution="resolved", operations=("G02:1",))
    result = replay_claims([], inspection, inspection, rules_by_id={"G02": rule},
                           facts=_facts(), obligations=ledger,
                           targets_by_id={"T": target.as_dict()})
    assert result["obligations"]["G02:B1:O2"]["state"] == "PASS"
    assert result["obligations"]["G02:B1:O2"]["artifact_status"] == "ALREADY_SATISFIED"


def test_g06_deletes_the_source_bound_class_conversion_prohibition(tmp_path):
    from fundlex.recall import Target, plan_target
    quote = "本基金不同基金份额类别之间不得互相转换。"
    text = "基金份额类别的具体设置由基金管理人确定。" + quote + "投资者可自行选择。"
    inspection = _inspection(tmp_path, "g06", [text])
    start = text.index(quote)
    target = Target("T", "G06", "p0", quote, "needs_edit", "",
                    start=start, end=start + len(quote), resolution="resolved",
                    operations=("G06:0",))
    edit = plan_target(_guide_rule("G06"), target, inspection, _facts(), None,
                       compose_budget=[0])
    assert edit is not None and edit.replacement == ""
    assert edit.start == start and edit.end == start + len(quote)
    assert edit.content_source == "g06_prohibition_delete"


def test_class_operation_recovery_binds_an_already_satisfied_group(tmp_path):
    from fundlex.recall import Target
    from fundlex.targeting import target_for_missing_operation
    text = ("基金份额类别：本基金将基金份额分为A类和C类不同的类别。"
            "在投资者认购、申购基金时收取认购、申购费用，不计提销售服务费的，称为A类基金份额；"
            "在投资者认购、申购基金份额时不收取认购、申购费用，而是从本类别基金资产中计提销售服务费的，称为C类基金份额。")
    inspection = _inspection(tmp_path, "class-recovery", [text])
    rule = _guide_rule("G02")
    record = SimpleNamespace(spec=SimpleNamespace(
        rule_id="G02", operation_id="G02:1", operation_index=1))
    section = SimpleNamespace(id="s1", locations=["p0"])
    target = target_for_missing_operation(
        record, [], {"G02": rule}, _facts(), inspection,
        {"G02": ["s1"]}, [section])
    assert target is not None and target.state == "satisfied"
    assert target.location == "p0" and target.quote == "C类基金份额"
    assert target.start is not None and target.end is not None


def test_g03_recovery_stays_in_scope_and_prefers_class_definition(tmp_path):
    """景宁 G03 p1195 is an income-distribution summary copy, not a G03 site."""
    from fundlex.recall import Target
    from fundlex.rule_scope import RuleScope
    from fundlex.targeting import target_for_missing_operation
    body = ("本基金根据申购费用收取方式的不同，将基金份额分为不同的类别。"
            "在投资者申购基金份额时收取申购费用，而不计提销售服务费的，称为A类基金份额；"
            "在投资者申购基金份额时不收取申购费用，而是从本类别基金资产中计提销售服务费的，称为C类基金份额。")
    summary = ("3、由于本基金A类基金份额不收取销售服务费，而C类基金份额收取销售服务费，"
               "各基金份额类别对应的可分配收益将有所不同。")
    inspection = _inspection(tmp_path, "g03-scope", [body, summary])
    record = SimpleNamespace(spec=SimpleNamespace(
        rule_id="G03", operation_id="G03:0", operation_index=0))
    body_section = SimpleNamespace(id="s1", title="第三部分基金的基本情况",
                                   zone="body", locations=["p0"])
    summary_section = SimpleNamespace(id="s2", title="第二十三部分基金合同内容摘要",
                                      zone="summary", locations=["p1"])
    scope = RuleScope(
        rule_id="G03", leaf_ids=frozenset({"s1:L01"}),
        locations=frozenset({"p0"}), zone="body",
        include_summary_copies=False, copy_keys=frozenset({"份额类别"}),
        resolution="resolved", evidence=())
    target = target_for_missing_operation(
        record, [], {"G03": _guide_rule("G03")}, _facts(), inspection,
        {"G03": ["s1", "s2"]}, [body_section, summary_section],
        scopes={"G03": scope})
    assert target is not None and target.location == "p0"
    assert target.site_candidates == ("p0",)
    assert not any(site == "p1" for site in target.site_candidates)


def test_g12_source_shape_selects_its_own_one_of_variant(tmp_path):
    from fundlex.recall import Target, _selected_operation_indices
    text = ("赎回费用由赎回基金份额的基金份额持有人承担，在基金份额持有人赎回基金份额时收取。"
            "不低于赎回费总额的25%应归基金财产，其余用于支付登记费和其他必要的手续费。"
            "对于持续持有期少于 7 日的投资者收取不低于 1.5%的赎回费，并将上述赎回费全额计入基金财产。")
    inspection = _inspection(tmp_path, "g12", [text])
    target = Target("T", "G12", "p0", text, "needs_edit", "",
                    start=0, end=len(text), resolution="resolved", operations=("G12:0",))
    assert _selected_operation_indices(_guide_rule("G12"), target, inspection, _facts()) == [1]


def test_satisfied_target_trims_quote_prefix_that_lives_in_the_previous_split(tmp_path):
    from fundlex.recall import Target
    from fundlex.recall_pipeline import _rebind_unresolved_targets
    inspection = _inspection(tmp_path, "g04-split", [
        "本基金根据申购费用收取方式的不同，将基金份额分为不同的类别。在投资",
        "者申购基金份额时收取申购费用，而不计提销售服务费的，称为 A 类基金份额；"
        "在投资者申购基金份额时不收取申购费用，而是从本类别基金资产中计提销售服务费的，"
        "称为 C 类基金份额。",
    ])
    quote = ("在投资者申购基金份额时收取申购费用，而不计提销售服务费的，"
             "称为 A 类基金份额")
    target = Target("T", "G04", "p1", quote, "satisfied", "")
    target.resolution = "quote_not_found"
    _rebind_unresolved_targets([target], {}, inspection,
                               rules_by_id={"G04": {}}, facts=_facts())
    assert target.location == "p1" and target.resolution == "resolved"
    assert target.start == 0 and target.quote.startswith("者申购基金份额")
    assert target.quote.endswith("称为 A 类基金份额")


def test_survey_satisfied_composite_duplicate_yields_to_actionable_target(tmp_path):
    from fundlex.recall import Target, dedupe_operation_targets
    texts = ["销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告",
             "费、促销活动费、基金份额持有人服务费等。"]
    inspection = _inspection(tmp_path, "g17", texts)
    composite = ("p0:0:32", "p1:0:20")
    actionable = Target("T", "G17", "p0", texts[0], "needs_edit", "",
                        resolution="composite", composite=composite, operations=("G17:0",))
    survey = Target("Y", "G17", "p0", texts[0], "satisfied", "",
                    resolution="composite", composite=composite, operations=("G17:0",))
    targets, _ = dedupe_operation_targets([actionable, survey],
                                          {"G17": _guide_rule("G17")}, inspection)
    assert actionable.out_of_scope is False
    assert survey.out_of_scope is True
    assert "DUPLICATE_OPERATION_TARGET:T" in survey.reason


def test_satisfied_target_rebinds_only_across_an_unterminated_split(tmp_path):
    from fundlex.recall import Target
    from fundlex.recall_pipeline import _rebind_unresolved_targets
    inspection = _inspection(tmp_path, "split", [
        "本基金根据申购费用收取方式的不同，将基金份额分为不同的类别。在投资",
        "者申购基金份额时收取申购费用，而不计提销售服务费的，称为 A 类基金份额；"
        "在投资者申购基金份额时不收取申购费用，而是从本类别基金资产中计提销售服务费的，"
        "称为 C 类基金份额。",
        "在每一份额类别内，本基金根据申购、赎回所使用货币的不同，再分为人民币份额和美元现汇份额。",
    ])
    target = Target("T1", "G04", "p0", "A 类基金份额", "satisfied", "")
    target.resolution = "quote_not_found"
    origins = {}
    records = _rebind_unresolved_targets(
        [target], origins, inspection, rules_by_id={"G04": {}}, facts=_facts())
    assert target.location == "p1" and target.resolution == "resolved"
    assert target.segment_ok is True
    assert origins["T1"] == "rebind_split_paragraph"
    assert records[0]["source"] == "split_paragraph"

    # A completed sentence is not a split; the quote in the next paragraph is
    # a different physical site and must not be borrowed.
    terminal = _inspection(tmp_path, "terminal", [
        "本基金将基金份额分为不同的类别。", "称为 A 类基金份额；称为 C 类基金份额。"])
    other = Target("T2", "G04", "p0", "A 类基金份额", "satisfied", "")
    other.resolution = "quote_not_found"
    _rebind_unresolved_targets([other], {}, terminal, rules_by_id={"G04": {}}, facts=_facts())
    assert other.location == "p0" and other.resolution == "quote_not_found"

    # Two adjacent paragraphs carrying the same quote are ambiguous.
    ambiguous = _inspection(tmp_path, "ambiguous", [
        "称为 A 类基金份额。", "本基金根据费用收取方式的不同，将基金份额分为不同的类别。在投资",
        "者申购基金份额时称为 A 类基金份额。"])
    ambiguous_target = Target("T3", "G04", "p1", "A 类基金份额", "satisfied", "")
    ambiguous_target.resolution = "quote_not_found"
    _rebind_unresolved_targets([ambiguous_target], {}, ambiguous,
                               rules_by_id={"G04": {}}, facts=_facts())
    assert ambiguous_target.location == "p1"
    assert ambiguous_target.resolution == "quote_not_found"


def test_satisfied_excerpt_cannot_pick_first_of_multiple_occurrences(tmp_path, monkeypatch):
    from fundlex.recall import SurveyHit
    import fundlex.recall_pipeline as pipeline
    rule = _g18()
    source = _inspection(tmp_path, "ambiguous-excerpt", [
        SATISFIED + "另由于基金费用的不同，本基金A类基金份额和C类基金份额另行披露。"])
    excerpt = "由于基金费用的不同，本基金A类基金份额和C"
    monkeypatch.setattr(pipeline, "survey_targets", lambda *a, **kw: [SurveyHit(
        "G18", "satisfied", ("p0",), excerpt, 0, 0, 1.0,
        operation_keys=("G18:0",))])
    synthetic, gaps = [], []
    _run_survey(None, [rule], {"G18": rule}, _facts(), source, [], synthetic,
                {}, [], gaps, [])
    assert synthetic == []
    assert any(gap["code"] == "SURVEY_QUOTE_UNRESOLVED" for gap in gaps)


def test_satisfied_class_excerpt_binds_whole_unique_policy_paragraph(tmp_path, monkeypatch):
    from fundlex.recall import SurveyHit
    import fundlex.recall_pipeline as pipeline
    rule = _guide_rule("G02")
    source = _inspection(tmp_path, "class-excerpt", [
        "本基金A类基金份额收取申购费用且不计提销售服务费，C类基金份额不收取申购费用而计提销售服务费。"
    ])
    monkeypatch.setattr(pipeline, "survey_targets", lambda *a, **kw: [SurveyHit(
        "G02", "satisfied", ("p0",), "C类基金份额", 0, 0, 1.0,
        operation_keys=("G02:0", "G02:1"))])
    synthetic, gaps = [], []
    pipeline._run_survey(None, [rule], {"G02": rule}, _facts(), source, [], synthetic,
                         {}, [], gaps, [])
    assert not gaps
    assert len(synthetic) == 1
    assert synthetic[0].location == "p0"
    assert synthetic[0].start == 0
    assert synthetic[0].end == len(source.by_location("p0").text)
