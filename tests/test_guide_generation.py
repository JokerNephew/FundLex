from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import pytest
import jsonschema
from docx import Document

from fundlex.guide_generation.corpus_catalog import parse_contract_filename, sync_corpus
from fundlex.guide_generation.corpus_segmenter import segment_contract
from fundlex.guide_generation.discovery import (
    DiscoveryConfig,
    _validate_sweep,
    path_b_contract_sweep,
    resolve_uncertain_batches,
)
from fundlex.guide_generation.lexical import LexicalIndex, deterministic_query
from fundlex.guide_generation.authoring import induce_families
from fundlex.guide_generation.coverage import run_coverage_gate
from fundlex.guide_generation.quality import run_quality_gate
from fundlex.guide_generation.prompts import QUALITY_SCHEMA
from fundlex.guide_generation.residual import ResidualConfig, run_residual_mining
from fundlex.guide_generation.models import (
    ClauseUnit,
    ContractVersion,
    CorpusSnapshot,
    GuideItem,
    ImpactCandidate,
    LocalGuideCandidate,
    PolicyRequirement,
    RegulationBlock,
    RegulationDocument,
)
from fundlex.guide_generation.requirements import extract_requirements
from fundlex.regulatory.docx_parser import parse_regulation_docx
from fundlex.regulatory.ingestion import ingest_regulation
from fundlex.regulatory.pdf_parser import parse_regulation_pdf
from fundlex.store import RunStore
from fundlex.guide_generation.store import GuideGenerationStore

from guide_generation_fakes import FakeGuideRouter


def _contract(path: Path, fund: str, date: str, *, service_fee: str = "0.40%"):
    doc = Document()
    doc.add_paragraph("基金管理人：测试基金管理有限公司")
    doc.add_paragraph("基金托管人：测试银行")
    doc.add_paragraph(f"{fund}基金合同")
    doc.add_paragraph("第一部分  前言")
    doc.add_paragraph("订立本基金合同的目的是保护投资人合法权益。")
    doc.add_paragraph("第二部分  基金费用与税收")
    doc.add_paragraph("一、基金费用的种类")
    doc.add_paragraph("本基金的管理费按前一日基金资产净值的1.2%年费率计提。")
    doc.add_paragraph("二、基金销售服务费")
    doc.add_paragraph("本基金A类基金份额不收取销售服务费，C类基金份额的销售服务费按前一日C类基金份额资产净值的"
                      f"{service_fee}年费率计提。")
    doc.add_paragraph("C类基金份额的销售服务费每日计提，按月支付。")
    doc.add_paragraph("第三部分  基金份额的申购与赎回")
    doc.add_paragraph("六、申购和赎回的价格、费用及其用途")
    doc.add_paragraph("本基金A类基金份额的申购费率由基金管理人决定，并在招募说明书中列示。")
    doc.add_paragraph("赎回费用由赎回基金份额的基金份额持有人承担，在基金份额持有人赎回基金份额时收取。")
    doc.add_paragraph("第四部分  基金合同")
    doc.add_paragraph("本基金合同自基金份额持有人、基金管理人和基金托管人盖章之日起生效。")
    doc.save(path)


def _regulation_docx(path: Path):
    doc = Document()
    doc.add_paragraph("测试销售费用管理规定")
    doc.add_paragraph("第一条")
    doc.add_paragraph("基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。")
    doc.add_paragraph("第二条")
    doc.add_paragraph("对于投资者持续持有期限少于七日的基金份额，应当收取不低于赎回金额1.5%的赎回费，"
                      "并将上述赎回费全额计入基金财产。")
    doc.add_paragraph("第三条")
    doc.add_paragraph("基金管理人应当在基金合同、基金销售协议等法律文件中与相关主体约定上述费用的收取标准。")
    doc.save(path)


def test_docx_regulation_parse(tmp_path):
    path = tmp_path / "reg.docx"
    _regulation_docx(path)
    document = parse_regulation_docx(path)
    assert document.source_type == "docx"
    assert any(block.article_no == "第一条" for block in document.blocks)
    assert any("不得收取" in block.text for block in document.blocks)
    assert document.document_sha256


def test_reference_pdf_parse():
    path = Path("references/source/附件1：公开募集证券投资基金销售费用管理规定.pdf")
    if not path.is_file():
        pytest.skip("reference regulation pdf is not present")
    document = parse_regulation_pdf(path)
    assert document.page_count == 10
    assert len(document.blocks) > 50
    assert any(block.article_no == "第一条" for block in document.blocks)
    assert "销售费用" in document.title
    assert any("不得收取认" in block.text for block in document.blocks)


def test_contract_segmentation_and_sections(tmp_path):
    path = tmp_path / "fund_20240101.docx"
    _contract(path, "测试基金", "20240101")
    segmented = segment_contract(path, contract_version_id="CV-test")
    assert segmented.clauses
    assert any("基金费用与税收" in " ".join(clause.section_path) for clause in segmented.clauses)
    fee = next(clause for clause in segmented.clauses if "销售服务费" in clause.text)
    assert fee.paragraph_start >= 0
    assert fee.paragraph_end >= fee.paragraph_start
    assert fee.source_anchor.get("paragraph_start") is not None


def test_corpus_versioning(tmp_path):
    root = tmp_path / "corpus"
    root.mkdir()
    _contract(root / "测试基金基金合同_20230101.docx", "测试基金", "20230101")
    _contract(root / "测试基金基金合同_20240101.docx", "测试基金", "20240101")
    _contract(root / "另一基金基金合同_20240101.docx", "另一基金", "20240101")
    store = RunStore(tmp_path / "run")
    gg = GuideGenerationStore(store)
    result = sync_corpus(gg, root)
    versions = gg.list_contract_versions()
    assert len(versions) == 3
    assert result.snapshot.counts["latest"] == 2
    same_fund = [item for item in versions if item.fund_identity == "测试基金"]
    assert len(same_fund) == 2
    assert same_fund[0].contract_id == same_fund[1].contract_id
    assert any(item.previous_version_id for item in same_fund)
    assert any(item.latest_known_version for item in same_fund)
    assert not any(item.latest_known_version for item in same_fund if item.version_date == "20230101")
    clauses = gg.get_clauses(same_fund[-1].contract_version_id)
    assert clauses


def test_lexical_index_finds_rare_fee_clause(tmp_path):
    path = tmp_path / "c.docx"
    _contract(path, "测试基金", "20240101")
    segmented = segment_contract(path, contract_version_id="CV-test")
    index = LexicalIndex(segmented.clauses)
    requirement = PolicyRequirement(
        requirement_id="R-test", source_block_ids=["B1"], source_quote="不得收取销售服务费",
        article="第一条", norm_type="prohibition", subject="基金管理人", object="销售服务费",
        requirement="基金管理人不得收取销售服务费", conditions=[], exceptions=[],
        candidate_contract_domains=["基金费用与税收"], disposition="modify_contract")
    hits = index.search(deterministic_query(requirement))
    assert hits
    assert any("销售服务费" in segmented.clauses[hit.clause_index].text for hit in hits[:20])


def test_requirement_extraction_quote_validation(tmp_path):
    path = tmp_path / "reg.docx"
    _regulation_docx(path)
    document = ingest_regulation(path, original_filename=path.name)
    router = FakeGuideRouter()
    result = extract_requirements(router, document)
    assert result.requirements
    assert all(item.disposition for item in result.requirements)
    assert all(item.source_quote for item in result.requirements)
    by_id = {block.block_id: block.text for block in document.blocks}
    for requirement in result.requirements:
        joined = "\n".join(by_id[block_id] for block_id in requirement.source_block_ids)
        assert requirement.source_quote[:20] in joined or joined[:20] in requirement.source_quote

    class InventedRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            value = super()._value(task_type, context)
            if task_type == "guide_req_extract":
                for item in value.get("requirements", []):
                    item["source_quote"] = "这是模型编造的法规原文"
            return value

    invented = extract_requirements(InventedRouter(), document)
    assert not any("编造" in item.source_quote for item in invented.requirements)
    assert all(item.disposition for item in invented.requirements)
    assert any(item.adjudication_status == "unresolved_article" for item in invented.requirements)
    assert any(item["severity"] == "critical" for item in invented.findings)


def test_requirement_review_is_batched_for_long_regulation():
    blocks = [
        RegulationBlock(f"B{index}", index, "paragraph", f"第一条规则内容 {index}。" + "费用条款。" * 700,
                        article_no=f"第{index + 1}条")
        for index in range(3)
    ]
    document = RegulationDocument("REG-REVIEW-BATCH", "pdf", "sha", "长规定", blocks)

    class ReviewBatchRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.review_contexts = []

        def call(self, *, task_type, bounded_context, **kwargs):
            if task_type == "guide_req_review":
                self.review_contexts.append(bounded_context)
            return super().call(task_type=task_type, bounded_context=bounded_context, **kwargs)

    router = ReviewBatchRouter()
    result = extract_requirements(router, document)
    assert result.requirements
    assert len(router.review_contexts) == 3
    assert all("Review batch:" in context for context in router.review_contexts)
    assert all(len(context) < 11000 for context in router.review_contexts)


def test_requirement_extraction_uses_small_output_safe_batches():
    blocks = [
        RegulationBlock(f"B{index}", index, "paragraph", "基金销售机构应当保存销售业务资料。" * 8,
                        article_no=f"第{index + 1}条")
        for index in range(20)
    ]
    document = RegulationDocument("REG-EXTRACT-BATCH", "pdf", "sha", "长规定", blocks)

    class ExtractBatchRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.extract_contexts = []

        def call(self, *, task_type, bounded_context, **kwargs):
            if task_type == "guide_req_extract":
                self.extract_contexts.append((bounded_context, kwargs.get("prompt_version")))
            return super().call(task_type=task_type, bounded_context=bounded_context, **kwargs)

    router = ExtractBatchRouter()
    result = extract_requirements(router, document)
    assert result.requirements
    initial_contexts = [context for context, version in router.extract_contexts
                        if version == "guide_req_extract_v2"]
    assert len(initial_contexts) >= 2
    assert all(len(context) < 3000 for context in initial_contexts)


def test_requirement_source_supplement_only_repairs_cited_partial_span():
    blocks = [
        RegulationBlock("B0001", 0, "heading", "第一条", article_no="第一条", heading_level=3),
        RegulationBlock("B0002", 1, "paragraph", "前端收费方式可以按认（申）购金额（份额）数量适用不同费率。",
                        article_no="第一条"),
        RegulationBlock("B0003", 2, "paragraph", "后端收费方式可以按持有期限适用不同费率标准。",
                        article_no="第一条"),
        RegulationBlock("B0004", 3, "paragraph", "基金销售机构应当保存销售业务资料。",
                        article_no="第一条"),
    ]
    document = RegulationDocument("REG-SUPPLEMENT", "docx", "sha", "测试规定", blocks)

    class PartialSpanRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_req_extract":
                return {
                    "requirements": [{
                        "source_block_ids": ["B0002", "B0003"],
                        "source_quote": blocks[1].text,
                        "article": "第一条", "norm_type": "permission", "subject": "基金销售机构",
                        "object": "收费方式", "requirement": blocks[1].text,
                        "conditions": [], "exceptions": [], "applicable_scope": {},
                        "effective_time": "", "candidate_contract_domains": ["费用条款"],
                        "external_dependencies": [], "disposition": "modify_contract",
                        "disposition_reason": "partial span",
                    }],
                    "article_dispositions": [{
                        "article": "第一条", "disposition": "modify_contract", "reason": "partial span",
                    }],
                    "notes": "",
                }
            if task_type == "guide_req_review":
                return {"missing_requirements": [], "corrections": [],
                        "article_assessment": [], "notes": ""}
            return super()._value(task_type, context)

    result = extract_requirements(PartialSpanRouter(), document)
    supplement = next(item for item in result.requirements
                       if item.source_block_ids == ["B0003"])
    assert supplement.payload["source_supplement"] is True
    assert supplement.source_quote == blocks[2].text
    # B0004 was never cited by the model and must remain on normal coverage
    # paths instead of being silently manufactured as a source supplement.
    assert not any(item.source_block_ids == ["B0004"] for item in result.requirements)


def test_authoring_impact_context_carries_coverage_topic_reason():
    from fundlex.guide_generation.authoring import _impact_context

    impact = ImpactCandidate(
        impact_id="IC-topic-context", requirement_id="R-topic-context",
        contract_version_id="CV-topic-context", clause_ids=["CL-topic-context"],
        source_text="基金管理人收取申购费。", context="费用",
        impact_status="possible", discovery_sources=["coverage_topic"],
        impact_reason="Coverage gap: 客户维护费计提与上限",
    )
    clause = ClauseUnit("CL-topic-context", "CV-topic-context", ["费用"], 1, 1,
                        "基金管理人收取申购费。")
    text = _impact_context([impact], {clause.clause_id: clause}, 5000)
    assert "Coverage gap" in text
    assert "客户维护费计提与上限" in text


def test_direct_sale_prohibition_qualifies_historical_fee_pattern():
    requirement = PolicyRequirement(
        requirement_id="R-direct", source_block_ids=["B1"],
        source_quote="基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。",
        article="第七条", norm_type="prohibition", subject="基金管理人",
        object="认（申）购费、销售服务费",
        requirement="基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。",
        disposition="modify_contract",
    )
    impact = ImpactCandidate(
        impact_id="IC-direct", requirement_id="R-direct", contract_version_id="CV-direct",
        clause_ids=["CL-direct"], source_text="A类及D类基金份额收取申购费用。",
        context="费用条款", impact_status="definite", discovery_sources=["lexical"],
    )
    candidate = LocalGuideCandidate(
        local_id="LG-direct", impact_ids=["IC-direct"], requirement_ids=["R-direct"],
        needs_modification="yes", business_goal="调整费用条款",
        new_text=("基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。"
                  "本基金A类及D类基金份额收取申购费用，不计提销售服务费。"),
    )
    family_result = induce_families(
        FakeGuideRouter(), [candidate], [requirement],
        {"CL-direct": ClauseUnit("CL-direct", "CV-direct", ["基金费用"], 1, 1,
                                  impact.source_text)},
    )
    family = family_result.families[0]
    assert "基金管理人销售其管理的基金，不得收取" in family["new_text"]
    assert "通过其他基金销售机构销售" in family["new_text"]
    assert any("直销" in value for value in family["conditions"])


def test_repair_reapplies_direct_sale_channel_boundary():
    from fundlex.guide_generation.repair import repair_items

    requirement = PolicyRequirement(
        requirement_id="R-direct-repair", source_block_ids=["B-direct-repair"],
        source_quote="基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。",
        article="第七条", norm_type="prohibition", subject="基金管理人",
        object="认（申）购费、销售服务费",
        requirement="基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。",
        disposition="modify_contract",
    )
    impact = ImpactCandidate(
        impact_id="IC-direct-repair", requirement_id="R-direct-repair",
        contract_version_id="CV-direct-repair", clause_ids=["CL-direct-repair"],
        source_text="A类基金份额收取申购费用。", context="费用",
        impact_status="definite", discovery_sources=["lexical"],
    )
    item = GuideItem(
        guide_item_id="G-direct-repair", requirement_ids=["R-direct-repair"],
        impact_ids=["IC-direct-repair"], clause_family_ids=[], section="费用",
        title="直销收费", business_requirement="调整收费条款", instruction="调整收费条款",
        new_text="基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。",
        new_text_origin="model_drafted",
    )
    finding = {
        "finding_id": "F-direct-repair", "guide_item_id": "G-direct-repair",
        "category": "regulation_fidelity", "severity": "critical", "status": "open",
        "description": "changed subject", "evidence_quote": item.new_text,
    }

    class RepairRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_repair":
                return {
                    "title": "直销收费", "business_requirement": "调整收费条款",
                    "instruction": "调整收费条款", "applicability": "费用条款",
                    "conditions": [], "exceptions": [], "old_text_patterns": [],
                    "new_text": ("基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。"
                                 "A类基金份额收取申购费用。"),
                    "new_text_origin": "model_drafted",
                }
            return super()._value(task_type, context)

    repaired, fixes, calls = repair_items(
        RepairRouter(), [item], [requirement], [impact], {}, [finding],
    )
    assert calls == 1
    assert repaired
    assert "其他基金销售机构" in repaired[0].new_text


def test_quality_gate_blocks_unqualified_direct_sale_fee_pattern():
    requirement = PolicyRequirement(
        requirement_id="R-quality-direct", source_block_ids=["B1"],
        source_quote="基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。",
        article="第七条", norm_type="prohibition", subject="基金管理人",
        object="认（申）购费、销售服务费",
        requirement="基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。",
        disposition="modify_contract",
    )
    impact = ImpactCandidate(
        impact_id="IC-quality-direct", requirement_id="R-quality-direct",
        contract_version_id="CV-quality-direct", clause_ids=["CL-quality-direct"],
        source_text="A类基金份额收取申购费用。", context="费用条款",
        impact_status="definite", discovery_sources=["lexical"],
    )
    item = GuideItem(
        guide_item_id="G-quality-direct", requirement_ids=["R-quality-direct"],
        impact_ids=["IC-quality-direct"], clause_family_ids=["CF-quality-direct"],
        section="基金费用", title="直销收费", business_requirement="调整收费条款",
        instruction="调整收费条款", new_text=(
            "基金管理人销售其管理的基金，不得收取认（申）购费、销售服务费。"
            "A类基金份额收取申购费用，不计提销售服务费。"
        ), new_text_origin="contract_pattern_adapted",
    )
    result = run_quality_gate(
        FakeGuideRouter(), [item], [requirement], [impact],
        {"CL-quality-direct": ClauseUnit("CL-quality-direct", "CV-quality-direct",
                                          ["基金费用"], 1, 1, impact.source_text)},
    )
    assert result.run_status == "fail"
    assert any(value["severity"] == "critical" and "渠道边界" in value["description"]
               for value in result.findings)


def test_effective_or_repeal_requirement_is_not_a_contract_modification():
    block = RegulationBlock(
        block_id="B0001", ordinal=0, type="paragraph",
        text="本规定自2026年1月1日起施行，证监会公告〔2013〕26号同时废止。",
        article_no="第二十九条",
    )
    document = RegulationDocument(document_id="REG-EFF", source_type="docx",
                                  source_sha256="x", title="测试规定", blocks=[block])

    class EffectiveRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_req_extract":
                return {"requirements": [{
                    "source_block_ids": ["B0001"],
                    "source_quote": "本规定自2026年1月1日起施行，证监会公告〔2013〕26号同时废止。",
                    "article": "第二十九条", "norm_type": "transition",
                    "subject": "本规定", "object": "施行与废止",
                    "requirement": "本规定自2026年1月1日起施行，原公告同时废止。",
                    "conditions": [], "exceptions": [],
                    "applicable_scope": {"fund_types": [], "channels": [],
                                         "share_classes": [], "investor_types": [], "notes": ""},
                    "effective_time": "2026-01-01", "candidate_contract_domains": [],
                    "external_dependencies": [], "disposition": "modify_contract",
                    "disposition_reason": "fake",
                }], "article_dispositions": [{
                    "article": "第二十九条", "source_block_ids": ["B0001"],
                    "disposition": "modify_contract", "reason": "fake",
                }], "notes": ""}
            return super()._value(task_type, context)

    result = extract_requirements(EffectiveRouter(), document)
    assert result.requirements
    assert all(item.disposition == "out_of_contract_scope"
               for item in result.requirements)


def test_article_disposition_prevents_silent_omission(tmp_path):
    blocks = [
        RegulationBlock(block_id="B0001", ordinal=0, type="heading", text="第一条", article_no="第一条",
                                heading_level=3),
        RegulationBlock(block_id="B0002", ordinal=1, type="paragraph",
                                text="基金管理人应当在基金合同中约定销售费用。", article_no="第一条"),
        RegulationBlock(block_id="B0003", ordinal=2, type="heading", text="第二条", article_no="第二条",
                                heading_level=3),
        RegulationBlock(block_id="B0004", ordinal=3, type="paragraph",
                                text="中国证监会对基金销售业务实施监督管理。", article_no="第二条"),
    ]
    document = RegulationDocument(document_id="REG", source_type="docx", source_sha256="x",
                                  title="测试规定", blocks=blocks)

    class SparseRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_req_extract":
                return {
                    "requirements": [{
                        "source_block_ids": ["B0002"], "source_quote": "基金管理人应当在基金合同中约定销售费用。",
                        "article": "第一条", "norm_type": "obligation", "subject": "基金管理人",
                        "object": "销售费用", "requirement": "基金合同应约定销售费用",
                        "conditions": [], "exceptions": [],
                        "applicable_scope": {"fund_types": [], "channels": [], "share_classes": [],
                                             "investor_types": [], "notes": ""},
                        "effective_time": "", "candidate_contract_domains": ["基金费用与税收"],
                        "external_dependencies": [], "disposition": "modify_contract",
                        "disposition_reason": "fake",
                    }],
                    "article_dispositions": [{
                        "article": "第一条", "source_block_ids": ["B0001", "B0002"],
                        "disposition": "modify_contract", "reason": "fake",
                    }, {
                        "article": "第二条", "source_block_ids": ["B0003", "B0004"],
                        "disposition": "out_of_contract_scope", "reason": "监管职责",
                    }],
                    "notes": "",
                }
            return super()._value(task_type, context)

    result = extract_requirements(SparseRouter(), document)
    assert result.requirements
    assert all(item.disposition for item in result.requirements)
    assert not any(item["category"] == "requirement_coverage" and item["severity"] == "critical"
                   for item in result.findings)


def test_coverage_audit_batches_unresolved_impacts_and_keeps_unbound_covered_open():
    requirement = PolicyRequirement(
        requirement_id="R-batch", source_block_ids=["B1"], source_quote="应当调整费用。",
        article="第一条", norm_type="obligation", subject="基金管理人", object="费用",
        requirement="基金管理人应当调整费用。", disposition="modify_contract",
    )
    impacts = [ImpactCandidate(
        impact_id=f"IC-batch-{index}", requirement_id="R-batch", contract_version_id="CV-batch",
        clause_ids=[f"CL-batch-{index}"], source_text="费用条款", context="费用",
        impact_status="possible", discovery_sources=["lexical"],
    ) for index in range(81)]
    version = ContractVersion("CV-batch", "C-batch", "测试基金", "测试合同", latest_known_version=True)
    snapshot = CorpusSnapshot("S-batch", ["CV-batch"], {"CV-batch": "sha"},
                              ["CV-batch"], [], "parser", "index")

    class BatchAuditRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.contexts = []

        def call(self, *, bounded_context, **kwargs):
            self.contexts.append(bounded_context)
            return super().call(bounded_context=bounded_context, **kwargs)

        def _value(self, task_type, context):
            if task_type == "guide_coverage_audit":
                ids = re.findall(r"- (IC-batch-\d+) req=", context)
                results = []
                for impact_id in ids:
                    if impact_id == "IC-batch-0":
                        results.append({"impact_id": impact_id, "status": "covered",
                                        "guide_item_ids": [], "reason": "无绑定指引"})
                    else:
                        results.append({"impact_id": impact_id, "status": "false_positive",
                                        "guide_item_ids": [], "reason": "批量复核"})
                return {"requirement_results": [], "impact_results": results,
                        "missing_topics": [], "notes": "batch"}
            return super()._value(task_type, context)

    router = BatchAuditRouter()
    existing_item = GuideItem(
        guide_item_id="GI-existing", requirement_ids=["R-batch"], impact_ids=["IC-other"],
        clause_family_ids=[], section="费用", title="已有候选指引", business_requirement="调整费用",
        instruction="按要求复核费用条款", new_text="调整费用条款。",
    )
    result = run_coverage_gate(
        router, [requirement], impacts, [existing_item],
        {impact.clause_ids[0]: ClauseUnit(impact.clause_ids[0], "CV-batch", ["费用"], 1, 1,
                                          impact.source_text) for impact in impacts},
        [version], [], snapshot, semantic_applicability=False,
    )
    audit_calls = [record for record in router.records if record["task_type"] == "guide_coverage_audit"]
    assert len(audit_calls) == 4  # requirements + ceil(81 / 40) impact batches
    assert "GI-existing" in router.contexts[0]
    assert result.impact_summary["false_positive"] == 80
    assert result.impact_summary["unresolved"] == 1
    unresolved = next(row for row in result.rows
                      if row["level"] == "impact" and row["ref_id"] == "IC-batch-0")
    assert unresolved["status"] == "unresolved"


def test_guide_scope_comes_from_builtin_revision_guide():
    from fundlex.guide_generation.scope import load_builtin_guide_scope, section_in_scope

    scope = load_builtin_guide_scope()
    assert "基金费用与税收" in scope
    assert "基金份额的申购与赎回" in scope
    assert "基金财产的清算" in scope
    assert "基金的信息披露" not in scope
    assert section_in_scope("第十五部分  基金费用与税收", scope)
    assert section_in_scope("第三部分  基金的基本情况", scope)
    assert not section_in_scope("第十八部分  基金的信息披露", scope)
    assert not section_in_scope("第七部分  基金合同当事人及权利义务", scope)


def test_coverage_marks_out_of_scope_requirement_without_gap():
    requirements = [
        PolicyRequirement(
            requirement_id="R-scope-in", source_block_ids=["B-in"],
            source_quote="基金管理人应当约定销售费用。", article="第一条",
            norm_type="obligation", subject="基金管理人", object="销售费用",
            requirement="基金管理人应当约定销售费用。",
            candidate_contract_domains=["基金费用与税收"],
            disposition="modify_contract",
        ),
        PolicyRequirement(
            requirement_id="R-scope-out", source_block_ids=["B-out"],
            source_quote="基金管理人应当披露其他信息。", article="第二条",
            norm_type="obligation", subject="基金管理人", object="其他信息",
            requirement="基金管理人应当披露其他信息。",
            candidate_contract_domains=["基金的信息披露"],
            disposition="modify_contract",
        ),
    ]
    version = ContractVersion("CV-scope-gate", "C-scope-gate", "测试基金",
                              "测试合同", latest_known_version=True)
    snapshot = CorpusSnapshot("S-scope-gate", ["CV-scope-gate"],
                              {"CV-scope-gate": "sha"}, ["CV-scope-gate"],
                              [], "parser", "index")

    class EmptyAuditRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_coverage_audit":
                return {"requirement_results": [], "impact_results": [],
                        "missing_topics": [], "notes": ""}
            return super()._value(task_type, context)

    result = run_coverage_gate(
        EmptyAuditRouter(), requirements, [], [], {}, [version], [], snapshot,
        semantic_applicability=False, scope_sections=["基金费用与税收"],
    )
    rows = {row["ref_id"]: row for row in result.rows if row["level"] == "requirement"}
    assert rows["R-scope-in"]["status"] == "gap"
    assert rows["R-scope-out"]["status"] == "out_of_scope"
    assert result.requirement_summary["gap"] == 1
    assert result.requirement_summary["out_of_scope"] == 1
    critical_refs = {item.get("description", "") for item in result.findings
                     if item.get("severity") == "critical"}
    assert any("R-scope-in" in value for value in critical_refs)
    assert not any("R-scope-out" in value for value in critical_refs)


def test_coverage_requirement_audit_is_batched_and_indexes_every_guide_item():
    requirements = [
        PolicyRequirement(
            requirement_id=f"R-batch-audit-{index}", source_block_ids=[f"B-{index}"],
            source_quote="基金管理人应当调整费用。", article="第一条",
            norm_type="obligation", subject="基金管理人", object="费用",
            requirement=f"第{index}项费用要求。", disposition="modify_contract",
        ) for index in range(40)
    ]
    items = [
        GuideItem(
            guide_item_id=f"G-batch-audit-{index}",
            requirement_ids=[f"R-batch-audit-{index}"], impact_ids=[],
            clause_family_ids=[], section="费用", title=f"费用指引{index}",
            business_requirement="调整费用", instruction="调整费用条款",
            new_text="调整费用条款。",
        ) for index in range(40)
    ]
    version = ContractVersion("CV-batch-audit", "C-batch-audit", "测试基金", "测试合同",
                              latest_known_version=True)
    snapshot = CorpusSnapshot("S-batch-audit", ["CV-batch-audit"], {"CV-batch-audit": "sha"},
                              ["CV-batch-audit"], [], "parser", "index")

    class AuditRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.contexts = []

        def call(self, *, bounded_context, **kwargs):
            self.contexts.append(bounded_context)
            return super().call(bounded_context=bounded_context, **kwargs)

        def _value(self, task_type, context):
            if task_type == "guide_coverage_audit":
                return {"requirement_results": [], "impact_results": [],
                        "missing_topics": [], "notes": ""}
            return super()._value(task_type, context)

    router = AuditRouter()
    result = run_coverage_gate(
        router, requirements, [], items, {}, [version], [], snapshot,
        semantic_applicability=False,
    )
    requirement_contexts = [context for context in router.contexts
                            if "[ALL_GUIDE_ITEM_INDEX]" in context]
    assert len(requirement_contexts) == 3  # ceil(40 / 16)
    assert all(len(context) <= 24000 for context in requirement_contexts)
    joined = "\n".join(requirement_contexts)
    assert all(f"G-batch-audit-{index}" in joined for index in range(40))
    assert result.requirement_summary["covered"] == 40


def test_topic_coverage_is_a_registered_model_task():
    from fundlex.model import QUOTE_SELF_VALIDATED_TASKS, TASK_TYPES

    assert "guide_topic_coverage" in TASK_TYPES
    assert "guide_topic_coverage" in QUOTE_SELF_VALIDATED_TASKS


def test_contract_requirement_coverage_is_scoped_to_batch_window():
    requirements = [
        PolicyRequirement(
            requirement_id=f"R-scope-{index}", source_block_ids=[f"B-{index}"],
            source_quote="基金管理人应当调整费用。", article="第一条",
            norm_type="obligation", subject="基金管理人", object="费用",
            requirement=f"第{index}项费用要求。", disposition="modify_contract",
        ) for index in range(2)
    ]
    version = ContractVersion("CV-scope", "C-scope", "测试基金", "测试合同",
                              latest_known_version=True)
    snapshot = CorpusSnapshot("S-scope", ["CV-scope"], {"CV-scope": "sha"},
                              ["CV-scope"], [], "parser", "index")

    class ScopeAuditRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_coverage_audit":
                return {"requirement_results": [], "impact_results": [],
                        "missing_topics": [], "notes": ""}
            return super()._value(task_type, context)

    result = run_coverage_gate(
        ScopeAuditRouter(), requirements, [], [], {}, [version], [
            {"contract_version_id": "CV-scope", "clause_ids": ["CL-1"],
             "requirement_ids": ["R-scope-0"], "status": "affected"},
            {"contract_version_id": "CV-scope", "clause_ids": ["CL-2"],
             "requirement_ids": ["R-scope-1"], "status": "not_affected"},
        ], snapshot, semantic_applicability=False,
    )
    rows = {row["ref_id"]: row for row in result.rows
            if row["level"] == "contract_requirement"}
    assert rows["CV-scope:R-scope-1"]["status"] == "not_affected"
    assert rows["CV-scope:R-scope-0"]["status"] == "uncertain"


def test_sweep_coverage_tail_covers_every_requirement_per_contract():
    requirements = [
        PolicyRequirement(
            requirement_id=f"R-tail-{index}", source_block_ids=[f"B-{index}"],
            source_quote="基金管理人应当调整费用。", article="第一条",
            norm_type="obligation", subject="基金管理人", object="费用",
            requirement=f"第{index}项费用要求。", disposition="modify_contract",
        ) for index in range(20)
    ]
    versions = [
        ContractVersion(f"CV-tail-{name}", f"C-tail-{name}", "测试基金", "测试合同",
                        latest_known_version=True) for name in ("a", "b")
    ]
    clauses_by_version = {
        "CV-tail-a": [ClauseUnit("CL-tail-a", "CV-tail-a", ["费用"], 1, 1,
                                 "基金管理人收取申购费。")],
        "CV-tail-b": [ClauseUnit("CL-tail-b", "CV-tail-b", ["费用"], 1, 1,
                                 "基金管理人收取赎回费。")],
    }

    class NotAffectedRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_impact_sweep":
                return {"findings": [], "uncovered_suspicions": [],
                        "section_assessment": "not_affected"}
            return super()._value(task_type, context)

    impacts, suspicions, evidence, calls = path_b_contract_sweep(
        NotAffectedRouter(), None, requirements, versions, clauses_by_version,
        DiscoveryConfig(path_b_batch_chars=5000, path_b_max_requirements=6,
                        path_b_max_calls=100, path_c_enabled=False), max_workers=1)
    assert impacts == []
    for version in versions:
        covered = set()
        for batch in evidence:
            if batch.get("contract_version_id") == version.contract_version_id:
                covered.update(batch.get("requirement_ids") or [])
        assert covered == {item.requirement_id for item in requirements}
    assert calls == len(evidence)


def test_residual_circuit_open_waits_and_retries_once():
    from fundlex.model import ModelError

    requirement = PolicyRequirement(
        requirement_id="R-circuit", source_block_ids=["B-circuit"],
        source_quote="基金管理人应当调整费用。", article="第一条",
        norm_type="obligation", subject="基金管理人", object="费用",
        requirement="基金管理人应当调整费用。", disposition="modify_contract",
    )
    version = ContractVersion("CV-circuit", "C-circuit", "测试基金", "测试合同",
                              latest_known_version=True)
    clause = ClauseUnit("CL-circuit", "CV-circuit", ["费用"], 1, 1,
                        "基金管理人收取申购费。")

    class CircuitRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.sweep_calls = 0

        def _value(self, task_type, context):
            if task_type == "guide_residual_sweep":
                self.sweep_calls += 1
                if self.sweep_calls == 1:
                    raise ModelError("MODEL_CIRCUIT_OPEN")
                return {"uncovered": [], "guide_conflicts": [], "notes": ""}
            return super()._value(task_type, context)

    router = CircuitRouter()
    result = run_residual_mining(
        router, [requirement], [], [], [version],
        {"CV-circuit": [clause]}, {"CL-circuit": clause},
        [{"contract_version_id": "CV-circuit", "clause_ids": ["CL-circuit"],
          "status": "affected"}],
        config=ResidualConfig(max_rounds=1, required_clean_rounds=1,
                              batch_chars=5000, max_calls_per_round=2,
                              circuit_wait_s=0.01),
    )
    assert router.sweep_calls == 2
    assert result.saturated is True


def test_sweep_retry_does_not_leave_stale_model_error_evidence():
    requirement = PolicyRequirement(
        requirement_id="R-retry", source_block_ids=["B-retry"],
        source_quote="基金管理人应当调整费用。", article="第一条",
        norm_type="obligation", subject="基金管理人", object="费用",
        requirement="基金管理人应当调整费用。", disposition="modify_contract",
    )
    version = ContractVersion("CV-retry", "C-retry", "测试基金", "测试合同",
                              latest_known_version=True)
    clause = ClauseUnit("CL-retry", "CV-retry", ["费用"], 1, 1,
                        "基金管理人收取申购费。")

    class FlakyRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.sweep_calls = 0

        def _value(self, task_type, context):
            if task_type == "guide_impact_sweep":
                self.sweep_calls += 1
                if self.sweep_calls == 1:
                    raise RuntimeError("transient")
                return {"findings": [], "uncovered_suspicions": [],
                        "section_assessment": "not_affected"}
            return super()._value(task_type, context)

    impacts, suspicions, evidence, calls = path_b_contract_sweep(
        FlakyRouter(), None, [requirement], [version], {"CV-retry": [clause]},
        DiscoveryConfig(path_b_batch_chars=5000, path_b_max_requirements=2,
                        path_b_max_calls=5, path_c_enabled=False), max_workers=1)
    assert calls == 2
    assert [item["status"] for item in evidence] == ["not_affected"]
    assert all(item["status"] != "model_error" for item in evidence)


def test_topic_coverage_judge_closes_only_covered_topics():
    requirement = PolicyRequirement(
        requirement_id="R-topic", source_block_ids=["B-topic"],
        source_quote="基金管理人应当约定销售费用。", article="第一条",
        norm_type="obligation", subject="基金管理人", object="销售费用",
        requirement="基金管理人应当约定销售费用。", disposition="modify_contract",
    )
    item = GuideItem(
        guide_item_id="G-topic", requirement_ids=["R-topic"], impact_ids=[],
        clause_family_ids=[], section="费用", title="销售费用约定",
        business_requirement="约定销售费用", instruction="补充销售费用约定",
        new_text="基金合同应当约定销售费用。",
    )
    version = ContractVersion("CV-topic", "C-topic", "测试基金", "测试合同",
                              latest_known_version=True)
    snapshot = CorpusSnapshot("S-topic", ["CV-topic"], {"CV-topic": "sha"},
                              ["CV-topic"], [], "parser", "index")

    class TopicRouter(FakeGuideRouter):
        def __init__(self, verdict):
            super().__init__()
            self.verdict = verdict

        def _value(self, task_type, context):
            if task_type == "guide_coverage_audit":
                return {"requirement_results": [], "impact_results": [],
                        "missing_topics": [{
                            "description": "销售费用约定主题",
                            "requirement_ids": ["R-topic"],
                            "evidence_quote": "基金管理人应当约定销售费用。",
                        }], "notes": ""}
            if task_type == "guide_topic_coverage":
                topic_id = re.findall(r"\[TOPIC id=([^\]]+)\]", context)[0]
                return {"results": [{
                    "topic_id": topic_id, "verdict": self.verdict,
                    "guide_item_ids": ["G-topic"] if self.verdict == "covered" else [],
                    "reason": "测试判定",
                }]}
            return super()._value(task_type, context)

    covered = run_coverage_gate(
        TopicRouter("covered"), [requirement], [], [item], {}, [version], [], snapshot,
        semantic_applicability=False,
    )
    assert covered.missing_topics == []
    assert not any(value["category"] == "semantic_coverage"
                   for value in covered.findings)
    assert any(row["level"] == "topic" and row["status"] == "covered"
               for row in covered.rows)

    uncovered = run_coverage_gate(
        TopicRouter("not_covered"), [requirement], [], [item], {}, [version], [], snapshot,
        semantic_applicability=False,
    )
    assert len(uncovered.missing_topics) == 1
    assert any(value["category"] == "semantic_coverage" and value["severity"] == "critical"
               for value in uncovered.findings)
    assert any(row["level"] == "topic" and row["status"] == "unresolved"
               for row in uncovered.rows)


def test_impact_sweep_rejects_unknown_or_overbound_requirement_ids():
    clause = ClauseUnit("CL-sweep", "CV-sweep", ["费用"], 1, 1, "基金管理人收取申购费。")
    valid = {
        "findings": [{"clause_ids": ["CL-sweep"], "requirement_ids": ["R-1"],
                      "impact_status": "possible", "reason": "费用",
                      "evidence_quote": "基金管理人收取申购费。"}],
        "uncovered_suspicions": [], "section_assessment": "affected",
    }
    findings, _, assessment, error = _validate_sweep(
        valid, [clause], {"R-1", "R-2"}, max_requirement_ids=2)
    assert len(findings) == 1
    assert assessment == "affected"
    assert error == ""

    invalid = {
        "findings": [
            {"clause_ids": ["CL-sweep"], "requirement_ids": ["R-unknown"],
             "impact_status": "possible", "reason": "未知要求",
             "evidence_quote": "基金管理人收取申购费。"},
            {"clause_ids": ["CL-sweep"], "requirement_ids": ["R-1", "R-2", "R-3"],
             "impact_status": "possible", "reason": "绑定过多",
             "evidence_quote": "基金管理人收取申购费。"},
        ],
        "uncovered_suspicions": [], "section_assessment": "affected",
    }
    findings, _, _, error = _validate_sweep(
        invalid, [clause], {"R-1", "R-2"}, max_requirement_ids=2)
    assert findings == []
    assert error == "INVALID_REQUIREMENT_BINDINGS:2"


def test_uncertain_batch_resolution_closes_only_explicit_not_affected():
    requirements = [
        PolicyRequirement(
            requirement_id=f"R-resolve-{index}", source_block_ids=[f"B-{index}"],
            source_quote="基金管理人收取申购费。", article="第一条",
            norm_type="prohibition", subject="基金管理人", object="申购费",
            requirement="基金管理人不得收取申购费。", disposition="modify_contract",
        ) for index in range(2)
    ]
    version = ContractVersion("CV-resolve", "C-resolve", "测试基金", "测试合同",
                              latest_known_version=True)
    clauses = [
        ClauseUnit("CL-resolve-1", "CV-resolve", ["费用"], 1, 1, "基金管理人收取申购费。"),
        ClauseUnit("CL-resolve-2", "CV-resolve", ["费用"], 2, 2, "基金托管人收取管理费。"),
    ]

    class ResolveRouter(FakeGuideRouter):
        def __init__(self, results):
            super().__init__()
            self.results = results

        def _value(self, task_type, context):
            if task_type == "guide_batch_resolve":
                return {"results": self.results}
            return super()._value(task_type, context)

    def run(results):
        batch = {"contract_version_id": "CV-resolve",
                 "clause_ids": ["CL-resolve-1", "CL-resolve-2"], "status": "uncertain"}
        impacts, calls = resolve_uncertain_batches(
            ResolveRouter(results), requirements, [version],
            {"CV-resolve": clauses}, [batch],
            DiscoveryConfig(path_b_max_requirements=2, path_b_max_calls=5),
        )
        return batch, impacts, calls

    batch, impacts, calls = run([
        {"requirement_id": "R-resolve-0", "status": "not_affected", "reason": "无关"},
        {"requirement_id": "R-resolve-1", "status": "not_affected", "reason": "无关"},
    ])
    assert calls == 1
    assert impacts == []
    assert batch["status"] == "not_affected"

    batch, impacts, _ = run([
        {"requirement_id": "R-resolve-0", "status": "affected",
         "clause_ids": ["CL-resolve-1"], "evidence_quote": "基金管理人收取申购费。",
         "reason": "费用条款"},
        {"requirement_id": "R-resolve-1", "status": "not_affected", "reason": "无关"},
    ])
    assert batch["status"] == "affected"
    assert len(impacts) == 1
    assert impacts[0].requirement_id == "R-resolve-0"
    assert impacts[0].clause_ids == ["CL-resolve-1"]

    batch, impacts, _ = run([
        {"requirement_id": "R-resolve-0", "status": "not_affected", "reason": "无关"},
    ])
    assert batch["status"] == "uncertain"
    assert impacts == []


def test_impact_sweep_emits_one_physical_impact_per_bound_requirement():
    requirements = [
        PolicyRequirement(
            requirement_id=f"R-{index}", source_block_ids=[f"B-{index}"],
            source_quote="基金管理人收取申购费。", article="第一条",
            norm_type="prohibition", subject="基金管理人", object="申购费",
            requirement="基金管理人不得收取申购费。", disposition="modify_contract",
        ) for index in range(2)
    ]
    version = ContractVersion("CV-multi", "C-multi", "测试基金", "测试合同",
                              latest_known_version=True)
    clause = ClauseUnit("CL-multi", "CV-multi", ["费用"], 1, 1, "基金管理人收取申购费。")

    class MultiBindingRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_impact_sweep":
                return {"findings": [{
                    "clause_ids": ["CL-multi"],
                    "requirement_ids": ["R-0", "R-1"],
                    "impact_status": "possible", "reason": "共同费用条款",
                    "evidence_quote": "基金管理人收取申购费。",
                }], "uncovered_suspicions": [], "section_assessment": "affected"}
            return super()._value(task_type, context)

    impacts, _, evidence, calls = path_b_contract_sweep(
        MultiBindingRouter(), None, requirements, [version], {"CV-multi": [clause]},
        DiscoveryConfig(path_b_batch_chars=5000, path_b_max_requirements=2,
                        path_b_max_calls=5, path_c_enabled=False), max_workers=1)
    assert calls == 1
    assert {item.requirement_id for item in impacts} == {"R-0", "R-1"}
    assert len(impacts) == 2
    assert all(item.clause_ids == ["CL-multi"] for item in impacts)
    assert evidence[0]["finding_count"] == 1


def test_residual_reuses_existing_physical_pattern_when_model_description_changes():
    requirement = PolicyRequirement(
        requirement_id="R-residual", source_block_ids=["B-residual"],
        source_quote="基金管理人收取申购费。", article="第一条",
        norm_type="prohibition", subject="基金管理人", object="申购费",
        requirement="基金管理人不得收取申购费。", disposition="modify_contract",
    )
    version = ContractVersion("CV-residual", "C-residual", "测试基金", "测试合同",
                              latest_known_version=True)
    clause = ClauseUnit("CL-residual", "CV-residual", ["费用"], 1, 1,
                        "基金管理人收取申购费。")
    existing = ImpactCandidate(
        impact_id="IC-residual-existing", requirement_id="R-residual",
        contract_version_id="CV-residual", clause_ids=["CL-residual"],
        source_text=clause.text, context="费用", impact_status="possible",
        discovery_sources=["contract_semantic_sweep"],
    )

    class RewordingResidualRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.round = 0

        def _value(self, task_type, context):
            if task_type == "guide_residual_sweep":
                self.round += 1
                return {"uncovered": [{
                    "clause_ids": ["CL-residual"], "requirement_ids": ["R-residual"],
                    "reason": f"换一种说明 {self.round}",
                    "pattern_description": f"不同描述 {self.round}",
                    "evidence_quote": "基金管理人收取申购费。",
                }], "guide_conflicts": [], "notes": ""}
            return super()._value(task_type, context)

    result = run_residual_mining(
        RewordingResidualRouter(), [requirement], [], [existing], [version],
        {"CV-residual": [clause]}, {"CL-residual": clause},
        [{"contract_version_id": "CV-residual", "clause_ids": ["CL-residual"],
          "status": "affected"}],
        config=ResidualConfig(max_rounds=2, required_clean_rounds=2,
                              batch_chars=5000, max_calls_per_round=2),
    )
    assert result.saturated is True
    assert result.new_impacts == []
    assert [item.status for item in result.rounds] == ["clean", "clean"]


def test_residual_context_indexes_every_guide_item():
    requirements = [
        PolicyRequirement(
            requirement_id=f"R-residual-index-{index}", source_block_ids=[f"B-{index}"],
            source_quote="基金管理人应当调整费用。", article="第一条",
            norm_type="obligation", subject="基金管理人", object="费用",
            requirement=f"第{index}项费用要求。", disposition="modify_contract",
        ) for index in range(40)
    ]
    items = [
        GuideItem(
            guide_item_id=f"G-residual-index-{index}",
            requirement_ids=[f"R-residual-index-{index}"], impact_ids=[],
            clause_family_ids=[], section="费用", title=f"费用指引{index}",
            business_requirement="调整费用", instruction="调整费用条款",
            new_text=f"第{index}项费用条款调整。",
        ) for index in range(40)
    ]
    version = ContractVersion("CV-residual-index", "C-residual-index", "测试基金",
                              "测试合同", latest_known_version=True)
    clause = ClauseUnit("CL-residual-index", "CV-residual-index", ["费用"], 1, 1,
                        "基金管理人收取申购费。")

    class ResidualIndexRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.contexts = []

        def call(self, *, bounded_context, **kwargs):
            self.contexts.append(bounded_context)
            return super().call(bounded_context=bounded_context, **kwargs)

        def _value(self, task_type, context):
            if task_type == "guide_residual_sweep":
                return {"uncovered": [], "guide_conflicts": [], "notes": ""}
            return super()._value(task_type, context)

    router = ResidualIndexRouter()
    result = run_residual_mining(
        router, requirements, items, [], [version],
        {"CV-residual-index": [clause]}, {"CL-residual-index": clause},
        [{"contract_version_id": "CV-residual-index",
          "clause_ids": ["CL-residual-index"], "status": "affected"}],
        config=ResidualConfig(max_rounds=1, required_clean_rounds=1,
                              batch_chars=5000, max_calls_per_round=1,
                              max_requirements=40),
    )
    assert len(router.contexts) == 1
    context = router.contexts[0]
    assert len(context) <= 32000
    assert "[ALL_GUIDE_ITEM_INDEX]" in context
    assert all(f"G-residual-index-{index}" in context for index in range(40))
    assert result.saturated is True


def test_residual_context_stays_within_limit_for_many_items():
    requirements = [
        PolicyRequirement(
            requirement_id=f"R-residual-many-{index}", source_block_ids=[f"B-{index}"],
            source_quote="基金管理人应当调整费用。", article="第一条",
            norm_type="obligation", subject="基金管理人", object="费用",
            requirement=f"第{index}项费用要求。", disposition="modify_contract",
        ) for index in range(150)
    ]
    items = [
        GuideItem(
            guide_item_id=f"G-residual-many-{index}",
            requirement_ids=[f"R-residual-many-{index}"], impact_ids=[],
            clause_family_ids=[], section="费用", title=f"费用指引{index}",
            business_requirement="调整费用", instruction="调整费用条款",
            new_text=f"第{index}项费用条款调整。",
        ) for index in range(150)
    ]
    version = ContractVersion("CV-residual-many", "C-residual-many", "测试基金",
                              "测试合同", latest_known_version=True)
    clause = ClauseUnit("CL-residual-many", "CV-residual-many", ["费用"], 1, 1,
                        "基金管理人收取申购费。")

    class ManyRouter(FakeGuideRouter):
        def __init__(self):
            super().__init__()
            self.contexts = []

        def call(self, *, bounded_context, **kwargs):
            self.contexts.append(bounded_context)
            return super().call(bounded_context=bounded_context, **kwargs)

        def _value(self, task_type, context):
            if task_type == "guide_residual_sweep":
                return {"uncovered": [], "guide_conflicts": [], "notes": ""}
            return super()._value(task_type, context)

    router = ManyRouter()
    result = run_residual_mining(
        router, requirements, items, [], [version],
        {"CV-residual-many": [clause]}, {"CL-residual-many": clause},
        [{"contract_version_id": "CV-residual-many",
          "clause_ids": ["CL-residual-many"], "status": "affected"}],
        config=ResidualConfig(max_rounds=1, required_clean_rounds=1,
                              batch_chars=5000, max_calls_per_round=1),
    )
    context = router.contexts[0]
    assert len(context) <= 30000
    assert all(f"G-residual-many-{index}" in context for index in range(150))
    assert result.saturated is True


def test_router_waits_for_circuit_and_retries(monkeypatch):
    from types import SimpleNamespace

    from fundlex.guide_generation import model_router as router_module
    from fundlex.model import ModelError

    monkeypatch.setattr(router_module.time, "sleep", lambda _seconds: None)

    class Gateway:
        model = "fake"
        provider = "fake"
        budget = None
        max_context_chars = 32000
        max_completion_tokens = 8000
        request_options = {}

        def __init__(self):
            self.calls = 0

        def call(self, **kwargs):
            self.calls += 1
            if self.calls == 1:
                raise ModelError("MODEL_CIRCUIT_OPEN")
            return SimpleNamespace(value={}, model="fake", latency_ms=1,
                                   prompt_tokens=1, completion_tokens=1, total_tokens=2,
                                   request_id=None, attempts=1, estimated_cost_usd=0.0)

        def close(self):
            pass

    gateway = Gateway()
    router = router_module.GuideModelRouter({"primary": gateway})
    result = router.call(task_type="guide_residual_sweep", bounded_context="x",
                         response_schema={"type": "object"}, instruction="i",
                         role="coverage_auditor")
    assert gateway.calls == 2
    assert result.model == "fake"


def test_router_queues_calls_to_budget_concurrency():
    import threading
    import time
    from types import SimpleNamespace

    from fundlex.guide_generation.model_router import GuideModelRouter

    class Budget:
        concurrency = 1

    class Gateway:
        model = "fake"
        provider = "fake"
        max_context_chars = 32000
        max_completion_tokens = 8000
        request_options = {}

        def __init__(self):
            self.budget = Budget()
            self.active = 0
            self.max_active = 0
            self.lock = threading.Lock()

        def call(self, **kwargs):
            with self.lock:
                self.active += 1
                self.max_active = max(self.max_active, self.active)
            time.sleep(0.03)
            with self.lock:
                self.active -= 1
            return SimpleNamespace(value={}, model="fake", latency_ms=1, prompt_tokens=1,
                                   completion_tokens=1, total_tokens=2, request_id=None,
                                   attempts=1, estimated_cost_usd=0.0)

        def close(self):
            pass

    gateway = Gateway()
    router = GuideModelRouter({"primary": gateway})
    errors: list[Exception] = []

    def work():
        try:
            router.call(task_type="guide_impact_sweep", bounded_context="x",
                        response_schema={"type": "object"}, instruction="i", role="discovery")
        except Exception as exc:  # pragma: no cover - asserted empty below
            errors.append(exc)

    threads = [threading.Thread(target=work) for _ in range(3)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert not errors
    assert gateway.max_active == 1
    assert len(router.records) == 3


def test_quality_review_keeps_critical_when_one_of_two_quotes_matches():
    requirement = PolicyRequirement(
        requirement_id="R-quality-partial", source_block_ids=["B-quality-partial"],
        source_quote="基金管理人应当约定申购费收取标准。", article="第一条",
        norm_type="obligation", subject="基金管理人", object="申购费收取标准",
        requirement="基金管理人应当约定申购费收取标准。", disposition="modify_contract",
    )
    impact = ImpactCandidate(
        impact_id="IC-quality-partial", requirement_id="R-quality-partial",
        contract_version_id="CV-quality-partial", clause_ids=["CL-quality-partial"],
        source_text="本基金申购费由基金管理人决定。", context="费用",
        impact_status="possible", discovery_sources=["lexical"],
    )
    item = GuideItem(
        guide_item_id="GI-quality-partial", requirement_ids=["R-quality-partial"],
        impact_ids=["IC-quality-partial"], clause_family_ids=[], section="费用",
        title="申购费收取标准", business_requirement="约定申购费标准",
        instruction="补充申购费收取标准", new_text="本基金申购费由基金管理人决定。",
    )

    class PartialQuoteRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_quality_review":
                return {"verdict": "fail", "findings": [{
                    "category": "regulation_fidelity", "severity": "critical",
                    "description": "changed subject: the duty is placed on the wrong party",
                    "evidence_quote": "本基金申购费由基金管理人决定。",
                    "evidence_quote_2": "这条第二引文并不存在于提交的上下文中",
                }], "notes": ""}
            return super()._value(task_type, context)

    result = run_quality_gate(
        PartialQuoteRouter(), [item], [requirement], [impact],
        {"CL-quality-partial": ClauseUnit("CL-quality-partial",
                                            "CV-quality-partial", ["费用"], 1, 1,
                                            impact.source_text)},
    )
    finding = next(value for value in result.findings
                   if value.get("source") == "guide_reviewer")
    assert finding["evidence_verified"] is True
    assert finding["evidence_matched_quote_indexes"] == [1]
    assert finding["severity"] == "critical"
    assert result.run_status == "fail"


def test_quality_review_accepts_two_bound_evidence_quotes_and_rejects_unknown_fields():
    requirement = PolicyRequirement(
        requirement_id="R-quality-evidence", source_block_ids=["B-quality-evidence"],
        source_quote="基金管理人应当约定申购费收取标准。", article="第一条",
        norm_type="obligation", subject="基金管理人", object="申购费收取标准",
        requirement="基金管理人应当约定申购费收取标准。", disposition="modify_contract",
    )
    impact = ImpactCandidate(
        impact_id="IC-quality-evidence", requirement_id="R-quality-evidence",
        contract_version_id="CV-quality-evidence", clause_ids=["CL-quality-evidence"],
        source_text="本基金申购费由基金管理人决定。", context="费用",
        impact_status="possible", discovery_sources=["lexical"],
    )
    item = GuideItem(
        guide_item_id="GI-quality-evidence", requirement_ids=["R-quality-evidence"],
        impact_ids=["IC-quality-evidence"], clause_family_ids=[], section="费用",
        title="申购费收取标准", business_requirement="约定申购费标准",
        instruction="补充申购费收取标准", new_text="本基金申购费由基金管理人决定。",
    )

    class TwoQuoteRouter(FakeGuideRouter):
        def _value(self, task_type, context):
            if task_type == "guide_quality_review":
                return {"verdict": "pass_with_findings", "findings": [{
                    "category": "other", "severity": "minor", "description": "需要人工确认",
                    "evidence_quote": "本基金申购费由基金管理人决定。",
                    "evidence_quote_2": "基金管理人应当约定申购费收取标准。",
                    "recommendation": "人工确认",
                }], "notes": ""}
            return super()._value(task_type, context)

    result = run_quality_gate(
        TwoQuoteRouter(), [item], [requirement], [impact],
        {"CL-quality-evidence": ClauseUnit("CL-quality-evidence",
                                             "CV-quality-evidence", ["费用"], 1, 1,
                                             impact.source_text)},
    )
    reviewer_finding = next(value for value in result.findings
                            if value.get("source") == "guide_reviewer")
    assert reviewer_finding["evidence_verified"] is True
    assert reviewer_finding["evidence_quotes"] == [
        "本基金申购费由基金管理人决定。", "基金管理人应当约定申购费收取标准。"
    ]

    valid = {"verdict": "pass", "findings": [], "notes": ""}
    jsonschema.validate(valid, QUALITY_SCHEMA)
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate({**valid, "unexpected": True}, QUALITY_SCHEMA)
