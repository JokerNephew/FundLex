from fundlex.compare.models import CompareBlock, CompareDocument
from fundlex.compare.parser import parse_compare_document
from fundlex.compare.render import render_compare_document
from docx import Document
from docx.shared import Pt, RGBColor
from io import BytesIO


def test_compare_renderer_preserves_common_word_paragraph_hierarchy():
    document = CompareDocument((
        CompareBlock("L000000", 0, "p0", "paragraph", raw_text="合同标题", style="Title"),
        CompareBlock("L000001", 1, "p1", "paragraph", raw_text="第一条", style="Heading 1"),
        CompareBlock("L000002", 2, "p2", "paragraph", raw_text="适用范围", style="List Paragraph"),
        CompareBlock("L000003", 3, "p3", "paragraph", raw_text="中文标题", style="标题"),
    ))

    html = render_compare_document(document, [], side="left")

    assert '<h1 class="compare-block compare-style-title compare-heading"' in html
    assert '<h1 class="compare-block compare-style-heading-1 compare-heading"' in html
    assert 'class="compare-block compare-style-list-paragraph"' in html
    assert 'compare-list-label' not in html
    assert 'class="compare-block compare-style-title compare-heading"' in html


def test_compare_renderer_preserves_merged_table_geometry():
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "合并单元格"
    table.cell(0, 2).text = "右侧"
    table.cell(1, 0).text = "下方"
    table.cell(1, 1).text = "中间"
    table.cell(1, 2).text = "末列"
    table.cell(2, 0).text = "底部"
    table.cell(2, 1).text = "底部中间"
    table.cell(2, 2).text = "底部末列"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 2).merge(table.cell(2, 2))
    data = BytesIO()
    document.save(data)

    parsed = parse_compare_document(data.getvalue(), filename="merged.docx", side="L")
    html = render_compare_document(parsed, [], side="left")

    assert 'colspan="2"' in html
    assert 'rowspan="2"' in html
    assert html.count("合并单元格") == 1


def test_compare_parser_and_renderer_project_real_run_formatting():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    paragraph.paragraph_format.space_after = Pt(12)
    title = paragraph.add_run("标题")
    title.bold = True
    title.font.size = Pt(16)
    title.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    paragraph.add_run(" 普通")
    data = BytesIO()
    document.save(data)

    parsed = parse_compare_document(data.getvalue(), filename="formatted.docx", side="L")
    block = parsed.blocks[0]
    html = render_compare_document(parsed, [], side="left")

    assert ("text-align", "center") in block.paragraph_format
    assert block.runs[0].bold is True
    assert block.runs[0].font_size == "16pt"
    assert block.runs[0].color == "#123456"
    assert "text-align:center" in html
    assert "font-weight:bold" in html
    assert "font-size:16pt" in html
    assert "color:#123456" in html
