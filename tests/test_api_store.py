"""Local queue and access-control integration checks with a minimal DOCX fixture."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from docx import Document

from fundlex.api import create_app
from fundlex.jobs import Worker
from fundlex.model import ModelResult
from fundlex.pipeline import process
from fundlex.store import Conflict, RunStore


def docx_bytes(label: str = "x") -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", f"<document>{label}</document>")
    return stream.getvalue()


def submit(client: TestClient, data: bytes, *, owner: str = "alice", key: str = "key-1"):
    return client.post("/api/batches", headers={"X-FundLex-User": owner, "Idempotency-Key": key},
                       data={"ruleset_version": "v1"}, files=[("files", ("contract.docx", data))])


def test_api_auth_scope_idempotency_and_download(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("FUNDLEX_AUTH_MODE", "dev")
    app = create_app(tmp_path)
    client = TestClient(app)
    assert client.get("/").status_code == 200
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/api/batches").status_code == 401
    assert submit(client, b"%PDF").status_code == 415
    unsupported = client.post("/api/batches", headers={"X-FundLex-User": "alice"},
                              data={"ruleset_version": "v999"},
                              files=[("files", ("contract.docx", docx_bytes()))])
    assert unsupported.status_code == 422
    stream = io.BytesIO()
    contract = Document()
    contract.add_paragraph("Independent unchanged-artifact transport fixture")
    contract.save(stream)
    data = stream.getvalue()
    first = submit(client, data)
    assert first.status_code == 201
    batch_id = first.json()["id"]
    assert len(first.json()["ruleset_sha256"]) == 64
    doc_id = first.json()["documents"][0]["document_id"]
    assert submit(client, data).json()["id"] == batch_id
    assert submit(client, docx_bytes("other")).status_code == 409
    assert client.get(f"/api/batches/{batch_id}", headers={"X-FundLex-User": "bob"}).status_code == 404

    def processor(source: Path, output: Path, rules: str):
        assert rules == "v1"
        assert source.read_bytes() == data
        from fundlex.docx import validate_edit_plan, generate_unchanged_artifacts
        plan = validate_edit_plan(source, [], {})
        validation = generate_unchanged_artifacts(source, plan, output / "redline.docx", output / "final.docx")
        # This exercises verified unchanged DOCX transport, not business rule approval.
        return {"qualification_status": "eligible", "delivery_status": "complete", "actions": [],
                "output_validation": {"kind": "validated_unchanged", "checks": validation.checks,
                                      "final_sha256": validation.final_sha256, "redline_sha256": validation.redline_sha256}}

    assert Worker(app.state.store, processor).run_once()
    detail = client.get(f"/api/batches/{batch_id}", headers={"X-FundLex-User": "alice"}).json()
    assert detail["status"] == "completed" and detail["completed"] == 1
    assert detail["documents"][0]["downloads"]["final"]
    assert client.get(f"/api/batches/{batch_id}/documents/{doc_id}/report", headers={"X-FundLex-User": "bob"}).status_code == 404
    url = detail["documents"][0]["downloads"]["final"]
    assert client.get(url, headers={"X-FundLex-User": "alice"}).content == data
    assert client.get(url, headers={"X-FundLex-User": "bob"}).status_code == 404
    summary = client.get(f"/api/batches/{batch_id}/report", headers={"X-FundLex-User": "alice"})
    assert summary.status_code == 200 and "eligible" in summary.text and "complete" in summary.text


def test_single_document_failure_does_not_block_batch(tmp_path: Path):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1",
                               uploads=[("one.docx", docx_bytes("bad")), ("two.docx", docx_bytes("good"))])

    def processor(source: Path, output: Path, rules: str):
        if b"bad" in source.read_bytes():
            raise ValueError("semantic failure")
        return {"qualification_status": "excluded_by_fund_type", "delivery_status": "not_generated"}

    assert Worker(store, processor).run_until_idle() == 2
    result = store.get_batch("alice", batch["id"])
    assert result["status"] == "completed"
    assert [d["status"] for d in result["documents"]] == ["failed", "succeeded"]
    assert [d["delivery_status"] for d in result["documents"]] == ["not_generated", "not_generated"]


def test_expired_lease_reclaim_rejects_old_worker(tmp_path: Path):
    store = RunStore(tmp_path)
    store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    first = store.claim(lease_seconds=0)
    second = store.claim(lease_seconds=30)
    assert second["fencing_token"] > first["fencing_token"]
    with pytest.raises(Conflict):
        store.finish(first["id"], first["fencing_token"], report={"delivery_status": "not_generated"}, artifacts={})
    store.finish(second["id"], second["fencing_token"], report={"delivery_status": "not_generated"}, artifacts={})


def test_pause_resume_and_input_hash_guard(tmp_path: Path):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    store.pause("alice", batch["id"])
    assert store.claim() is None
    store.resume("alice", batch["id"])
    input_key = store.claim()["input_key"]
    store.object_path(input_key).write_bytes(b"corrupt")
    # Reclaim after lease expiry and fail closed without publishing a contract.
    with store._connect() as db:
        db.execute("UPDATE documents SET lease_until=0 WHERE batch_id=?", (batch["id"],))
    assert Worker(store, lambda *_: {}).run_once()
    result = store.get_batch("alice", batch["id"])["documents"][0]
    assert result["error_code"] == "INPUT_INTEGRITY_FAILED"
    assert result["delivery_status"] == "not_generated"
    assert "final" not in result["artifacts"]


def test_expired_max_attempts_becomes_terminal_report(tmp_path: Path):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    for _ in range(2):
        lease = store.claim(lease_seconds=0, max_attempts=2)
        assert lease is not None
    assert store.claim(max_attempts=2) is None
    result = store.get_batch("alice", batch["id"])
    assert result["status"] == "completed"
    doc = result["documents"][0]
    assert doc["error_code"] == "WORKER_ATTEMPTS_EXHAUSTED"
    assert store.get_report("alice", batch["id"], doc["id"])["delivery_status"] == "not_generated"


def test_complete_without_both_verified_contracts_fails_closed(tmp_path: Path):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])

    def incomplete_processor(source: Path, output: Path, rules: str):
        (output / "final.docx").write_bytes(docx_bytes())
        return {"qualification_status": "eligible", "delivery_status": "complete"}

    Worker(store, incomplete_processor).run_once()
    doc = store.get_batch("alice", batch["id"])["documents"][0]
    assert doc["error_code"] == "MISSING_VERIFIED_ARTIFACTS"
    assert "final" not in doc["artifacts"]


def test_ruleset_hash_change_blocks_processing(tmp_path: Path):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", ruleset_sha256="0" * 64,
                               uploads=[("a.docx", docx_bytes())])
    Worker(store, lambda *_: pytest.fail("processor must not run")).run_once()
    doc = store.get_batch("alice", batch["id"])["documents"][0]
    assert doc["error_code"] == "RULESET_HASH_MISMATCH"
    assert doc["delivery_status"] == "not_generated"


def test_http_upload_worker_pipeline_report_flow(tmp_path: Path):
    source = tmp_path / "real.docx"
    contract = Document()
    for line in ("示例证券投资基金基金合同", "第四部分 基金的基本情况", "二、基金的类别",
                 "混合型证券投资基金", "三、基金的运作方式", "契约型开放式", "份额类别设置",
                 "第十二部分 基金的投资", "本基金投资于债券的比例不低于基金资产的80%。"):
        contract.add_paragraph(line)
    contract.save(source)

    class Gateway:
        model = "test-gateway"
        calls = 0

        def extract_fund_type(self, context: str):
            self.calls += 1
            return ModelResult({"quote": context.splitlines()[0], "fund_type": "ordinary"},
                               self.model, 3, 10, 5, 15, None, 1)

    app = create_app(tmp_path / "data")
    client = TestClient(app)
    response = submit(client, source.read_bytes())
    assert response.status_code == 201
    batch_id = response.json()["id"]
    gateway = Gateway()
    worker = Worker(app.state.store, lambda source, output, version: process(source, output, version, gateway))
    assert worker.run_once()
    detail = client.get(f"/api/batches/{batch_id}", headers={"X-FundLex-User": "alice"}).json()
    doc = detail["documents"][0]
    assert detail["status"] == "completed"
    assert doc["eligibility"] == "eligible"
    assert doc["delivery_status"] != "complete"
    assert "final" not in doc["downloads"] and "redline" not in doc["downloads"]
    report = client.get(doc["report_url"], headers={"X-FundLex-User": "alice"}).json()
    assert report["model"]["calls"] == 1
    assert report["edit_plan"]["count"] == 0
    assert report["release_gate"]["status"] == "closed"
    report_md = client.get(doc["downloads"]["report_md"], headers={"X-FundLex-User": "alice"})
    assert report_md.status_code == 200 and "不完整" in report_md.text


def test_expired_lease_cannot_be_revived_or_retried(tmp_path):
    store = RunStore(tmp_path)
    store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    job = store.claim(lease_seconds=0)
    with pytest.raises(Conflict):
        store.heartbeat(job["id"], job["fencing_token"])
    with pytest.raises(Conflict):
        store.retry(job["id"], job["fencing_token"], error_code="TEMPORARY")


def test_worker_discards_expired_attempt_and_survives_for_reclaim(tmp_path):
    store=RunStore(tmp_path)
    batch=store.create_batch(owner='alice',ruleset_version='v1',uploads=[('a.docx',docx_bytes())])
    def interrupted(*_):
        with store._connect() as db:
            db.execute('UPDATE documents SET lease_until=0 WHERE batch_id=?',(batch['id'],))
        return {'qualification_status':'excluded_by_fund_type','delivery_status':'not_generated'}
    assert Worker(store,interrupted).run_once()
    doc=store.get_batch('alice',batch['id'])['documents'][0]
    assert doc['status'] in {'leased','running'} and not doc['artifacts']
    assert Worker(store,lambda *_:{'qualification_status':'excluded_by_fund_type','delivery_status':'not_generated'}).run_once()
    doc=store.get_batch('alice',batch['id'])['documents'][0]
    assert doc['attempt']==2 and doc['status']=='succeeded'


def test_long_processor_keeps_lease_and_excludes_second_worker(tmp_path):
    import time
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    def processor(*_):
        with pytest.raises(Conflict, match="one worker"):
            Worker(RunStore(tmp_path)).run_once()
        time.sleep(1.3)
        assert store.claim(lease_seconds=1) is None
        return {"qualification_status": "excluded_by_fund_type", "delivery_status": "not_generated"}
    assert Worker(store, processor, lease_seconds=1).run_once()
    assert store.get_batch("alice", batch["id"])["documents"][0]["attempt"] == 1


def test_transient_error_waits_then_terminal_report(tmp_path):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    def processor(*_):
        raise OSError("storage unavailable")
    worker = Worker(store, processor)
    for attempt in range(1, 4):
        assert worker.run_once()
        doc = store.get_batch("alice", batch["id"])["documents"][0]
        assert doc["attempt"] == attempt
        if attempt < 3:
            assert doc["status"] == "retry_wait"
            assert worker.run_once() is False
            with store._connect() as db:
                db.execute("UPDATE documents SET next_attempt_at=0 WHERE id=?", (doc["id"],))
    assert doc["status"] == "failed"
    assert store.get_report("alice", batch["id"], doc["id"])["error_code"] == "TEMPORARY_STORAGE_FAILURE"


def test_model_outage_pauses_and_manual_resume_preserves_attempt(tmp_path):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    worker = Worker(store, lambda *_: {"error_code": "MODEL_RETRY_EXHAUSTED"})
    for _ in range(2):
        assert worker.run_once()
        with store._connect() as db:
            db.execute("UPDATE documents SET next_attempt_at=0 WHERE batch_id=?", (batch["id"],))
    assert store.get_batch("alice", batch["id"])["status"] == "paused"
    assert worker.run_once() is False
    store.resume("alice", batch["id"])
    assert Worker(store, lambda *_: {"qualification_status": "excluded_by_fund_type"}).run_once()
    assert store.get_batch("alice", batch["id"])["documents"][0]["attempt"] == 3


def test_engine_change_invalidates_idempotency_and_queued_work(tmp_path, monkeypatch):
    import fundlex.store as module
    import fundlex.jobs as jobs
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())], idempotency_key="same")
    monkeypatch.setattr(module, "engine_build", lambda: "different-build")
    with pytest.raises(Conflict):
        store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())], idempotency_key="same")
    monkeypatch.setattr(jobs, "engine_build", lambda: "different-build")
    Worker(store, lambda *_: pytest.fail("stale engine must not execute")).run_once()
    assert store.get_batch("alice", batch["id"])["documents"][0]["error_code"] == "ENGINE_SCHEMA_MISMATCH"


def test_corrupt_published_artifact_and_production_auth_fail_closed(tmp_path, monkeypatch):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    Worker(store, lambda *_: {"qualification_status": "excluded_by_fund_type"}).run_once()
    doc = store.get_batch("alice", batch["id"])["documents"][0]
    path, _ = store.get_artifact("alice", batch["id"], doc["id"], "report_json")
    path.write_bytes(b"corrupt")
    with pytest.raises(Conflict, match="integrity"):
        store.get_artifact("alice", batch["id"], doc["id"], "report_json")
    monkeypatch.setenv("FUNDLEX_AUTH_MODE", "oidc")
    client = TestClient(create_app(tmp_path))
    assert client.get("/api/batches", headers={"X-FundLex-User": "alice"}).status_code == 503


@pytest.mark.parametrize("code", ["MODEL_CIRCUIT_OPEN", "MODEL_GLOBAL_CALL_BUDGET_EXCEEDED", "MODEL_COST_BUDGET_EXCEEDED"])
def test_global_model_stop_pauses_remaining_queue(tmp_path, code):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes()), ("b.docx", docx_bytes())])
    worker = Worker(store, lambda *_: {"error_code": code})
    assert worker.run_once()
    result = store.get_batch("alice", batch["id"])
    assert result["status"] == "paused"
    assert [d["status"] for d in result["documents"]] == ["retry_wait", "queued"]
    assert not worker.run_once()


@pytest.mark.parametrize("validation,error", [
    (None, "MISSING_VERIFIED_ARTIFACTS"),
    ({"kind": "zero_edit_evaluation_only", "checks": ["package", "byte_identity"]}, "MISSING_VERIFIED_ARTIFACTS"),
    ({"kind": "validated_unchanged", "checks": ["package", "zero_edit", "byte_identity"], "final_sha256": "wrong", "redline_sha256": "wrong"}, "ARTIFACT_VERIFICATION_MISMATCH")])
def test_complete_without_independent_validation_fails_closed(tmp_path, validation, error):
    store = RunStore(tmp_path)
    batch = store.create_batch(owner="alice", ruleset_version="v1", uploads=[("a.docx", docx_bytes())])
    def processor(source, output, _):
        for kind in ("final", "redline"):
            (output / f"{kind}.docx").write_bytes(source.read_bytes())
        return {"delivery_status": "complete", "output_validation": validation}
    Worker(store, processor).run_once()
    doc = store.get_batch("alice", batch["id"])["documents"][0]
    assert doc["error_code"] == error
    assert "final" not in doc["artifacts"]
