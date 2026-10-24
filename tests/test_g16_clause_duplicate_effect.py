"""Source-scoped cross-paragraph G16 proof; actual DOCX artifacts (FL-D192)."""
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from fundlex.docx import inspect_docx
from fundlex.generation.claim_replay import replay_claims
from fundlex.obligations import compile_obligations
from fundlex.recall import Target, _g16_clause_locations, rebind_g16_clause_targets
from test_recall_core import _facts


HEAD = "基金销售服务费每日计提，逐日累计至每个月月末，按月支付。"
DELAY = "若遇法定节假日、休息日或不可抗力致使无法按时支付等，支付日期顺延。"
USAGE = "销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费。"
TAIL = "销售服务费使用范围不包括基金募集期间的上述费用。"


def _write(path, paragraphs):
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(path)
    return inspect_docx(path, full=True)


@pytest.mark.parametrize("with_ledger", [True, False])
def test_replay_blocks_cross_paragraph_duplicate_even_when_every_insertion_is_authorized(
        tmp_path, with_ledger):
    guide = json.loads((Path(__file__).resolve().parents[1] /
                        "guides/builtin_guide_v1.json").read_text())
    rule = next(rule for rule in guide["groups"] if rule["rule_id"] == "G16")
    policy = rule["operations"][0]["new_text"]
    paragraphs = [HEAD, DELAY + USAGE, TAIL]
    before = _write(tmp_path / "source.docx", paragraphs)
    # Reproduce D18: complete policy at both the payment head and usage tail.
    after = _write(tmp_path / "duplicate.docx",
                   [HEAD + policy, DELAY + USAGE, TAIL + policy])
    edits, targets = [], {}
    for location in ("p0", "p2"):
        source = before.by_location(location).text
        edits.append(dict(rule_id="G16", target_id=location, location=location,
                          start=len(source), end=len(source), replacement=policy,
                          operation="insert", content_source="tier1_template",
                          evidence={"operation_keys": ["G16:0"], "fixed_text": policy}))
        targets[location] = Target(location, "G16", location, source, "needs_edit", "",
                                   start=0, end=len(source), resolution="resolved",
                                   operations=("G16:0",)).as_dict()
    ledger = compile_obligations({"groups": [rule]}) if with_ledger else None
    if ledger:
        for record in ledger.records:
            record.applicability = SimpleNamespace(status="TRUE")
    replay = replay_claims(edits, before, after, rules_by_id={"G16": rule},
                           facts=_facts(), targets_by_id=targets, obligations=ledger)
    assert not replay["verified"]
    failures = [f for f in replay["failures"] if f["code"] == "CLAIM_DUPLICATE_EFFECT"]
    assert any(f.get("locations") == ["p0", "p1", "p2"] for f in failures)
    if with_ledger:
        proof = next(iter(replay["obligations"].values()))
        assert proof["state"] == "FAIL"
        assert proof["failure_code"] == "CLAIM_DUPLICATE_EFFECT"
        assert all(site["state"] == "FAIL" for site in proof["site_effects"])


@pytest.mark.parametrize("paragraphs, expected", [
    ([HEAD, DELAY, USAGE, TAIL], [("p0", "p1", "p2", "p3")]),
    ([HEAD, DELAY + USAGE, TAIL, HEAD, USAGE, TAIL],
     [("p0", "p1", "p2"), ("p3", "p4", "p5")]),
    # An unterminated paragraph must not absorb an independent payment mechanism.
    ([HEAD.rstrip("。") + "，", HEAD, USAGE, TAIL],
     [("p0",), ("p1", "p2", "p3")]),
    ([HEAD, "基金管理费若遇法定节假日，支付日期顺延。" + USAGE, TAIL], [("p0",)]),
    ([HEAD, "投资者服务。" + USAGE, TAIL], [("p0",)]),
    ([HEAD, "基金的管理费按前一日基金资产净值计提。", USAGE, TAIL], [("p0",)]),
    ([HEAD, "二十、基金合同内容摘要", USAGE, TAIL], [("p0",)]),
])
def test_clause_boundaries_preserve_independent_mechanisms_and_do_not_cross_unrelated_text(
        tmp_path, paragraphs, expected):
    source = _write(tmp_path / "source.docx", paragraphs)
    assert _g16_clause_locations(source) == expected
    targets = [Target(str(i), "G16", p.location, p.text, "needs_edit", "",
                       start=0, end=len(p.text)) for i, p in enumerate(source.paragraphs)]
    rebind_g16_clause_targets(targets, source)
    for clause in expected:
        assert all(target.location == clause[-1]
                    for target in targets if target.target_id in
                    {loc.removeprefix("p") for loc in clause})


def test_split_usage_tail_appends_before_independent_expense_sentence(tmp_path, monkeypatch):
    from fundlex.obligations import apply_facts
    from fundlex.recall import append_site_ok, ensure_delete_targets
    from fundlex.recall_pipeline import _plan_and_generate
    from test_g17_usage_site_closure import NoModel
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    guide = json.loads((Path(__file__).resolve().parents[1] /
                        "guides/builtin_guide_v1.json").read_text())
    rules = {r['rule_id']:r for r in guide['groups'] if r['rule_id'] in {'G16','G17'}}
    usage = ('销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、'
             '促销活动费、基金份额持有人服务费等。')
    suffix = '上述费用根据有关法规及相应协议规定，按费用实际支出金额列入或摊入当期费用。 '
    texts = [HEAD, usage[:32], usage[32:] + suffix, '独立的下一条。'] * 2
    source = tmp_path/'source.docx'
    original = _write(source, texts)
    targets = [Target(f'g16-{i}', 'G16', f'p{i}', texts[i], 'needs_edit', 'inventory',
        start=0, end=len(texts[i]), resolution='resolved', operations=('G16:0',)) for i in (2,6)]
    rebind_g16_clause_targets(targets, original)
    assert all(t.end == len(usage[32:]) for t in targets)
    assert all(append_site_ok('G16', original.by_location(t.location), original) for t in targets)
    targets.extend(ensure_delete_targets(targets, rules, original))
    facts = _facts(); ledger = compile_obligations({'groups':list(rules.values())});apply_facts(ledger,facts)
    plan, _, _, _, actual = _plan_and_generate(NoModel(), rules, targets, original,
        facts, source, tmp_path, max_compose_calls=0, obligations=ledger)
    policy = rules['G16']['operations'][0]['new_text']
    for i in (0,4):
        assert actual.by_location(f'p{i}').text == HEAD
        assert actual.by_location(f'p{i+1}').text == ''
        assert actual.by_location(f'p{i+2}').text == policy + suffix
        assert actual.by_location(f'p{i+3}').text == texts[i+3]
    proof = replay_claims(plan['edits'], original, actual, rules_by_id=rules, facts=facts,
        obligations=ledger, targets_by_id={t.target_id:t.as_dict() for t in targets})
    assert proof['verified'], proof['failures']
