"""Deterministic guide parsing tests: DOCX extraction, validation, verification and repair.

The scripted gateway replaces only the network transport. Paragraph extraction,
verbatim binding, condition checking, quote validation and the bounded repair loop
remain the real implementation under test.
"""
from __future__ import annotations

import copy
import hashlib
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.shared import RGBColor

from fundlex.guide import (GUIDE_PARSE_PROMPT_VERSION, GUIDE_SCHEMA_VERSION, GUIDE_VERIFY_PROMPT_VERSION,
                           GuideError, build_guide_document, extract_guide_source, guide_document_sha256,
                           load_guide_document, parse_guide, save_guide_document, validate_parsed_guide,
                           verify_guide)
from fundlex.model import ModelError


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_docx(path: Path, paragraphs: list[tuple[str, str]]) -> Path:
    """A synthetic guide: (text, emphasis) with plain, bold and red paragraphs."""
    document = Document()
    for text, emphasis in paragraphs:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        if emphasis == "red":
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        elif emphasis == "bold":
            run.bold = True
    document.save(path)
    return path


def patch_document_xml(path: Path, transform) -> Path:
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["word/document.xml"] = transform(entries["word/document.xml"].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return path


CLASS_DEFINITION = (
    "基金份额类别：本基金将基金份额分为A类和C类不同的类别。在投资者认购、申购基金时收取认购、申购费用"
    "而不计提销售服务费的，称为A类基金份额；在投资者认购、申购基金份额时不收取认购、申购费用而是从本类别"
    "基金资产中计提销售服务费的，称为C类基金份额。"
)
DIRECT_SALES_NEW = (
    "投资者通过直销机构（含基金管理人及其销售子公司，下同）申购本基金A类基金份额时，不收取认购/申购费用；"
    "通过直销机构申购的C类基金份额，不收取销售服务费。"
)
CLASS_CHANGE_OLD = (
    "在符合法律法规、基金合同的约定且不损害已有基金份额持有人权益的情况下，经与基金托管人协商一致，"
    "基金管理人在履行适当程序后可以增加新的基金份额类别、或者在法律法规和基金合同规定的范围内变更现有"
    "基金份额类别的申购费率、调低赎回费率或变更收费方式、或者停止现有基金份额类别的销售等，调整前基金"
    "管理人需及时公告并报中国证监会备案。"
)
SERVICE_FEE_BRANCH = (
    "如果有收取销售服务费的份额类别，则将其中的“调低赎回费率”改为“调低销售服务费率”，否则直接删除"
    "“、调低赎回费率”"
)
SERVICE_FEE_DELETE_OLD = (
    "C类基金份额的销售服务费每日计提，按月支付。由基金管理人向基金托管人发送基金销售服务费划付指令，"
    "经基金托管人复核后于次月首日起10个工作日内从基金财产中一次性支付给登记机构，由登记机构代付给销售"
    "机构。若遇法定节假日、公休日等，支付日期顺延。"
)

GUIDE_PARAGRAPHS: list[tuple[str, str]] = [
    ("释义-基金份额类别定义：", "bold"),
    ("原条款，参考如下描述，实际份额类别以具体产品为准：", "plain"),
    (CLASS_DEFINITION, "plain"),
    ("新条款，在原有条款后增加如下内容，实际份额类别以具体产品为准：", "plain"),
    (DIRECT_SALES_NEW, "red"),
    ("基金的基本情况-份额类别设置：", "bold"),
    ("原条款，参考以下描述:", "plain"),
    (CLASS_CHANGE_OLD, "plain"),
    ("新条款，按照以下方式调整：", "plain"),
    (SERVICE_FEE_BRANCH, "plain"),
    ("删除“并报中国证监会备案”。", "plain"),
    ("基金费用与税收-3、基金销售服务费（有多个份额类别时适用）", "bold"),
    (SERVICE_FEE_DELETE_OLD, "plain"),
    ("新条款，删除原条款。", "plain"),
]


def guide_source(tmp_path: Path, name: str = "guide.docx"):
    path = build_docx(tmp_path / name, GUIDE_PARAGRAPHS)
    return extract_guide_source(path), path


def valid_parsed() -> dict:
    result = {
        "groups": [
            {
                "rule_id": "G01",
                "title": "释义-基金份额类别定义",
                "section_hint": "释义",
                "requirement": "在释义的份额类别定义后追加直销机构申购的收费安排。",
                "operations": [{"op": "append", "new_text": DIRECT_SALES_NEW,
                                "condition": {"kind": "none", "value": None}}],
                "old_text_variants": [CLASS_DEFINITION],
                "slot_tokens": ["A类基金份额", "C类基金份额"],
                "class_substitution": True,
                "references": [],
                "source_paragraphs": [0, 1, 2, 3, 4],
                "notes": "",
            },
            {
                "rule_id": "G02",
                "title": "基金的基本情况-份额类别设置",
                "section_hint": "基金的基本情况",
                "requirement": "有销售服务费类别时把调低赎回费率改为调低销售服务费率，否则删除；并删除备案句。",
                "operations": [
                    {"op": "replace", "old_text": "调低赎回费率", "new_text": "调低销售服务费率",
                     "condition": {"kind": "has_service_fee", "value": True}},
                    {"op": "delete", "old_text": "、调低赎回费率",
                     "condition": {"kind": "has_service_fee", "value": False}},
                    {"op": "delete", "old_text": "并报中国证监会备案",
                     "condition": {"kind": "none", "value": None}},
                ],
                "old_text_variants": [CLASS_CHANGE_OLD],
                "slot_tokens": [],
                "class_substitution": False,
                "references": [],
                "source_paragraphs": [5, 6, 7, 8, 9, 10],
                "notes": "",
            },
            {
                "rule_id": "G03",
                "title": "基金费用与税收-销售服务费",
                "section_hint": "基金费用与税收",
                "requirement": "删除销售服务费支付路径条款。",
                "operations": [{"op": "delete", "old_text": SERVICE_FEE_DELETE_OLD,
                                "condition": {"kind": "has_service_fee", "value": True}}],
                "old_text_variants": [],
                "slot_tokens": [],
                "class_substitution": False,
                "references": ["同前述说明"],
                "source_paragraphs": [11, 12, 13],
                "notes": "",
            },
        ],
        "notes": "synthetic guide",
    }

    for group in result["groups"]:
        for operation in group["operations"]:
            operation["group_kind"] = "INDEPENDENT"
            operation["one_of_group"] = ""
    return result


class ScriptedGateway:
    """Pops one scripted response per guide call; ModelError entries are raised."""

    def __init__(self, responses: list[tuple[str, object]], *, max_context_chars: int = 12000):
        self.responses = list(responses)
        self.model = "scripted-guide"
        self.max_context_chars = max_context_chars
        self.calls = 0
        self.events: list[dict] = []
        self.contexts: list[tuple[str, str]] = []

    def call(self, *, task_type, bounded_context, response_schema, instruction, prompt_version="v1"):
        if not self.responses:
            raise AssertionError(f"unexpected {task_type} call")
        expected, outcome = self.responses.pop(0)
        assert expected == task_type, f"expected {expected}, got {task_type}"
        self.calls += 1
        self.contexts.append((task_type, bounded_context))
        if isinstance(outcome, ModelError):
            self.events.append({"task_type": task_type, "prompt_version": prompt_version,
                                "error_code": outcome.code})
            raise outcome
        event = {"task_type": task_type, "prompt_version": prompt_version, "actual_model": self.model,
                 "request_id": f"req-{self.calls}", "latency_ms": 1, "http_status": 200, "error_code": None,
                 "prompt_tokens": 10, "completion_tokens": 10, "estimated_cost_usd": 0.0}
        self.events.append(event)
        return SimpleNamespace(value=outcome, model=self.model, latency_ms=1, prompt_tokens=10,
                               completion_tokens=10, total_tokens=20, request_id=event["request_id"],
                               attempts=1, estimated_cost_usd=0.0, task_type=task_type)

    def metrics(self) -> dict:
        return {"logical_calls": self.calls, "http_attempts": self.calls, "checkpoint_reuses": 0,
                "prompt_tokens": 10 * self.calls, "completion_tokens": 10 * self.calls,
                "total_tokens": 20 * self.calls, "latency_ms": self.calls, "rate_limit_count": 0,
                "error_count": 0, "estimated_cost_usd": 0.0, "events": list(self.events)}


def clean_verification() -> dict:
    return {"findings": [], "coverage_explanation": "全部段落均被分组覆盖。"}


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------

def test_extraction_paragraphs_emphasis_and_full_text(tmp_path):
    source, path = guide_source(tmp_path)
    assert source.source_sha256 == sha256_file(path)
    assert len(source.paragraphs) == len(GUIDE_PARAGRAPHS)
    assert [paragraph.index for paragraph in source.paragraphs] == list(range(len(GUIDE_PARAGRAPHS)))
    assert source.paragraphs[0].emphasis == "bold"
    assert source.paragraphs[4].emphasis == "red"
    assert source.paragraphs[2].emphasis == "plain"
    assert source.paragraphs[2].text == CLASS_DEFINITION
    lines = source.full_text.split("\n")
    assert lines[2] == f"[2] {CLASS_DEFINITION}"
    assert lines[0] == "[0] 释义-基金份额类别定义："


def test_extraction_emphasis_majority_and_plain_default(tmp_path):
    document = Document()
    first = document.add_paragraph()
    first.add_run("红红红").font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    first.add_run("普通")
    second = document.add_paragraph()
    second.add_run("红").font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    second.add_run("普通文本内容")
    third = document.add_paragraph()
    third.add_run("加粗内容").bold = True
    third.add_run("普通")
    path = tmp_path / "emphasis.docx"
    document.save(path)
    source = extract_guide_source(path)
    assert source.paragraphs[0].emphasis == "red"
    assert source.paragraphs[1].emphasis == "plain"
    assert source.paragraphs[2].emphasis == "bold"


def test_extraction_skips_instructions_and_deleted_text(tmp_path):
    path = build_docx(tmp_path / "fields.docx", [("正文内容", "plain")])
    patch_document_xml(path, lambda xml: xml.replace(
        "</w:p>",
        '<w:r><w:instrText xml:space="preserve">PAGE</w:instrText></w:r>'
        '<w:del w:id="1" w:author="a"><w:r><w:t>已删除</w:t></w:r></w:del>'
        '<w:r><w:delText>删除域</w:delText></w:r></w:p>', 1))
    source = extract_guide_source(path)
    assert source.paragraphs[0].text == "正文内容"


def test_extraction_failure_codes(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(GuideError) as unreadable:
        extract_guide_source(broken)
    assert unreadable.value.code == "GUIDE_SOURCE_UNREADABLE"
    empty = build_docx(tmp_path / "empty.docx", [("", "plain"), ("   ", "plain")])
    with pytest.raises(GuideError) as empty_error:
        extract_guide_source(empty)
    assert empty_error.value.code == "GUIDE_SOURCE_EMPTY"
    with pytest.raises(GuideError) as missing:
        extract_guide_source(tmp_path / "missing.docx")
    assert missing.value.code == "GUIDE_SOURCE_UNREADABLE"


# ---------------------------------------------------------------------------
# Deterministic validation
# ---------------------------------------------------------------------------

def test_validation_accepts_faithful_parse(tmp_path):
    source, _ = guide_source(tmp_path)
    result = validate_parsed_guide(source, valid_parsed())
    assert result["errors"] == []
    assert result["coverage_gaps"] == []
    assert result["stats"]["groups"] == 3
    assert result["stats"]["operations_by_kind"] == {"append": 1, "replace": 1, "delete": 3, "noop": 0}
    assert result["stats"]["conditions_by_kind"]["has_service_fee"] == 3


def test_validation_rejects_non_verbatim_text(tmp_path):
    source, _ = guide_source(tmp_path)
    parsed = valid_parsed()
    parsed["groups"][0]["operations"][0]["new_text"] = "投资者通过直销机构申购本基金A类基金份额、C类基金份额时不收取申购费用。"
    result = validate_parsed_guide(source, parsed)
    codes = [error["code"] for error in result["errors"]]
    assert codes == ["GUIDE_TEXT_NOT_VERBATIM"]
    assert result["errors"][0]["rule_id"] == "G01"


def test_validation_rejects_bad_conditions_and_operations(tmp_path):
    source, _ = guide_source(tmp_path)
    parsed = valid_parsed()
    parsed["groups"][1]["operations"][0]["condition"] = {"kind": "class_count", "value": 2}
    parsed["groups"][1]["operations"][1]["condition"] = {"kind": "none", "value": True}
    parsed["groups"][1]["operations"].append({"op": "move", "new_text": DIRECT_SALES_NEW,
                                              "condition": {"kind": "none", "value": None},
                                              "group_kind": "INDEPENDENT", "one_of_group": ""})
    parsed["groups"][1]["operations"].append({"op": "append", "condition": {"kind": "none", "value": None},
                                              "group_kind": "INDEPENDENT", "one_of_group": ""})
    result = validate_parsed_guide(source, parsed)
    codes = [error["code"] for error in result["errors"]]
    assert codes == ["GUIDE_GROUP_INVALID"] * 4
    details = " | ".join(error["detail"] for error in result["errors"])
    assert "class_count requires value 1" in details
    assert "none requires value null" in details
    assert "op 'move'" in details
    assert "new_text must be a non-empty string" in details


def test_validation_rejects_rule_ids_and_ranges_and_slots(tmp_path):
    source, _ = guide_source(tmp_path)
    parsed = valid_parsed()
    parsed["groups"][1]["rule_id"] = "G01"
    parsed["groups"][2]["source_paragraphs"] = [11, 99]
    parsed["groups"][0]["slot_tokens"].append("D类基金份额")
    result = validate_parsed_guide(source, parsed)
    codes = sorted(error["code"] for error in result["errors"])
    assert codes == ["GUIDE_PARAGRAPH_RANGE_INVALID", "GUIDE_RULE_ID_INVALID", "GUIDE_SLOT_NOT_IN_TEXT"]
    assert result["stats"]["coverage_gaps"] == 2


def test_validation_reports_coverage_gaps_and_sequential_ids(tmp_path):
    source, _ = guide_source(tmp_path)
    parsed = valid_parsed()
    parsed["groups"] = parsed["groups"][:2]
    parsed["groups"][1]["rule_id"] = "G03"
    result = validate_parsed_guide(source, parsed)
    codes = [error["code"] for error in result["errors"]]
    assert codes == ["GUIDE_RULE_ID_INVALID"]
    assert "not sequential" in result["errors"][0]["detail"]
    assert result["coverage_gaps"] == [12, 13]


def test_validation_rejects_empty_document(tmp_path):
    source, _ = guide_source(tmp_path)
    result = validate_parsed_guide(source, {"notes": "nothing"})
    assert [error["code"] for error in result["errors"]] == ["GUIDE_GROUP_INVALID"]


# ---------------------------------------------------------------------------
# Parse and verify calls
# ---------------------------------------------------------------------------

def test_parse_guide_passes_guide_context_and_prompt_version(tmp_path):
    source, _ = guide_source(tmp_path)
    gateway = ScriptedGateway([("guide_parse", valid_parsed())])
    parsed = parse_guide(source, gateway)
    assert parsed["groups"][0]["rule_id"] == "G01"
    task_type, context = gateway.contexts[0]
    assert task_type == "guide_parse"
    assert source.full_text in context
    assert gateway.events[0]["prompt_version"] == GUIDE_PARSE_PROMPT_VERSION


def test_parse_guide_wraps_model_error(tmp_path):
    source, _ = guide_source(tmp_path)
    gateway = ScriptedGateway([("guide_parse", ModelError("MODEL_HTTP_ERROR"))])
    with pytest.raises(GuideError) as failure:
        parse_guide(source, gateway)
    assert failure.value.code == "GUIDE_MODEL_FAILED:MODEL_HTTP_ERROR"
    assert failure.value.details["model_error"] == "MODEL_HTTP_ERROR"


def test_verify_guide_drops_findings_with_invalid_evidence(tmp_path):
    source, _ = guide_source(tmp_path)
    value = {"findings": [
        {"kind": "wrong_text", "description": "备案句应删除。", "evidence_quote": "删除“并报中国证监会备案”。",
         "paragraphs": [10]},
        {"kind": "wrong_condition", "description": "条件不忠实。", "evidence_quote": "这段文字不在指引中",
         "paragraphs": [9]},
        {"kind": "missing_group", "description": "段落不存在。", "evidence_quote": "新条款，删除原条款。",
         "paragraphs": [77]},
        {"kind": "unknown_kind", "description": "种类非法。", "evidence_quote": "新条款，删除原条款。",
         "paragraphs": [13]},
    ], "coverage_explanation": "第 12—13 段未覆盖。"}
    gateway = ScriptedGateway([("guide_verify", value)])
    result = verify_guide(source, valid_parsed(), gateway)
    assert result["invalid_findings"] == 3
    assert [finding["kind"] for finding in result["findings"]] == ["wrong_text"]
    assert result["findings"][0]["paragraphs"] == [10]
    assert result["coverage_explanation"] == "第 12—13 段未覆盖。"
    assert gateway.events[0]["prompt_version"] == GUIDE_VERIFY_PROMPT_VERSION
    assert "PARSED_GUIDE_JSON" in gateway.contexts[0][1]


def test_verify_guide_normalizes_whitespace_in_quotes(tmp_path):
    source, _ = guide_source(tmp_path)
    quoted = SERVICE_FEE_BRANCH[:12]
    value = {"findings": [{"kind": "wrong_text", "description": "逐字不一致。",
                           "evidence_quote": f" {quoted}\n", "paragraphs": [9]}],
             "coverage_explanation": ""}
    gateway = ScriptedGateway([("guide_verify", value)])
    result = verify_guide(source, valid_parsed(), gateway)
    assert result["invalid_findings"] == 0
    assert len(result["findings"]) == 1


# ---------------------------------------------------------------------------
# Build loop
# ---------------------------------------------------------------------------

def test_build_guide_document_repairs_after_validation_error(tmp_path):
    source, _ = guide_source(tmp_path)
    broken = valid_parsed()
    broken["groups"][0]["operations"][0]["new_text"] = "投资者通过直销机构申购本基金A类基金份额、C类基金份额时不收取申购费用。"
    gateway = ScriptedGateway([
        ("guide_parse", broken),
        ("guide_verify", clean_verification()),
        ("guide_parse", valid_parsed()),
        ("guide_verify", clean_verification()),
    ])
    document = build_guide_document(source, gateway)
    assert document["validation"]["clean"] is True
    assert document["validation"]["rounds"] == 2
    assert [round_["errors"] for round_ in document["validation"]["rounds_detail"]] == [1, 0]
    assert document["schema_version"] == GUIDE_SCHEMA_VERSION
    assert document["source_sha256"] == source.source_sha256
    assert document["source_path_hint"] == "guide.docx"
    assert document["prompt_versions"] == {"guide_parse": GUIDE_PARSE_PROMPT_VERSION,
                                            "guide_verify": GUIDE_VERIFY_PROMPT_VERSION}
    assert document["model"]["parse_rounds"] == 2
    assert document["model"]["logical_calls"] == 4
    second_context = gateway.contexts[2][1]
    assert "REPAIR_CONTEXT" in second_context
    assert "GUIDE_TEXT_NOT_VERBATIM" in second_context


def test_build_guide_document_repairs_after_verifier_finding(tmp_path):
    source, _ = guide_source(tmp_path)
    finding = {"kind": "ambiguous_branch", "description": "单类别分支表述有歧义。",
               "evidence_quote": "否则直接删除“、调低赎回费率”", "paragraphs": [9]}
    gateway = ScriptedGateway([
        ("guide_parse", valid_parsed()),
        ("guide_verify", {"findings": [finding], "coverage_explanation": "有歧义。"}),
        ("guide_parse", valid_parsed()),
        ("guide_verify", clean_verification()),
    ])
    document = build_guide_document(source, gateway)
    assert document["validation"]["clean"] is True
    assert document["validation"]["rounds"] == 2
    assert document["validation"]["rounds_detail"][0]["findings"] == 1
    assert "INDEPENDENT_VERIFIER_FINDINGS" in gateway.contexts[2][1]


def test_build_guide_document_stops_at_max_repair_rounds(tmp_path):
    source, _ = guide_source(tmp_path)
    broken = valid_parsed()
    broken["groups"][2]["operations"][0]["old_text"] = "不存在的旧文。"
    gateway = ScriptedGateway([
        ("guide_parse", broken),
        ("guide_verify", clean_verification()),
        ("guide_parse", broken),
        ("guide_verify", clean_verification()),
    ])
    document = build_guide_document(source, gateway, max_repair_rounds=1)
    assert document["validation"]["clean"] is False
    assert document["validation"]["rounds"] == 2
    assert [error["code"] for error in document["validation"]["errors"]] == ["GUIDE_TEXT_NOT_VERBATIM"]
    with pytest.raises(GuideError) as invalid:
        build_guide_document(source, gateway, max_repair_rounds=-1)
    assert invalid.value.code == "GUIDE_BUILD_ARGUMENT_INVALID"


def test_guide_document_hash_is_stable_across_rebuilds(tmp_path):
    source, _ = guide_source(tmp_path)
    gateway = ScriptedGateway([("guide_parse", valid_parsed()), ("guide_verify", clean_verification())])
    document = build_guide_document(source, gateway)
    other = copy.deepcopy(document)
    other["built_at"] = "2026-10-15T00:00:00Z"
    other["model"] = {"configured_model": "other"}
    assert guide_document_sha256(document) == guide_document_sha256(other)
    other["groups"][0]["requirement"] = "改动了要求。"
    assert guide_document_sha256(document) != guide_document_sha256(other)


def test_save_and_load_guide_document(tmp_path):
    source, _ = guide_source(tmp_path)
    gateway = ScriptedGateway([("guide_parse", valid_parsed()), ("guide_verify", clean_verification())])
    document = build_guide_document(source, gateway)
    target = tmp_path / "out" / "builtin_guide_v1.json"
    save_guide_document(document, target)
    assert load_guide_document(target) == document
    assert json.loads(target.read_text(encoding="utf-8"))["schema_version"] == GUIDE_SCHEMA_VERSION
    unsupported = dict(document, schema_version="guide-rules-v0")
    with pytest.raises(GuideError) as save_error:
        save_guide_document(unsupported, tmp_path / "bad.json")
    assert save_error.value.code == "GUIDE_DOCUMENT_SCHEMA_UNSUPPORTED"
    with pytest.raises(GuideError) as load_error:
        load_guide_document(tmp_path / "missing.json")
    assert load_error.value.code == "GUIDE_DOCUMENT_UNREADABLE"
