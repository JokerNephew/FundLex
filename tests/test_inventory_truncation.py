"""A truncated inventory splits its source window instead of repeating the request."""
from types import SimpleNamespace

import pytest
from docx import Document

from fundlex.docx import inspect_docx
from fundlex.model import ModelError
from fundlex.recall import Window
from fundlex.recall_pipeline import _inventory, _WindowState


def test_truncated_inventory_splits_and_preserves_each_bound_target(tmp_path):
    source = tmp_path / "contract.docx"
    document = Document()
    for text in ("甲业务旧句。", "乙业务旧句。"):
        document.add_paragraph(text)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    locations = tuple(span.location for span in inspection.paragraphs)
    entry = _WindowState(Window("W", ("S1",), locations,
        "\n".join(f"[{span.location}] {span.text}" for span in inspection.paragraphs)),
        ({"rule_id": "G99", "title": "业务条款", "operations": []},))
    contexts = []

    class Gateway:
        max_context_chars = 10000

        def call(self, **kwargs):
            context = kwargs["bounded_context"]
            assert context not in contexts  # no unchanged request can be retried
            contexts.append(context)
            if len(contexts) == 1:
                raise ModelError("MODEL_OUTPUT_TRUNCATED")
            span = inspection.paragraphs[len(contexts) - 2]
            return SimpleNamespace(value={"targets": [{"rule_id": "G99",
                "location": span.location, "quote": span.text, "state": "needs_edit",
                "reason": "exact business clause", "operation_id": "G99:0"}]})

    _inventory(entry, None, SimpleNamespace(as_ledger=lambda: {}), Gateway(), inspection, round_number=0)
    assert len(contexts) == 3
    assert {target.location for target in entry.targets} == set(locations)
    assert all(target.resolution == "resolved" and target.operations == ("G99:0",)
               for target in entry.targets)
    assert entry.ledgers[0]["reason"] == "MODEL_OUTPUT_TRUNCATED"


def test_single_paragraph_truncation_still_fails_closed(tmp_path):
    source = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("业务旧句。")
    document.save(source)
    inspection = inspect_docx(source, full=True)
    span = inspection.paragraphs[0]
    entry = _WindowState(Window("W", ("S1",), (span.location,), f"[{span.location}] {span.text}"),
        ({"rule_id": "G99", "title": "业务条款", "operations": []},))

    class Gateway:
        max_context_chars = 10000
        calls = 0

        def call(self, **kwargs):
            self.calls += 1
            raise ModelError("MODEL_OUTPUT_TRUNCATED")

    gateway = Gateway()
    with pytest.raises(ModelError, match="MODEL_OUTPUT_TRUNCATED"):
        _inventory(entry, None, SimpleNamespace(as_ledger=lambda: {}), gateway, inspection, round_number=0)
    assert gateway.calls == 1
    assert not entry.targets
