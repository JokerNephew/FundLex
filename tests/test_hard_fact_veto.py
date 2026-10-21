"""FL-D106: the model judges; the code vetoes only on hard facts.

The engine used to silence a model's semantic judgment with deterministic tests that
were themselves semantic — a reference text that did not appear verbatim, a class group
the product lacked, a fee-homogeneous class set. Those are exactly the cases where a
real edit is most likely needed, so dropping them causes misses (可以错杀不能漏判).

These tests pin the new split:

* ``hard_fact_topic_absent`` — the only veto that may discard a model judgment, and only
  when the contract provably never mentions the rule's topic;
* ``class_substitution_letter_stable`` — a hard fact (instantiated replacement equals the
  reference, so no change is possible) and therefore still a legitimate drop reason;
* ``class_substitution_noop`` / ``slot_group_missing`` / ``weak_rule_signals`` — the
  semantic signals, which flag for review but no longer discard.
"""

from fundlex.recall import (
    ClassFact,
    Facts,
    class_substitution_letter_stable,
    class_substitution_noop,
    fee_homogeneous_classes,
    hard_fact_topic_absent,
)


def facts(classes, floating=False):
    return Facts(
        fund_type="equity", fund_type_quote="", fund_type_basis="",
        classes=tuple(classes), floating_management_fee=floating,
        floating_quote="", floating_basis="", notes="",
        quotes_verified={}, problems=(), fact_status={"share_class_structure":"VERIFIED", "subscription_fee_group":"VERIFIED", "service_fee_group":"VERIFIED"}, set_complete_verified=True,
    )


AC_FACTS = facts([
    ClassFact(name="A", subscription_fee=True, service_fee=False, current=True, rate=None, quote="q"),
    ClassFact(name="C", subscription_fee=False, service_fee=True, current=True, rate=None, quote="q"),
])

# A/Y product: both classes charge a subscription fee and no service fee — no C group.
AY_FACTS = facts([
    ClassFact(name="A", subscription_fee=True, service_fee=False, current=True, rate=None, quote="q"),
    ClassFact(name="Y", subscription_fee=True, service_fee=False, current=True, rate=None, quote="q"),
])

G02_RULE = {
    "rule_id": "G02", "title": "释义-基金份额类别定义",
    "requirement": "按产品实际类别对份额类别描述分组替换。",
    "operations": [
        {"op": "replace", "old_text": "A类基金份额", "new_text": "A类、B类基金份额",
         "condition": {"kind": "none", "value": None}},
        {"op": "replace", "old_text": "C类基金份额", "new_text": "C类、D类基金份额",
         "condition": {"kind": "none", "value": None}},
    ],
    "class_substitution": True,
}


def test_hard_fact_topic_absent_only_when_topic_never_mentioned():
    contract = "本基金的申购费率由基金管理人决定，并在招募说明书中列示。"
    assert hard_fact_topic_absent({"rule_id": "G15"}, contract) is True   # no 销售服务费/支付
    contract2 = contract + "销售服务费每日计提，按月支付。"
    assert hard_fact_topic_absent({"rule_id": "G15"}, contract2) is False  # topic present


def test_hard_fact_topic_absent_unknown_rule_never_vetoes():
    assert hard_fact_topic_absent({"rule_id": "G99"}, "") is False
    assert hard_fact_topic_absent({"rule_id": "G99"}, "任意文本") is False


def test_letter_stable_is_a_hard_fact_for_ac_product():
    # A->A (+B maps onto the product's own classes => A), C->C: nothing changes.
    assert class_substitution_letter_stable(G02_RULE, AC_FACTS) is True
    assert class_substitution_noop(G02_RULE, AC_FACTS) is True


def test_fee_homogeneous_is_semantic_not_a_letter_stable_fact():
    # A/Y share one fee profile: the grouping selects nothing (semantic), so it must not
    # be reported as a hard "letters already match" fact.
    assert fee_homogeneous_classes(AY_FACTS) is True
    assert class_substitution_letter_stable(G02_RULE, AY_FACTS) is False
    # class_substitution_noop stays True for backwards compatibility ...
    assert class_substitution_noop(G02_RULE, AY_FACTS) is True


def test_weak_rule_signals_flags_semantic_cases():
    from fundlex.recall_pipeline import weak_rule_signals
    ops = [op for op in G02_RULE["operations"]]
    signals = weak_rule_signals(G02_RULE, AY_FACTS, ops)
    assert "fee_homogeneous_classes" in signals
    assert "class_substitution_noop" in signals


def test_weak_rule_signals_empty_for_stable_ac_product():
    from fundlex.recall_pipeline import weak_rule_signals
    ops = [op for op in G02_RULE["operations"]]
    # The AC product is a hard-fact no-op, not a "weak/uncertain" signal.
    signals = weak_rule_signals(G02_RULE, AC_FACTS, ops)
    assert "class_substitution_noop" not in signals


def test_sweep_keeps_class_candidate_until_operation_site_proof(tmp_path):
    """Even a letter-stable product needs the shared artifact proof at this site."""
    from types import SimpleNamespace
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall_pipeline import _Chunk, _sweep_batch
    text = ("本基金将基金份额分为A类和Y类不同的类别。在投资者申购基金份额时收取申购费用，"
            "而不计提销售服务费的，称为A类基金份额。")
    path = tmp_path / "classes.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    actual = inspect_docx(path, full=True)
    location = actual.paragraphs[0].location
    class Gateway:
        def call(self, **kwargs):
            return SimpleNamespace(value={"findings": [{"rule_id": "G02", "location": location,
                "quote": text, "issue": "unaddressed_requirement"}]})
    for profile in (AY_FACTS, AC_FACTS):
        result = _sweep_batch(Gateway(), {"G02": G02_RULE}, "class substitution",
            [_Chunk("C1", (location,), f"[{location}] {text}")], actual,
            round_number=0, counter=[0], facts=profile)
        assert len(result["findings"]) == 1
        assert not result["dropped"]
