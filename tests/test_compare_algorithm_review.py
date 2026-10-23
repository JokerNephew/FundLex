"""Adversarial checks found during the independent compare review.

These tests deliberately exercise cases where a permissive parser or a
direction-dependent alignment can hide a real textual change.  They are kept
separate from the implementation tests so future changes cannot silently
remove the high-recall and symmetry gates.
"""

from docx import Document
from docx.oxml import OxmlElement

from fundlex.compare import compare_documents, parse_compare_document
from fundlex.compare.models import CompareBlock, CompareDocument
from fundlex.compare.normalize import normalize_diff, normalize_match


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _paragraph(text: str):
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    paragraph.append(run)
    return paragraph


def _sdt(*children):
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    for child in children:
        content.append(child)
    control.append(content)
    return control


def _doc_with_blocks(values, side):
    return CompareDocument(
        tuple(
            CompareBlock(
                f"{side}{i:06d}", i, f"p{i}", "paragraph",
                raw_text=value, match_text=normalize_match(value),
                diff_text=normalize_diff(value),
            )
            for i, value in enumerate(values)
        )
    )


def _semantic(result, reverse=False):
    """Return text-level changes with the sides put into one orientation."""
    rows = []
    for item in result.diffs:
        left = "".join(span.text for span in item.left_spans)
        right = "".join(span.text for span in item.right_spans)
        kind = item.type
        if reverse:
            left, right = right, left
            kind = {"insert": "delete", "delete": "insert"}.get(kind, kind)
        rows.append((kind, left, right))
    return sorted(rows)


def test_nested_body_sdt_text_is_not_silently_dropped(tmp_path):
    document = Document()
    body = document._element.body
    body.insert(0, _sdt(_sdt(_paragraph("嵌套控件文字"))))
    path = tmp_path / "nested-sdt.docx"
    document.save(path)

    parsed = parse_compare_document(path)

    assert "嵌套控件文字" in [block.raw_text for block in parsed.blocks]
    assert parsed.coverage.status == "complete"


def test_table_cell_sdt_text_preserves_cell_scope(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    # Keep the ordinary paragraph and add a content control as a sibling in
    # the same cell.  The text must remain in the table cell scope.
    cell._tc.append(_sdt(_paragraph("单元格控件文字")))
    path = tmp_path / "cell-sdt.docx"
    document.save(path)

    parsed = parse_compare_document(path)
    blocks = [block for block in parsed.blocks if block.raw_text == "单元格控件文字"]
    assert len(blocks) == 1
    assert blocks[0].container == "table"
    assert blocks[0].table_index == 0
    assert blocks[0].row_index == 0
    assert blocks[0].cell_index == 0
    assert parsed.coverage.status == "complete"


def test_direction_swap_preserves_all_textual_change_polarity():
    # SequenceMatcher has direction-dependent tie breaking.  This small
    # repeated-character case used to produce a different set of changed text
    # when the two contracts were swapped, which violates the compare contract
    # even though both runs are deterministic.
    left = _doc_with_blocks(["甲甲乙"], "L")
    right = _doc_with_blocks(["甲乙甲甲"], "R")

    forward = compare_documents(left, right)
    backward = compare_documents(right, left)

    assert _semantic(forward) == _semantic(backward, reverse=True)


def test_reordering_unique_blocks_is_direction_symmetric():
    left = _doc_with_blocks(["甲", "乙"], "L")
    right = _doc_with_blocks(["乙", "甲"], "R")
    forward = compare_documents(left, right)
    backward = compare_documents(right, left)
    assert _semantic(forward) == _semantic(backward, reverse=True)


def test_split_into_four_blocks_is_structural_only():
    left = _doc_with_blocks(["甲乙丙丁"], "L")
    right = _doc_with_blocks(["甲", "乙", "丙", "丁"], "R")
    assert compare_documents(left, right).diffs == ()


def test_adjacent_similar_blocks_keep_both_numeric_changes():
    left = _doc_with_blocks(
        ["管理费率为0.40%，按日计提。", "托管费率为0.10%，按日计提。"], "L"
    )
    right = _doc_with_blocks(
        ["管理费率为0.30%，按日计提。", "托管费率为0.20%，按日计提。"], "R"
    )

    result = compare_documents(left, right)
    changed_left = "".join(span.text for item in result.diffs for span in item.left_spans)
    changed_right = "".join(span.text for item in result.diffs for span in item.right_spans)
    assert "0.40%" in changed_left and "0.30%" in changed_right
    assert "0.10%" in changed_left and "0.20%" in changed_right


def test_w_sym_is_reported_as_coverage_gap_instead_of_zero_diff(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("符号")
    run = OxmlElement("w:r")
    symbol = OxmlElement("w:sym")
    symbol.set(f"{{{W}}}font", "Wingdings")
    symbol.set(f"{{{W}}}char", "F0FC")
    run.append(symbol)
    paragraph._p.append(run)
    path = tmp_path / "symbol.docx"
    document.save(path)

    parsed = parse_compare_document(path)

    assert parsed.coverage.status == "incomplete"
    assert any("sym" in warning for warning in parsed.coverage.warnings)


def test_unknown_table_wrapper_with_visible_text_is_not_silent(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=1)
    wrapper = OxmlElement("w:customXml")
    wrapper.append(_paragraph("被包装的单元格文字"))
    table.cell(0, 0)._tc.append(wrapper)
    path = tmp_path / "unknown-table-wrapper.docx"
    document.save(path)

    parsed = parse_compare_document(path)

    # The first implementation only walked direct w:p/w:sdt children of a
    # cell.  A customXml wrapper therefore discarded visible text while still
    # claiming complete coverage.  Until such wrappers are supported they
    # must produce an incomplete coverage warning.
    assert parsed.coverage.status == "incomplete"
    assert any("table" in warning for warning in parsed.coverage.warnings)


def test_visible_header_text_is_reported_as_out_of_scope(tmp_path):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "页眉费率0.40%"
    path = tmp_path / "header.docx"
    document.save(path)
    parsed = parse_compare_document(path)
    assert parsed.coverage.status == "incomplete"
    assert "UNSUPPORTED_VISIBLE_CONTENT:header" in parsed.coverage.warnings
