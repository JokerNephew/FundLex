"""Independent source-layout and semantic controls for bounded rule recognition."""
from pathlib import Path
import json
import hashlib
import pytest
from docx import Document
from fundlex.docx import inspect_docx
from fundlex.rules import evaluate_candidates,load_ruleset


def action_result(path,aid,profile=None):
    return next(a for a in evaluate_candidates(inspect_docx(path),load_ruleset('v1'),profile)
                if a['action_id']==aid)


def document(path,heading,subsection,text):
    doc=Document();doc.add_paragraph(heading);doc.add_paragraph(subsection)
    doc.add_paragraph(text);doc.save(path)


@pytest.mark.parametrize('title',['基金份额类别设置','基金份额的分类','基金份额类别','基金的份额类别设置'])
def test_class_header_alias_and_exact_phrase_boundaries(tmp_path,title):
    path=tmp_path/'classes.docx'
    text='本基金可调整份额类别，并报中国证监会备案，无需召开大会。'
    document(path,'三、基金的基本情况','（五）'+title,text)
    candidate=action_result(path,'R03-A3')['candidates'][0]
    assert candidate['target_verified'] and candidate['quote']=='并报中国证监会备案'
    assert text[candidate['char_start']:candidate['char_end']]==candidate['quote']


def test_rate_variant_has_source_host_and_ownership_paragraph_is_excluded(tmp_path):
    path=tmp_path/'rates.docx'
    doc=Document();doc.add_paragraph('六、基金份额的申购、赎回与转换')
    doc.add_paragraph('（七）申购和赎回的费用及其用途')
    variant='对于持续持有期不足7天的基金份额持有人收取1.5%的赎回费，并全部计入基金财产。'
    doc.add_paragraph('实际执行的赎回费率在招募说明书中载明。'+variant)
    doc.add_paragraph('赎回费用由持有人承担。'+variant)
    doc.save(path)
    candidates=action_result(path,'R11-A1')['candidates']
    assert len(candidates)==2
    assert candidates[0]['target_verified'] and candidates[0]['quote']==variant
    evidence=candidates[0]['host_evidence'][0]
    span=inspect_docx(path).by_location(evidence['paragraph_id'])
    assert span.text[evidence['char_start']:evidence['char_end']]==evidence['quote']
    assert candidates[1]['status']=='excluded' and candidates[1]['reason']=='SHORT_HOLDING_NON_TARGET_HOST'


def test_missing_rate_host_fails_closed_and_adjacent_subsection_is_not_borrowed(tmp_path):
    path=tmp_path/'missing-host.docx'
    doc=Document();doc.add_paragraph('第六部分 基金份额的申购与赎回')
    doc.add_paragraph('六、其他业务');doc.add_paragraph('基金管理人可调整费率。')
    doc.add_paragraph('七、申购和赎回的价格、费用及其用途')
    doc.add_paragraph('本基金对持续持有期少于7日的投资者收取不低于1.50%的赎回费，并将上述赎回费全额计入基金财产。')
    doc.save(path)
    c=action_result(path,'R11-A1')['candidates'][0]
    assert not c.get('target_verified') and c['reason']=='SHORT_HOLDING_HOST_UNVERIFIED'


def test_cross_paragraph_phrase_is_reported_but_never_an_edit(tmp_path):
    path=tmp_path/'split.docx'
    doc=Document();doc.add_paragraph('第四部分 基金的基本情况')
    doc.add_paragraph('基金份额类别设置')
    doc.add_paragraph('调整前需公告并');doc.add_paragraph('报中国证监会备案。')
    doc.save(path)
    c=action_result(path,'R03-A3')['candidates'][0]
    assert c['reason']=='CROSS_PARAGRAPH_TARGET_UNSUPPORTED'
    assert c['status']=='failed' and not c['mapping_safe'] and not c.get('target_verified')
    assert [s['quote'] for s in c['evidence_spans']]==['并','报中国证监会备案']


def test_conversion_sentence_with_extra_fee_obligation_is_not_verified(tmp_path):
    path=tmp_path/'extra-duty.docx'
    document(path,'第四部分 基金的基本情况','基金份额类别设置',
        '本基金不同基金份额类别之间不得互相转换，并应支付相关费用。')
    profile={'share_class_count':2,'share_class_count_verified':True,'share_class_count_min':2}
    c=action_result(path,'R04-A1',profile)['candidates'][0]
    assert c['reason']=='CONVERSION_SCOPE_UNVERIFIED' and not c.get('target_verified')


def test_alternate_headings_do_not_unlock_directory_page_entries(tmp_path):
    path=tmp_path/'toc.docx'
    doc=Document();doc.add_paragraph('目录')
    from docx.enum.style import WD_STYLE_TYPE
    doc.styles.add_style('TOC1',WD_STYLE_TYPE.PARAGRAPH)
    p=doc.add_paragraph('三、基金的基本情况');p.style='TOC1'
    doc.add_paragraph('1.11 基金的投资 ...... 18')
    doc.add_paragraph('本基金投资于债券。');doc.save(path)
    inspection=inspect_docx(path)
    assert all(p.zone=='toc' and not p.editable for p in inspection.paragraphs)


def test_numbered_summary_keeps_its_own_region(tmp_path):
    path=tmp_path/'summary.docx'
    document(path,'二十六、基金合同的内容摘要','第四部分 基金的基本情况','本基金的类别为混合型。')
    inspection=inspect_docx(path)
    assert all(p.zone=='summary' for p in inspection.paragraphs)


def test_recognition_update_does_not_change_confirmed_template_literals():
    current=json.loads(Path('rules/rules_v1.json').read_text(encoding='utf-8'))
    digest=hashlib.sha256(json.dumps(current['templates'],sort_keys=True,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
    assert digest=='04da6f3a7a3960b78b171d885bbea4d4ddc4fe1b0082b598442e0b92c9b58596'
    assert all(not a['auto_write'] and not a['approved'] for a in current['actions'])
