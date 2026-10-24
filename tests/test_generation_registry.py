"""Phase 1 acceptance tests: Strategy Registry + Router (plan v4 §17, §18, §22).

The registry must cover every rule group of the compiled guide; routing must
honour the plan's two hard rules: fact failures block (never fall through to
Tier3), and only text-structure reasons may degrade a tier.
"""
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from fundlex.generation import (
    OPERATION_REGISTRY,
    PolicyComponents,
    STRATEGY_REGISTRY,
    GenerationTier,
    operation_for,
    operation_registry_coverage_gaps,
    registry_coverage_gaps,
    route_operation,
    strategy_for,
    strategy_mode,
)
from fundlex.generation.registry import register_guide_operations

GUIDE_PATH = Path(__file__).resolve().parents[1] / "guides" / "builtin_guide_v1.json"


def _guide() -> dict:
    return json.loads(GUIDE_PATH.read_text(encoding="utf-8"))


def test_registry_covers_every_guide_rule():
    guide = _guide()
    assert registry_coverage_gaps(guide) == []
    rule_ids = {group["rule_id"] for group in guide["groups"]}
    assert set(STRATEGY_REGISTRY) == rule_ids


def test_operation_registry_covers_every_guide_operation():
    guide = _guide()
    assert operation_registry_coverage_gaps(guide) == []
    assert len(OPERATION_REGISTRY) == sum(
        len(group.get("operations") or []) for group in guide["groups"])


def test_sibling_operations_do_not_share_strategy():
    first = operation_for("G11:0")
    second = operation_for("G11:1")
    assert first.operation_key == "G11:0"
    assert second.operation_key == "G11:1"
    assert first.default_tier is GenerationTier.TEMPLATE
    assert second.default_tier is GenerationTier.SPLICE
    assert first.executor == "template"
    assert second.executor == "splice"
    assert strategy_for("G11:0") is first
    assert strategy_for("G11:1") is second


def test_operation_policy_components_are_scoped_to_operation():
    spec = operation_for("G16:0")
    assert isinstance(spec.policy_components, PolicyComponents)
    assert spec.policy_components.operation_key == "G16:0"
    assert spec.policy_components.fact_dependencies == spec.required_fact_status
    assert {item.id for item in spec.policy_components.must_add} == {
        "daily_accrual", "direct_refund", "non_direct_sales_policy",
        "over_one_year_refund",
    }


def test_registry_specs_are_internally_consistent():
    for spec in STRATEGY_REGISTRY.values():
        assert spec.default_tier in spec.allowed_tiers
        assert spec.mode in {"replace", "append", "delete", "class_substitution", "noop"}
        assert all(isinstance(name, str) and name for name in spec.required_fact_status)


def test_unregistered_rule_falls_back_to_compose():
    spec = strategy_for("G99:0")
    assert spec.default_tier is GenerationTier.COMPOSE
    assert "未登记" in spec.note


def test_unknown_fact_blocks_instead_of_falling_through_to_tier3():
    decision = route_operation("G08:0", fact_status={"share_class_structure": "UNKNOWN"})
    assert decision.blocked is True
    assert decision.actual_tier is None
    assert decision.selected_executor is None
    assert "share_class_structure" in decision.fact_blocking
    assert dict(decision.fact_status)["share_class_structure"] == "UNKNOWN"
    assert decision.fallback_reason.startswith("FACT_NOT_VERIFIED")
    # CONFLICT behaves like UNKNOWN; only VERIFIED passes the gate.
    assert route_operation(
        "G08:0", fact_status={"share_class_structure": "CONFLICT",
                              "subscription_fee_group": "VERIFIED",
                              "service_fee_group": "VERIFIED"}).blocked is True
    assert route_operation(
        "G08:0", fact_status={"share_class_structure": "VERIFIED",
                              "subscription_fee_group": "VERIFIED",
                               "service_fee_group": "VERIFIED"}).resolved is True


def test_route_decision_records_executor_and_serializable_tiers():
    decision = route_operation(
        "G11:1",
        fact_status={"share_class_structure": "VERIFIED",
                     "subscription_fee_group": "VERIFIED",
                     "service_fee_group": "VERIFIED"},
        structure="unsafe",
    )
    assert decision.requested_tier is GenerationTier.SPLICE
    assert decision.actual_tier is GenerationTier.COMPOSE
    assert decision.selected_executor == "compose"
    assert decision.fallback_reason == "structure_unsafe_for_splice"
    assert decision.as_dict()["actual_tier"] == GenerationTier.COMPOSE.value


def test_custom_explicit_operation_requires_registered_registry_entry():
    guide = {"groups": [{"rule_id": "G99", "operations": [
        {"operation_id": "G99:CUSTOM", "op": "replace"},
    ]}]}
    assert operation_registry_coverage_gaps(guide) == ["G99:CUSTOM"]
    assert register_guide_operations(guide) == ["G99:CUSTOM"]


def test_routed_planner_fails_closed_for_missing_custom_operation_registry_entry(
        monkeypatch):
    """A routed planner must not silently borrow a rule-level or sibling executor."""
    from fundlex.recall import RecallError, Target, plan_target

    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    rule = {
        "rule_id": "G99",
        "operations": [{
            "operation_id": "G99:CUSTOM_MISSING",
            "op": "replace",
            "old_text": "旧文。",
            "new_text": "新文。",
            "condition": {"kind": "none", "value": None},
        }],
    }
    target = Target(
        target_id="T-CUSTOM-MISSING",
        rule_id="G99",
        location="p0",
        quote="旧文。",
        state="needs_edit",
        reason="registry integration regression",
        start=0,
        end=3,
        resolution="resolved",
        operations=("G99:CUSTOM_MISSING",),
    )
    facts = SimpleNamespace(status_of=lambda _name: "VERIFIED")

    with pytest.raises(RecallError) as excinfo:
        plan_target(rule, target, None, facts, None, compose_budget=[0])

    assert excinfo.value.code == "RECALL_OPERATION_REGISTRY_MISSING"
    assert excinfo.value.detail == {
        "operation": "G99:CUSTOM_MISSING",
        "rule": "G99",
    }


def test_legacy_operation_without_identity_remains_outside_routed_registry():
    guide = {"groups": [{"rule_id": "G99", "operations": [
        {"op": "replace"},
    ]}]}
    assert operation_registry_coverage_gaps(guide) == ["G99:0"]


def test_unknown_fact_can_never_select_compose_even_for_compose_requested_rule():
    decision = route_operation(
        "G16:0",
        fact_status={"share_class_structure": "VERIFIED",
                     "service_fee_group": "CONFLICT"},
        structure="unsafe",
    )
    assert decision.requested_tier is GenerationTier.COMPOSE
    assert decision.actual_tier is None
    assert decision.selected_executor is None
    assert decision.blocked_facts == ("service_fee_group",)


def test_route_degrades_only_on_structure():
    verified = {"share_class_structure": "VERIFIED"}
    mixed = route_operation("G14:0", fact_status=verified, structure="mixed")
    assert mixed.actual_tier is GenerationTier.SPLICE
    assert mixed.fallback_reason == "complex_mixed_clause"
    unsafe = route_operation("G14:0", fact_status=verified, structure="unsafe")
    assert unsafe.actual_tier is GenerationTier.COMPOSE
    assert unsafe.fallback_reason == "structure_unsafe_for_splice"
    clean = route_operation("G08:0",
                            fact_status={**verified,
                                         "subscription_fee_group": "VERIFIED",
                                         "service_fee_group": "VERIFIED"},
                            structure="clean")
    assert clean.actual_tier is GenerationTier.TEMPLATE
    assert clean.fallback_reason == "clean_single_sentence"


def test_delete_rules_keep_their_deterministic_path():
    decision = route_operation("G15:0", structure="unsafe")
    assert decision.actual_tier is GenerationTier.DELETE


def test_g09_clean_clause_keeps_implemented_splice_executor():
    decision = route_operation("G09:0", structure="clean")
    assert decision.selected_executor == "splice"
    assert decision.actual_tier is GenerationTier.SPLICE
    assert route_operation("G09:0", structure="unsafe").blocked


def test_g09_routed_planner_writes_both_phrases_without_compose(tmp_path, monkeypatch):
    from docx import Document
    from fundlex.docx import inspect_docx, validate_edit_plan, generate_artifacts
    from fundlex.recall import Target, plan_target, build_docx_edits, content_check_factory
    from test_recall_core import _facts

    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    rule = next(group for group in _guide()["groups"] if group["rule_id"] == "G09")
    old = "赎回金额为按实际确认的有效赎回份额乘以当日该类基金份额净值并扣除相应的费用，赎回金额单位为元。"
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("六、基金份额的申购与赎回")
    doc.add_paragraph(old)
    doc.save(source)
    inspection = inspect_docx(source, full=True)
    target = Target("T-G09", "G09", "p1", old, "needs_edit", "variant",
                    0, len(old), "resolved", operations=("G09:0",))
    edit = plan_target(rule, target, inspection, _facts(), None, compose_budget=[0])
    assert edit.evidence["selected_executor"] == "splice"
    assert edit.operation_keys == ("G09:0",)
    plan = validate_edit_plan(source, build_docx_edits([edit], inspection), {},
        content_check=content_check_factory([edit.as_dict()], inspection))
    final = tmp_path / "final.docx"
    generate_artifacts(source, plan, tmp_path / "redline.docx", final)
    text = inspect_docx(final, full=True).by_location("p1").text
    assert "当日该类基金份额的基金份额净值并扣除或加上相应的费用" in text


def test_flag_default_is_registry_and_unknown_values_fall_back(monkeypatch):
    monkeypatch.delenv("FUNDLEX_GENERATION_STRATEGY", raising=False)
    assert strategy_mode() == "registry"
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "bogus")
    assert strategy_mode() == "registry"
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry_shadow")
    assert strategy_mode() == "registry_shadow"


def test_generation_mix_classification_matches_plan_buckets():
    from tools.evaluation.generation_mix import classify
    assert classify("fragment_replace") == "deterministic"
    assert classify("whole_paragraph_delete") == "deterministic"
    assert classify("composed") == "model"
    assert classify("composed_delete") == "model"
    assert classify("merged") == "merge"
    assert classify("unknown-source") == "other"
