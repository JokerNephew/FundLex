"""Actual DOCX and replay checks for a numbered fee-burden clause (FL-D215)."""
import json
from pathlib import Path

import pytest
from docx import Document

from fundlex.docx import inspect_docx
from fundlex.generation.claim_replay import replay_claims
from fundlex.generation.template import plan_policy_replacement
from fundlex.obligations import apply_facts, compile_obligations
from fundlex.recall import Target
from fundlex.recall_pipeline import _plan_and_generate
from test_recall_core import _facts


class NoModel:
    max_context_chars = 60000

    def call(self, **kwargs):
        raise AssertionError("clean numbered fee clause requires no model")


@pytest.mark.parametrize("prefix", ["4、", " ４、 ", "（四）", "(12) ", "4. ", "四、"])
def test_numbered_fee_clause_survives_docx_and_replay(tmp_path, monkeypatch, prefix):
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    source = tmp_path / "source.docx"
    before = prefix + "A 类基金份额的申购费用由投资人承担，不列入基金财产；C 类基金份额不收取申购费用。"
    doc = Document()
    doc.add_paragraph(before)
    doc.add_paragraph("5、相邻赎回条款及原费率必须保留。")
    doc.save(source)
    original = inspect_docx(source, full=True)
    guide = json.loads((Path(__file__).resolve().parents[1] / "guides/builtin_guide_v1.json").read_text())
    rule = next(g for g in guide["groups"] if g["rule_id"] == "G11")
    facts = _facts()
    ledger = compile_obligations({"groups": [rule]})
    apply_facts(ledger, facts)
    target = Target("T", "G11", "p0", before, "needs_edit", "inventory",
                    start=0, end=len(before), resolution="resolved", operations=("G11:0",))
    planned, _, _, _, actual = _plan_and_generate(NoModel(), {"G11": rule}, [target],
        original, facts, source, tmp_path, max_compose_calls=0, obligations=ledger)
    expected = prefix + rule["operations"][0]["new_text"]
    assert actual.by_location("p0").text == expected
    assert actual.by_location("p1").text == original.by_location("p1").text
    proof = replay_claims(planned["edits"], original, actual, rules_by_id={"G11": rule},
        facts=facts, obligations=ledger, targets_by_id={"T": target.as_dict()})
    assert proof["verified"], proof["failures"]
    # A complete policy does not authorize dropping or renumbering the clause.
    for forged_prefix in ("", "9、"):
        replacement = forged_prefix + rule["operations"][0]["new_text"]
        forged = {**planned["edits"][0], "replacement": replacement, "components": []}
        changed = Document()
        changed.add_paragraph(replacement)
        changed.add_paragraph(original.by_location("p1").text)
        changed.save(tmp_path / "forged.docx")
        rejected = replay_claims([forged], original, inspect_docx(tmp_path / "forged.docx", full=True),
            rules_by_id={"G11": rule}, facts=facts, obligations=ledger,
            targets_by_id={"T": target.as_dict()})
        assert not rejected["verified"]
        assert any(f["code"] == "CLAIM_OPERATION_SPAN_UNAUTHORIZED" for f in rejected["failures"])


def test_numbered_mixed_clause_still_requires_bounded_planning():
    mixed = "4、申购费用由投资人承担，不列入基金财产。5、赎回费用由投资者承担。"
    assert plan_policy_replacement(mixed, "新政策。", "G11") is None


def test_numbered_g08_keeps_literal_ordinal():
    before = "（二）本基金的申购费率由基金管理人决定，并在招募说明书中列示。"
    result = plan_policy_replacement(before, "通过直销机构申购本基金A类基金份额不收取申购费。", "G08")
    assert result.replacement.startswith("（二）通过直销机构")
