"""Corpus and scoring checks independent from production alignment code."""
from pathlib import Path
from docx import Document
from tools.evaluation.compare_gold import gold_pair
from tools.evaluation.compare_evaluate import score_pair
from tools.evaluation.compare_evaluate import _semantic
from fundlex.compare.parser import parse_compare_document

def document(path, texts):
    d=Document()
    for text in texts:d.add_paragraph(text)
    d.save(path);return path

def test_gold_keeps_exact_offset_for_repeated_text(tmp_path):
    a=document(tmp_path/'a.docx',['重复条款','基金管理人应当披露','重复条款','基金管理人应当披露','结束'])
    b=document(tmp_path/'b.docx',['重复条款','基金管理人应当披露','重复条款','基金管理人可以披露','结束'])
    gold=gold_pair(a,b)
    assert gold['operations']
    assert gold['operations'][0]['left']['ordinal']==3
    score=score_pair({'gold':gold,'base_path':str(a),'path':str(b)})
    assert score['covered']==score['gold_ops'] and score['gold_ops']>0

def test_zero_gold_respects_paragraph_split(tmp_path):
    a=document(tmp_path/'a.docx',['基金管理人应当披露'])
    b=document(tmp_path/'b.docx',['基金管理人','应当披露'])
    gold=gold_pair(a,b,expected_zero=True)
    assert gold['expected_zero']
    assert score_pair({'gold':gold,'base_path':str(a),'path':str(b)})['diff_count']==0

def test_real_diff_gold_demands_location(tmp_path):
    a=document(tmp_path/'a.docx',['费率0.40%','管理人应当履约'])
    b=document(tmp_path/'b.docx',['费率0.30%','管理人可以履约'])
    gold=gold_pair(a,b)
    assert len(gold['operations'])==2
    score=score_pair({'gold':gold,'base_path':str(a),'path':str(b)})
    assert score['covered']==2

def test_zero_gold_is_derived_from_visible_stream_not_manifest_flag(tmp_path):
    a=document(tmp_path/'a.docx',['基金管理人应当披露'])
    b=document(tmp_path/'b.docx',['基金管理人','应当披露'])
    gold=gold_pair(a,b,expected_zero=False)
    assert gold['expected_zero'] is True
    assert gold['declared_expected_zero'] is False
    assert gold['zero_declaration_matches'] is False

def test_symmetry_reverses_insert_delete_polarity_once(tmp_path):
    a=document(tmp_path/'a.docx',['原始条款'])
    b=document(tmp_path/'b.docx',['原始条款','新增条款'])
    from fundlex.compare.text_diff import compare_docx
    ab=compare_docx(a,b).to_dict(); ba=compare_docx(b,a).to_dict()
    assert any(d['type']=='insert' for d in ab['diffs'])
    assert _semantic(ab)==_semantic(ba,reverse=True)

def test_symmetry_rejects_wrong_polarity_and_wrong_block():
    ab={'diffs':[{'type':'insert','right_spans':[{'block_id':'R000001','text':'新增','start':0,'end':2}]}]}
    wrong={'diffs':[{'type':'insert','left_spans':[{'block_id':'L000001','text':'新增','start':0,'end':2}]}]}
    assert _semantic(ab)!=_semantic(wrong,reverse=True)
    wrong['diffs'][0]['type']='delete'
    assert _semantic(ab)==_semantic(wrong,reverse=True)
    wrong['diffs'][0]['left_spans'][0]['block_id']='L000002'
    assert _semantic(ab)!=_semantic(wrong,reverse=True)

def test_gold_does_not_import_compare_runtime():
    import ast
    from pathlib import Path
    tree=ast.parse(Path('tools/evaluation/compare_gold.py').read_text(encoding='utf-8'))
    imported=[n.module or '' for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)]
    names=[alias.name for n in ast.walk(tree) if isinstance(n, ast.Import) for alias in n.names]
    assert not any(name.startswith('fundlex.compare') for name in imported+names)

def test_drawingml_move_to_is_not_word_revision(tmp_path):
    # DrawingML has a moveTo element of its own.  Namespace-aware revision
    # detection must not reject a normal chart/shape path as an existing w:moveTo.
    d=Document(); d.add_paragraph('合同正文'); p=tmp_path/'shape.docx'; d.save(p)
    with __import__('zipfile').ZipFile(p) as z: parts={n:z.read(n) for n in z.namelist()}
    from lxml import etree
    root=etree.fromstring(parts['word/document.xml']); drawing=etree.SubElement(root, '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r'); dml=etree.SubElement(drawing, '{http://schemas.openxmlformats.org/drawingml/2006/main}moveTo'); etree.SubElement(dml, '{http://schemas.openxmlformats.org/drawingml/2006/main}pt')
    parts['word/document.xml']=etree.tostring(root)
    with __import__('zipfile').ZipFile(p,'w') as z:
        for n,v in parts.items():z.writestr(n,v)
    assert parse_compare_document(p).blocks
