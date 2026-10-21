from __future__ import annotations

from types import SimpleNamespace

import pytest

from fundlex.fact_evidence import verify_class_set
from fundlex.frozen import FrozenError, ModelCassette
from fundlex.guide import GuideError, validate_runtime_guide
from fundlex.generation.effects import evaluate_operation_effect
from fundlex.obligations import apply_facts, compile_obligations, project


class Facts:
    fact_status = {"share_class_structure": "VERIFIED", "subscription_fee_group": "VERIFIED",
                   "service_fee_group": "VERIFIED", "fund_type": "VERIFIED"}
    quotes_verified = {"floating": True}
    class_count = 2
    has_service_fee = True
    floating_management_fee = False

    def status_of(self, name):
        return self.fact_status.get(name, "UNKNOWN")

    def as_ledger(self):
        return {"fact_status": self.fact_status}


def _guide(op=None):
    return {"groups": [{"rule_id": "G01", "title": "G01", "operations": [op or {
        "operation_id": "G01:0", "op": "delete", "old_text": "备案句",
        "condition": {"kind": "none"}, "group_kind": "INDEPENDENT"}]}]}


def test_true_operation_never_becomes_not_applicable_when_site_is_missing():
    ledger = compile_obligations(_guide())
    apply_facts(ledger, Facts())
    project(ledger, {"targets": [], "plan": {"ledger": []}, "artifact_claim_replay": {}})
    record = ledger.records[0]
    assert record.applicability.status == "TRUE"
    assert record.completion != "NOT_APPLICABLE"


def test_class_set_quote_rejects_partial_a_c_for_a_c_e():
    inspection = SimpleNamespace(paragraphs=[SimpleNamespace(text="本基金将基金份额分为A类、C类和E类不同的类别")])
    classes = [SimpleNamespace(name="A", current=True), SimpleNamespace(name="C", current=True)]
    proof = verify_class_set(inspection, classes, "本基金将基金份额分为A类、C类和E类不同的类别")
    assert proof["quote_verified"] is True
    assert proof["set_complete_verified"] is False


def test_effect_proof_is_site_scoped():
    original = SimpleNamespace(paragraphs=[SimpleNamespace(location="p1", text="旧句"),
                                           SimpleNamespace(location="p2", text="旧句")])
    artifact = SimpleNamespace(paragraphs=[SimpleNamespace(location="p1", text="新句"),
                                           SimpleNamespace(location="p2", text="旧句")])
    artifact.by_location = lambda location: next(x for x in artifact.paragraphs if x.location == location)
    original.by_location = lambda location: next(x for x in original.paragraphs if x.location == location)
    record = SimpleNamespace(obligation_id="G01:B1:O1", applicability=SimpleNamespace(status="TRUE"),
                            spec=SimpleNamespace(rule_id="G01", operation_id="G01:0", op="replace",
                                                 old_text="旧句", new_text="新句"))
    proof = evaluate_operation_effect(record, original, artifact,
                                      [{"location": "p2", "quote": "旧句"}], [], Facts())
    assert proof.state == "FAIL"
    assert proof.site_effects[0]["failure_code"] == "CLAIM_OPERATION_NOT_PLANNED"


def test_cassette_is_immutable_and_replayable(tmp_path):
    cassette = ModelCassette(tmp_path)
    request = {"task_type": "recall_facts", "prompt_version": "v1", "configured_model": "m",
               "context_sha256": "a" * 64, "request": {"bounded_context": "x"}}
    schema = {"type": "object", "properties": {"ok": {"type": "boolean"}}, "required": ["ok"]}
    cassette.save(request, {"ok": True}, {"task_type": "recall_facts", "prompt_version": "v1",
                                           "context_sha256": "a" * 64, "configured_model": "m",
                                           "actual_model": "m", "http_status": 200})
    assert cassette.load(request, schema)[0] == {"ok": True}
    with pytest.raises(FrozenError):
        cassette.save(request, {"ok": False}, {})


def test_guide_gate_rejects_uncertified_document():
    with pytest.raises(GuideError) as error:
        validate_runtime_guide(_guide())
    assert error.value.code == "GUIDE_NOT_CERTIFIED"


def test_gateway_replay_never_uses_network_or_budget(tmp_path):
    import httpx
    from fundlex.model import ModelGateway, ModelError
    calls = []
    def record(request):
        calls.append(request)
        return httpx.Response(200, json={"model": "m", "choices": [
            {"message": {"content": '{"ok":true}'}, "finish_reason": "stop"}]})
    kwargs = {"task_type": "recall_facts", "bounded_context": "public contract",
              "response_schema": {"type": "object", "required": ["ok"]},
              "instruction": "extract", "prompt_version": "v1"}
    with ModelGateway(base_url="https://example.invalid", model="m", api_key="test-secret",
                      mode="record", cassette_dir=tmp_path,
                      client=httpx.Client(transport=httpx.MockTransport(record))) as gateway:
        assert gateway.call(**kwargs).value == {"ok": True}
    def forbidden(request):
        raise AssertionError("replay attempted HTTP")
    class Budget:
        def reserve(self, **kwargs):
            raise AssertionError("replay attempted budget reservation")
    with ModelGateway(base_url="", model="m", api_key="", mode="replay", cassette_dir=tmp_path,
                      budget=Budget(), client=httpx.Client(transport=httpx.MockTransport(forbidden))) as gateway:
        assert gateway.call(**kwargs).value == {"ok": True}
        assert gateway.metrics()["http_attempts"] == 0
        with pytest.raises(ModelError, match="MODEL_REPLAY_MISS"):
            gateway.call(**{**kwargs, "bounded_context": "changed"})
    assert len(calls) == 1
    assert all("test-secret" not in path.read_text() for path in tmp_path.rglob("*.json"))


@pytest.mark.parametrize("scope", [None, [], ["S1"]])
@pytest.mark.parametrize("target_state", [None, "absent", "satisfied", "needs_edit"])
def test_true_completion_invariant_across_target_and_scope_states(scope, target_state):
    ledger = compile_obligations(_guide())
    apply_facts(ledger, Facts())
    report = {"targets": [] if target_state is None else [{"target_id": "T1", "rule_id": "G01",
              "state": target_state, "operations": ["G01:0"], "location": "p1"}],
              "plan": {"ledger": []}, "artifact_claim_replay": {}}
    if scope is not None:
        report["scope_mapping"] = {"mappings": {"G01": scope}}
    project(ledger, report, artifact_text="no old text")
    assert ledger.records[0].applicability.status == "TRUE"
    assert ledger.records[0].completion in {"ARTIFACT_VERIFIED", "ALREADY_SATISFIED", "UNRESOLVED"}


def test_primary_etf_misclassification_cannot_skip_regular_fund(tmp_path):
    from fundlex.qualification import certify_exclusion
    from fundlex.recall import facts_from_value
    from fundlex.recall_pipeline import _qualification_status
    from docx import Document
    from fundlex.docx import inspect_docx
    document = Document()
    document.add_paragraph("本基金为普通开放式基金，可以投资ETF。")
    path = tmp_path / "regular.docx"
    document.save(path)
    source = inspect_docx(path, full=True)
    facts = facts_from_value(source, {"fund_type": {"value": "etf", "quote": source.paragraphs[0].text}})
    class Gateway:
        max_context_chars = 10000
        def call(self, **kwargs):
            return SimpleNamespace(value={"fund_type": "open_regular", "quote": source.paragraphs[0].text})
    certified = certify_exclusion(source, facts, Gateway(), ["etf"])
    assert certified.fund_type_certification["status"] == "CONFLICT"
    assert _qualification_status(certified, ["etf"])[1] is False


def test_final_ledger_cannot_borrow_same_operation_edit_from_another_site():
    from fundlex.recall import _reconcile_final_plan_ledger
    ledger = [{"target_id": target, "status": "planned", "operation_keys": ["G05:2"]}
              for target in ("BODY", "SUMMARY")]
    records = [{"rule_id": "G05", "target_id": "BODY", "location": "p10",
                "operation_keys": ["G05:2"], "components": []}]
    _reconcile_final_plan_ledger(ledger, records, [], [], [])
    assert ledger[0]["status"] == "planned"
    assert ledger[1]["status"] == "unresolved"
    assert ledger[1]["reason"] == "RECALL_PLAN_COMPONENT_DROPPED"


def test_final_ledger_requires_every_selected_sibling_component():
    from fundlex.recall import _reconcile_final_plan_ledger
    ledger = [{"target_id": "T", "status": "planned", "operation_keys": ["G05:0", "G05:2"]}]
    _reconcile_final_plan_ledger(ledger, [{"rule_id": "G05", "target_id": "T",
                                          "operation_keys": ["G05:0"]}], [], [], [])
    assert ledger[0]["status"] == "unresolved"
    assert ledger[0]["operation_effects"]["G05:2"]["status"] == "unresolved"


@pytest.mark.parametrize("code,has_target,has_plan,expected", [
    ("CLAIM_OPERATION_NOT_PLANNED", False, False, "TARGET"),
    ("CLAIM_OPERATION_NOT_PLANNED", True, False, "GENERATION"),
    ("CLAIM_PLANNED_EFFECT_MISSING", True, True, "GENERATION"),
    ("CLAIM_EXECUTION_EFFECT_MISSING", True, True, "EXECUTION"),
])
def test_claim_root_cause_routes_the_missing_layer(code, has_target, has_plan, expected):
    from fundlex.recall import Target
    from fundlex.recall_pipeline import _claim_failure_context
    target = Target("T", "G05", "p1", "调低赎回费率", "needs_edit", "",
                    operations=("G05:0",))
    edits = [{"target_id": "T", "location": "p1", "operation_keys": ["G05:0"]}] if has_plan else []
    context = _claim_failure_context({"code": code, "rule_id": "G05", "location": "p1",
                                       "operation_id": "G05:0"}, [target] if has_target else [], edits)
    assert context.resolved_class.value == expected
    assert context.operation_id == "G05:0" and context.site_id == "p1"


def _inspection(tmp_path, *paragraphs):
    from docx import Document
    from fundlex.docx import inspect_docx
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    path = tmp_path / 'source.docx'
    document.save(path)
    return inspect_docx(path, full=True)


def test_audit_missed_target_deduplicates_each_operation_at_the_site(tmp_path):
    from fundlex.recall import Target
    from fundlex.recall_pipeline import _absorb_missed_targets
    source = _inspection(tmp_path, "旧业务甲句。旧业务乙句。")
    location = source.paragraphs[0].location
    rule = {"rule_id": "G99", "operations": [
        {"op": "delete", "old_text": "旧业务甲句。", "condition": {"kind": "none"}},
        {"op": "delete", "old_text": "旧业务乙句。", "condition": {"kind": "none"}}]}
    existing = Target("T", "G99", location, "旧业务甲句。", "needs_edit", "", operations=("G99:0",))
    state = {"missed": {"G99": [{"location": location, "quote": "旧业务乙句。",
                                  "operation_id": "G99:1"}]}}
    synthetic, origins = [], {}
    assert _absorb_missed_targets(state, [existing], synthetic, origins, source,
        rules_by_id={"G99": rule}, facts=Facts()) == 1
    assert synthetic[0].operations == ("G99:1",)
    assert _absorb_missed_targets(state, [existing], synthetic, origins, source,
        rules_by_id={"G99": rule}, facts=Facts()) == 0


def test_sweep_corrects_scope_for_verified_business_clause(tmp_path):
    from fundlex.recall_pipeline import _Chunk, _sweep_batch
    quote = "销售服务费主要用于本基金的市场推广。"
    source = _inspection(tmp_path, quote)
    location = source.paragraphs[0].location
    rule = {"rule_id": "G17", "operations": [{"op": "delete", "old_text": quote}]}
    class Gateway:
        def call(self, **kwargs):
            return SimpleNamespace(value={"findings": [{"rule_id": "G17", "location": location,
                "quote": quote, "issue": "residual_old_text", "reason": "exact residual"}]})
    result = _sweep_batch(Gateway(), {"G17": rule}, "delete old policy",
        [_Chunk("C1", (location,), f"[{location}] {quote}")], source, round_number=0,
        counter=[0], rule_locations={"G17": set()}, scope_by_chunk={"C1": set()})
    assert result["findings"][0]["routing_status"] == "SCOPE_CORRECTION_TARGET"


def test_delete_scan_keeps_duplicate_sites_and_rejects_unrelated_filing(tmp_path):
    from fundlex.generation.claim_replay import check_delete_residuals
    old = "并报中国证监会备案"
    source = _inspection(tmp_path, "调整基金份额类别并报中国证监会备案。",
        "信息披露计划并报中国证监会备案。", "调整基金份额类别并报中国证监会备案。")
    locations = [span.location for span in source.paragraphs]
    rules = {"G05": {"rule_id": "G05", "operations": [{"op": "delete", "old_text": old}]}}
    findings = check_delete_residuals(rules, source, original=source,
        targets=[{"location": locations[0], "operations": ["G05:0"]}])
    assert {item["location"] for item in findings} == {locations[0], locations[2]}
    assert findings[0]["scan"] == "known_site"


def test_same_paragraph_independent_operations_use_one_atomic_edit_without_model(tmp_path):
    from test_recall_core import _facts
    facts = _facts()
    from fundlex.recall import Target, plan_edits, build_docx_edits, content_check_factory
    from fundlex.docx import validate_edit_plan
    from fundlex.generation.claim_replay import replay_claims
    before = "旧业务甲句。旧业务乙句。保留句。"
    after = "新业务甲句。保留句。"
    source = _inspection(tmp_path, before)
    location = source.paragraphs[0].location
    rule = {"rule_id": "G99", "operations": [
        {"op": "replace", "old_text": "旧业务甲句。", "new_text": "新业务甲句。",
         "group_kind": "INDEPENDENT", "condition": {"kind": "none"}},
        {"op": "delete", "old_text": "旧业务乙句。", "group_kind": "INDEPENDENT",
         "condition": {"kind": "none"}}]}
    ledger = compile_obligations({"groups": [rule]})
    apply_facts(ledger, facts)
    targets = [Target(f"T{index}", "G99", location, before, "needs_edit", "",
        start=0, end=len(before), segment_ok=True, resolution="resolved",
        operations=(f"G99:{index}",)) for index in range(2)]
    class Gateway:
        def call(self, **kwargs):
            raise AssertionError("independent deterministic transforms called a model")
    planned = plan_edits({"G99": rule}, targets, source, facts, Gateway(), obligations=ledger)
    assert planned["conflicts"] == []
    assert len(planned["edits"]) == 1
    edit = planned["edits"][0]
    assert edit["content_source"] == "atomic_paragraph" and edit["replacement"] == after
    assert [unit["operation_keys"] for unit in edit["components"]] == [["G99:0"], ["G99:1"]]
    assert all(entry["status"] == "planned" for entry in planned["ledger"])
    validate_edit_plan(tmp_path / "source.docx", build_docx_edits(planned["edits"], source), {},
                       content_check=content_check_factory(planned["edits"], source))
    artifact = _inspection(tmp_path, after)
    replay = replay_claims(planned["edits"], source, artifact, rules_by_id={"G99": rule},
        targets_by_id={target.target_id: target.as_dict() for target in targets},
        obligations=ledger, facts=facts)
    assert replay["verified"] is True, replay["failures"]
    assert all(proof["state"] == "PASS" for proof in replay["obligations"].values())


def test_sweep_cannot_borrow_delete_proof_from_another_site(tmp_path):
    from fundlex.generation.claim_replay import replay_claims
    from fundlex.recall_pipeline import _apply_artifact_proof_to_findings
    rule = {"rule_id": "G99", "operations": [{"op": "delete", "old_text": "旧业务删除句。",
        "group_kind": "INDEPENDENT", "condition": {"kind": "none"}}]}
    ledger = compile_obligations({"groups": [rule]})
    apply_facts(ledger, Facts())
    source = _inspection(tmp_path, "业务条款已满足。", "另一个业务条款。")
    first, second = [span.location for span in source.paragraphs]
    replay = replay_claims([], source, source, rules_by_id={"G99": rule}, facts=Facts(),
        obligations=ledger, targets_by_id={"T": {"target_id": "T", "location": first,
            "quote": "业务条款已满足。", "operations": ["G99:0"]}})
    assert replay["verified"] is True
    finding = {"status": "pending", "rule_id": "G99", "location": second,
        "quote": "另一个业务条款。", "issue": "unaddressed_requirement"}
    _apply_artifact_proof_to_findings([finding], replay, source, ledger, {"G99": rule})
    assert finding["status"] == "pending"
    assert finding["proof_kind"] == "SEMANTIC_UNCERTAIN"


def test_generation_recovery_replans_only_dirty_sites(tmp_path, monkeypatch):
    from copy import deepcopy
    from fundlex import recall
    from test_recall_core import _facts
    facts = _facts()
    text = '旧业务条款。保留句。'
    source = _inspection(tmp_path, text, text)
    rule = {'rule_id': 'G99', 'operations': [{'op': 'replace', 'old_text': '旧业务条款。',
        'new_text': '新业务条款。', 'group_kind': 'INDEPENDENT', 'condition': {'kind': 'none'}}]}
    targets = [recall.Target(f'T{index}', 'G99', span.location, text, 'needs_edit', '',
        start=0, end=len(text), segment_ok=True, resolution='resolved', operations=('G99:0',))
        for index, span in enumerate(source.paragraphs)]
    class Gateway:
        def call(self, **kwargs):
            raise AssertionError('deterministic recovery called a model')
    original = recall.plan_edits({'G99': rule}, targets, source, facts, Gateway())
    retained = deepcopy(original['edits'][1])
    broken = deepcopy(original)
    broken['edits'][0]['replacement'] = '错误候选。'
    called = []
    planner = recall.plan_target
    def traced(rule, target, *args, **kwargs):
        called.append(target.target_id)
        return planner(rule, target, *args, **kwargs)
    monkeypatch.setattr(recall, 'plan_target', traced)
    repaired = recall.replan_sites(broken, {source.paragraphs[0].location}, {'G99': rule},
        targets, source, facts, Gateway(), max_compose_calls=1)
    assert called == ['T0']
    assert retained in repaired['edits']
    assert {edit['location']: edit['replacement'] for edit in repaired['edits']} == {
        span.location: '新业务条款。保留句。' for span in source.paragraphs}
    assert all(entry['status'] == 'planned' for entry in repaired['ledger'])
