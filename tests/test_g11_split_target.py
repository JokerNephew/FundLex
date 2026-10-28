"""D06 G11: the bearer sentence can cross a DOCX paragraph boundary."""
from pathlib import Path

from docx import Document

from fundlex.docx import inspect_docx
from fundlex.recall import g11_composite_source_parts


def test_g11_split_bearer_parts_stop_before_c_class_policy(tmp_path: Path):
    guide = __import__("json").loads(
        (Path(__file__).resolve().parents[1] / "guides/builtin_guide_v1.json").read_text())
    rule = next(item for item in guide["groups"] if item["rule_id"] == "G11")
    document = Document()
    document.add_paragraph("本基金 A 类、Y 类基金份额的申购费用由该类基金份额的投资人承担，")
    document.add_paragraph("不列入基金财产。本基金 C 类基金份额、E 类基金份额和 F 类基金份额不收取申购费。")
    path = tmp_path / "g11-split.docx"
    document.save(path)
    inspection = inspect_docx(path, full=True)

    parts = g11_composite_source_parts(inspection, rule["operations"][0], "p0")
    assert parts == ("p0:0:31", "p1:0:8")
