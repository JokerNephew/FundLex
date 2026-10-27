"""V1 native Compare selection: browser contracts and real authenticated DOCX flow."""
from __future__ import annotations

import hashlib
import io
import json
import os
import socket
import threading
import time
from pathlib import Path
from zipfile import ZipFile

import pytest

pytestmark = pytest.mark.skipif(os.getenv("FUNDLEX_BROWSER_TESTS") != "1", reason="opt-in V1 Chromium suite")


@pytest.fixture(scope="module")
def browser():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as playwright:
        instance = playwright.chromium.launch(headless=True)
        yield instance
        instance.close()


@pytest.fixture
def live_app(tmp_path, monkeypatch):
    import uvicorn
    from fundlex.api import create_app
    monkeypatch.setenv("FUNDLEX_AUTH_MODE", "dev")
    monkeypatch.setenv("FUNDLEX_FRONTEND_VERSION", "v1")
    monkeypatch.setenv("FUNDLEX_DATA_DIR", str(tmp_path / "module-default"))
    app = create_app(tmp_path / "data", auth_db=str(tmp_path / "auth.sqlite3"))
    server = uvicorn.Server(uvicorn.Config(app, log_level="critical"))
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    base = f"http://127.0.0.1:{sock.getsockname()[1]}"
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10
    while not server.started and thread.is_alive() and time.monotonic() < deadline:
        time.sleep(.02)
    assert server.started
    yield app, base
    server.should_exit = True
    thread.join(timeout=10)
    sock.close()


@pytest.fixture
def password_app(tmp_path, monkeypatch):
    import uvicorn
    import secrets
    monkeypatch.setenv("FUNDLEX_AUTH_MODE", "password")
    monkeypatch.setenv("FUNDLEX_FRONTEND_VERSION", "v1")
    monkeypatch.setenv("FUNDLEX_DATA_DIR", str(tmp_path / "module-default"))
    from fundlex.api import create_app
    app = create_app(tmp_path / "data", auth_db=str(tmp_path / "auth.sqlite3"))
    password = secrets.token_urlsafe(24)
    app.state.users.create_user("compare-reader", password, "Compare Reader")
    app.state.users.create_user("compare-other", password, "Other Reader")
    server = uvicorn.Server(uvicorn.Config(app, log_level="critical"))
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    base = f"http://127.0.0.1:{sock.getsockname()[1]}"
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10
    while not server.started and thread.is_alive() and time.monotonic() < deadline:
        time.sleep(.02)
    assert server.started
    yield app, base, password
    server.should_exit = True
    thread.join(timeout=10)
    sock.close()


def mock_compare(page, *, structural=False, diff_count=2):
    pair = {"id": "pair-one", "label": "第一组", "ordinal": 0, "status": "succeeded",
            "left": {"name": "left.docx"}, "right": {"name": "right.docx"},
            "summary": {"total": 2, "modify": 2, "insert": 0, "delete": 0}}
    task = {"id": "selection-task", "name": "V1 人工选择", "status": "completed", "total": 2,
            "pairs": [pair, {**pair, "id": "pair-two", "label": "第二组", "ordinal": 1}]}
    diffs = [{"id": f"D{i}", "type": "modify", "section": "费率",
              "left": {"text": "旧文"}, "right": {"text": "新文"}} for i in range(1, diff_count + 1)]
    if structural:
        diffs[0]["type"] = "insert"
        diffs[1]["type"] = "delete"
    states = {p["id"]: {"selection_epoch": 0, "decisions": {d["id"]: "pending" for d in diffs}}
              for p in task["pairs"]}
    controls = {"conflict": False, "generations": [], "downloads": [], "held": None, "hold": False}

    def payload(pair_id):
        selection = states[pair_id]
        pending = [key for key, choice in selection["decisions"].items() if choice == "pending"]
        return {"pair": {"id": pair_id}, "result_sha256": "a" * 64, "left_source_sha256": "b" * 64,
                "right_source_sha256": "c" * 64, "selection": selection, "diffs": diffs,
                "assessment": {"pending_ids": pending, "issues": [], "can_generate": not pending}, "artifacts": []}

    def handle(route):
        request = route.request
        path = request.url.split("/api", 1)[-1].split("?", 1)[0]
        if path == "/compare/tasks":
            return route.fulfill(json={"items": [task]})
        if path == "/compare/tasks/selection-task":
            return route.fulfill(json=task)
        if not path.startswith("/compare/pairs/"):
            return route.continue_()
        pair_id = path.split("/")[3]
        if path.endswith("/view"):
            html = "<p>" + "</p><p>".join(f"合同条款 {i}" for i in range(60)) + "</p>"
            html += "<p><span data-diff-id='D1'>旧文</span><span data-diff-id='D2'>新文</span></p>"
            return route.fulfill(json={"pair": {"id": pair_id, "label": pair_id}, "left": {"html": html},
                                       "right": {"html": html}, "summary": pair["summary"],
                                       "coverage": {"status": "complete"}, "diffs": diffs})
        if path.endswith("/selection"):
            if request.method == "PUT":
                body = request.post_data_json
                if controls["conflict"]:
                    controls["conflict"] = False
                    states[pair_id]["selection_epoch"] += 1
                    return route.fulfill(status=409, json={"detail": "SELECTION_EPOCH_CONFLICT"})
                assert body["expected_epoch"] == states[pair_id]["selection_epoch"]
                states[pair_id] = {"selection_epoch": body["expected_epoch"] + 1, "decisions": body["decisions"]}
            return route.fulfill(json=payload(pair_id))
        if path.endswith("/preview"):
            return route.fulfill(json=payload(pair_id))
        if path.endswith("/generate"):
            controls["generations"].append(request.headers["idempotency-key"])
            result = {"version": 1, "artifacts": [{"kind": "selected_final", "status": "current", "version": 1,
                       "sha256": "d" * 64, "download_url": f"/api/compare/pairs/{pair_id}/selection-artifacts/1/selected_final"}]}
            if controls["hold"]:
                controls["held"] = (route, result)
                return
            return route.fulfill(json=result)
        if "/selection-artifacts/" in path:
            controls["downloads"].append(request.headers)
            return route.fulfill(body=b"browser-download-fixture", headers={
                "Content-Disposition": 'attachment; filename="selected.docx"'})
        return route.continue_()

    page.route("**/api/**", handle)
    return controls, states


def open_compare(page, base):
    page.goto(base + "/#compare", wait_until="networkidle")
    page.locator('[data-compare-task="selection-task"]').click()
    page.locator('[data-selection-choice]').first.wait_for(state="visible")
    page.wait_for_function("!document.querySelector('[data-selection-choice]').disabled")


@pytest.mark.parametrize("width", [390, 768, 1440])
def test_v1_pending_choices_refresh_generation_download_and_mobile(browser, live_app, width):
    from playwright.sync_api import expect
    _, base = live_app
    context = browser.new_context(viewport={"width": width, "height": 960}, accept_downloads=True)
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v1-selection-check')")
    page = context.new_page()
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    controls, states = mock_compare(page)
    try:
        open_compare(page, base)
        expect(page.locator('h1:visible')).to_have_count(1)
        expect(page.locator('#compare-task-name')).to_have_text('V1 人工选择')
        expect(page.locator('#breadcrumb [aria-current="page"]')).to_have_text('任务详情')
        expect(page.locator('.topbar h1')).to_have_count(0)
        expect(page.locator('#compare-selection-generate')).to_be_disabled()
        expect(page.locator('#compare-selection-progress')).to_contain_text('待选择 2')
        page.locator('#compare-selection-preview').click()
        expect(page.locator('#compare-selection-status')).to_contain_text('仍有 2 处待选择')
        page.locator('input[data-diff-id="D1"][value="left"]').check()
        expect(page.locator('#compare-selection-status')).to_have_text('已保存')
        page.locator('input[data-diff-id="D2"][value="right"]').check()
        expect(page.locator('#compare-selection-generate')).to_be_enabled()
        assert states['pair-one']['decisions'] == {"D1": "left", "D2": "right"}
        page.reload(wait_until="networkidle")
        expect(page.locator('input[data-diff-id="D1"][value="left"]')).to_be_checked()
        expect(page.locator('input[data-diff-id="D2"][value="right"]')).to_be_checked()
        page.locator('#compare-selection-generate').click()
        expect(page.locator('#compare-selection-status')).to_contain_text('已生成 v1')
        with page.expect_download() as pending_download:
            page.locator('[data-selection-download]').click()
        assert pending_download.value.suggested_filename == 'selected.docx'
        assert controls['downloads'][0]['x-fundlex-user'] == 'v1-selection-check'
        page.locator('#compare-selection-generate').click()
        expect(page.locator('#compare-selection-status')).to_contain_text('已生成 v1')
        assert len(controls['generations']) == 2
        assert len(set(controls['generations'])) == 1
        page.locator('#compare-selection-left-all').click()
        expect(page.locator('#compare-selection-status')).to_have_text('已保存')
        assert page.locator('[data-selection-download]').count() == 0
        assert states['pair-one']['decisions'] == {"D1": "left", "D2": "left"}
        if width == 390:
            page.locator('[data-compare-side="right"]').click()
            assert page.locator('#compare-viewer').get_attribute('data-mobile-side') == 'right'
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        assert not errors
    finally:
        context.close()


@pytest.mark.parametrize('width', [390, 768, 1280, 1440])
def test_v1_review_rail_geometry_independent_scroll_and_form(browser, live_app, width):
    from playwright.sync_api import expect
    _, base = live_app
    context = browser.new_context(viewport={'width': width, 'height': 960})
    page = context.new_page()
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    mock_compare(page, diff_count=60)
    try:
        open_compare(page, base)
        geometry = page.evaluate('''() => {
          const doc=document.querySelector('.compare-document-area').getBoundingClientRect();
          const rail=document.querySelector('.compare-diff-rail').getBoundingClientRect();
          const list=document.querySelector('#compare-diff-list');
          return {doc:{x:doc.x,right:doc.right,top:doc.top},rail:{x:rail.x,bottom:rail.bottom,top:rail.top},
                  scroll:list.scrollHeight>list.clientHeight,overflow:getComputedStyle(list).overflowY};
        }''')
        if width > 760:
            assert geometry['rail']['x'] >= geometry['doc']['right']
            assert abs(geometry['rail']['top'] - geometry['doc']['top']) < 2
        else:
            assert geometry['rail']['bottom'] <= geometry['doc']['top']
        assert geometry['scroll'] and geometry['overflow'] == 'auto'
        page.locator('#compare-next-diff').scroll_into_view_if_needed()
        if os.getenv('FUNDLEX_UX_EVIDENCE'):
            out = Path(os.environ['FUNDLEX_UX_EVIDENCE'])
            out.mkdir(parents=True, exist_ok=True)
            page.screenshot(path=str(out / f'local-reader-{width}.png'))
        before = page.locator('.main').evaluate('(n)=>n.scrollTop')
        for _ in range(8):
            page.locator('#compare-next-diff').click()
        assert page.locator('#compare-diff-list').evaluate('(n)=>n.scrollTop') > 0
        assert page.locator('.main').evaluate('(n)=>n.scrollTop') == before
        page.locator('.sidebar [data-nav="compare"]').click()
        page.locator('#compare-new').click()
        page.locator('#compare-name').fill('统一表单测试')
        assert page.locator('#compare-name').evaluate('(n)=>n.clientHeight') >= 40
        with page.expect_file_chooser() as chooser:
            page.locator('[data-compare-choose="left"]').click()
        chooser.value.set_files({'name': '左侧真实名称.docx', 'mimeType': 'application/octet-stream', 'buffer': b'docx'})
        expect(page.locator('[data-file-side="left"] .compare-file-name')).to_have_text('左侧真实名称.docx')
        expect(page.locator('[data-compare-choose="left"]')).to_have_text('更换文件')
        if os.getenv('FUNDLEX_UX_EVIDENCE'):
            page.screenshot(path=str(Path(os.environ['FUNDLEX_UX_EVIDENCE']) / f'local-form-{width}.png'))
        page.locator('#compare-add-pair').click()
        expect(page.locator('.compare-pair-input')).to_have_count(2)
        page.locator('[data-compare-remove="1"]').click()
        expect(page.locator('[data-file-side="left"] .compare-file-name')).to_have_text('左侧真实名称.docx')
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        assert not errors
    finally:
        context.close()


@pytest.mark.parametrize('width', [390, 1440])
def test_v1_feishu_page_states_expiry_confirmation_and_reset(browser, live_app, width):
    from playwright.sync_api import expect
    _, base = live_app
    context = browser.new_context(viewport={'width': width, 'height': 960})
    page = context.new_page()
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    state = {'bound': False, 'deletes': 0, 'starts': 0}
    png = 'data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='

    def binding(route):
        if route.request.method == 'DELETE':
            state['deletes'] += 1
            state['bound'] = False
        route.fulfill(json={'bound': state['bound']})

    def start(route):
        state['starts'] += 1
        route.fulfill(json={'authorization_url': 'https://accounts.feishu.cn/authorize',
                            'qr_png': png, 'expires_at': page.evaluate('Date.now()/1000+600')})

    page.route('**/api/integrations/feishu/binding', binding)
    page.route('**/api/integrations/feishu/oauth/start', start)
    try:
        page.goto(base + '/#settings', wait_until='networkidle')
        expect(page.locator('#breadcrumb [aria-current="page"]')).to_have_text('飞书绑定')
        expect(page.locator('#feishu-binding-badge')).to_have_text('未绑定')
        assert not page.locator('#feishu-binding-code-btn').count()
        assert '实际运行能力' not in page.locator('#view-settings').inner_text()
        assert '本地开发身份' not in page.locator('#view-settings').inner_text()
        page.locator('#feishu-oauth-btn').click()
        expect(page.locator('#feishu-oauth-qr img')).to_be_visible()
        expect(page.locator('#feishu-oauth-link')).to_have_attribute('href', 'https://accounts.feishu.cn/authorize')
        state['bound'] = True
        page.locator('#feishu-binding-refresh').click()
        expect(page.locator('#feishu-binding-badge')).to_have_text('已绑定')
        expect(page.locator('#feishu-oauth-panel')).to_be_hidden()
        page.locator('#feishu-unbind-btn').click()
        expect(page.locator('#feishu-unbind-confirm')).to_be_visible()
        assert state['deletes'] == 0
        page.locator('[data-action="feishu-unbind-cancel"]').click()
        expect(page.locator('#feishu-unbind-confirm')).to_be_hidden()
        page.locator('#feishu-unbind-btn').click()
        page.locator('[data-action="feishu-unbind-confirm"]').click()
        expect(page.locator('#feishu-binding-badge')).to_have_text('未绑定')
        assert state['deletes'] == 1
        page.clock.install()
        page.locator('#feishu-oauth-btn').click()
        expect(page.locator('#feishu-oauth-qr img')).to_be_visible()
        page.clock.fast_forward(601_000)
        expect(page.locator('#feishu-oauth-message')).to_contain_text('已过期')
        assert page.locator('#feishu-oauth-qr img').count() == 0
        expect(page.locator('#feishu-oauth-link')).to_be_hidden()
        page.locator('#feishu-oauth-btn').click()
        expect(page.locator('#feishu-oauth-qr img')).to_be_visible()
        page.locator('.sidebar [data-nav="compare"]').click()
        assert page.locator('#feishu-oauth-qr img').count() == 0
        page.locator('[data-nav="settings"]').click()
        expect(page.locator('#feishu-binding-badge')).to_have_text('未绑定')
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        assert not errors
    finally:
        context.close()


def test_v1_conflict_retains_choices_and_late_generation_cannot_cross_pairs(browser, live_app):
    from playwright.sync_api import expect
    _, base = live_app
    context = browser.new_context(viewport={"width": 1440, "height": 960})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v1-selection-check')")
    page = context.new_page()
    controls, states = mock_compare(page)
    try:
        open_compare(page, base)
        controls['conflict'] = True
        page.locator('#compare-selection-right-all').click()
        expect(page.locator('#compare-selection-status')).to_contain_text('本地选择已保留')
        expect(page.locator('input[data-diff-id="D1"][value="right"]')).to_be_checked()
        expect(page.locator('#compare-selection-generate')).to_be_disabled()
        page.locator('#compare-selection-save').click()
        expect(page.locator('#compare-selection-status')).to_have_text('已保存')
        assert states['pair-one']['decisions'] == {"D1": "right", "D2": "right"}
        controls['hold'] = True
        page.locator('#compare-selection-generate').click()
        expect(page.locator('#compare-selection-left-all')).to_be_disabled()
        page.locator('#compare-next-pair').click()
        expect(page.locator('#compare-pair-title')).to_have_text('pair-two')
        route, result = controls['held']
        route.fulfill(json=result)
        expect(page.locator('#compare-selection-progress')).to_contain_text('待选择 2')
        assert page.locator('[data-selection-download]').count() == 0
        assert '已生成' not in page.locator('#compare-selection-status').inner_text()
    finally:
        context.close()


def test_v1_feishu_errors_and_late_qr_cannot_survive_navigation(browser, live_app):
    from playwright.sync_api import expect
    _, base = live_app
    context = browser.new_context()
    page = context.new_page()
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    controls = {'fail': True, 'held': None}

    def binding(route):
        if controls['fail']:
            return route.fulfill(status=503, json={'detail': 'temporarily unavailable'})
        return route.fulfill(json={'bound': False})

    def start(route):
        controls['held'] = route

    page.route('**/api/integrations/feishu/binding', binding)
    page.route('**/api/integrations/feishu/oauth/start', start)
    try:
        page.goto(base + '/#settings', wait_until='networkidle')
        expect(page.locator('#feishu-binding-badge')).to_have_text('读取失败')
        expect(page.locator('#feishu-oauth-btn')).to_be_disabled()
        controls['fail'] = False
        page.locator('#feishu-binding-refresh').click()
        expect(page.locator('#feishu-oauth-btn')).to_be_enabled()
        page.locator('#feishu-oauth-btn').click()
        expect(page.locator('#feishu-oauth-btn')).to_be_disabled()
        page.locator('.sidebar [data-nav="compare"]').click()
        assert controls['held']
        controls['held'].fulfill(json={'authorization_url': 'https://accounts.feishu.cn/authorize',
                                      'qr_png': 'data:image/png;base64,late', 'expires_at': time.time() + 600})
        page.locator('[data-nav="settings"]').click()
        expect(page.locator('#feishu-binding-badge')).to_have_text('未绑定')
        expect(page.locator('#feishu-oauth-panel')).to_be_hidden()
        assert not page.locator('#feishu-oauth-qr img').count()
        assert not errors
    finally:
        context.close()


def test_v1_insert_delete_choices_remain_explicit_but_generation_is_blocked(browser, live_app):
    from playwright.sync_api import expect
    _, base = live_app
    context = browser.new_context()
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v1-selection-check')")
    page = context.new_page()
    controls, states = mock_compare(page, structural=True)
    try:
        open_compare(page, base)
        page.locator('#compare-selection-right-all').click()
        expect(page.locator('#compare-selection-status')).to_have_text('已保存')
        assert states['pair-one']['decisions'] == {"D1": "right", "D2": "right"}
        expect(page.locator('#compare-selection-generate')).to_be_disabled()
        expect(page.locator('#compare-selection-issues')).to_contain_text('暂不支持')
        assert not controls['generations']
    finally:
        context.close()


def test_v1_password_login_upload_worker_selection_and_verified_downloads(browser, password_app, tmp_path):
    from docx import Document
    from lxml import etree
    from playwright.sync_api import expect
    from fundlex.compare.worker import CompareWorker

    app, base, password = password_app
    paths = []
    for side, rate in (("left", "0.40%"), ("right", "0.30%")):
        document = Document()
        document.add_paragraph(f"基金管理费按前一日基金资产净值的{rate}年费率计提。")
        document.add_paragraph("基金托管人与基金管理人应当履行各自职责。")
        path = tmp_path / f"{side}.docx"
        document.save(path)
        paths.append(path)
    source_hashes = [hashlib.sha256(path.read_bytes()).hexdigest() for path in paths]
    context = browser.new_context(viewport={"width": 1440, "height": 960}, accept_downloads=True)
    page = context.new_page()
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    try:
        page.goto(base + "/#compare", wait_until="domcontentloaded")
        page.locator('#login-view').wait_for(state='visible')
        page.locator('#login-username').fill('compare-reader')
        page.locator('#login-password').fill(password)
        page.locator('#login-submit').click()
        expect(page.locator('#login-view')).to_be_hidden()
        page.locator('#compare-new').click()
        page.locator('#compare-name').fill('真实 Worker 人工选择浏览器回归')
        page.locator('[data-compare-left]').set_input_files(str(paths[0]))
        page.locator('[data-compare-right]').set_input_files(str(paths[1]))
        with page.expect_response(lambda response: '/api/compare/tasks/' in response.url and response.url.endswith('/seal')) as sealed:
            page.locator('#compare-submit').click()
        assert sealed.value.ok
        assert CompareWorker(app.state.compare_store).run_once()
        page.reload(wait_until="networkidle")
        page.wait_for_function("!!document.querySelector('[data-selection-choice]') && !document.querySelector('[data-selection-choice]').disabled")
        expect(page.locator('#compare-selection-generate')).to_be_disabled()
        page.locator('[data-selection-choice="right"]').check()
        expect(page.locator('#compare-selection-generate')).to_be_enabled()
        page.reload(wait_until="networkidle")
        expect(page.locator('[data-selection-choice="right"]')).to_be_checked()
        page.locator('#compare-selection-generate').click()
        expect(page.locator('#compare-selection-status')).to_contain_text('已生成 v1')
        downloaded = {}
        for kind in ('selected_final', 'selected_redline', 'selection_report_json', 'selection_report_markdown'):
            button = page.locator(f'[data-selection-download$="/{kind}"]')
            with page.expect_download() as pending_download:
                button.click()
            download = pending_download.value
            saved = tmp_path / download.suggested_filename
            download.save_as(saved)
            downloaded[kind] = saved.read_bytes()
        report = json.loads(downloaded['selection_report_json'])
        for kind in ('selected_final', 'selected_redline'):
            assert hashlib.sha256(downloaded[kind]).hexdigest() == report['outputs'][kind]['sha256']
        assert [hashlib.sha256(path.read_bytes()).hexdigest() for path in paths] == source_hashes
        assert source_hashes == [report['left_source_sha256'], report['right_source_sha256']]
        final = Document(io.BytesIO(downloaded['selected_final']))
        assert [p.text for p in final.paragraphs] == [
            '基金管理费按前一日基金资产净值的0.30%年费率计提。',
            '基金托管人与基金管理人应当履行各自职责。',
        ]
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with ZipFile(io.BytesIO(downloaded['selected_redline'])) as archive:
            root = etree.fromstring(archive.read('word/document.xml'))
        assert root.xpath('//w:ins', namespaces=ns) and root.xpath('//w:del', namespaces=ns)
        for deletion in root.xpath('//w:del', namespaces=ns):
            deletion.getparent().remove(deletion)
        accepted = [''.join(p.xpath('.//w:t/text()', namespaces=ns)) for p in root.xpath('//w:body/w:p', namespaces=ns)]
        assert accepted == [p.text for p in final.paragraphs]
        assert b'0.30%' in downloaded['selection_report_markdown']
        final_url = page.locator('[data-selection-download$="/selected_final"]').get_attribute('data-selection-download')
        assert context.request.get(base + final_url).status == 401
        login = context.request.post(base + '/api/auth/login', data=json.dumps({'username': 'compare-other', 'password': password}),
                                     headers={'Content-Type': 'application/json'})
        token = login.json()['token']
        assert context.request.get(base + final_url, headers={'Authorization': f'Bearer {token}'}).status == 404
        assert not errors
    finally:
        context.close()
