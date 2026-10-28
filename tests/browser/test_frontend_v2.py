"""Browser smoke coverage for the React/Vite workbench entry.

This suite checks the shipped React shell and deep-link refresh behavior. It
does not start the Worker or call a model; the real document delivery path is
covered by the separate live end-to-end evidence.
"""
from __future__ import annotations

import io
import json
import os
import re
import socket
import threading
import time

import pytest


pytestmark = pytest.mark.skipif(
    os.getenv("FUNDLEX_BROWSER_TESTS") != "1",
    reason="opt-in React V2 Chromium smoke suite",
)


@pytest.fixture(scope="module")
def browser():
    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        instance = playwright.chromium.launch(
            headless=os.getenv("FUNDLEX_BROWSER_HEADLESS", "1") != "0")
        yield instance
        instance.close()


@pytest.fixture
def live_app(tmp_path, monkeypatch):
    import uvicorn
    from fundlex.api import create_app

    monkeypatch.setenv("FUNDLEX_AUTH_MODE", "dev")
    monkeypatch.setenv("FUNDLEX_FRONTEND_VERSION", "v2")
    monkeypatch.setenv("FUNDLEX_DATA_DIR", str(tmp_path / "module-default"))
    app = create_app(tmp_path / "data", auth_db=str(tmp_path / "auth.sqlite3"))
    server = uvicorn.Server(uvicorn.Config(app, log_level="critical"))
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    base = f"http://127.0.0.1:{sock.getsockname()[1]}"
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10
    while not server.started and thread.is_alive() and time.monotonic() < deadline:
        time.sleep(0.02)
    assert server.started
    yield base
    server.should_exit = True
    thread.join(timeout=10)
    sock.close()


@pytest.fixture
def prefixed_live_app(tmp_path, monkeypatch):
    """Run the React shell behind the same ``/test/`` prefix as isolation."""
    import uvicorn
    from fundlex.api import create_app

    monkeypatch.setenv("FUNDLEX_AUTH_MODE", "dev")
    monkeypatch.setenv("FUNDLEX_FRONTEND_VERSION", "v2")
    monkeypatch.setenv("FUNDLEX_FRONTEND_BASE_PATH", "/test")
    monkeypatch.setenv("FUNDLEX_DATA_DIR", str(tmp_path / "module-default"))
    app = create_app(tmp_path / "data", auth_db=str(tmp_path / "auth.sqlite3"))
    server = uvicorn.Server(uvicorn.Config(app, log_level="critical"))
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    base = f"http://127.0.0.1:{sock.getsockname()[1]}/test"
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
    thread.start()
    deadline = time.monotonic() + 10
    while not server.started and thread.is_alive() and time.monotonic() < deadline:
        time.sleep(0.02)
    assert server.started
    yield base
    server.should_exit = True
    thread.join(timeout=10)
    sock.close()


def test_react_routes_and_deep_links(browser, live_app):
    context = browser.new_context(viewport={"width": 1440, "height": 900})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v2-browser-check')")
    page = context.new_page()
    page_errors: list[str] = []
    console_errors: list[str] = []
    api_404s: list[str] = []
    missing_resources: list[str] = []
    page.on("pageerror", lambda error: page_errors.append(str(error)))
    page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
    def collect_404(response):
        if response.status != 404:
            return
        if "/api/" in response.url:
            api_404s.append(response.url)
        else:
            missing_resources.append(response.url)

    page.on("response", collect_404)
    try:
        for path, heading in (("/tasks", "任务"), ("/tasks/new", "新建任务"), ("/guides", "指引配置"), ("/settings", "设置"), ("/compare", "合同对比")):
            response = page.goto(live_app + path, wait_until="networkidle")
            assert response and response.ok
            page.locator("#root").wait_for(state="visible")
            assert heading in page.locator("body").inner_text()
        response = page.goto(live_app + "/tasks/nonexistent/documents/doc/review", wait_until="networkidle")
        assert response and response.ok
        page.locator("#root").wait_for(state="visible")
        assert not page_errors
        unexpected_console_errors = [
            message for message in console_errors
            if message != "Failed to load resource: the server responded with a status of 404 (Not Found)"
        ]
        assert not unexpected_console_errors, f"console errors: {unexpected_console_errors}; API 404s: {api_404s}; resources: {missing_resources}"
        assert not missing_resources, f"404 resources: {missing_resources}; expected API 404s: {api_404s}"
    finally:
        context.close()


def test_built_entry_and_assets_are_cacheable(browser, live_app):
    context = browser.new_context()
    try:
        response = context.request.get(live_app + "/")
        assert response.ok
        assert "no-store" in response.headers.get("cache-control", "")
        html = response.text()
        assert "/assets/" in html
        assets = re.findall(r'(?:src|href)="(/assets/[^"]+)"', html)
        assert assets
        for asset in assets:
            asset_response = context.request.get(live_app + asset)
            assert asset_response.ok
            assert "immutable" in asset_response.headers.get("cache-control", "")
    finally:
        context.close()


def test_prefixed_entry_routes_assets_api_and_logo_same_origin(browser, prefixed_live_app):
    context = browser.new_context(viewport={"width": 390, "height": 844})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v2-prefix-browser-check');")
    page = context.new_page()
    page_errors: list[str] = []
    console_errors: list[str] = []
    failed_resources: list[str] = []
    page.on("pageerror", lambda error: page_errors.append(str(error)))
    page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
    page.on("requestfailed", lambda request: failed_resources.append(f"{request.url}: {request.failure}"))
    try:
        response = page.goto(prefixed_live_app + "/tasks/new", wait_until="networkidle")
        assert response and response.ok
        assert "新建任务" in page.locator("body").inner_text()
        logo_src = page.locator("img[alt='广发基金']").get_attribute("src")
        assert logo_src and logo_src.startswith("/test/static/")
        assert all(url.startswith(prefixed_live_app + "/") for url in page.evaluate("performance.getEntriesByType('resource').map(entry => entry.name)"))
        metrics = page.evaluate("""() => ({ viewport: innerWidth, documentWidth: document.documentElement.scrollWidth, bodyWidth: document.body.scrollWidth })""")
        assert metrics["documentWidth"] <= metrics["viewport"]
        assert metrics["bodyWidth"] <= metrics["viewport"]
        assert not page_errors, page_errors
        assert not console_errors, console_errors
        assert not failed_resources, failed_resources
    finally:
        context.close()


def test_real_task_document_deep_links_use_task_scoped_api(browser, live_app):
    from docx import Document

    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("React V2 浏览器深链测试合同。")
    document.save(stream)
    data = stream.getvalue()
    identity = "v2-document-check"
    context = browser.new_context(viewport={"width": 1440, "height": 900})
    context.add_init_script(f"localStorage.setItem('fundlex-dev-user', {json.dumps(identity)});")
    page = context.new_page()
    requests: list[str] = []
    page_errors: list[str] = []
    page.on("request", lambda request: requests.append(request.url))
    page.on("pageerror", lambda error: page_errors.append(str(error)))
    try:
        headers = {"X-FundLex-User": identity}
        created = context.request.post(
            live_app + "/api/tasks",
            headers=headers,
            data={
                "name": "React V2 深链任务",
                "engine_id": "recall_first",
                "files": [{"client_id": "browser-doc-1", "name": "深链合同.docx", "size": len(data)}],
            },
        )
        assert created.ok, created.text()
        task_id = created.json()["id"]
        uploaded = context.request.post(
            live_app + f"/api/tasks/{task_id}/uploads",
            headers=headers,
            multipart={
                "client_ids": json.dumps(["browser-doc-1"]),
                "files": {
                    "name": "深链合同.docx",
                    "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "buffer": data,
                },
            },
        )
        assert uploaded.ok, uploaded.text()
        documents = context.request.get(live_app + f"/api/tasks/{task_id}/documents", headers=headers)
        assert documents.ok, documents.text()
        document_id = documents.json()["items"][0]["document_id"]

        response = page.goto(live_app + f"/tasks/{task_id}/documents/{document_id}/detail", wait_until="networkidle")
        assert response and response.ok
        assert "深链合同.docx" in page.locator("body").inner_text()
        assert "报告读取失败" in page.locator("body").inner_text()

        response = page.goto(live_app + f"/tasks/{task_id}/documents/{document_id}/review", wait_until="networkidle")
        assert response and response.ok
        assert "修订呈现读取失败" in page.locator("body").inner_text()
        assert any(f"/api/tasks/{task_id}/documents/{document_id}" in url for url in requests)
        assert not page_errors, page_errors
    finally:
        context.close()


def test_review_renders_void_document_tags_without_page_error(browser, live_app):
    """A real revision payload may contain HTML void tags such as ``<hr>``."""
    task_id = "review-task"
    document_id = "review-document"
    context = browser.new_context(viewport={"width": 390, "height": 844})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','review-browser-check');")
    page = context.new_page()
    page_errors: list[str] = []
    page.on("pageerror", lambda error: page_errors.append(str(error)))

    def json_route(route, body):
        route.fulfill(status=200, content_type="application/json", body=json.dumps(body))

    document = {
        "task_id": task_id,
        "document_id": document_id,
        "name": "Review void tag.docx",
        "job_status": "succeeded",
        "delivery_status": "complete",
        "downloads": {"redline": f"/api/tasks/{task_id}/documents/{document_id}/download/redline"},
    }
    revisions = {
        "document_id": document_id,
        "html": "<p>原文</p><hr><p><ins data-rev='2'>新文</ins></p>",
        "revisions": [{"del_id": 1, "ins_id": 2, "old_text": "原文", "new_text": "新文"}],
    }

    def handle(route):
        request = route.request
        path = request.url.split("/api", 1)[-1].split("?", 1)[0]
        if path == f"/tasks/{task_id}/documents/{document_id}":
            return json_route(route, document)
        if path == f"/tasks/{task_id}/documents":
            return json_route(route, {"items": [document], "total": 1, "page": 1, "page_size": 200})
        if path == f"/tasks/{task_id}/documents/{document_id}/revisions":
            return json_route(route, revisions)
        return route.continue_()

    page.route("**/api/**", handle)
    try:
        response = page.goto(live_app + f"/tasks/{task_id}/documents/{document_id}/review", wait_until="networkidle")
        assert response and response.ok
        assert "Review void tag.docx" in page.locator("body").inner_text()
        metrics = page.evaluate("""() => ({ viewport: innerWidth, documentWidth: document.documentElement.scrollWidth, bodyWidth: document.body.scrollWidth })""")
        assert metrics["documentWidth"] <= metrics["viewport"]
        assert metrics["bodyWidth"] <= metrics["viewport"]
        left_tab = page.get_by_role("tab", name="合同呈现")
        right_tab = page.get_by_role("tab", name="修订清单")
        assert left_tab.get_attribute("aria-selected") == "true"
        right_tab.click()
        page.get_by_role("heading", name="修订清单").wait_for(state="visible")
        assert right_tab.get_attribute("aria-selected") == "true"
        left_tab.click()
        assert left_tab.get_attribute("aria-selected") == "true"
        assert not page_errors, page_errors
    finally:
        context.close()


def test_compare_result_view_supports_pair_navigation_and_mobile_switching(browser, live_app):
    """The React Compare route keeps result navigation usable at phone width."""
    task_id = "compare-browser-task"
    pair_one = "compare-browser-pair-1"
    pair_two = "compare-browser-pair-2"
    task = {
        "id": task_id,
        "name": "React Compare 结果测试",
        "status": "completed",
        "total": 2,
        "pairs": [
            {"id": pair_one, "label": "第一组", "ordinal": 0, "status": "succeeded", "left": {"name": "a-1.docx"}, "right": {"name": "b-1.docx"}, "summary": {"total": 1, "modify": 1, "insert": 0, "delete": 0}},
            {"id": pair_two, "label": "第二组", "ordinal": 1, "status": "succeeded", "left": {"name": "a-2.docx"}, "right": {"name": "b-2.docx"}, "summary": {"total": 1, "modify": 0, "insert": 1, "delete": 0}},
        ],
    }
    views = {
        pair_one: {"pair": {"id": pair_one, "label": "第一组", "status": "succeeded"}, "left": {"name": "a-1.docx", "html": "<p>" + "</p><p>".join(f"合同段落 {index}" for index in range(80)) + "</p><p>旧文 <span class='compare-diff-modify' data-diff-id='D1'>A</span></p>"}, "right": {"name": "b-1.docx", "html": "<p>" + "</p><p>".join(f"合同段落 {index}" for index in range(80)) + "</p><p>新文 <span class='compare-diff-modify' data-diff-id='D1'>B</span> <span class='compare-diff-insert' data-diff-id='D3'>新增</span></p>"}, "summary": {"total": 2, "modify": 1, "insert": 1, "delete": 0}, "coverage": {"status": "complete"}, "diffs": [{"id": "D1", "type": "modify", "section": "正文", "left": {"block_id": "p1", "start": 0, "end": 1, "text": "A"}, "right": {"block_id": "p1", "start": 0, "end": 1, "text": "B"}}, {"id": "D3", "type": "insert", "section": "正文", "right": {"block_id": "p1", "start": 2, "end": 4, "text": "新增"}}]},
        pair_two: {"pair": {"id": pair_two, "label": "第二组", "status": "succeeded"}, "left": {"name": "a-2.docx", "html": "<p>原文 <span class='compare-diff-insert' data-diff-id='D2'>新增</span></p>"}, "right": {"name": "b-2.docx", "html": "<p>修订 <span class='compare-diff-insert' data-diff-id='D2'>新增</span></p>"}, "summary": {"total": 1, "modify": 0, "insert": 1, "delete": 0}, "coverage": {"status": "complete"}, "diffs": [{"id": "D2", "type": "insert", "section": "正文", "left": {"block_id": "p1", "start": 0, "end": 1, "text": ""}, "right": {"block_id": "p1", "start": 0, "end": 1, "text": "新增"}}]},
    }
    context = browser.new_context(viewport={"width": 1440, "height": 900})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','compare-browser-check');")
    page = context.new_page()
    page_errors: list[str] = []
    console_errors: list[str] = []
    page.on("pageerror", lambda error: page_errors.append(str(error)))
    page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)

    def json_route(route, body):
        route.fulfill(status=200, content_type="application/json", body=json.dumps(body))

    def handle(route):
        path = route.request.url.split("/api", 1)[-1].split("?", 1)[0]
        if path == "/compare/tasks":
            return json_route(route, {"items": [task], "total": 1})
        if path == f"/compare/tasks/{task_id}":
            return json_route(route, task)
        if path.startswith("/compare/pairs/") and path.endswith("/view"):
            return json_route(route, views[path.split("/")[3]])
        return route.continue_()

    page.route("**/api/**", handle)
    try:
        response = page.goto(live_app + "/compare", wait_until="networkidle")
        assert response and response.ok
        page.get_by_role("heading", name="合同对比").first.wait_for(state="visible")
        page.get_by_text("共 2 处差异 · 修改 1 · 新增 1 · 删除 0").wait_for(state="visible")
        assert page.locator('.compare-splitter').count() == 1
        assert page.locator('.compare-splitter .ant-splitter-panel').count() == 2
        assert page.locator('.compare-splitter .ant-splitter-bar-dragger').count() == 1
        page.locator('.compare-lines path').first.wait_for(state="attached")
        left_panel = page.locator('.compare-splitter .ant-splitter-panel').first
        dragger = page.locator('.compare-splitter .ant-splitter-bar-dragger')
        dragger.scroll_into_view_if_needed()
        width_before = left_panel.bounding_box()["width"]
        dragger_box = dragger.bounding_box()
        page.mouse.move(dragger_box["x"] + dragger_box["width"] / 2, dragger_box["y"] + dragger_box["height"] / 2)
        page.mouse.down()
        page.mouse.move(dragger_box["x"] + dragger_box["width"] / 2 - 120, dragger_box["y"] + dragger_box["height"] / 2, steps=8)
        page.mouse.up()
        page.wait_for_timeout(300)
        assert left_panel.bounding_box()["width"] < width_before - 40
        assert page.locator('.compare-lines path').count() >= 1
        page.set_viewport_size({"width": 390, "height": 844})
        page.wait_for_function("() => !document.querySelector('.compare-splitter')")
        page.locator('.compare-mobile-switch').wait_for(state="visible")
        assert page.locator('[data-diff-id="D1"]').count() == 2
        assert page.locator('.compare-diff-rail').count() == 1
        outer_scroll_before = page.evaluate("() => window.scrollY")
        page.locator('.compare-diff-item').first.evaluate("node => node.click()")
        page.wait_for_function("() => document.querySelector('.compare-pane-wrap')?.scrollTop > 0")
        assert page.evaluate("() => window.scrollY") == outer_scroll_before
        assert page.locator('.compare-pane-wrap').first.evaluate("pane => { const target = pane.querySelector('[data-diff-id=\\\"D1\\\"]'); const a = pane.getBoundingClientRect(); const b = target.getBoundingClientRect(); return b.top < a.bottom && b.bottom > a.top }")
        assert page.locator('[data-diff-id="D3"]').count() == 1
        page.locator('.compare-diff-item').nth(1).evaluate("node => node.click()")
        page.locator('.compare-viewer[data-mobile-side="right"]').wait_for(state="attached")
        page.wait_for_function("() => { const pane = document.querySelector('.compare-viewer[data-mobile-side=\"right\"] .compare-pane-wrap:last-of-type'); const target = pane?.querySelector('[data-diff-id=\"D3\"]'); if (!pane || !target) return false; const a = pane.getBoundingClientRect(); const b = target.getBoundingClientRect(); return b.top < a.bottom && b.bottom > a.top }")
        metrics = page.evaluate("""() => ({ viewport: innerWidth, documentWidth: document.documentElement.scrollWidth, bodyWidth: document.body.scrollWidth })""")
        assert metrics["documentWidth"] <= metrics["viewport"]
        assert metrics["bodyWidth"] <= metrics["viewport"]

        page.get_by_role("button", name="查看对比版").click()
        assert page.locator('.compare-viewer[data-mobile-side="right"]').count() == 1
        next_pair = page.get_by_role("button", name="下一组", exact=True)
        next_pair.wait_for(state="visible")
        next_pair.click()
        page.get_by_text("共 1 处差异 · 修改 0 · 新增 1 · 删除 0").wait_for(state="visible")
        assert page.locator('[data-diff-id="D2"]').count() == 2
        assert page.locator('.compare-viewer[data-mobile-side="left"]').count() == 1
        assert not page_errors, page_errors
        assert not console_errors, console_errors
    finally:
        context.close()


def test_compare_first_upload_seals_and_refetches_explicit_task(browser, live_app, tmp_path):
    """The first Compare upload never refetches an undefined task."""
    from docx import Document

    left_path = tmp_path / "原版.docx"
    right_path = tmp_path / "修订版.docx"
    for path, text in ((left_path, "原版合同"), (right_path, "修订版合同")):
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    task_id = "compare-upload-browser-task"
    task = {
        "id": task_id, "name": "首次上传回归", "status": "sealed", "total": 1,
        "pairs": [{"id": "compare-upload-browser-pair", "label": "合同 1", "status": "pending",
                    "left": {"name": "原版.docx"}, "right": {"name": "修订版.docx"}}],
    }
    request_urls: list[str] = []
    context = browser.new_context(viewport={"width": 1440, "height": 900})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','compare-first-upload-check');")
    page = context.new_page()

    def json_route(route, body, status=200):
        route.fulfill(status=status, content_type="application/json", body=json.dumps(body))

    def handle(route):
        request_urls.append(route.request.url)
        path = route.request.url.split("/api", 1)[-1].split("?", 1)[0]
        if path == "/compare/tasks" and route.request.method == "GET":
            return json_route(route, {"items": [task], "total": 1})
        if path == "/compare/tasks" and route.request.method == "POST":
            return json_route(route, task, status=201)
        if path == f"/compare/tasks/{task_id}" and route.request.method == "GET":
            return json_route(route, task)
        if path == f"/compare/tasks/{task_id}/uploads":
            return json_route(route, {"receipts": [{"client_id": "left", "status": "accepted"}, {"client_id": "right", "status": "accepted"}]})
        if path == f"/compare/tasks/{task_id}/seal":
            return json_route(route, task)
        return route.continue_()

    page.route("**/api/**", handle)
    try:
        response = page.goto(live_app + "/compare", wait_until="networkidle")
        assert response and response.ok
        page.get_by_role("heading", name="合同对比").first.wait_for(state="visible")
        page.get_by_role("button", name="＋ 新建对比任务").click()
        file_inputs = page.locator(".compare-side-inputs input[type=file]")
        file_inputs.nth(0).set_input_files(str(left_path))
        file_inputs.nth(1).set_input_files(str(right_path))
        page.get_by_role("button", name="开始上传并对比").click()
        page.get_by_text("已提交，对比任务正在处理。").wait_for(state="visible")
        assert not any("/api/compare/tasks/undefined" in url for url in request_urls)
        assert any(f"/api/compare/tasks/{task_id}" in url for url in request_urls)
        assert "上传中断" not in page.locator("body").inner_text()
    finally:
        context.close()


def test_new_task_manifest_search_filter_pagination_and_compare_entry(browser, live_app, tmp_path):
    """The React upload list keeps the original manifest interactions reachable."""
    from docx import Document

    files = []
    for index in range(11):
        path = tmp_path / f"清单合同-{index:02d}.docx"
        document = Document()
        document.add_paragraph(f"合同 {index}")
        document.save(path)
        files.append(path)
    invalid = tmp_path / "说明.txt"
    invalid.write_text("invalid", encoding="utf-8")

    context = browser.new_context(viewport={"width": 1440, "height": 900})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v2-manifest-check');")
    page = context.new_page()
    try:
        response = page.goto(live_app + "/tasks/new", wait_until="networkidle")
        assert response and response.ok
        page.locator("input[type=file]").first.set_input_files([str(path) for path in files] + [str(invalid)])
        page.get_by_text("文件清单").wait_for()
        assert "第 1 / 2 页" in page.locator("body").inner_text()
        assert page.locator(".upload-manifest-table tbody tr").count() == 10

        page.get_by_role("button", name="下一页").click()
        assert "清单合同-10.docx" in page.locator("body").inner_text()
        page.get_by_placeholder("搜索当前待提交文件").fill("10.docx")
        assert page.locator(".upload-manifest-table tbody tr").count() == 1
        page.get_by_placeholder("搜索当前待提交文件").fill("")
        page.get_by_role("combobox").last.select_option("issue")
        assert "说明.txt" in page.locator("body").inner_text()
        assert page.locator(".upload-manifest-table tbody tr").count() == 1

        compare = page.locator('a[href="/compare"]')
        assert compare.is_visible()
        compare.click()
        page.wait_for_url(re.compile(r"/compare$"))
        page.get_by_role("heading", name="合同对比").first.wait_for()
        page.get_by_role("button", name="＋ 新建对比任务").wait_for()
        assert "原版 → 对比版" in page.locator("body").inner_text()
    finally:
        context.close()


def test_new_task_invalid_replacement_removes_old_upload_row(browser, live_app, tmp_path):
    """A failed replacement cannot leave the previous file silently submittable."""
    from docx import Document

    original = tmp_path / "原合同.docx"
    document = Document()
    document.add_paragraph("原合同")
    document.save(original)
    invalid = tmp_path / "替换说明.txt"
    invalid.write_text("invalid", encoding="utf-8")

    context = browser.new_context(viewport={"width": 1440, "height": 900})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v2-replace-check');")
    page = context.new_page()
    try:
        response = page.goto(live_app + "/tasks/new", wait_until="networkidle")
        assert response and response.ok
        page.locator("input[type=file]").first.set_input_files(str(original))
        page.get_by_text("原合同.docx").wait_for()
        page.get_by_role("button", name="替换").click()
        page.wait_for_timeout(100)
        page.locator('input[type=file][accept*="docx"]').last.set_input_files(str(invalid))
        page.get_by_text("只支持 DOCX 文件。", exact=False).wait_for()
        assert "原合同.docx" not in page.locator(".upload-manifest-table").inner_text()
        assert "替换说明.txt" in page.locator(".upload-manifest-table").inner_text()
        assert page.get_by_role("button", name="创建任务并上传").is_disabled()
    finally:
        context.close()


def test_shell_shrinks_when_viewport_changes_from_desktop_to_medium(browser, live_app):
    """Responsive layout also stays bounded when a browser is resized in place."""
    context = browser.new_context(viewport={"width": 1440, "height": 900})
    context.add_init_script("localStorage.setItem('fundlex-dev-user','v2-resize-check');")
    page = context.new_page()
    try:
        response = page.goto(live_app + "/tasks", wait_until="networkidle")
        assert response and response.ok
        page.set_viewport_size({"width": 768, "height": 900})
        page.wait_for_timeout(150)
        metrics = page.evaluate("""() => ({ viewport: innerWidth, documentWidth: document.documentElement.scrollWidth, bodyWidth: document.body.scrollWidth })""")
        assert metrics["documentWidth"] <= metrics["viewport"]
        assert metrics["bodyWidth"] <= metrics["viewport"]
    finally:
        context.close()
