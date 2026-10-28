"""Independent edge cases for the compare parser/alignment/diff contracts."""
from hashlib import sha256
import io
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
import pytest

from fundlex.compare import compare_docx, compare_documents, parse_compare_document
from fundlex.compare.revisions import prepare_compare_input
from fundlex.compare.normalize import normalize_match, normalize_diff
from fundlex.docx import DOCUMENT, Q, DocxError


def doc(tmp_path, name, paras):
    d = Document()
    for text in paras:
        d.add_paragraph(text)
    p = tmp_path / name
    d.save(p)
    return p


def texts(result, side):
    return ''.join(s.text for d in result.diffs for s in getattr(d, f'{side}_spans'))


def mutate_xml(path, output, change):
    with ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = etree.fromstring(parts[DOCUMENT]); change(root)
    parts[DOCUMENT] = etree.tostring(root)
    with ZipFile(output, 'w', ZIP_DEFLATED) as z:
        for n, data in parts.items(): z.writestr(n, data)
    return output


@pytest.mark.parametrize(('a', 'b'), [('0.40%', '0.30%'), ('应当', '可以'), ('不得', '可以'), ('。', '；'), ('10日', '15日'), ('ABC', 'abc'), ('基金管理人', '基金 管理人')])
def test_visible_changes_never_disappear(tmp_path, a, b):
    left = doc(tmp_path, 'l.docx', ['前文锚点', f'本基金{a}履行义务。', '后文锚点'])
    right = doc(tmp_path, 'r.docx', ['前文锚点', f'本基金{b}履行义务。', '后文锚点'])
    r = compare_docx(left, right)
    assert r.summary.total > 0
    assert texts(r, 'left') or texts(r, 'right')


def test_run_and_formatting_differences_are_not_text_differences(tmp_path):
    left = doc(tmp_path, 'l.docx', ['基金管理人应当履行义务'])
    d = Document(); p = d.add_paragraph()
    for text in ['基金', '管理人', '应当', '履行义务']:
        r = p.add_run(text); r.bold = True; r.font.name = '黑体'
    right = tmp_path / 'r.docx'; d.save(right)
    assert compare_docx(left, right).summary.total == 0


@pytest.mark.parametrize(('a', 'b'), [(['ABCDEF'], ['ABC', 'DEF']), (['ABC', 'DEF'], ['ABCDEF']), (['ABCDEF'], ['AB', 'CD', 'EF']), (['AB', 'CD', 'EF'], ['ABCDEF']), (['A  B'], ['A    B'])])
def test_paragraph_split_merge_and_space_counts(tmp_path, a, b):
    assert compare_docx(doc(tmp_path, 'l.docx', a), doc(tmp_path, 'r.docx', b)).summary.total == 0


def test_split_with_true_edit_keeps_only_changed_token(tmp_path):
    left = doc(tmp_path, 'l.docx', ['基金管理费为0.40%，基金托管费为0.10%。'])
    right = doc(tmp_path, 'r.docx', ['基金管理费为0.30%，', '基金托管费为0.10%。'])
    r = compare_docx(left, right)
    assert texts(r, 'left') == '0.40%'
    assert texts(r, 'right') == '0.30%'
    assert r.summary.modify == 1


def test_repeated_paragraphs_and_insert_do_not_shift_following_differences(tmp_path):
    a = ['标题', '具体以届时公告为准', '原管理费0.40%', '具体以届时公告为准', '基金管理人应当履行义务', '尾部']
    b = ['标题', '具体以届时公告为准', '新增段落', '原管理费0.30%', '具体以届时公告为准', '基金管理人可以履行义务', '尾部']
    r = compare_docx(doc(tmp_path, 'l.docx', a), doc(tmp_path, 'r.docx', b))
    assert '0.40%' in texts(r, 'left') and '0.30%' in texts(r, 'right')
    assert '应当' in texts(r, 'left') and '可以' in texts(r, 'right')
    assert '新增段落' in texts(r, 'right')


def test_table_row_insert_does_not_mark_unchanged_rows(tmp_path):
    def table_doc(name, rows):
        d = Document(); t = d.add_table(rows=0, cols=2)
        for a, b in rows:
            cells = t.add_row().cells; cells[0].text = a; cells[1].text = b
        p = tmp_path / name; d.save(p); return p
    a = [('A', '甲'), ('B', '乙'), ('C', '丙')]
    b = [('A', '甲'), ('X', '新'), ('B', '乙'), ('C', '丙')]
    r = compare_docx(table_doc('l.docx', a), table_doc('r.docx', b))
    assert texts(r, 'left') == ''
    assert texts(r, 'right') == 'X新'


def test_table_cell_boundaries_are_semantic(tmp_path):
    def tdoc(name, values):
        d = Document(); t = d.add_table(rows=1, cols=2)
        for c, v in zip(t.rows[0].cells, values): c.text = v
        p = tmp_path/name; d.save(p); return p
    r = compare_docx(tdoc('l.docx',['ABC','DEF']), tdoc('r.docx',['ABCDEF','']))
    assert r.summary.total > 0
    assert 'DEF' in texts(r, 'left') and 'DEF' in texts(r, 'right')


def test_hyperlinks_sdt_and_merged_cells_are_read(tmp_path):
    d = Document(); p = d.add_paragraph('前')
    link = OxmlElement('w:hyperlink'); run = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text='链接文字'; run.append(t); link.append(run); p._p.append(link)
    sdt=OxmlElement('w:sdt'); content=OxmlElement('w:sdtContent'); p=OxmlElement('w:p');run=OxmlElement('w:r');t=OxmlElement('w:t');t.text='控件文字';run.append(t);p.append(run);content.append(p);sdt.append(content);d._element.body.insert(1,sdt)
    table=d.add_table(rows=1, cols=2);table.cell(0,0).text='合并文字';table.cell(0,0).merge(table.cell(0,1))
    path=tmp_path/'source.docx';d.save(path)
    result=parse_compare_document(path)
    assert ''.join(b.raw_text for b in result.blocks)=='前链接文字控件文字合并文字'
    assert result.coverage.status=='complete'


def test_visible_unsupported_content_never_claims_complete(tmp_path):
    path=doc(tmp_path,'a.docx',['普通正文'])
    def add(root):
        p=root.find('.//'+Q('p'));r=OxmlElement('w:r');drawing=OxmlElement('w:drawing');r.append(drawing);p.append(r)
    bad=mutate_xml(path,tmp_path/'drawing.docx',add)
    parsed=parse_compare_document(bad)
    assert parsed.coverage.status=='incomplete'
    assert compare_docx(bad,bad).coverage.status=='incomplete'


def test_existing_revisions_rejected(tmp_path):
    path=doc(tmp_path,'a.docx',['原文字'])
    def add(root): root.find('.//'+Q('p')).append(OxmlElement('w:ins'))
    bad=mutate_xml(path,tmp_path/'revision.docx',add)
    with pytest.raises(DocxError, match='EXISTING_REVISION_UNSUPPORTED'): parse_compare_document(bad)


def test_compare_input_accepts_existing_revisions_without_overwriting_source(tmp_path):
    path = doc(tmp_path, 'source.docx', ['原文字'])

    def add(root):
        paragraph = root.find('.//' + Q('p'))
        run = paragraph.find(Q('r'))
        text = run.find(Q('t'))
        deleted = OxmlElement('w:del')
        deleted.set(Q('id'), '1')
        deleted_run = OxmlElement('w:r')
        deleted_text = OxmlElement('w:delText')
        deleted_text.text = text.text
        deleted_run.append(deleted_text)
        deleted.append(deleted_run)
        inserted = OxmlElement('w:ins')
        inserted.set(Q('id'), '2')
        inserted_run = OxmlElement('w:r')
        inserted_text = OxmlElement('w:t')
        inserted_text.text = '新文字'
        inserted_run.append(inserted_text)
        inserted.append(inserted_run)
        paragraph.remove(run)
        paragraph.extend([deleted, inserted])

    revised = mutate_xml(path, tmp_path / 'revised.docx', add)
    original = revised.read_bytes()
    prepared = prepare_compare_input(original)

    assert prepared.revisions_accepted == 3
    assert prepared.original_size == len(original)
    assert prepared.original_sha256 == sha256(original).hexdigest()
    assert revised.read_bytes() == original
    assert b'<w:ins' not in prepared.data
    assert b'<w:del' not in prepared.data
    parsed = parse_compare_document(prepared.data, filename='revised.docx')
    assert ''.join(block.raw_text for block in parsed.blocks) == '新文字'
    repeated = prepare_compare_input(original)
    assert repeated.data == prepared.data
    assert repeated.revisions_accepted == prepared.revisions_accepted


def test_compare_input_accepts_moved_runs_and_range_markers(tmp_path):
    path = doc(tmp_path, 'source.docx', ['甲位移文字乙'])

    def add(root):
        paragraph = root.find('.//' + Q('p'))
        run = paragraph.find(Q('r'))
        text = run.find(Q('t')).text
        paragraph.remove(run)
        for tag in ('moveFromRangeStart', 'moveToRangeStart'):
            marker = OxmlElement('w:' + tag)
            marker.set(Q('id'), tag)
            marker.set(Q('name'), tag)
            paragraph.append(marker)
        for tag, wrapper, text_tag in (
                ('moveFromRangeEnd', 'moveFrom', 'delText'),
                ('moveToRangeEnd', 'moveTo', 't')):
            moved = OxmlElement('w:' + wrapper)
            moved.set(Q('id'), wrapper)
            moved_run = OxmlElement('w:r')
            moved_text = OxmlElement('w:' + text_tag)
            moved_text.text = text
            moved_run.append(moved_text)
            moved.append(moved_run)
            paragraph.append(moved)
            marker = OxmlElement('w:' + tag)
            marker.set(Q('id'), tag)
            paragraph.append(marker)

    revised = mutate_xml(path, tmp_path / 'moved.docx', add)
    prepared = prepare_compare_input(revised.read_bytes())
    assert prepared.revisions_accepted == 7
    assert all(tag not in prepared.data for tag in (
        b'<w:moveFrom', b'<w:moveTo', b'<w:delText', b'<w:moveFromRangeStart'))
    parsed = parse_compare_document(prepared.data, filename='moved.docx')
    assert [block.raw_text for block in parsed.blocks] == ['甲位移文字乙']


def test_compare_input_accepts_deleted_and_moved_paragraph_marks(tmp_path):
    path = doc(tmp_path, 'source.docx', ['第一段', '第二段', '第三段'])

    def add(root):
        for index, tag in enumerate(('del', 'moveFrom')):
            paragraph = root.findall('.//' + Q('p'))[index]
            p_pr = paragraph.find(Q('pPr'))
            if p_pr is None:
                p_pr = OxmlElement('w:pPr')
                paragraph.insert(0, p_pr)
            r_pr = p_pr.find(Q('rPr'))
            if r_pr is None:
                r_pr = OxmlElement('w:rPr')
                p_pr.append(r_pr)
            marker = OxmlElement('w:' + tag)
            marker.set(Q('id'), str(index + 10))
            r_pr.append(marker)

    revised = mutate_xml(path, tmp_path / 'paragraph-marks.docx', add)
    prepared = prepare_compare_input(revised.read_bytes())
    assert prepared.revisions_accepted == 2
    parsed = parse_compare_document(prepared.data, filename='paragraph-marks.docx')
    assert [block.raw_text for block in parsed.blocks] == ['第一段第二段第三段']


def test_compare_input_accepts_deleted_rows_cells_and_cell_merges(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    for row_index, values in enumerate((('A', 'B'), ('C', 'D'))):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    path = tmp_path / 'table.docx'
    document.save(path)

    def remove(root):
        rows = root.findall('.//' + Q('tr'))
        tr_pr = rows[1].find(Q('trPr'))
        if tr_pr is None:
            tr_pr = OxmlElement('w:trPr')
            rows[1].insert(0, tr_pr)
        deleted_row = OxmlElement('w:del')
        deleted_row.set(Q('id'), 'row')
        tr_pr.append(deleted_row)
        cell = rows[0].findall(Q('tc'))[1]
        tc_pr = cell.find(Q('tcPr'))
        if tc_pr is None:
            tc_pr = OxmlElement('w:tcPr')
            cell.insert(0, tc_pr)
        deleted_cell = OxmlElement('w:cellDel')
        deleted_cell.set(Q('id'), 'cell')
        tc_pr.append(deleted_cell)

    revised = mutate_xml(path, tmp_path / 'table-deleted.docx', remove)
    prepared = prepare_compare_input(revised.read_bytes())
    assert prepared.revisions_accepted == 2
    assert b'<w:cellDel' not in prepared.data and b'<w:del' not in prepared.data
    parsed = parse_compare_document(prepared.data, filename='table-deleted.docx')
    assert [block.raw_text for block in parsed.blocks] == ['A']

    def merge(root):
        cell = root.findall('.//' + Q('tc'))[0]
        tc_pr = cell.find(Q('tcPr'))
        if tc_pr is None:
            tc_pr = OxmlElement('w:tcPr')
            cell.insert(0, tc_pr)
        change = OxmlElement('w:cellMerge')
        change.set(Q('id'), 'merge')
        change.set(Q('vMerge'), 'rest')
        tc_pr.append(change)

    merged = mutate_xml(path, tmp_path / 'table-merged.docx', merge)
    prepared = prepare_compare_input(merged.read_bytes())
    assert prepared.revisions_accepted == 1
    with ZipFile(io.BytesIO(prepared.data)) as archive:
        root = etree.fromstring(archive.read(DOCUMENT))
    value = root.find('.//' + Q('vMerge'))
    assert value is not None and value.get(Q('val')) == 'restart'
    assert b'<w:cellMerge' not in prepared.data


def test_compare_input_rejects_invalid_cell_merge(tmp_path):
    path = doc(tmp_path, 'source.docx', ['表格'])

    def add(root):
        table = OxmlElement('w:tbl')
        row = OxmlElement('w:tr')
        cell = OxmlElement('w:tc')
        tc_pr = OxmlElement('w:tcPr')
        tc_pr.append(OxmlElement('w:cellMerge'))
        cell.append(tc_pr)
        row.append(cell)
        table.append(row)
        root.find(Q('body')).append(table)

    revised = mutate_xml(path, tmp_path / 'invalid-merge.docx', add)
    with pytest.raises(DocxError, match='REVISION_CELL_MERGE_UNSUPPORTED'):
        prepare_compare_input(revised.read_bytes())


def test_compare_input_counts_revisions_inside_property_snapshots(tmp_path):
    path = doc(tmp_path, 'source.docx', ['属性修订'])
    original = path.read_bytes()

    def add(root):
        paragraph = root.find('.//' + Q('p'))
        p_pr = paragraph.find(Q('pPr'))
        if p_pr is None:
            p_pr = OxmlElement('w:pPr')
            paragraph.insert(0, p_pr)
        change = OxmlElement('w:pPrChange')
        change.set(Q('id'), 'change')
        old = OxmlElement('w:pPr')
        old_r_pr = OxmlElement('w:rPr')
        inserted = OxmlElement('w:ins')
        inserted.set(Q('id'), 'old-insert')
        old_r_pr.append(inserted)
        old.append(old_r_pr)
        change.append(old)
        p_pr.append(change)

    revised = mutate_xml(path, tmp_path / 'property-change.docx', add)
    prepared = prepare_compare_input(revised.read_bytes())
    assert prepared.original_sha256 == sha256(revised.read_bytes()).hexdigest()
    assert prepared.revisions_accepted == 2
    assert b'<w:pPrChange' not in prepared.data and b'<w:ins' not in prepared.data
    assert prepared.data != original
    parsed = parse_compare_document(prepared.data, filename='property-change.docx')
    assert [block.raw_text for block in parsed.blocks] == ['属性修订']


def test_compare_input_ignores_drawingml_move_to(tmp_path):
    path = doc(tmp_path, 'source.docx', ['合同正文'])

    def add(root):
        run = etree.SubElement(root, Q('r'))
        drawing = etree.SubElement(
            run, '{http://schemas.openxmlformats.org/drawingml/2006/main}moveTo')
        etree.SubElement(
            drawing, '{http://schemas.openxmlformats.org/drawingml/2006/main}pt')

    revised = mutate_xml(path, tmp_path / 'shape.docx', add)
    original = revised.read_bytes()
    prepared = prepare_compare_input(original)
    assert prepared.revisions_accepted == 0
    assert prepared.data == original
    assert parse_compare_document(prepared.data, filename='shape.docx').blocks


def test_compare_input_rejects_unsafe_original_even_when_revision_is_present(tmp_path):
    path = doc(tmp_path, 'source.docx', ['原文字'])

    def add(root):
        paragraph = root.find('.//' + Q('p'))
        paragraph.append(OxmlElement('w:ins'))
        protection = OxmlElement('w:documentProtection')
        root.insert(0, protection)

    revised = mutate_xml(path, tmp_path / 'unsafe.docx', add)
    with pytest.raises(DocxError, match='DOCUMENT_PROTECTED'):
        prepare_compare_input(revised.read_bytes())


def test_bytes_path_identity_and_source_hash(tmp_path):
    path=doc(tmp_path,'source.docx',['合同'])
    a=parse_compare_document(path);b=parse_compare_document(path.read_bytes(),filename=path.name)
    assert a == b
    assert a.source_sha256 == sha256(path.read_bytes()).hexdigest()


def test_insert_delete_virtual_anchors_are_bound_to_real_blocks(tmp_path):
    a=doc(tmp_path,'a.docx',['开头','尾部'])
    b=doc(tmp_path,'b.docx',['开头','新增文字','尾部'])
    r=compare_docx(a,b)
    ids={x.block_id for x in r.left.blocks}
    for item in r.diffs:
        if not item.left_spans: assert item.left_anchor.block_id in ids


def test_normalization_preserves_significant_characters():
    assert normalize_match('  A  0.40%！  ') == 'A 0.40%！'
    assert normalize_diff('ABC') != normalize_diff('abc')
    assert normalize_diff('基金管理人') != normalize_diff('基金 管理人')


def test_deterministic_and_symmetric(tmp_path):
    a=doc(tmp_path,'a.docx',['重复','费率0.40%','重复','不得赎回','删除内容','终点'])
    b=doc(tmp_path,'b.docx',['重复','费率0.30%','新增内容','重复','可以赎回','终点'])
    runs=[compare_docx(a,b).to_dict()['diffs'] for _ in range(5)]
    assert all(r == runs[0] for r in runs)
    r=compare_docx(a,b);reverse=compare_docx(b,a)
    def sig(result, rev=False):
        out=[]
        for d in result.diffs:
            typ={'insert':'delete','delete':'insert','modify':'modify'}[d.type] if rev else d.type
            left=''.join(s.text for s in d.left_spans);right=''.join(s.text for s in d.right_spans)
            out.append((typ,right,left) if rev else (typ,left,right))
        return sorted(out)
    assert sig(r)==sig(reverse,True)
