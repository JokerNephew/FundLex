"""G17's complete usage-exclusion sentence remains required (FL-D214)."""
import json
from pathlib import Path

import pytest
from docx import Document

from fundlex.docx import inspect_docx
from fundlex.generation.claim_replay import replay_claims
from fundlex.generation.effects import operation_site_bound
from fundlex.obligations import apply_facts, compile_obligations
from fundlex.recall import Target, guard_g17_deletes_against_g16, plan_target
from fundlex.recall_pipeline import _plan_and_generate
from test_recall_core import _facts


HEAD = ("销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、"
        "促销活动费、基金份额持有人服务费等。")
TAIL = "销售服务费的使用范围不包括基金募集期间的上述费用。"


class NoModel:
    max_context_chars = 60000

    def call(self, **kwargs):
        raise AssertionError("a complete authorized sentence needs no model")


def _setup(tmp_path, tail=TAIL, rule=None):
    if rule is None:
        guide = json.loads((Path(__file__).resolve().parents[1] /
                            "guides/builtin_guide_v1.json").read_text())
        rule = next(g for g in guide["groups"] if g["rule_id"] == "G17")
    source = tmp_path / "source.docx"
    _write(source, [HEAD, tail, HEAD, tail])
    original = inspect_docx(source, full=True)
    facts = _facts()
    ledger = compile_obligations({"groups": [rule]})
    apply_facts(ledger, facts)
    targets = [Target(f"T{i}", "G17", f"p{i}", text, "needs_edit", "source",
                      start=0, end=len(text), resolution="resolved", operations=("G17:0",))
               for i, text in enumerate([HEAD, tail, HEAD, tail])]
    return source, original, facts, ledger, rule, targets


def _write(path, texts):
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    document.save(path)


@pytest.mark.parametrize("reported", [True, False])
def test_retained_tail_fails_actual_docx_replay_even_after_head_deleted(tmp_path, reported):
    _, original, facts, ledger, rule, targets = _setup(tmp_path)
    final = tmp_path / "actual.docx"
    _write(final, ["", TAIL, "", TAIL])
    edits = [dict(target_id=t.target_id, rule_id="G17", location=t.location,
                  start=0, end=len(HEAD), replacement="", operation="replace",
                  content_source="variant_delete", operation_keys=["G17:0"])
             for t in (targets[0], targets[2])]
    for edit in edits:
        edit["components"] = [{**edit, "expected_old_text": HEAD}]
    proof = replay_claims(edits, original, inspect_docx(final, full=True),
        rules_by_id={"G17": rule}, facts=facts, obligations=ledger,
        targets_by_id={t.target_id: t.as_dict() for t in targets
                       if reported or t.location in {"p0", "p2"}})
    assert not proof["verified"]
    effect = proof["obligations"]["G17:B1:O1"]
    assert set(effect["sites"]) == {"p0", "p1", "p2", "p3"}
    assert {s["location"] for s in effect["site_effects"] if s["state"] == "FAIL"} == {"p1", "p3"}
    assert proof["failures"] and all(f.get("operation_id") == "G17:0"
                                     and f["location"] in {"p1", "p3"}
                                     for f in proof["failures"])


@pytest.mark.parametrize("tail", [TAIL, TAIL.replace("的使用", "使用"),
    "C 类基金份额的" + TAIL, "基金" + TAIL])
def test_exact_usage_exclusion_variant_plans_and_deletes_all_docx_sites(tmp_path, monkeypatch, tail):
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    source, original, facts, ledger, rule, targets = _setup(tmp_path, tail)
    planned, _, _, _, actual = _plan_and_generate(NoModel(), {"G17": rule}, targets,
        original, facts, source, tmp_path, max_compose_calls=0, obligations=ledger)
    assert len(planned["edits"]) == 4, planned["ledger"]
    assert all(not actual.by_location(f"p{i}").text.strip() for i in range(4))
    proof = replay_claims(planned["edits"], original, actual, rules_by_id={"G17": rule},
        facts=facts, obligations=ledger, targets_by_id={t.target_id: t.as_dict() for t in targets})
    assert proof["verified"], proof["failures"]
    assert set(proof["obligations"]["G17:B1:O1"]["sites"]) == {"p0", "p1", "p2", "p3"}
    assert (tmp_path / "redline.docx").is_file()


@pytest.mark.parametrize("tail", [TAIL.replace("不包括", "包括"),
    TAIL.replace("销售服务费", "管理费"), TAIL.replace("募集期间", "存续期间"),
    TAIL + "该项费用应另行核对。"])
def test_usage_sentence_binding_rejects_different_meaning_or_extra_sentence(tmp_path, tail):
    _, original, facts, ledger, _, _ = _setup(tmp_path, tail)
    assert not operation_site_bound(ledger.record("G17:0").spec, tail,
                                    original, "p1", facts)


def test_head_only_operation_cannot_borrow_guide_tail_authorization(tmp_path, monkeypatch):
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    _, _, _, _, rule, _ = _setup(tmp_path)
    rule["operations"][0]["old_text"] = HEAD
    _, original, facts, ledger, rule, targets = _setup(tmp_path, rule=rule)
    assert not operation_site_bound(ledger.record("G17:0").spec, TAIL,
                                    original, "p1", facts)
    from fundlex.recall import RecallError
    with pytest.raises(RecallError):
        plan_target(rule, targets[1], original, facts, NoModel(), compose_budget=[0], obligations=ledger)


@pytest.mark.parametrize("stale", [False, True])
def test_g17_delete_is_blocked_when_same_clause_g16_has_no_plan(tmp_path, stale):
    source, original, facts, ledger, g17, _ = _setup(tmp_path)
    head = "销售服务费每日计提，按月支付。"
    tail = TAIL
    _write(source, [head, tail])
    original = inspect_docx(source, full=True)
    g16 = Target("g16", "G16", "p0", head, "needs_edit", "source",
                 start=0, end=len(head), resolution="resolved", out_of_scope=stale)
    g17 = Target("g17", "G17", "p1", tail, "needs_edit", "source",
                 start=0, end=len(tail), resolution="resolved")
    planned = {"edits": [{"target_id": "g17", "rule_id": "G17", "location": "p1",
                           "operation": "replace", "replacement": "", "status": "planned"}],
               "ledger": [{"target_id": "g17", "status": "planned"}]}
    blocked = guard_g17_deletes_against_g16(planned, [g16, g17], original)
    assert blocked and planned["edits"] == []
    assert planned["ledger"][0]["status"] == "unresolved"


@pytest.mark.parametrize("loss", [False, True])
def test_merged_g16_g17_actual_docx_and_before_loss_audit(tmp_path, monkeypatch, loss):
    """Authorized G17 removal must not conceal missing unrelated source (FL-D224)."""
    from types import SimpleNamespace
    from fundlex.recall_pipeline import _run_audit, _merge_audit_state, _audit_issues
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    guide = json.loads((Path(__file__).resolve().parents[1] /
                        "guides/builtin_guide_v1.json").read_text())
    rules = {g["rule_id"]: g for g in guide["groups"] if g["rule_id"] in {"G16", "G17"}}
    suffix = "费用计算结果应当逐项核对并保留记录。"
    tail = TAIL.replace("的使用", "使用")
    texts = ["销售服务费每日计提，按月支付。", HEAD, tail + suffix]
    source = tmp_path / "source.docx"
    _write(source, texts)
    original, facts = inspect_docx(source, full=True), _facts()
    ledger = compile_obligations({"groups": list(rules.values())})
    apply_facts(ledger, facts)
    targets = [Target("g16", "G16", "p2", texts[2], "needs_edit", "source",
        start=0, end=len(texts[2]), resolution="resolved", operations=("G16:0",)),
        Target("usage", "G17", "p1", HEAD, "needs_edit", "source",
        start=0, end=len(HEAD), resolution="resolved", operations=("G17:0",)),
        Target("scope", "G17", "p2", tail, "needs_edit", "source",
        start=0, end=len(tail), resolution="resolved", operations=("G17:0",))]
    planned, _, _, _, actual = _plan_and_generate(NoModel(), rules, targets,
        original, facts, source, tmp_path, max_compose_calls=0, obligations=ledger)
    proof = replay_claims(planned["edits"], original, actual, rules_by_id=rules,
        facts=facts, obligations=ledger, targets_by_id={t.target_id: t.as_dict() for t in targets})
    assert proof["verified"], proof["failures"]
    assert all(proof["obligations"][f"{rid}:B1:O1"]["state"] == "PASS" for rid in rules)
    assert suffix in actual.by_location("p2").text and tail not in actual.by_location("p2").text
    assert not actual.by_location("p1").text.strip()
    if loss:
        artifact = tmp_path / "lost-suffix.docx"
        document = Document(tmp_path / "final.docx")
        document.paragraphs[2].text = document.paragraphs[2].text.replace(suffix, "")
        document.save(artifact)
        actual = inspect_docx(artifact, full=True)
    quote = suffix if loss else tail
    class Gateway:
        max_context_chars = 100000
        def call(self, **kwargs):
            return SimpleNamespace(value={"results": [dict(target_id=t.target_id,
                verdict="wrong", failure_class="GENERATION", evidence_quote=quote,
                reason="原文丢失") for t in targets if t.location == "p2"], "missed_targets": []})
    result = _run_audit(Gateway(), rules, targets, planned, original, actual,
                        facts=facts, obligations=ledger)
    assert len(result["results"]) == 2
    for item in result["results"]:
        assert item["evidence_verified"] and item["evidence_in_before"]
        assert not item["evidence_in_after"]
        if loss:
            assert not item.get("artifact_proven")
            assert item.get("review_status") != "REVIEW_UNPROVEN"
        else:
            assert item["artifact_proven"]
            assert item["review_evidence_status"] == "AUTHORIZED_REMOVAL_QUOTED_AS_ERROR"
    state = {"results": {}, "missed": {}}
    _merge_audit_state(state, result, targets)
    issues = [issue for issue in _audit_issues(state, targets) if issue["location"] == "p2"]
    assert len(issues) == (2 if loss else 0)


@pytest.mark.parametrize("text", [TAIL + "投资者可以随时免责。",
    TAIL.replace("不包括", "包括"), TAIL.replace("销售服务费", "管理费"),
    TAIL.replace("募集期间", "存续期间"), "销售服务费不主要用于支付销售机构佣金。"])
def test_g17_forged_provenance_cannot_authorize_source_loss(tmp_path, text):
    """Plan/artifact agreement and planner quotes cannot grant removal rights."""
    _, _, _, _, rule, _ = _setup(tmp_path)
    source, final = tmp_path / "forged-source.docx", tmp_path / "forged-final.docx"
    _write(source, [text])
    _write(final, [""])
    edit = dict(target_id="T", rule_id="G17", location="p0", start=0, end=len(text),
        replacement="", operation="replace", content_source="whole_paragraph_delete",
        operation_keys=["G17:0"], expected_old_text=text,
        evidence={"deleted_paragraph": text, "deleted_quote": text,
                  "delete_texts": [text], "fragments": [{"old_text": text}]})
    proof = replay_claims([edit], inspect_docx(source, full=True), inspect_docx(final, full=True),
                          rules_by_id={"G17": rule})
    assert not proof["verified"]
    assert any(f["code"] == "CLAIM_OPERATION_SPAN_UNAUTHORIZED" for f in proof["failures"])


def test_g17_planner_rejects_whole_paragraph_fallback_with_extra_sentence(tmp_path, monkeypatch):
    import fundlex.recall as recall
    from fundlex.recall import PlannedEdit, RecallError
    _, _, facts, _, rule, _ = _setup(tmp_path)
    text = TAIL + "投资者可以随时免责。"
    source = tmp_path / "forged.docx"
    _write(source, [text])
    original = inspect_docx(source, full=True)
    target = Target("T", "G17", "p0", text, "needs_edit", "source", start=0,
                    end=len(text), operations=("G17:0",))
    monkeypatch.setattr(recall, "_tier1_enabled", lambda: False)
    monkeypatch.setattr(recall, "_plan_target_inner", lambda *a, **kw: PlannedEdit(
        target_id="T", rule_id="G17", location="p0", start=0, end=len(text),
        replacement="", operation="replace", content_source="whole_paragraph_delete", evidence={}))
    with pytest.raises(RecallError, match="RECALL_DELETE_SCOPE_UNPROVEN"):
        plan_target(rule, target, original, facts, NoModel(), compose_budget=[0])


@pytest.mark.parametrize("text", [HEAD.replace("、以及", "，及"),
    HEAD.replace("基金份额持有人服务费", "持有人服务费"),
    TAIL.replace("的使用范围", ""), HEAD[:-1] + "；"])
def test_g17_finite_source_variants_allow_only_complete_sentence(tmp_path, text):
    from fundlex.recall import _g17_source_effect_owned
    _, _, _, _, rule, _ = _setup(tmp_path)
    source, final = tmp_path / "variant.docx", tmp_path / "actual.docx"
    suffix = "基金管理人可根据基金发展情况下调销售服务费率。"
    _write(source, [text + suffix])
    _write(final, [suffix])
    original = inspect_docx(source, full=True)
    assert _g17_source_effect_owned(original, "p0", 0, len(text), "", rule["operations"])
    assert not _g17_source_effect_owned(original, "p0", 0, len(text + suffix), "", rule["operations"])
    for changed in (text.replace("主要用于", "不主要用于"),
                    text.replace("不包括", "包括"), text.replace("销售服务费", "管理费")):
        if changed == text:
            continue
        _write(source, [changed])
        assert not _g17_source_effect_owned(inspect_docx(source, full=True), "p0", 0,
                                            len(changed), "", rule["operations"])
    edit = dict(target_id="T", rule_id="G17", location="p0", start=0, end=len(text),
                replacement="", operation="replace", operation_keys=["G17:0"])
    proof = replay_claims([edit], original, inspect_docx(final, full=True), rules_by_id={"G17": rule})
    assert proof["verified"], proof["failures"]
