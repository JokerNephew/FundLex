"""Source-bound append after a definition split within the class letter (FL-D207)."""
import pytest
from docx import Document

from fundlex.docx import generate_artifacts, inspect_docx, validate_edit_plan
from fundlex.recall import (RecallError, Target, dedupe_operation_targets, plan_target,
                            rebind_split_append_targets, build_docx_edits, content_check_factory)
from test_operation_execution_closure import Facts


HEAD = "基金份额类别：本基金根据销售所适用账户的不同，将基金份额分为 A"
MIDDLE = "类基金份额和 Y 类基金份额 。A 类基金份额不得对个人养老金账户进行销售，"
END = "Y 类基金份额仅面向个人养老金账户销售。 "


def source(tmp_path, lines):
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    path = tmp_path / "split.docx"
    document.save(path)
    return inspect_docx(path, full=True)


@pytest.mark.parametrize("rule_id", ["G01", "G03"])
def test_append_after_entire_wrapped_definition_preserves_account_restrictions(tmp_path, rule_id):
    inspection = source(tmp_path, [HEAD, MIDDLE, END, "其他独立条款。"])
    rule = {"rule_id": rule_id, "operations": [dict(op="append", new_text="直销申购新政策。",
                                                  condition={"kind": "none"})]}
    target = Target("T", rule_id, "p0", HEAD, "needs_edit", "definition",
                    start=0, end=len(HEAD), resolution="resolved", operations=(rule_id + ":0",))
    edit = plan_target(rule, target, inspection, Facts(), None, compose_budget=[0])
    assert edit.operation == "insert"
    assert edit.location == "p2"
    assert edit.start == edit.end in {len(END.rstrip()), len(END)}
    assert inspection.by_location("p0").text == HEAD
    assert inspection.by_location("p1").text == MIDDLE
    plan = validate_edit_plan(tmp_path / "split.docx", build_docx_edits([edit], inspection), {},
                             content_check=content_check_factory([edit.as_dict()], inspection))
    final = tmp_path / "final.docx"
    generate_artifacts(tmp_path / "split.docx", plan, tmp_path / "redline.docx", final)
    actual = inspect_docx(final, full=True)
    assert actual.by_location("p0").text == HEAD
    assert actual.by_location("p1").text == MIDDLE
    assert "直销申购新政策。" in actual.by_location("p2").text
    assert actual.by_location("p3").text == "其他独立条款。"


def test_head_and_tail_discoveries_collapse_to_one_append_site(tmp_path):
    inspection = source(tmp_path, [HEAD, MIDDLE, END])
    targets = [Target("T" + str(i), "G01", "p" + str(i), text, "needs_edit", "definition",
        start=0, end=len(text), resolution="resolved", operations=("G01:0",))
        for i, text in ((0, HEAD), (2, END))]
    rebind_split_append_targets(targets, inspection)
    kept, removed = dedupe_operation_targets(targets, {"G01": {"operations": [{"op": "append"}]}}, inspection)
    active = [target for target in kept if not target.out_of_scope]
    assert len(active) == 1 and len(removed) == 1
    assert active[0].location == "p2"


@pytest.mark.parametrize("follower", ["其他条款与基金份额无关。", "类别。其他内容。"])
def test_unproven_continuation_blocks_append_inside_class_name(tmp_path, follower):
    inspection = source(tmp_path, [HEAD, follower])
    target = Target("T", "G01", "p0", HEAD, "needs_edit", "definition",
                    start=0, end=len(HEAD), resolution="resolved", operations=("G01:0",))
    with pytest.raises(RecallError, match="RECALL_APPEND_CLAUSE_INCOMPLETE"):
        plan_target({"rule_id": "G01", "operations": [dict(op="append", new_text="新政策。",
            condition={"kind": "none"})]}, target, inspection, Facts(), None, compose_budget=[0])


@pytest.mark.parametrize("rule_id", ["G01", "G03"])
def test_nav_formula_is_kept_with_definition_before_append(tmp_path, rule_id):
    head = ("本基金将基金份额分为A类、C类和D类三种不同的类别。"
            "各类基金份额将分别计算基金份额净值。计算公式为： ")
    formula = "计算日某类基金份额净值＝计算日该类基金份额的基金资产净值/计算日该类基金份额余额总数 "
    inspection = source(tmp_path, [head, formula, "投资者可自行选择基金份额类别。"])
    target = Target("T", rule_id, "p0", head, "needs_edit", "definition",
                    start=0, end=len(head), resolution="resolved", operations=(rule_id + ":0",))
    edit = plan_target({"rule_id": rule_id, "operations": [dict(op="append",
        new_text="直销申购新政策。", condition={"kind": "none"})]},
        target, inspection, Facts(), None, compose_budget=[0])
    assert edit.location == "p1" and edit.operation == "insert"
    plan = validate_edit_plan(tmp_path / "split.docx", build_docx_edits([edit], inspection), {},
                             content_check=content_check_factory([edit.as_dict()], inspection))
    generate_artifacts(tmp_path / "split.docx", plan, tmp_path / "redline.docx", tmp_path / "final.docx")
    actual = inspect_docx(tmp_path / "final.docx", full=True)
    assert actual.by_location("p0").text == head
    assert actual.by_location("p1").text.startswith(formula.rstrip())
    assert "直销申购新政策。" in actual.by_location("p1").text
    assert actual.by_location("p2").text == "投资者可自行选择基金份额类别。"


@pytest.mark.parametrize("follower", ["基金份额净值＝未知公式", "其他段落。"])
def test_formula_lead_does_not_authorize_unproven_follower(tmp_path, follower):
    head = "本基金将基金份额分为A类和C类。计算公式为："
    inspection = source(tmp_path, [head, follower])
    target = Target("T", "G03", "p0", head, "needs_edit", "definition",
                    start=0, end=len(head), resolution="resolved", operations=("G03:0",))
    with pytest.raises(RecallError, match="RECALL_APPEND_CLAUSE_INCOMPLETE"):
        plan_target({"rule_id": "G03", "operations": [dict(op="append", new_text="新政策。",
            condition={"kind": "none"})]}, target, inspection, Facts(), None, compose_budget=[0])


@pytest.mark.parametrize('split', [1, 2, 3, 4, 5, 7, 10])
def test_append_follows_fee_definition_broken_within_investor_intro(tmp_path, split):
    from fundlex.generation.claim_replay import replay_claims
    opening = '在投资者申购基金份额时'
    lead = '本基金根据申购费用收取方式的不同，将基金份额分为不同的类别。'
    continuation = '收取申购费用，而不计提销售服务费的，称为A类基金份额；不收取申购费用而计提销售服务费的，称为C类基金份额。 '
    head, tail = lead + opening[:split], opening[split:] + continuation
    inspection = source(tmp_path, [head, tail, '独立下一条。'])
    rule = {'rule_id':'G03','operations':[dict(op='append',new_text='直销申购新政策。',condition={'kind':'none'})]}
    target = Target('T','G03','p0',head,'needs_edit','definition',start=0,end=len(head),resolution='resolved',operations=('G03:0',))
    edit = plan_target(rule,target,inspection,Facts(),None,compose_budget=[0])
    assert edit.location == 'p1'
    plan = validate_edit_plan(tmp_path/'split.docx',build_docx_edits([edit],inspection),{},
        content_check=content_check_factory([edit.as_dict()],inspection))
    generate_artifacts(tmp_path/'split.docx',plan,tmp_path/'redline.docx',tmp_path/'final.docx')
    actual = inspect_docx(tmp_path/'final.docx',full=True)
    assert actual.by_location('p0').text == head
    assert actual.by_location('p1').text.startswith(tail.rstrip())
    assert actual.by_location('p2').text == '独立下一条。'
    proof = replay_claims([edit.as_dict()],inspection,actual,rules_by_id={'G03':rule},facts=Facts(),
        targets_by_id={'T':target.as_dict()})
    assert proof['verified'],proof['failures']
    wrong_path = tmp_path/'wrong.docx'
    doc = Document(tmp_path/'split.docx')
    doc.paragraphs[0].text = lead + '直销申购新政策。' + opening[:split]
    doc.save(wrong_path)
    wrong_edit = dict(edit.as_dict(),location='p0',start=len(lead),end=len(lead),replacement='直销申购新政策。')
    wrong = replay_claims([wrong_edit],inspection,inspect_docx(wrong_path,full=True),rules_by_id={'G03':rule},facts=Facts())
    assert not wrong['verified']
    assert any(f['code']=='CLAIM_OPERATION_SPAN_UNAUTHORIZED' for f in wrong['failures'])
