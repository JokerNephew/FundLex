import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from fundlex.compare import compare_docx, generate_selection_artifacts
from fundlex.docx import DOCUMENT, DocxError


def _save_doc(path: Path, paragraphs: list[str], *, table: tuple[str, str] | None = None) -> None:
    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    if table is not None:
        grid = document.add_table(rows=1, cols=2)
        grid.cell(0, 0).text = table[0]
        grid.cell(0, 1).text = table[1]
    document.save(path)


def _result(left: Path, right: Path) -> tuple[dict, str]:
    payload = compare_docx(left, right).to_dict()
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    return payload, hashlib.sha256(encoded).hexdigest()


def _generate(tmp_path: Path, left_text: list[str], right_text: list[str], decisions: dict[str, str], *, tables=None):
    left, right = tmp_path / "left.docx", tmp_path / "right.docx"
    _save_doc(left, left_text, table=tables[0] if tables else None)
    _save_doc(right, right_text, table=tables[1] if tables else None)
    result, result_hash = _result(left, right)
    choices = {diff["id"]: decisions.get(diff["id"], next(iter(decisions.values())))
               for diff in result["diffs"]}
    final, redline = tmp_path / "final.docx", tmp_path / "redline.docx"
    validation = generate_selection_artifacts(
        left, right, result, choices, final, redline,
        selection_epoch=3, result_sha256=result_hash)
    return left, right, final, redline, result, validation


def test_modify_choice_right_writes_real_final_and_redline(tmp_path):
    left, right, final, redline, result, validation = _generate(
        tmp_path, ["前文", "费率为0.40%", "后文"], ["前文", "费率为0.30%", "后文"], {"D0001": "right"})
    assert validation.final_sha256 == hashlib.sha256(final.read_bytes()).hexdigest()
    assert validation.redline_sha256 == hashlib.sha256(redline.read_bytes()).hexdigest()
    reopened = Document(final)
    assert [paragraph.text for paragraph in reopened.paragraphs] == ["前文", "费率为0.30%", "后文"]
    with ZipFile(redline) as archive:
        xml = archive.read(DOCUMENT)
    assert b"w:ins" in xml and b"w:del" in xml
    assert validation.manifest[0]["selected_text"] == "0.30%"


def test_consecutive_modify_paragraphs_are_written_independently(tmp_path):
    left, right, final, redline, result, validation = _generate(
        tmp_path,
        ["第一段费率为0.40%", "第二段费率为0.10%"],
        ["第一段费率为0.30%", "第二段费率为0.08%"],
        {"D0001": "right", "D0002": "left"},
    )
    # The two adjacent modified paragraphs share one alignment group; each
    # selected diff must still map to its own corresponding left paragraph.
    assert [diff["group_id"] for diff in result["diffs"]] == ["G0001", "G0001"]
    assert [paragraph.text for paragraph in Document(final).paragraphs] == [
        "第一段费率为0.30%", "第二段费率为0.10%"]
    assert len(validation.manifest) == 1
    with ZipFile(redline) as archive:
        xml = archive.read(DOCUMENT)
    assert xml.count(b"<w:ins ") == 1
    assert xml.count(b"<w:del ") == 1


def test_multiple_selected_modifications_in_one_group_are_written(tmp_path):
    _, _, final, _, _, validation = _generate(
        tmp_path,
        ["第一段费率为0.40%", "第二段费率为0.10%"],
        ["第一段费率为0.30%", "第二段费率为0.08%"],
        {"D0001": "right", "D0002": "right"},
    )
    assert [paragraph.text for paragraph in Document(final).paragraphs] == [
        "第一段费率为0.30%", "第二段费率为0.08%"]
    assert len(validation.manifest) == 2


def test_modify_choice_left_keeps_left_and_has_no_redline_for_unmodified_selection(tmp_path):
    left, right, final, redline, result, validation = _generate(
        tmp_path, ["费率为0.40%"], ["费率为0.30%"], {"D0001": "left"})
    assert Document(final).paragraphs[0].text == "费率为0.40%"
    # A selection that explicitly keeps the left side has no physical edit;
    # the redline artifact still reports that no right-side operation exists.
    assert validation.manifest == ()
    with ZipFile(redline) as archive:
        assert b"w:ins" not in archive.read(DOCUMENT)


def test_structural_insert_delete_is_blocked_before_output(tmp_path):
    left = tmp_path / "left.docx"
    right = tmp_path / "right.docx"
    _save_doc(left, ["前文", "尾部"])
    _save_doc(right, ["前文", "新增整段", "尾部"])
    result, result_hash = _result(left, right)
    decisions = {diff["id"]: "right" for diff in result["diffs"]}
    with pytest.raises(DocxError, match="SELECTION_STRUCTURAL_DIFF_UNSUPPORTED"):
        generate_selection_artifacts(left, right, result, decisions,
                                     tmp_path / "final.docx", tmp_path / "redline.docx",
                                     selection_epoch=1, result_sha256=result_hash)


def test_table_cell_selection_is_written_and_reopened(tmp_path):
    _, _, final, _, _, validation = _generate(
        tmp_path, [], [], {"D0001": "right"}, tables=(("条款内容：左侧", "固定"), ("条款内容：右侧", "固定")))
    assert Document(final).tables[0].cell(0, 0).text == "条款内容：右侧"
    assert Document(final).tables[0].cell(0, 1).text == "固定"
    assert validation.checks[-1] == "selected_text"


def test_pending_and_source_or_result_tampering_are_blocked(tmp_path):
    left = tmp_path / "left.docx"
    right = tmp_path / "right.docx"
    _save_doc(left, ["旧"],)
    _save_doc(right, ["新"],)
    result, result_hash = _result(left, right)
    with pytest.raises(DocxError, match="PENDING_SELECTIONS"):
        generate_selection_artifacts(left, right, result, {}, tmp_path / "f.docx", tmp_path / "r.docx",
                                     selection_epoch=1, result_sha256=result_hash)
    changed = dict(result)
    changed["diffs"] = [dict(result["diffs"][0], section="changed")]
    with pytest.raises(DocxError, match="SELECTION_RESULT_HASH_MISMATCH"):
        generate_selection_artifacts(left, right, changed, {"D0001": "right"}, tmp_path / "f.docx", tmp_path / "r.docx",
                                     selection_epoch=1, result_sha256=result_hash)


def test_structural_diff_is_blocked_before_output(tmp_path):
    left, right = tmp_path / "left.docx", tmp_path / "right.docx"
    _save_doc(left, ["前文", "尾部"])
    _save_doc(right, ["前文", "新增整段", "尾部"])
    result, result_hash = _result(left, right)
    decisions = {diff["id"]: "right" for diff in result["diffs"]}
    with pytest.raises(DocxError, match="SELECTION_STRUCTURAL_DIFF_UNSUPPORTED"):
        generate_selection_artifacts(left, right, result, decisions, tmp_path / "f.docx", tmp_path / "r.docx",
                                     selection_epoch=1, result_sha256=result_hash)
