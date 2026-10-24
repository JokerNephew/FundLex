"""Sweep repairs preserve target/plan ownership instead of restarting discovery."""
import pytest

from fundlex.recall import Target
from fundlex.recall_pipeline import _absorb_sweep_findings, _sweep_recovery_contexts


def test_bound_sweep_site_uses_generation_without_inventory_discovery():
    target = Target("T1", "G17", "p8", "旧用途句", "needs_edit", "",
                    start=0, end=4, resolution="resolved", operations=("G17:0",))
    contexts = _sweep_recovery_contexts([
        {"finding_id": "SW1", "rule_id": "G17", "location": "p8",
         "issue": "residual_old_text", "reason": "exact residual"}],
        [target], {"edits": [{"operation_keys": ["G17:0"], "location": "p8"}]},
        {"failures": []})
    assert len(contexts) == 1
    assert contexts[0].stage == "sweep"
    assert contexts[0].operation_id == "G17:0"
    assert contexts[0].site_id == "p8"
    assert contexts[0].resolved_class.value == "GENERATION"


def test_unbound_sweep_site_routes_to_target_discovery():
    contexts = _sweep_recovery_contexts([
        {"finding_id": "SW2", "rule_id": "G17", "location": "p9",
         "issue": "missing_new_text", "reason": "no bound target"}],
        [], {"edits": []}, {"failures": []})
    assert len(contexts) == 1
    assert contexts[0].resolved_class.value == "TARGET"
    assert contexts[0].error_code == "SWEEP_TARGET_UNBOUND"


def test_sweep_execution_failure_preserves_existing_final_plan_route():
    target = Target("T1", "G17", "p8", "旧用途句", "needs_edit", "",
                    start=0, end=4, resolution="resolved", operations=("G17:0",))
    contexts = _sweep_recovery_contexts([
        {"finding_id": "SW1", "rule_id": "G17", "location": "p8",
         "issue": "residual_old_text", "reason": "execution missed planned delete"}],
        [target], {"edits": [{"operation_keys": ["G17:0"], "location": "p8"}]},
        {"failures": [{"code": "CLAIM_EXECUTION_EFFECT_MISSING", "rule_id": "G17",
                       "operation_id": "G17:0", "location": "p8"}]})
    assert len(contexts) == 1
    assert contexts[0].resolved_class.value == "EXECUTION"


@pytest.mark.parametrize("code", ["CLAIM_OPERATION_NOT_PLANNED", "CLAIM_OPERATION_UNPROVEN"])
def test_sweep_claim_no_plan_with_existing_target_routes_generation(code):
    target = Target("T1", "G05", "p0", "旧句", "needs_edit", "",
                    start=0, end=2, resolution="resolved", operations=("G05:0",))
    contexts = _sweep_recovery_contexts([
        {"finding_id": "SW3", "rule_id": "G05", "location": "p0",
         "issue": "missing_new_text", "reason": "claim did not plan operation"}],
        [target], {"edits": []},
        {"failures": [{"code": code, "rule_id": "G05",
                        "operation_id": "G05:0", "location": "p0",
                        "target_id": "T1"}]})
    assert len(contexts) == 1
    assert contexts[0].resolved_class.value == "GENERATION"
    assert contexts[0].metadata["root_cause"] == "TARGET_EXISTS_NO_PLAN"


def test_sweep_variant_reactivates_absent_target_with_unique_delete_operation(tmp_path):
    """A verified G15 wording variant must become an executable repair target."""
    from docx import Document
    from fundlex.docx import inspect_docx

    text = ("基金销售服务费每日计提，逐日累计至每月月末，按月支付，经基金管理人与基金托管人核对一致后，"
            "由基金托管人于次月前 5 个工作日内从基金财产中一次性划出，由基金管理人代收，"
            "基金管理人收到后按相关合同规定支付给基金销售机构等。若遇法定节假日、公休假或不可抗力等，支付日期顺延。")
    source = tmp_path / "variant.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    old = Target("T1", "G15", "p0", text, "absent", "inventory variant", operations=("G15:0",))
    synthetic = []
    origins = {}
    rule = {"rule_id": "G15", "operations": [{"op": "delete", "old_text": "C类基金份额的销售服务费每日计提，按月支付。由基金管理人向基金托管人发送基金销售服务费划付指令，经基金托管人复核后于次月首日起10个工作日内从基金财产中一次性支付给登记机构，由登记机构代付给销售机构。若遇法定节假日、公休日等，支付日期顺延。", "condition": {"kind": "none"}}]}
    pending = [{"finding_id": "SW-G15", "rule_id": "G15", "location": "p0",
                "quote": text, "issue": "residual_old_text"}]

    added = _absorb_sweep_findings(
        pending, [old], synthetic, origins, inspection,
        rules_by_id={"G15": rule}, facts=type("Facts", (), {})())

    assert added == 1
    assert old.out_of_scope is True
    assert synthetic[0].state == "needs_edit"
    assert synthetic[0].operations == ("G15:0",)
    assert synthetic[0].resolution == "resolved"


def test_sweep_variant_does_not_rebind_absent_target_to_inactive_state(tmp_path):
    """A short matching fingerprint must also create an active repair target."""
    from docx import Document
    from fundlex.docx import inspect_docx

    text = "基金销售服务费每日计提，逐日累计至每月月末，按月支付。"
    source = tmp_path / "short-variant.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    old = Target("T1", "G15", "p0", text, "absent", "inventory variant", operations=("G15:0",))
    synthetic = []
    rule = {"rule_id": "G15", "operations": [{"op": "delete", "old_text": "基金销售服务费每日计提，按月支付。由基金管理人向基金托管人发送基金销售服务费划付指令。", "condition": {"kind": "none"}}]}

    added = _absorb_sweep_findings(
        [{"finding_id": "SW-G15-short", "rule_id": "G15", "location": "p0",
          "quote": text, "issue": "residual_old_text"}],
        [old], synthetic, {}, inspection, rules_by_id={"G15": rule},
        facts=type("Facts", (), {})())

    assert added == 1
    assert old.state == "absent"
    assert synthetic[0].state == "needs_edit"
    assert synthetic[0].resolution == "resolved"


def test_sweep_variant_with_multiple_operations_stays_fail_closed(tmp_path):
    """A location alone cannot choose between sibling operations."""
    from docx import Document
    from fundlex.docx import inspect_docx

    text = "基金销售服务费每日计提，逐日累计至每月月末，按月支付。"
    source = tmp_path / "ambiguous-variant.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    old = Target("T1", "G15", "p0", text, "absent", "inventory variant", operations=("G15:0",))
    synthetic = []
    rule = {"rule_id": "G15", "operations": [
        {"op": "delete", "old_text": "第一种销售服务费每日计提，按月支付。", "condition": {"kind": "none"}},
        {"op": "delete", "old_text": "第二种销售服务费每日计提，按月支付。", "condition": {"kind": "none"}},
    ]}

    added = _absorb_sweep_findings(
        [{"finding_id": "SW-G15-ambiguous", "rule_id": "G15", "location": "p0",
          "quote": text, "issue": "residual_old_text"}],
        [old], synthetic, {}, inspection, rules_by_id={"G15": rule},
        facts=type("Facts", (), {})())

    assert added == 0
    assert synthetic == []
    assert old.out_of_scope is False
