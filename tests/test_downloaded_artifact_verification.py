"""Downloaded files must prove the report and revision round trip independently."""
from hashlib import sha256
import json

from docx import Document

from fundlex.docx import generate_artifacts, inspect_docx, validate_edit_plan
from tools.evaluation.verify_downloaded_artifacts import verify_case


def test_download_verification_detects_file_and_report_drift(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("旧条款及不改内容。")
    document.add_paragraph("其他完整保留。")
    document.save(source)
    case = tmp_path / "D01"
    case.mkdir()
    inspected = inspect_docx(source, full=True)
    span = inspected.by_location("p0")
    edit = {"edit_id": "e1", "location": "p0", "zone": span.zone,
            "start": 0, "end": 3, "expected_old_text": "旧条款",
            "expected_old_hash": sha256("旧条款".encode()).hexdigest(),
            "replacement": "新条款", "template_id": "T", "verified": True,
            "fingerprint": span.fingerprint, "mapping_version": span.mapping_version,
            "source_sha256": inspected.source_sha256}
    plan = validate_edit_plan(source, [edit], {"T": "新条款"})
    validation = generate_artifacts(source, plan, case / "redline.docx", case / "final.docx")
    report = {"source_sha256": inspected.source_sha256, "delivery_status": "complete",
              "plan": {"edits": [edit]}, "output_validation": {
                  "final_sha256": validation.final_sha256,
                  "redline_sha256": validation.redline_sha256}}
    (case / "report.json").write_text(json.dumps(report), encoding="utf-8")
    checked = verify_case(case, source)
    assert checked["mechanical_pass"]
    assert checked["changed_paragraphs"][0]["after"] == "新条款及不改内容。"
    assert checked["business_complete"] == "not assessed by this tool"

    report["plan"]["edits"][0]["replacement"] = "未经写入的文字"
    (case / "report.json").write_text(json.dumps(report), encoding="utf-8")
    checked = verify_case(case, source)
    assert not checked["mechanical_pass"]
    assert checked["checks"]["final_hash"]
    assert not checked["checks"]["plan_matches_actual"]

    (case / "final.docx").write_bytes(source.read_bytes())
    checked = verify_case(case, source)
    assert not checked["checks"]["final_hash"]
    assert not checked["checks"]["accepted_matches_final"]
