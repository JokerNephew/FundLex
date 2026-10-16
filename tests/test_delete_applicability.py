"""A delete the model inventory judged "absent" must not silence the survey.

G17's reference is two sentences with a class prefix
("C类基金份额的销售服务费主要用于支付销售机构佣金、…。C类基金份额的销售服务费使用范围不
包括基金募集期间的上述费用。"), but contracts carry the clause without the prefix and
with a differently worded second sentence — sometimes with no second sentence at all.
The model therefore reports ``absent`` and records a target at the very location the
reference text still occupies, which made ``survey_targets`` skip the hit and dropped
the deletion (D04/D08/D17/D19/D20, FL-D101).

These tests pin the override in both directions: a rule whose clause is genuinely
present gets a target even though an ``absent`` target covers that location, and a
rule that already has a resolved ``needs_edit`` target is not re-surfaced.
"""

from pathlib import Path

from docx import Document

from fundlex.docx import inspect_docx
from fundlex.recall import ClassFact, Facts, is_delete_only_rule, survey_targets

G17_OLD = ("C类基金份额的销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、"
           "促销活动费、基金份额持有人服务费等。C类基金份额的销售服务费使用范围不包括基金募集"
           "期间的上述费用。")

G17_RULE = {
    "rule_id": "G17",
    "title": "基金费用与税收-3、基金销售服务费（有多个份额类别时适用）",
    "requirement": f"删除原条款“{G17_OLD}”。",
    "operations": [{"op": "delete", "old_text": G17_OLD,
                    "condition": {"kind": "none", "value": None}}],
}

# The contract's own wording: no class prefix, a differently worded (or absent) second
# sentence, exactly as the real dev set showed.
CONTRACT_CLAUSE = ("销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、"
                   "促销活动费、基金份额持有人服务费等。")


def _facts() -> Facts:
    return Facts(fund_type="open_regular", fund_type_quote="", fund_type_basis="",
                 classes=(ClassFact("A类基金份额", True, False, True, None, ""),
                          ClassFact("C类基金份额", False, True, True, None, "")),
                 floating_management_fee=False, floating_quote="", floating_basis="",
                 notes="", quotes_verified={}, problems=())


def _document(tmp_path: Path) -> Path:
    path = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("第五部分  基金费用与税收")
    document.add_paragraph(CONTRACT_CLAUSE)
    document.add_paragraph("基金管理人可以在基金合同约定的范围内调整费率。")
    document.save(path)
    return path


def test_delete_only_rule_detection():
    assert is_delete_only_rule(G17_RULE, _facts()) is True
    # A replace operation makes it not delete-only.
    assert is_delete_only_rule(
        {"rule_id": "G18", "operations": [
            {"op": "replace", "old_text": "a", "new_text": "b",
             "condition": {"kind": "none", "value": None}}]}, _facts()) is False
    # A noop branch alongside a delete is still delete-only.
    assert is_delete_only_rule(
        {"rule_id": "G17", "operations": [
            {"op": "delete", "old_text": G17_OLD, "condition": {"kind": "none", "value": None}},
            {"op": "noop", "condition": {"kind": "class_count", "value": 1}}]}, _facts()) is True


def test_absent_target_does_not_suppress_the_delete(tmp_path: Path):
    """The inventory's "absent" target sits at the clause's own location.

    Without the override the survey skips that (rule, location) pair and the deletion
    is dropped; with it the reference still resolves and a target is produced.
    """
    inspection = inspect_docx(_document(tmp_path), full=True)
    located = [span.location for span in inspection.paragraphs
               if "销售服务费主要用于" in span.text]
    assert located, "fixture must carry the contract clause"
    known = [("G17", location) for location in located]      # the model's "absent" targets

    suppressed = survey_targets(inspection, [G17_RULE], _facts(), known=known)
    assert not [hit for hit in suppressed if hit.rule_id == "G17"]

    recovered = survey_targets(inspection, [G17_RULE], _facts(), known=known,
                               override_rules=frozenset({"G17"}))
    hits = [hit for hit in recovered if hit.operation == "survey"]
    assert [hit.rule_id for hit in hits] == ["G17"]
    assert hits[0].locations[0] in located


def test_resolved_needs_edit_target_keeps_its_coverage(tmp_path: Path):
    """A rule already being edited must not be re-surfaced by the override.

    ``_run_survey`` only overrides rules that have no resolved ``needs_edit`` target, so
    a rule that is genuinely being edited is unaffected by the widened set.
    """
    inspection = inspect_docx(_document(tmp_path), full=True)
    located = [span.location for span in inspection.paragraphs
               if "销售服务费主要用于" in span.text]

    class _Target:
        rule_id = "G17"
        state = "needs_edit"
        resolution = "resolved"
        composite: tuple[str, ...] = ()

        def __init__(self, location):
            self.location = location

    targets = [_Target(location) for location in located]
    resolved_needs_edit = {target.rule_id for target in targets
                           if target.state == "needs_edit" and target.resolution == "resolved"}
    override = frozenset({"G17"}) - resolved_needs_edit
    assert override == frozenset()

    hits = survey_targets(inspection, [G17_RULE], _facts(),
                          known=[("G17", location) for location in located],
                          override_rules=override)
    assert not hits
