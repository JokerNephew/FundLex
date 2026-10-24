"""Composite slice replacements must not reinsert executor-preserved text (FL-D210)."""
from dataclasses import replace

import pytest
import json
from pathlib import Path
from docx import Document

from fundlex.docx import generate_artifacts, inspect_docx, validate_edit_plan
from fundlex.generation.claim_replay import replay_claims
from fundlex.obligations import apply_facts, compile_obligations
from fundlex.recall import (Target, _component_ok, build_docx_edits, content_check_factory,
                            find_old_text_all, plan_composite_target)
from test_recall_core import _facts


PREFIX = "若遇法定节假日、休息日或不可抗力致使无法按时支付等，支付日期顺延。"
HEAD = "销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、促销活动费、基金份额持有人服务费等。"
TAIL = "销售服务费使用范围不包括基金募集期间的上述费用。"


@pytest.mark.parametrize("semicolon", [False, True])
@pytest.mark.parametrize("stale", [False, True])
def test_g17_survey_and_residual_keep_every_split_copy_narrow(tmp_path, monkeypatch, semicolon, stale):
    """Actual survey→targets→planner→DOCX covers all copies without suffix loss."""
    from fundlex.recall import (assign_target_operations, dedupe_operation_targets,
                                ensure_delete_targets, survey_targets)
    from fundlex.recall_pipeline import _plan_and_generate, _survey_targets_from_hits
    from test_g17_usage_site_closure import NoModel
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "registry")
    rule = next(g for g in json.loads((Path(__file__).resolve().parents[1] /
        "guides/builtin_guide_v1.json").read_text())["groups"] if g["rule_id"] == "G17")
    head = HEAD[:-1] + "；" if semicolon else HEAD
    split = 34 if semicolon else 32
    suffix = ("基金管理人可根据基金发展情况下调销售服务费率或为促进销售特定时间内对销售服务费率进行打折优惠。"
              if semicolon else "上述费用按费用实际支出金额列入或摊入当期费用。")
    texts = [head[:split], head[split:] + suffix + " ", TAIL, "副本之间的独立记录。"] * 2
    source = tmp_path / "source.docx"
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(source)
    original = inspect_docx(source, full=True)
    facts = _facts()
    ledger = compile_obligations({"groups": [rule]})
    apply_facts(ledger, facts)
    if stale:
        targets = [Target(f"W{i}", "G17", f"p{i}", texts[i], "needs_edit", "wide_inventory",
            operations=("G17:0",), resolution="composite",
            composite=(f"p{i}:0:{len(texts[i])}", f"p{i+1}:0:{len(texts[i+1].strip())}"))
            for i in (0, 4)]
    else:
        hits = survey_targets(original, [rule], facts, known=[])
        targets, gaps = _survey_targets_from_hits(hits, {"G17": rule}, original, [], start_index=0)
        assert not gaps
        assert all(int(t.composite[-1].rsplit(":", 1)[1]) == len(head[split:])
                   for t in targets if t.composite)
    targets.extend(ensure_delete_targets(targets, {"G17": rule}, original))
    assign_target_operations(targets, {"G17": rule}, original, facts, ledger)
    targets, _ = dedupe_operation_targets(targets, {"G17": rule}, original)
    plan, _, _, _, actual = _plan_and_generate(NoModel(), {"G17": rule}, targets,
        original, facts, source, tmp_path, max_compose_calls=0, obligations=ledger)
    for i in (0, 4):
        assert not actual.by_location(f"p{i}").text.strip()
        assert actual.by_location(f"p{i+1}").text == suffix + " "
        assert not actual.by_location(f"p{i+2}").text.strip()
        assert actual.by_location(f"p{i+3}").text == texts[i+3]
    proof = replay_claims(plan["edits"], original, actual, rules_by_id={"G17": rule},
        facts=facts, obligations=ledger, targets_by_id={t.target_id: t.as_dict() for t in targets})
    assert proof["verified"], proof["failures"]


@pytest.mark.parametrize("new", ["", "新的费用用途约定。"])
def test_composite_actual_docx_keeps_prefix_suffix_once(tmp_path, new):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph(PREFIX + HEAD + " ")
    doc.add_paragraph(TAIL + "计算结果应当保留。")
    doc.save(source)
    original = inspect_docx(source, full=True)
    rule = {"rule_id": "G17", "operations": [dict(op="replace" if new else "delete",
        old_text=HEAD + TAIL, new_text=new, condition={"kind": "none"})]}
    parts = find_old_text_all(list(original.paragraphs), HEAD + TAIL)[0]
    target = Target("T", "G17", "p0", HEAD + TAIL, "needs_edit", "source",
        composite=tuple(f"{loc}:{start}:{end}" for loc, start, end in parts),
        operations=("G17:0",))
    edits = plan_composite_target(rule, target, original, _facts(), new_text=new)
    for edit in edits:
        edit.operation_keys = ("G17:0",)
    plan = validate_edit_plan(source, build_docx_edits(edits, original), {},
        content_check=content_check_factory([edit.as_dict() for edit in edits], original))
    assert len(plan.edits) == 2
    generate_artifacts(source, plan, tmp_path / "redline.docx", tmp_path / "final.docx")
    actual = inspect_docx(tmp_path / "final.docx", full=True)
    assert actual.by_location("p0").text == PREFIX + new + " "
    assert actual.by_location("p1").text == "计算结果应当保留。"
    ledger = compile_obligations({"groups": [rule]})
    apply_facts(ledger, _facts())
    rows = [edit.as_dict() for edit in edits]
    for row in rows:
        row["components"] = [{**row, "expected_old_text": original.by_location(row["location"]).text[
            row["start"]:row["end"]]}]
    proof = replay_claims(rows, original, actual, rules_by_id={"G17": rule},
                          targets_by_id={"T": target.as_dict()}, obligations=ledger, facts=_facts())
    assert proof["verified"], proof["failures"]
    if not new:
        # Reproduce the old producer error without altering source scope/parts.
        bad = replace(edits[0], replacement=PREFIX + " ")
        component = {**bad.as_dict(), "expected_old_text": HEAD}
        assert not _component_ok(component, original)
        rows[0] = {**component, "components": [component]}
        duplicate_doc = Document(tmp_path / "final.docx")
        duplicate_doc.paragraphs[0].text = PREFIX + PREFIX + "  "
        duplicate_doc.save(tmp_path / "bad.docx")
        bad_proof = replay_claims(rows, original, inspect_docx(tmp_path / "bad.docx", full=True),
            rules_by_id={"G17": rule}, targets_by_id={"T": target.as_dict()},
            obligations=ledger, facts=_facts())
        assert not bad_proof["verified"]
        assert any(failure["code"] == "CLAIM_OPERATION_SPAN_UNAUTHORIZED"
                    for failure in bad_proof["failures"])


@pytest.mark.parametrize("forged", [False, True])
def test_g17_sentence_split_parts_are_proved_from_original(tmp_path, forged):
    """A split sentence is authorized; counterfeit parts cannot absorb its suffix."""
    from fundlex.recall import RecallError
    source = tmp_path / "split.docx"
    doc = Document()
    split = len(HEAD) // 2
    doc.add_paragraph(PREFIX + HEAD[:split])
    doc.add_paragraph(HEAD[split:] + TAIL + "独立核对记录应保留。")
    doc.save(source)
    original = inspect_docx(source, full=True)
    rule = {"rule_id": "G17", "operations": [dict(op="delete", old_text=HEAD + TAIL,
             condition={"kind": "none"})]}
    parts = find_old_text_all(list(original.paragraphs), HEAD + TAIL)[0]
    target = Target("T", "G17", "p0", HEAD + TAIL, "needs_edit", "source",
        composite=tuple(f"{loc}:{start}:{end}" for loc, start, end in parts), operations=("G17:0",))
    edits = plan_composite_target(rule, target, original, _facts(), new_text="")
    for edit in edits:
        edit.operation_keys = ("G17:0",)
    if forged:
        tail = original.by_location("p1").text
        edits[1].end = len(tail)
        fake_parts = [target.composite[0], f"p1:0:{len(tail)}"]
        edits[1].evidence.update(parts=fake_parts, deleted_paragraph=tail)
        with pytest.raises(RecallError, match="RECALL_DELETE_SCOPE_UNPROVEN"):
            plan_composite_target(rule, replace(target, composite=tuple(fake_parts)),
                                  original, _facts(), new_text="")
    rows = [edit.as_dict() for edit in edits]
    for row in rows:
        row["expected_old_text"] = original.by_location(row["location"]).text[row["start"]:row["end"]]
    # Independently produce the claimed artifact, including the forged loss.
    doc.paragraphs[0].text = PREFIX
    doc.paragraphs[1].text = "" if forged else "独立核对记录应保留。"
    final = tmp_path / "actual.docx"
    doc.save(final)
    proof = replay_claims(rows, original, inspect_docx(final, full=True), rules_by_id={"G17": rule})
    assert proof["verified"] is not forged, proof["failures"]
    if forged:
        assert any(f["code"] == "CLAIM_OPERATION_SPAN_UNAUTHORIZED" for f in proof["failures"])
