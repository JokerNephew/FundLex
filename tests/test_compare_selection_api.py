"""Compare selection draft, generation, versioning and download API.

The browser and revision engine are not involved. These tests use real DOCX
files, the real Compare worker and the real OOXML selection writer, then re-read
the downloaded artifact through the HTTP API.
"""
from __future__ import annotations

import hashlib
import io
import json
import os

import pytest
from docx import Document
from fastapi.testclient import TestClient

from fundlex.api import create_app
from fundlex.compare.service import CompareService
from fundlex.compare.store import CompareStore
from fundlex.compare.worker import CompareWorker
from fundlex.db import Database
from fundlex.store import RunStore


def _docx(paragraphs: list[str]) -> bytes:
    stream = io.BytesIO()
    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    document.save(stream)
    return stream.getvalue()


def _pair(client: TestClient, left: list[str], right: list[str],
          user: str = "alice") -> tuple[str, dict]:
    headers = {"X-FundLex-User": user}
    left_bytes, right_bytes = _docx(left), _docx(right)
    payload = {"name": "selection", "pairs": [{
        "client_pair_id": "p", "label": "P",
        "left": {"client_id": "l", "name": "left.docx", "size": len(left_bytes)},
        "right": {"client_id": "r", "name": "right.docx", "size": len(right_bytes)},
    }]}
    created = client.post("/api/compare/tasks", headers=headers, json=payload)
    assert created.status_code == 201, created.text
    task_id = created.json()["id"]
    uploaded = client.post(f"/api/compare/tasks/{task_id}/uploads", headers=headers,
                           data={"client_ids": json.dumps(["l", "r"])},
                           files=[("files", ("left.docx", left_bytes)),
                                  ("files", ("right.docx", right_bytes))])
    assert uploaded.status_code == 200, uploaded.text
    assert client.post(f"/api/compare/tasks/{task_id}/seal", headers=headers).status_code == 200
    assert CompareWorker(client.app.state.compare_store).run_once()
    pair = client.get(f"/api/compare/tasks/{task_id}", headers=headers).json()["pairs"][0]
    assert pair["status"] == "succeeded", pair
    return pair["id"], headers


def _binding(state: dict) -> dict:
    return {"result_sha256": state["result_sha256"],
            "left_source_sha256": state["left_source_sha256"],
            "right_source_sha256": state["right_source_sha256"]}


def _save(client: TestClient, pair_id: str, headers: dict, state: dict,
          decisions: dict[str, str]) -> dict:
    response = client.put(f"/api/compare/pairs/{pair_id}/selection", headers=headers,
                          json={"expected_epoch": state["selection_epoch"],
                                **_binding(state), "decisions": decisions})
    assert response.status_code == 200, response.text
    return response.json()


def _generate(client: TestClient, pair_id: str, headers: dict, state: dict) -> dict:
    response = client.post(f"/api/compare/pairs/{pair_id}/selection/generate",
                           headers=headers,
                           json={"expected_epoch": state["selection_epoch"],
                                 "decision_hash": state["draft"]["decision_hash"],
                                 "result_sha256": state["result_sha256"]})
    assert response.status_code == 200, response.text
    return response.json()


def test_save_restore_generate_and_download_real_docx(tmp_path):
    client = TestClient(create_app(tmp_path))
    pair_id, headers = _pair(client, ["前文", "费率为0.40%", "后文"],
                             ["前文", "费率为0.30%", "后文"])
    state = client.get(f"/api/compare/pairs/{pair_id}/selection", headers=headers).json()
    assert state["selection_epoch"] == 0
    assert state["decisions"] == {"D0001": "pending"}
    assert state["can_generate"] is False
    assert state["diffs"][0]["left_text"] == "0.40%"
    assert state["artifacts"] == []

    state = _save(client, pair_id, headers, state, {"D0001": "right"})
    assert state["selection_epoch"] == 1
    assert state["decisions"] == {"D0001": "right"}
    assert state["can_generate"] is True
    assert state["assessment"]["counts"] == {"pending": 0, "left": 0, "right": 1}

    # Repeated identical saves are idempotent and keep the same epoch.
    again = _save(client, pair_id, headers, state, {"D0001": "right"})
    assert again["selection_epoch"] == 1
    # Omitted ids are merged from the draft rather than reset to pending.
    incremental = client.put(f"/api/compare/pairs/{pair_id}/selection", headers=headers,
                             json={"expected_epoch": 1, **_binding(state), "decisions": {}})
    assert incremental.status_code == 200
    assert incremental.json()["selection_epoch"] == 1
    assert incremental.json()["decisions"] == {"D0001": "right"}

    preview = client.post(f"/api/compare/pairs/{pair_id}/selection/preview",
                          headers=headers, json={}).json()
    assert preview["can_generate"] is True

    artifact = _generate(client, pair_id, headers, state)
    assert artifact["version"] == 1 and artifact["current"] is True
    final = artifact["files"]["final"]
    redline = artifact["files"]["redline"]
    assert len(final["sha256"]) == 64 and len(redline["sha256"]) == 64
    assert artifact["report"]["decision_hash"] == state["draft"]["decision_hash"]
    assert "final_docx_reopen" in artifact["report"]["checks"]

    # Retrying the same generation is idempotent: no second version is created.
    repeated = _generate(client, pair_id, headers, state)
    assert repeated["version"] == 1
    assert repeated["files"]["final"]["sha256"] == final["sha256"]

    download = client.get(final["download_url"], headers=headers)
    assert download.status_code == 200
    assert download.headers["X-Content-SHA256"] == final["sha256"]
    assert hashlib.sha256(download.content).hexdigest() == final["sha256"]
    reopened = Document(io.BytesIO(download.content))
    assert [p.text for p in reopened.paragraphs] == ["前文", "费率为0.30%", "后文"]
    assert client.get(final["download_url"], headers={"X-FundLex-User": "bob"}).status_code == 404

    # A changed decision advances the epoch and marks the old version stale.
    state = _save(client, pair_id, headers, state, {"D0001": "left"})
    assert state["selection_epoch"] == 2
    old = next(item for item in state["artifacts"] if item["version"] == 1)
    assert old["stale"] is True and old["current"] is False
    assert client.get(final["download_url"], headers=headers).status_code == 200

    stale_request = client.post(f"/api/compare/pairs/{pair_id}/selection/generate",
                                headers=headers,
                                json={"expected_epoch": 1,
                                      "decision_hash": old["decision_hash"],
                                      "result_sha256": state["result_sha256"]})
    assert stale_request.status_code == 409
    assert stale_request.json()["detail"] == "SELECTION_EPOCH_CONFLICT"

    current = _generate(client, pair_id, headers, state)
    assert current["version"] == 2
    download = client.get(current["files"]["final"]["download_url"], headers=headers)
    assert [p.text for p in Document(io.BytesIO(download.content)).paragraphs] == [
        "前文", "费率为0.40%", "后文"]


def test_invalid_selection_hashes_epochs_and_cross_user_are_rejected(tmp_path):
    client = TestClient(create_app(tmp_path))
    pair_id, headers = _pair(client, ["费率为0.40%"], ["费率为0.30%"])
    state = client.get(f"/api/compare/pairs/{pair_id}/selection", headers=headers).json()
    url = f"/api/compare/pairs/{pair_id}/selection"

    unknown = client.put(url, headers=headers, json={
        "expected_epoch": 0, **_binding(state), "decisions": {"D9999": "right"}})
    assert unknown.status_code == 422 and unknown.json()["detail"] == "UNKNOWN_DIFF_ID"

    changed_result = client.put(url, headers=headers, json={
        "expected_epoch": 0, "result_sha256": "0" * 64,
        "left_source_sha256": state["left_source_sha256"],
        "right_source_sha256": state["right_source_sha256"],
        "decisions": {"D0001": "right"}})
    assert changed_result.status_code == 409
    assert changed_result.json()["detail"] == "COMPARE_RESULT_CHANGED"

    changed_source = client.put(url, headers=headers, json={
        "expected_epoch": 0, "result_sha256": state["result_sha256"],
        "left_source_sha256": "0" * 64,
        "right_source_sha256": state["right_source_sha256"],
        "decisions": {"D0001": "right"}})
    assert changed_source.status_code == 409
    assert changed_source.json()["detail"] == "COMPARE_SOURCE_CHANGED"

    stale_epoch = client.put(url, headers=headers, json={
        "expected_epoch": 5, **_binding(state), "decisions": {"D0001": "right"}})
    assert stale_epoch.status_code == 409
    assert stale_epoch.json()["detail"] == "SELECTION_EPOCH_CONFLICT"

    state = _save(client, pair_id, headers, state, {"D0001": "pending"})
    pending = client.post(f"{url}/generate", headers=headers, json={
        "expected_epoch": state["selection_epoch"],
        "decision_hash": state["draft"]["decision_hash"],
        "result_sha256": state["result_sha256"]})
    assert pending.status_code == 422
    assert pending.json()["detail"] == "PENDING_SELECTIONS"

    other = {"X-FundLex-User": "bob"}
    assert client.get(url, headers=other).status_code == 404
    assert client.put(url, headers=other, json={
        "expected_epoch": 0, **_binding(state), "decisions": {}}).status_code == 404
    assert client.get(f"{url}-artifacts/1/final", headers=other).status_code == 404


MYSQL_URL = os.environ.get("FUNDLEX_COMPARE_MYSQL_TEST_URL", "").strip()


@pytest.fixture
def mysql_compare_store(tmp_path):
    if not MYSQL_URL:
        pytest.skip("FUNDLEX_COMPARE_MYSQL_TEST_URL is not configured")
    with Database(MYSQL_URL).connect() as db:
        for table in ("compare_selection_artifacts", "compare_selection_drafts",
                      "compare_artifacts", "compare_pair_files", "compare_pairs", "compare_tasks"):
            db.execute(f"DROP TABLE IF EXISTS {table}")
    yield CompareStore(RunStore(tmp_path, database_url=MYSQL_URL))


def test_mysql_selection_generation_and_stale_versions(mysql_compare_store):
    store = mysql_compare_store
    left, right = _docx(["费率为0.40%"]), _docx(["费率为0.30%"])
    task = store.create("alice", "mysql-selection", [{
        "client_pair_id": "p", "label": "P",
        "left": {"client_id": "l", "name": "left.docx", "size": len(left)},
        "right": {"client_id": "r", "name": "right.docx", "size": len(right)},
    }])
    store.upload("alice", task["id"], [
        dict(client_id="l", name="left.docx", data=left),
        dict(client_id="r", name="right.docx", data=right)])
    store.seal("alice", task["id"])
    job = store.claim()
    assert job is not None
    files = {item["side"]: item for item in store.pair_inputs(job["id"])}
    result = CompareService().compare(
        store.store.object_path(files["left"]["object_key"]),
        store.store.object_path(files["right"]["object_key"]),
        left_filename="left.docx", right_filename="right.docx")
    store.finish(job["id"], job["fencing_token"], result)
    pair = store.get_pair("alice", job["id"])
    assert pair["status"] == "succeeded"

    state = store.selection_state("alice", job["id"])
    decisions = {item["id"]: "right" for item in state["diffs"]}
    state = store.save_selection(
        "alice", job["id"], expected_epoch=0, result_sha256=state["result_sha256"],
        left_source_sha256=state["left_source_sha256"],
        right_source_sha256=state["right_source_sha256"], decisions=decisions)
    artifact = store.generate_selection(
        "alice", job["id"], expected_epoch=state["selection_epoch"],
        expected_decision_hash=state["draft"]["decision_hash"],
        result_sha256=state["result_sha256"])
    path, row = store.selection_artifact("alice", job["id"], artifact["version"], "final")
    assert hashlib.sha256(path.read_bytes()).hexdigest() == row["sha256"]
    assert [p.text for p in Document(path).paragraphs] == ["费率为0.30%"]

    changed = store.save_selection(
        "alice", job["id"], expected_epoch=state["selection_epoch"],
        result_sha256=state["result_sha256"],
        left_source_sha256=state["left_source_sha256"],
        right_source_sha256=state["right_source_sha256"],
        decisions={item["id"]: "left" for item in state["diffs"]})
    old = next(item for item in changed["artifacts"] if item["version"] == artifact["version"])
    assert old["stale"] is True


def test_structural_selection_and_tampered_artifact_fail_closed(tmp_path):
    client = TestClient(create_app(tmp_path))
    pair_id, headers = _pair(client, ["前文", "尾部"], ["前文", "新增整段", "尾部"])
    state = client.get(f"/api/compare/pairs/{pair_id}/selection", headers=headers).json()
    decisions = {item["id"]: "right" for item in state["diffs"]}
    state = _save(client, pair_id, headers, state, decisions)
    blocked = client.post(f"/api/compare/pairs/{pair_id}/selection/generate", headers=headers,
                          json={"expected_epoch": state["selection_epoch"],
                                "decision_hash": state["draft"]["decision_hash"],
                                "result_sha256": state["result_sha256"]})
    assert blocked.status_code == 422
    assert blocked.json()["detail"] == "SELECTION_STRUCTURAL_DIFF_UNSUPPORTED"

    # A tampered object is refused even though its database row still exists.
    pair_id, headers = _pair(client, ["费率为0.40%"], ["费率为0.30%"])
    state = client.get(f"/api/compare/pairs/{pair_id}/selection", headers=headers).json()
    state = _save(client, pair_id, headers, state, {"D0001": "right"})
    artifact = _generate(client, pair_id, headers, state)
    store = client.app.state.compare_store
    with store._connect() as db:
        row = db.execute("""SELECT object_key FROM compare_selection_artifacts
            WHERE pair_id=? AND version=1 AND kind='final'""", (pair_id,)).fetchone()
    store.store.object_path(row["object_key"]).write_bytes(b"tampered")
    response = client.get(artifact["files"]["final"]["download_url"], headers=headers)
    assert response.status_code == 409
    assert response.json()["detail"] == "COMPARE_SELECTION_ARTIFACT_INVALID"
