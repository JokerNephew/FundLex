import hashlib, io, json, zipfile
from docx import Document
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient
from lxml import etree
from fundlex.api import create_app
from fundlex.compare.revisions import prepare_compare_input
from fundlex.compare.worker import CompareWorker
from fundlex.docx import Q

def docx(text):
    f=io.BytesIO(); d=Document(); d.add_paragraph(text); d.save(f); return f.getvalue()


def revised_docx(old='原文', new='新文'):
    source = io.BytesIO()
    document = Document()
    document.add_paragraph(old)
    document.save(source)
    with zipfile.ZipFile(io.BytesIO(source.getvalue())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts['word/document.xml'])
    paragraph = root.find('.//' + Q('p'))
    run = paragraph.find(Q('r'))
    text = run.find(Q('t'))
    deleted = OxmlElement('w:del')
    deleted.set(Q('id'), '1')
    deleted_run = OxmlElement('w:r')
    deleted_text = OxmlElement('w:delText')
    deleted_text.text = text.text
    deleted_run.append(deleted_text)
    deleted.append(deleted_run)
    inserted = OxmlElement('w:ins')
    inserted.set(Q('id'), '2')
    inserted_run = OxmlElement('w:r')
    inserted_text = OxmlElement('w:t')
    inserted_text.text = new
    inserted_run.append(inserted_text)
    inserted.append(inserted_run)
    paragraph.remove(run)
    paragraph.extend([deleted, inserted])
    parts['word/document.xml'] = etree.tostring(root)
    revised = io.BytesIO()
    with zipfile.ZipFile(revised, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return revised.getvalue()

def test_compare_api_roundtrip_and_xss(tmp_path):
    a,b=docx('A <script>alert(1)</script>'),docx('B')
    c=TestClient(create_app(tmp_path)); h={'X-FundLex-User':'alice'}
    payload={'name':'cmp','pairs':[{'client_pair_id':'p','label':'P','left':{'client_id':'l','name':'a.docx','size':len(a)},'right':{'client_id':'r','name':'b.docx','size':len(b)}}]}
    created=c.post('/api/compare/tasks',headers=h,json=payload); assert created.status_code==201
    tid=created.json()['id']; c.post(f'/api/compare/tasks/{tid}/uploads',headers=h,data={'client_ids':'["l","r"]'},files=[('files',('a.docx',a)),('files',('b.docx',b))])
    c.post(f'/api/compare/tasks/{tid}/seal',headers=h); from fundlex.compare.worker import CompareWorker
    assert CompareWorker(c.app.state.compare_store).run_once(); pid=created.json()['pairs'][0]['id']
    view=c.get(f'/api/compare/pairs/{pid}/view',headers=h); assert view.status_code==200
    assert '<script>' not in view.json()['left']['html']
    assert c.get(f'/api/compare/pairs/{pid}',headers={'X-FundLex-User':'bob'}).status_code==404


def test_compare_upload_accepts_revisions_and_returns_audit_metadata(tmp_path):
    revised = revised_docx()
    normalized = prepare_compare_input(revised).data
    client = TestClient(create_app(tmp_path / 'store'))
    headers = {'X-FundLex-User': 'alice'}
    payload = {'name': 'cmp', 'pairs': [{'client_pair_id': 'p', 'label': 'P',
        'left': {'client_id': 'l', 'name': 'a.docx', 'size': len(revised)},
        'right': {'client_id': 'r', 'name': 'b.docx', 'size': len(normalized)}}]}
    created = client.post('/api/compare/tasks', headers=headers, json=payload)
    assert created.status_code == 201
    task_id = created.json()['id']
    uploaded = client.post(f'/api/compare/tasks/{task_id}/uploads', headers=headers,
                           data={'client_ids': '["l","r"]'},
                           files=[('files', ('a.docx', revised)), ('files', ('b.docx', normalized))])
    assert uploaded.status_code == 200, uploaded.text
    receipts = {item['client_id']: item for item in uploaded.json()['receipts']}
    normalized_sha256 = hashlib.sha256(normalized).hexdigest()
    assert receipts['l']['status'] == 'accepted'
    assert receipts['l']['revisions_accepted'] == 3
    assert receipts['l']['original_sha256'] == hashlib.sha256(revised).hexdigest()
    assert receipts['l']['original_size'] == len(revised)
    assert receipts['l']['normalized_sha256'] == normalized_sha256
    assert receipts['l']['normalized_size'] == len(normalized)
    assert receipts['r']['revisions_accepted'] == 0
    pair = uploaded.json()['task']['pairs'][0]
    assert pair['left']['revisions_accepted'] == 3
    assert pair['left']['original_size'] == len(revised)
    assert pair['left']['size'] == len(normalized)
    assert pair['left']['normalized_sha256'] == normalized_sha256

    # Retrying the same original upload is idempotent and must not replace the
    # stored normalized object with a differently-timestamped copy.
    retried = client.post(f'/api/compare/tasks/{task_id}/uploads', headers=headers,
                          data={'client_ids': '["l"]'},
                          files=[('files', ('a.docx', revised))])
    assert retried.status_code == 200, retried.text
    retry_receipt = retried.json()['receipts'][0]
    assert retry_receipt['status'] == 'already_accepted'
    assert retry_receipt['normalized_sha256'] == normalized_sha256
    assert retried.json()['task']['pairs'][0]['left']['size'] == len(normalized)


def test_compare_upload_rejects_unsafe_revision_input(tmp_path):
    revised = revised_docx()
    with zipfile.ZipFile(io.BytesIO(revised)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    settings = etree.fromstring(parts['word/settings.xml'])
    settings.insert(0, OxmlElement('w:documentProtection'))
    parts['word/settings.xml'] = etree.tostring(settings)
    protected = io.BytesIO()
    with zipfile.ZipFile(protected, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    protected = protected.getvalue()
    right = docx('新文')
    client = TestClient(create_app(tmp_path / 'invalid'))
    headers = {'X-FundLex-User': 'alice'}
    payload = {'name': 'invalid', 'pairs': [{'client_pair_id': 'p', 'label': 'P',
        'left': {'client_id': 'l', 'name': 'left.docx', 'size': len(protected)},
        'right': {'client_id': 'r', 'name': 'right.docx', 'size': len(right)}}]}
    created = client.post('/api/compare/tasks', headers=headers, json=payload)
    assert created.status_code == 201
    task_id = created.json()['id']
    uploaded = client.post(f'/api/compare/tasks/{task_id}/uploads', headers=headers,
                           data={'client_ids': '["l","r"]'},
                           files=[('files', ('left.docx', protected)),
                                  ('files', ('right.docx', right))])
    assert uploaded.status_code == 200, uploaded.text
    receipt = next(item for item in uploaded.json()['receipts']
                   if item['client_id'] == 'l')
    assert receipt['status'] == 'rejected'
    assert receipt['error_code'] == 'COMPARE_INVALID_DOCX'
    assert receipt['normalized_sha256'] is None and receipt['normalized_size'] is None
    assert receipt['original_sha256'] == hashlib.sha256(protected).hexdigest()
    assert receipt['revisions_accepted'] == 0
    assert uploaded.json()['task']['pairs'][0]['left']['upload_status'] == 'rejected'


def test_compare_legacy_doc_is_rejected_with_conversion_guidance(tmp_path):
    legacy = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"legacy bytes"
    right = docx('B')
    client = TestClient(create_app(tmp_path / 'legacy'))
    headers = {'X-FundLex-User': 'alice'}
    payload = {'name': 'legacy', 'pairs': [{'client_pair_id': 'p', 'label': 'P',
        'left': {'client_id': 'l', 'name': 'left.docx', 'size': len(legacy)},
        'right': {'client_id': 'r', 'name': 'right.docx', 'size': len(right)}}]}
    created = client.post('/api/compare/tasks', headers=headers, json=payload)
    assert created.status_code == 201
    task_id = created.json()['id']
    uploaded = client.post(f'/api/compare/tasks/{task_id}/uploads', headers=headers,
                           data={'client_ids': '["l","r"]'},
                           files=[('files', ('left.docx', legacy)),
                                  ('files', ('right.docx', right))])
    assert uploaded.status_code == 200, uploaded.text
    receipt = next(item for item in uploaded.json()['receipts']
                   if item['client_id'] == 'l')
    assert receipt['status'] == 'rejected'
    assert receipt['error_code'] == 'LEGACY_DOC_UNSUPPORTED'
    assert receipt['normalized_sha256'] is None and receipt['normalized_size'] is None

    bad_payload = {'name': 'legacy-manifest', 'pairs': [{'client_pair_id': 'p', 'label': 'P',
        'left': {'client_id': 'l', 'name': 'left.doc', 'size': len(legacy)},
        'right': {'client_id': 'r', 'name': 'right.docx', 'size': len(right)}}]}
    response = client.post('/api/compare/tasks', headers=headers, json=bad_payload)
    assert response.status_code == 422


def test_compare_revised_input_selection_binds_normalized_hash(tmp_path):
    revised = revised_docx('费率为0.40%', '费率为0.40%')
    normalized = prepare_compare_input(revised).data
    right = docx('费率为0.30%')
    client = TestClient(create_app(tmp_path / 'selection'))
    headers = {'X-FundLex-User': 'alice'}
    payload = {'name': 'selection', 'pairs': [{'client_pair_id': 'p', 'label': 'P',
        'left': {'client_id': 'l', 'name': 'left.docx', 'size': len(revised)},
        'right': {'client_id': 'r', 'name': 'right.docx', 'size': len(right)}}]}
    created = client.post('/api/compare/tasks', headers=headers, json=payload)
    assert created.status_code == 201, created.text
    task_id = created.json()['id']
    uploaded = client.post(f'/api/compare/tasks/{task_id}/uploads', headers=headers,
                           data={'client_ids': '["l","r"]'},
                           files=[('files', ('left.docx', revised)),
                                  ('files', ('right.docx', right))])
    assert uploaded.status_code == 200, uploaded.text
    assert client.post(f'/api/compare/tasks/{task_id}/seal', headers=headers).status_code == 200
    assert CompareWorker(client.app.state.compare_store).run_once()
    pair = client.get(f'/api/compare/tasks/{task_id}', headers=headers).json()['pairs'][0]
    assert pair['status'] == 'succeeded', pair
    normalized_sha256 = hashlib.sha256(normalized).hexdigest()
    assert pair['left']['normalized_sha256'] == normalized_sha256

    url = f"/api/compare/pairs/{pair['id']}/selection"
    state = client.get(url, headers=headers).json()
    assert state['left_source_sha256'] == normalized_sha256
    assert state['left_source_sha256'] != hashlib.sha256(revised).hexdigest()
    assert len(state['diffs']) == 1 and state['diffs'][0]['type'] == 'modify'
    saved = client.put(url, headers=headers, json={
        'expected_epoch': 0, 'result_sha256': state['result_sha256'],
        'left_source_sha256': state['left_source_sha256'],
        'right_source_sha256': state['right_source_sha256'],
        'decisions': {state['diffs'][0]['id']: 'right'}})
    assert saved.status_code == 200, saved.text
    state = saved.json()
    generated = client.post(f'{url}/generate', headers=headers, json={
        'expected_epoch': state['selection_epoch'],
        'decision_hash': state['draft']['decision_hash'],
        'result_sha256': state['result_sha256']})
    assert generated.status_code == 200, generated.text
    artifact = generated.json()
    assert artifact['report']['left_source_sha256'] == normalized_sha256
    downloaded = client.get(artifact['files']['final']['download_url'], headers=headers)
    assert downloaded.status_code == 200
    assert Document(io.BytesIO(downloaded.content)).paragraphs[0].text == '费率为0.30%'
