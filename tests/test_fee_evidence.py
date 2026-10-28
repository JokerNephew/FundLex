import pytest

from fundlex.fee_evidence import fee_value, verify_fee_evidence


@pytest.mark.parametrize("quote,field,name,expected", [
    ("H类基金份额收取申购费用，不计提销售服务费。", "sales_service_fee", "H", False),
    ("J 类及 K 类基金份额不收取申购费用，计提销售服务费。", "sales_service_fee", "K", True),
    ("H类基金份额收取申购费用，J类基金份额不收取申购费用。", "subscription_fee", "H", True),
    ("H类基金份额收取申购费用，J类基金份额不收取申购费用。", "subscription_fee", "J", False),
    ("H类基金份额收取申购费用，养老金通过直销机构申购时免收申购费用。", "subscription_fee", "H", True),
    ("H类基金份额收取申购费用。H类基金份额通过直销机构申购时免收申购费用。", "subscription_fee", "H", True),
    ("通过直销机构申购H类基金份额免收申购费用。", "subscription_fee", "H", None),
    ("H类基金份额通过电子交易申购时免收申购费用。", "subscription_fee", "H", None),
    ("H类基金份额收取申购费用。销售服务费率为0.40%。", "sales_service_fee", "H", None),
    ("H类基金份额收取申购费用，J类基金份额计提销售服务费。", "sales_service_fee", "H", None),
    ("H类基金份额销售服务费率为0.00%。", "sales_service_fee", "H", False),
    ("H类基金份额销售服务费率为0.40%。", "sales_service_fee", "H", True),
    ("H类基金份额销售服务费按年费率0.40%计算。", "sales_service_fee", "H", True),
    ("H类基金份额销售服务费率不超过0.40%。", "sales_service_fee", "H", None),
    ("H类基金份额收取申购费用，不收取申购费用。", "subscription_fee", "H", None),
    ("H类基金份额收取申购费用。H类基金份额不收取申购费用。", "subscription_fee", "H", None),
    ("收取申购费用且不计提销售服务费的基金份额称为H类基金份额。", "sales_service_fee", "H", False),
    ("H类和J类基金份额分别收取申购费用和销售服务费。", "subscription_fee", "H", None),
    ("H类基金份额不计提管理费。", "sales_service_fee", "H", None),
    ("对H类基金份额免收申购费。", "subscription_fee", "H", False),
    ("对特定投资者的H类基金份额免收申购费。", "subscription_fee", "H", None),
    ("H类基金份额可以收取申购费用。", "subscription_fee", "H", None),
    ("H类基金份额是否收取申购费用。", "subscription_fee", "H", None),
    ("H类基金份额不得收取申购费用。", "subscription_fee", "H", None),
    ("H类基金份额收取申购费用，但不收取销售服务费。", "sales_service_fee", "H", False),
    ("H类基金份额不收取认购和申购费。", "subscription_fee", "H", False),
])
def test_clause_bound_facts(quote, field, name, expected):
    assert fee_value(quote, field, name) is expected


def test_single_class_requires_verified_binding():
    quote = "本基金收取申购费用，不计提销售服务费。"
    assert fee_value(quote, "sales_service_fee", "__single__") is None
    assert fee_value(quote, "sales_service_fee", "__single__", single_bound=True) is False
    assert fee_value("H类基金份额计提销售服务费。", "sales_service_fee", "__single__", single_bound=True) is None


def test_exact_source_never_joins_or_normalizes_quotes():
    quote = "H类基金份额不收取申购费用。"
    assert verify_fee_evidence(quote, "subscription_fee", "H", ["前言" + quote + "后续"]) is False
    assert verify_fee_evidence(quote, "subscription_fee", "H", ["H类基金份额", "不收取申购费用。"]) is None
    assert verify_fee_evidence(quote, "subscription_fee", "H", ["H 类基金份额不收取申购费用。"]) is None
    assert verify_fee_evidence("", "subscription_fee", "H", [quote]) is None


@pytest.mark.parametrize("connector", ["及", "和", "与", "、"])
@pytest.mark.parametrize("name", ["H", "J", "K"])
def test_complete_names_with_explicit_joint_subject(connector, name):
    quote = connector.join(f"{c}类基金份额" for c in ("H", "J", "K")) + "不收取销售服务费。"
    assert fee_value(quote, "sales_service_fee", name) is False


@pytest.mark.parametrize("quote,name,expected", [
    ("H类基金份额销售服务费按前一日H类基金份额资产净值的0.50%年费率计提。", "H", True),
    ("H类基金份额销售服务费按前一日H类基金份额资产净值的0%年费率计提。", "H", False),
    ("H类基金份额销售服务费按前一日J类基金份额资产净值的0.50%年费率计提。", "H", None),
    ("H类基金份额销售服务费按前一日J类基金份额资产净值的0.50%年费率计提。", "J", None),
    ("H类基金份额销售服务费按前一日H类基金份额资产净值计算，管理费按0.50%年费率计提。", "H", None),
    ("H类基金份额销售服务费不按前一日H类基金份额资产净值的0.50%年费率计提。", "H", None),
    ("H类基金份额销售服务费按前一日H类基金份额资产净值的0.50%年费率不计提。", "H", None),
    ("特定投资者H类基金份额销售服务费按前一日H类基金份额资产净值的0.50%年费率计提。", "H", None),
    ("H类基金份额销售服务费率为0%，H类基金份额通过直销机构销售服务费率为0.50%。", "H", False),
    ("H类基金份额销售服务费率为0%，H类基金份额销售服务费率为0.50%。", "H", None),
    ("H类基金份额不收取销售服务费和J类基金份额计提销售服务费。", "H", False),
    ("H类基金份额不收取销售服务费和J类基金份额计提销售服务费。", "J", True),
    ("H类基金份额J类基金份额不收取销售服务费。", "H", None),
])
def test_repeated_subject_rates_keep_property_and_scope(quote, name, expected):
    assert fee_value(quote, "sales_service_fee", name) is expected


def test_repeated_class_subject_and_rate_with_de_shows_positive_service_fee():
    quote = ("本基金 A 类基金份额不收取销售服务费，C 类基金份额的销售服务费按前一日 "
             "C 类基金份额资产净值的 0.40%的年费率计提。")
    assert fee_value(quote, "sales_service_fee", "A") is False
    assert fee_value(quote, "sales_service_fee", "C") is True


def test_investor_bears_subscription_fee_is_positive_evidence():
    quote = "A类基金份额的申购费用由投资人承担，不列入基金财产。"
    assert fee_value(quote, "subscription_fee", "A") is True


@pytest.mark.parametrize("name,expected", [("H", False), ("J", True), ("K", True)])
def test_respective_same_fee_does_not_hide_previous_class(name, expected):
    quote = "本基金H类基金份额不收取销售服务费，J类基金份额和K类基金份额分别从本类别份额基金资产中计提销售服务费。"
    assert fee_value(quote, "sales_service_fee", name) is expected


@pytest.mark.parametrize("quote,field", [
    ("J类基金份额和K类基金份额分别收取申购费和销售服务费。", "sales_service_fee"),
    ("J类基金份额和K类基金份额分别收取申购费和销售服务费。", "subscription_fee"),
    ("J类基金份额和K类基金份额分别从本类别基金资产中计提销售服务费和管理费。", "sales_service_fee"),
    ("J类基金份额和K类基金份额分别从本类别基金资产中按0.20%和0.40%计提销售服务费。", "sales_service_fee"),
    ("J类基金份额和K类基金份额分别从本类别基金资产中计提销售服务费，申购费率为0%。", "sales_service_fee"),
    ("J类基金份额和K类基金份额分别从本类别基金资产中计提销售服务费和不计提销售服务费。", "sales_service_fee"),
    ("J类基金份额分别从本类别基金资产中计提销售服务费。", "sales_service_fee"),
])
@pytest.mark.parametrize("name", ["J", "K"])
def test_respective_allocation_remains_ambiguous(quote, field, name):
    assert fee_value(quote, field, name) is None


def test_joint_respective_special_investors_remain_unknown():
    quote = "特定投资者J类基金份额和K类基金份额分别从本类别基金资产中计提销售服务费。"
    assert fee_value(quote, "sales_service_fee", "J") is None
    assert fee_value(quote, "sales_service_fee", "K") is None


@pytest.mark.parametrize("prefix", ["不", "并不", "无需", "以往会", "计划"])
@pytest.mark.parametrize("name", ["J", "K"])
def test_joint_respective_narrative_is_not_an_explicit_current_fee(prefix, name):
    quote = "J类基金份额和K类基金份额" + prefix + "分别从本类别基金资产中计提销售服务费。"
    assert fee_value(quote, "sales_service_fee", name) is None


def test_multi_currency_class_labels_are_parsed_as_their_letter():
    """QDII classes such as A类人民币 belong to the same fee group as A."""
    subscription = ("本基金C类基金份额不收取申购费，A类人民币基金份额和A类美元现汇基金份额的"
                    "申购费率由基金管理人决定，并在招募说明书中列示。")
    assert fee_value(subscription, "subscription_fee", "A") is True
    assert fee_value(subscription, "subscription_fee", "C") is False
    service = ("本基金A类基金份额不收取销售服务费，C类基金份额的销售服务费按前一日"
               "C类基金资产净值0.40%的年费率计提。")
    assert fee_value(service, "sales_service_fee", "A") is False
    assert fee_value(service, "sales_service_fee", "C") is True


def test_class_set_accepts_full_multi_currency_enumeration(tmp_path):
    """Four full display names enumerate the set without an explicit count."""
    from docx import Document

    from fundlex.docx import inspect_docx
    from fundlex.fact_evidence import verify_class_set
    from fundlex.recall import ClassFact

    text = ("本基金A类人民币基金份额、A类美元基金份额、C类人民币基金份额、C类美元基金份额"
            "分别设置代码，由于基金费用的不同，本基金各类基金份额将分别计算基金资产净值。")
    source = tmp_path / "class-set.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    classes = [
        ClassFact("A类人民币", True, False, True, None, ""),
        ClassFact("A类美元", True, False, True, None, ""),
        ClassFact("C类人民币", False, True, True, None, ""),
        ClassFact("C类美元", False, True, True, None, ""),
    ]
    evidence = verify_class_set(inspection, classes, text)
    assert evidence["quote_verified"] is True
    assert evidence["set_complete_verified"] is True
    assert evidence["reason"] == "CLASS_SET_VERIFIED"


def test_capped_percentage_and_bracketed_nav_rate_are_fee_evidence():
    """QDII contracts cap the subscription rate and bracket the NAV percentage."""
    assert fee_value("本基金A类基金份额（含A类人民币和A类美元基金份额）的"
                     "申购费率最高不超过申购金额的5％", "subscription_fee", "A") is True
    assert fee_value("C类基金份额不收取申购费用", "subscription_fee", "C") is False
    assert fee_value("本基金A类基金份额不收取销售服务费", "sales_service_fee", "A") is False
    assert fee_value("C类基金份额的销售服务费按前一日C类基金份额资产净值的"
                     "[ 0.4% ]年费率计提", "sales_service_fee", "C") is True


def test_class_set_base_groups_cover_currency_variants(tmp_path):
    """A three-group enumeration covers A类人民币/A类美元现汇/C类/E类."""
    from docx import Document

    from fundlex.docx import inspect_docx
    from fundlex.fact_evidence import verify_class_set
    from fundlex.recall import ClassFact

    text = ("本基金根据申购费用、销售服务费收取方式的不同，将基金份额分为"
            "A类、C类和E类三种不同的类别。")
    source = tmp_path / "class-set-qdii.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    classes = [
        ClassFact("A类人民币", True, False, True, None, ""),
        ClassFact("A类美元现汇", True, False, True, None, ""),
        ClassFact("C类", False, True, True, None, ""),
        ClassFact("E类", False, True, True, None, ""),
    ]
    evidence = verify_class_set(inspection, classes, text)
    assert evidence["set_complete_verified"] is True
    assert evidence["reason"] == "CLASS_SET_VERIFIED"


def test_class_set_accepts_common_enumeration_verbs(tmp_path):
    """Real contracts say 现有/设 as well as 分为; the label set still must match."""
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.fact_evidence import verify_class_set
    from fundlex.recall import ClassFact

    classes = [ClassFact("A类", True, False, True, None, ""),
               ClassFact("C类", False, True, True, None, "")]
    for text in ("本基金现有A类与C类基金份额。",
                 "本基金设A类、C类基金份额。",
                 "本基金包括A类基金份额和C类基金份额。"):
        source = tmp_path / "contract.docx"
        document = Document()
        document.add_paragraph(text)
        document.save(source)
        evidence = verify_class_set(inspect_docx(source, full=True), classes, text)
        assert evidence["set_complete_verified"] is True, (text, evidence)
    # A different label set is still rejected even with the same verb.
    source = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("本基金设A类、C类和D类基金份额。")
    document.save(source)
    evidence = verify_class_set(inspect_docx(source, full=True), classes,
                                "本基金设A类、C类和D类基金份额。")
    assert evidence["set_complete_verified"] is False


def test_fee_value_accepts_shixing_rate_phrasing():
    from fundlex.fee_evidence import fee_value
    assert fee_value("C类基金份额销售服务费实行0.5%的年费率。",
                     "sales_service_fee", "C") is True
    assert fee_value("C类基金份额销售服务费实行0%的年费率。",
                     "sales_service_fee", "C") is False
