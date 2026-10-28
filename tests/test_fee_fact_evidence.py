"""Fee booleans require fee-bearing evidence, not merely a class quote."""
import pytest
from docx import Document

from fundlex.docx import inspect_docx
from fundlex.recall import facts_from_value
from test_recall_core import CLASS_DEFINITION


def test_management_fee_quote_cannot_certify_subscription_or_service_fee(tmp_path):
    path = tmp_path / "contract.docx"
    Document().save(path)
    doc = Document(path)
    doc.add_paragraph("本基金将基金份额分为A类和Y类不同的类别。")
    doc.add_paragraph("本基金 A 类基金份额的管理费按前一日该类基金份额资产净值的0.90%年费率计提。")
    doc.add_paragraph("本基金 Y 类基金份额的管理费按前一日该类基金份额资产净值的0.45%年费率计提。")
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    facts = facts_from_value(inspection, {
        "fund_type": {"value": "open_regular", "quote": "本基金将基金份额分为A类和Y类不同的类别。"},
        "classes": [
            {"name": "A类基金份额", "subscription_fee": True, "service_fee": False,
             "current": True, "quote": "本基金 A 类基金份额的管理费按前一日该类基金份额资产净值的0.90%年费率计提。"},
            {"name": "Y类基金份额", "subscription_fee": True, "service_fee": False,
             "current": True, "quote": "本基金 Y 类基金份额的管理费按前一日该类基金份额资产净值的0.45%年费率计提。"},
        ],
        "class_set_quote": "本基金将基金份额分为A类和Y类不同的类别。",
        "floating_management_fee": {"value": False, "quote": "本基金 A 类基金份额的管理费按前一日该类基金份额资产净值的0.90%年费率计提。"},
    })
    assert facts.status_of("share_class_structure") == "VERIFIED"
    assert facts.status_of("subscription_fee_group") == "UNKNOWN"
    assert facts.status_of("service_fee_group") == "UNKNOWN"
    assert not facts.value_verified
    assert facts.fee_polarities["A类基金份额"] == {"subscription_fee": None, "service_fee": None}


def test_product_wide_rate_discretion_does_not_prove_class_fee_polarities(tmp_path):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("本基金将基金份额分为A类和Y类不同的类别。")
    doc.add_paragraph("本基金的申购费率由基金管理人决定，并在招募说明书中列示。")
    doc.add_paragraph("申购费用由投资人承担，不列入基金财产。")
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    enum = "本基金将基金份额分为A类和Y类不同的类别。"
    quote = "本基金的申购费率由基金管理人决定，并在招募说明书中列示。"
    facts = facts_from_value(inspection, {
        "fund_type": {"value": "open_regular", "quote": enum},
        "classes": [
            {"name": name, "subscription_fee": True, "service_fee": False,
             "current": True, "quote": quote}
            for name in ("A类基金份额", "Y类基金份额")
        ],
        "class_set_quote": enum,
    })
    assert facts.status_of("subscription_fee_group") == "UNKNOWN"
    assert facts.status_of("service_fee_group") == "UNKNOWN"


@pytest.mark.parametrize("tamper", [False, True])
def test_separate_fee_quotes_are_verified_independently(tmp_path, tamper):
    enum = "本基金将基金份额分为A类和C类不同的类别。"
    subscription = "本基金A类基金份额在申购时收取申购费用，C类基金份额不收取申购费用。"
    service = "本基金A类基金份额不收取销售服务费，C类基金份额的销售服务费按前一日C类基金份额资产净值的0.20%年费率计提。"
    path = tmp_path / "contract.docx"
    doc = Document()
    for paragraph in (enum, subscription, service):
        doc.add_paragraph(paragraph)
    doc.save(path)
    facts = facts_from_value(inspect_docx(path, full=True), {
        "classes": [
            {"name": name, "subscription_fee": sub, "service_fee": serv, "current": True,
             "quote": enum, "subscription_quote": subscription,
             "service_quote": service + ("不存在的尾文" if tamper else "")}
            for name, sub, serv in (("A", True, False), ("C", False, True))
        ],
        "class_set_quote": enum,
    })
    assert facts.status_of("share_class_structure") == "VERIFIED"
    # The tampered quote is the *service* one; the subscription family is proven
    # independently and must stay VERIFIED (FL-D440).
    assert facts.status_of("subscription_fee_group") == "VERIFIED"
    assert facts.status_of("service_fee_group") == ("UNKNOWN" if tamper else "VERIFIED")
    assert facts.value_verified is (not tamper)
    assert facts.as_ledger()["classes"][0]["subscription_quote"] == subscription


def test_contracted_ren_shengou_definition_keeps_fee_meaning(tmp_path):
    quote = CLASS_DEFINITION.replace("认购、申购", "认申购")
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph(quote)
    doc.save(path)
    facts = facts_from_value(inspect_docx(path, full=True), {
        "classes": [
            {"name": name, "subscription_fee": sub, "service_fee": serv, "current": True, "quote": quote}
            for name, sub, serv in (("A", True, False), ("C", False, True), ("D", True, False))
        ],
        "class_set_quote": quote,
    })
    assert facts.status_of("subscription_fee_group") == "VERIFIED"


def test_definition_fee_quote_certifies_both_polarities(tmp_path):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph(CLASS_DEFINITION)
    doc.save(path)
    quote = CLASS_DEFINITION
    facts = facts_from_value(inspect_docx(path, full=True), {
        "fund_type": {"value": "open_regular", "quote": quote},
        "classes": [
            {"name": "A类基金份额", "subscription_fee": True, "service_fee": False, "current": True, "quote": quote},
            {"name": "C类基金份额", "subscription_fee": False, "service_fee": True, "current": True, "quote": quote},
            {"name": "D类基金份额", "subscription_fee": True, "service_fee": False, "current": True, "quote": quote},
        ],
        "class_set_quote": quote,
        "floating_management_fee": {"value": False, "quote": quote},
    })
    assert all(facts.status_of(name) == "VERIFIED" for name in
               ("share_class_structure", "subscription_fee_group", "service_fee_group"))


@pytest.mark.parametrize("subscription,service", [(False, False), (True, True)])
def test_literal_fee_contradiction_is_conflict(tmp_path, subscription, service):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph(CLASS_DEFINITION)
    doc.save(path)
    facts = facts_from_value(inspect_docx(path, full=True), {
        "classes": [
            {"name": "A类基金份额", "subscription_fee": subscription, "service_fee": service,
             "current": True, "quote": CLASS_DEFINITION},
            {"name": "C类基金份额", "subscription_fee": False, "service_fee": True,
             "current": True, "quote": CLASS_DEFINITION},
            {"name": "D类基金份额", "subscription_fee": True, "service_fee": False,
             "current": True, "quote": CLASS_DEFINITION},
        ],
        "class_set_quote": CLASS_DEFINITION,
    })
    assert facts.status_of("share_class_structure") == "VERIFIED"
    assert facts.status_of("subscription_fee_group") == "CONFLICT"
    assert facts.status_of("service_fee_group") == "CONFLICT"
    assert not facts.value_verified
    assert any(p.startswith("FEE_QUOTE_CONFLICT:") for p in facts.problems)


def test_numbered_fee_quote_repair_keeps_source_prefix_and_certifies_fact(tmp_path):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("本基金将基金份额分为A类和C类不同的类别。")
    doc.add_paragraph("4、A 类基金份额的申购费用由投资人承担，不列入基金财产；C 类基金份额不收取申购费用。")
    doc.add_paragraph("本基金 A 类基金份额不收取销售服务费，C 类基金份额的销售服务费按前一日 C 类基金份额资产净值的 0.40%的年费率计提。")
    doc.save(path)
    inspection = inspect_docx(path, full=True)
    facts = facts_from_value(inspection, {
        "classes": [
            {"name": "A类", "subscription_fee": True, "service_fee": False,
             "current": True, "quote": "本基金将基金份额分为A类和C类不同的类别。",
             "subscription_quote": "本基金 A 类基金份额的申购费用由投资人承担，不列入基金财产；C 类基金份额不收取申购费用。",
             "service_quote": "本基金 A 类基金份额不收取销售服务费"},
            {"name": "C类", "subscription_fee": False, "service_fee": True,
             "current": True, "quote": "本基金将基金份额分为A类和C类不同的类别。",
             "subscription_quote": "本基金 C 类基金份额不收取申购费用。",
             "service_quote": "C 类基金份额的销售服务费按前一日 C 类基金份额资产净值的 0.40%的年费率计提"},
        ],
        "class_set_quote": "本基金将基金份额分为A类和C类不同的类别。",
    })
    assert facts.classes[0].subscription_quote.startswith("4、A 类")
    assert facts.quotes_verified["subscription_fee:A类"]
    assert facts.status_of("subscription_fee_group") == "VERIFIED"
    assert facts.status_of("service_fee_group") == "VERIFIED"
