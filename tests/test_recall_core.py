"""Core recall-first engine tests: facts, mapping, inventory, planning, artifacts.

The gateway is scripted; no real model calls are made.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from fundlex.docx import generate_artifacts, inspect_docx, validate_edit_plan
from fundlex.recall import (
    RecallError, build_docx_edits, build_windows, class_label, condition_matches,
    content_check_factory, display_text, extract_facts, find_normalized,
    instantiate_slots, inventory_window, plan_edits, plan_windows, resolve_quote, scope_map,
    structure_index,
)

CLASS_DEFINITION = (
    "基金份额类别：本基金将基金份额分为A类、C类和D类不同的类别。在投资者认购、申购基金时收取认购、申购费用而不计提"
    "销售服务费的，称为A类基金份额；在投资者认购、申购基金份额时不收取认购、申购费用而是从本类别基金资产中计提销售服"
    "务费的，称为C类基金份额；在投资者申购时收取申购费用且不计提销售服务费的，称为D类基金份额。"
)
FEE_ADJUSTMENT = (
    "在符合法律法规、基金合同的约定且不损害已有基金份额持有人权益的情况下，经与基金托管人协商一致，基金管理人在履行"
    "适当程序后可以增加新的基金份额类别、或者在法律法规和基金合同规定的范围内变更现有基金份额类别的申购费率、调低赎回"
    "费率或变更收费方式、或者停止现有基金份额类别的销售等，调整前基金管理人需及时公告并报中国证监会备案。"
)
CONVERSION_CLAUSE = "除非基金管理人在未来条件成熟后另行公告开通相关业务，本基金不同基金份额类别之间不得互相转换。"
SUBSCRIPTION_FEE_CLAUSE = (
    "申购份额的计算及余额的处理方式：本基金申购份额的计算详见《招募说明书》。本基金的申购费率由基金管理人决定，并在招"
    "募说明书中列示。申购的有效份额为净申购金额除以当日的基金份额净值，有效份额单位为份。"
)
PROMO_CLAUSE = (
    "基金管理人可以在不违反法律法规规定及基金合同约定的情形下根据市场情况制定基金促销计划。在基金促销活动期间，按相关"
    "监管部门要求履行必要手续后，基金管理人可以适当调低基金申购费率、赎回费率和转换费率。"
)
SERVICE_FEE_CLAUSE = (
    "C类基金份额的销售服务费每日计提，按月支付。由基金管理人向基金托管人发送基金销售服务费划付指令，经基金托管人复核后"
    "于次月首日起10个工作日内从基金财产中一次性支付给登记机构，由登记机构代付给销售机构。若遇法定节假日、公休日等，支付"
    "日期顺延。"
)
LIQUIDATION_CLAUSE = (
    "依据基金财产清算的分配方案，将基金财产清算后的全部剩余资产扣除基金财产清算费用、交纳所欠税款并清偿基金债务后，按"
    "基金份额持有人持有的基金份额比例进行分配。"
)


def build_contract(path: Path, *, title: str = "某某灵活配置混合型证券投资基金基金合同") -> None:
    document = Document()
    document.add_paragraph(title)
    document.add_paragraph("一、前言")
    document.add_paragraph("本基金合同经中国证监会核准。")
    document.add_paragraph("二、释义")
    document.add_paragraph(CLASS_DEFINITION)
    document.add_paragraph("三、基金的基本情况")
    document.add_paragraph(CLASS_DEFINITION)
    document.add_paragraph(FEE_ADJUSTMENT)
    document.add_paragraph(CONVERSION_CLAUSE)
    document.add_paragraph("六、基金份额的申购与赎回")
    document.add_paragraph(SUBSCRIPTION_FEE_CLAUSE)
    document.add_paragraph(PROMO_CLAUSE)
    document.add_paragraph("七、基金费用与税收")
    document.add_paragraph(SERVICE_FEE_CLAUSE)
    document.add_paragraph("八、基金的收益与分配")
    document.add_paragraph("本基金同一类别的每一基金份额享有同等分配权。")
    document.add_paragraph("九、基金合同的变更、终止与基金财产的清算")
    document.add_paragraph(LIQUIDATION_CLAUSE)
    document.save(path)


def guide_rules() -> list[dict]:
    return [
        {
            "rule_id": "G01", "title": "释义-份额类别定义", "section_hint": "释义",
            "requirement": "在份额类别定义原条款后增加直销机构收费新条款。",
            "operations": [{"op": "append",
                            "new_text": "投资者通过直销机构（含基金管理人及其销售子公司，下同）申购本基金A类基金份额时，"
                                        "不收取认购/申购费用；通过直销机构申购的C类基金份额，不收取销售服务费。",
                            "condition": {"kind": "none", "value": None}}],
            "old_text_variants": [], "slot_tokens": ["A类基金份额", "C类基金份额"],
            "class_substitution": True, "references": [], "source_paragraphs": [0, 4],
        },
        {
            "rule_id": "G05", "title": "基本情况-类别调整句", "section_hint": "基金的基本情况",
            "requirement": "有销售服务费类别时把调低赎回费率改为调低销售服务费率，并删除备案句。",
            "operations": [
                {"op": "replace", "old_text": "调低赎回费率", "new_text": "调低销售服务费率",
                 "condition": {"kind": "has_service_fee", "value": True}},
                {"op": "delete", "old_text": "、调低赎回费率",
                 "condition": {"kind": "has_service_fee", "value": False}},
                {"op": "delete", "old_text": "并报中国证监会备案",
                 "condition": {"kind": "none", "value": None}},
            ],
            "old_text_variants": [], "slot_tokens": [], "class_substitution": False,
            "references": [], "source_paragraphs": [17],
        },
        {
            "rule_id": "G06", "title": "基本情况-类别转换句", "section_hint": "基金的基本情况",
            "requirement": "删除类别转换句。",
            "operations": [{"op": "delete", "old_text": CONVERSION_CLAUSE,
                            "condition": {"kind": "none", "value": None}}],
            "old_text_variants": [], "slot_tokens": [], "class_substitution": False,
            "references": [], "source_paragraphs": [23],
        },
        {
            "rule_id": "G08", "title": "申购份额计算-多类别", "section_hint": "基金份额的申购与赎回",
            "requirement": "多类别时替换申购费率句。",
            "operations": [{"op": "replace",
                            "old_text": "本基金的申购费率由基金管理人决定，并在招募说明书中列示。",
                            "new_text": "本基金A类基金份额在申购时收取申购费用，C类基金份额不收取申购费用。通过直销机构申购"
                                        "本基金A类基金份额不收取申购费。通过直销机构以外的其他销售机构申购本基金的申购费率由"
                                        "基金管理人决定，并在招募说明书中列示。",
                            "condition": {"kind": "class_count", "value": "multiple"}}],
            "old_text_variants": [], "slot_tokens": ["A类基金份额", "C类基金份额"],
            "class_substitution": True, "references": ["同前述说明"], "source_paragraphs": [32],
        },
        {
            "rule_id": "G14", "title": "促销费率调整句", "section_hint": "基金份额的申购与赎回",
            "requirement": "把促销费率句替换为多类别新文。",
            "operations": [{"op": "replace", "old_text": "基金管理人可以适当调低基金申购费率和赎回费率。",
                            "new_text": "基金管理人可以适当调低基金申购费率和销售服务费率。",
                            "condition": {"kind": "class_count", "value": "multiple"}}],
            "old_text_variants": ["基金管理人可以适当调低基金申购费率、赎回费率和转换费率。"],
            "slot_tokens": [], "class_substitution": False, "references": [], "source_paragraphs": [66],
        },
        {
            "rule_id": "G15", "title": "销售服务费-删除支付句", "section_hint": "基金费用与税收",
            "requirement": "删除销售服务费支付句。",
            "operations": [{"op": "delete", "old_text": SERVICE_FEE_CLAUSE,
                            "condition": {"kind": "none", "value": None}}],
            "old_text_variants": [], "slot_tokens": [], "class_substitution": False,
            "references": [], "source_paragraphs": [76],
        },
        {
            "rule_id": "G16", "title": "销售服务费-追加新条款", "section_hint": "基金费用与税收",
            "requirement": "在销售服务费条款后增加每日计提与返还新条款。",
            "operations": [{"op": "append",
                            "new_text": "销售服务费每日计提。对于投资者通过直销机构持有的C类基金份额计提的销售服务费，"
                                        "在投资者赎回C类基金份额或基金合同终止时，随赎回款（或清算款）一并返还给投资者。",
                            "condition": {"kind": "none", "value": None}}],
            "old_text_variants": [], "slot_tokens": ["C类基金份额"], "class_substitution": True,
            "references": [], "source_paragraphs": [78],
        },
        {
            "rule_id": "G22", "title": "清算-非浮动有服务费", "section_hint": "基金财产的清算",
            "requirement": "非浮动管理费产品收取销售服务费时替换清算分配句。",
            "operations": [{"op": "replace", "old_text": LIQUIDATION_CLAUSE,
                            "new_text": LIQUIDATION_CLAUSE[:-1] + "，并加上返还基金份额持有人的销售服务费（如有）。",
                            "condition": {"kind": "compound", "value": [
                                {"kind": "is_floating_fee", "value": False},
                                {"kind": "has_service_fee", "value": True},
                                {"kind": "class_count", "value": "multiple"}]}}],
            "old_text_variants": [], "slot_tokens": [], "class_substitution": False,
            "references": [], "source_paragraphs": [105],
        },
    ]


class ScriptedResult:
    def __init__(self, value):
        self.value = value


class ScriptedGateway:
    def __init__(self, responses: dict[str, object], *, max_context_chars: int = 60000):
        self.responses = responses
        # The validator hard gate runs for every composed edit; tests that do not care
        # about the independent review get a passing review by default.
        self.responses.setdefault("recall_verify",
                                  {"verdict": "pass", "reason": "scripted independent review"})
        self.max_context_chars = max_context_chars
        self.max_calls = 50
        self.calls: list[tuple[str, str]] = []

    def call(self, *, task_type, bounded_context, response_schema, instruction, prompt_version="v1"):
        self.calls.append((task_type, prompt_version))
        response = self.responses.get(task_type)
        if response is None:
            raise AssertionError(f"unexpected task_type {task_type}")
        value = response(bounded_context) if callable(response) else response
        return ScriptedResult(value)


def facts_response(context: str) -> dict:
    return {
        "fund_type": {"value": "open_regular", "quote": "某某灵活配置混合型证券投资基金基金合同", "basis": "标题"},
        "classes": [
            {"name": "A", "subscription_fee": True, "service_fee": False, "current": True,
             "quote": "在投资者认购、申购基金时收取认购、申购费用而不计提销售服务费的，称为A类基金份额"},
            {"name": "C", "subscription_fee": False, "service_fee": True, "current": True,
             "quote": "不收取认购、申购费用而是从本类别基金资产中计提销售服务费的，称为C类基金份额"},
            {"name": "D", "subscription_fee": True, "service_fee": False, "current": True,
             "quote": "在投资者申购时收取申购费用且不计提销售服务费的，称为D类基金份额"},
        ],
        "floating_management_fee": {"value": False, "quote": "本基金合同经中国证监会核准。", "basis": "无浮动安排"},
    }


def scope_response(context: str) -> dict:
    mapping = {"G01": ["S03", "S04"], "G05": ["S04"], "G06": ["S04"], "G08": ["S05"],
               "G14": ["S05"], "G15": ["S06"], "G16": ["S06"], "G22": ["S08"]}
    # The scope contract is complete: every rule is explicitly mapped, even when the
    # answer is "no section in this contract" (an omitted rule is a coverage gap).
    for rule in guide_rules():
        mapping.setdefault(rule["rule_id"], [])
    return {"mappings": [{"rule_id": rule_id, "section_ids": ids, "reason": "业务子节"}
                         for rule_id, ids in mapping.items()]}


def inventory_response(context: str) -> dict:
    targets = []
    if CLASS_DEFINITION in context:
        for location in ("p4", "p6"):
            if f"[{location}]" in context:
                targets.append({"rule_id": "G01", "location": location, "quote": CLASS_DEFINITION,
                                "state": "needs_edit", "reason": "定义句需追加直销新条款"})
                break
    if FEE_ADJUSTMENT in context:
        targets.append({"rule_id": "G05", "location": "p7", "quote": "调低赎回费率",
                        "state": "needs_edit", "reason": "需改为销售服务费率"})
        targets.append({"rule_id": "G05", "location": "p7", "quote": "并报中国证监会备案",
                        "state": "needs_edit", "reason": "删除备案句"})
    if CONVERSION_CLAUSE in context:
        targets.append({"rule_id": "G06", "location": "p8", "quote": CONVERSION_CLAUSE,
                        "state": "needs_edit", "reason": "删除转换限制"})
    if SUBSCRIPTION_FEE_CLAUSE in context:
        targets.append({"rule_id": "G08", "location": "p10",
                        "quote": "本基金的申购费率由基金管理人决定，并在招募说明书中列示。",
                        "state": "needs_edit", "reason": "多类别替换"})
    if PROMO_CLAUSE in context:
        targets.append({"rule_id": "G14", "location": "p11",
                        "quote": "基金管理人可以适当调低基金申购费率、赎回费率和转换费率。",
                        "state": "needs_edit", "reason": "变体句替换"})
    if SERVICE_FEE_CLAUSE in context:
        targets.append({"rule_id": "G15", "location": "p13", "quote": SERVICE_FEE_CLAUSE,
                        "state": "needs_edit", "reason": "删除支付句"})
        targets.append({"rule_id": "G16", "location": "p13", "quote": SERVICE_FEE_CLAUSE,
                        "state": "needs_edit", "reason": "追加每日计提条款"})
    if LIQUIDATION_CLAUSE in context:
        targets.append({"rule_id": "G22", "location": "p17", "quote": LIQUIDATION_CLAUSE,
                        "state": "needs_edit", "reason": "追加销售服务费返还"})
    return {"targets": targets}


@pytest.fixture()
def contract(tmp_path: Path) -> Path:
    path = tmp_path / "contract.docx"
    build_contract(path)
    return path


def test_text_helpers():
    assert find_normalized("申购费率由基金管理人决定，并在招募说明书中列示。",
                           "申购费率由基金管理人决定， 并在招募说明书中列示。") is not None
    cleaned, index = display_text("abc\ufff9def")
    assert cleaned == "abcdef" and index == [0, 1, 2, 4, 5, 6]


def test_class_label_and_slots():
    from fundlex.recall import Facts, ClassFact
    facts = Facts(fund_type="open_regular", fund_type_quote="", fund_type_basis="",
                  classes=(ClassFact("A", True, False, True, None, ""),
                           ClassFact("C", False, True, True, None, ""),
                           ClassFact("D", True, False, True, None, "")),
                  floating_management_fee=False, floating_quote="", floating_basis="",
                  notes="", quotes_verified={}, problems=())
    assert class_label(facts.subscription_group) == "A类、D类基金份额"
    text, issues = instantiate_slots("本基金A类基金份额与C类基金份额。", facts)
    assert text == "本基金A类、D类基金份额与C类基金份额。" and not issues
    assert condition_matches({"kind": "class_count", "value": "multiple"}, facts)
    assert not condition_matches({"kind": "class_count", "value": 1}, facts)
    assert condition_matches({"kind": "compound", "value": [
        {"kind": "has_service_fee", "value": True}, {"kind": "is_floating_fee", "value": False}]}, facts)


def test_class_name_suffixes_and_example_tokens():
    from fundlex.recall import Facts, ClassFact
    facts = Facts(fund_type="open_regular", fund_type_quote="", fund_type_basis="",
                  classes=(ClassFact("A类基金份额", True, False, True, None, ""),
                           ClassFact("C类基金份额", False, True, True, None, "")),
                  floating_management_fee=False, floating_quote="", floating_basis="",
                  notes="", quotes_verified={}, problems=())
    assert class_label(facts.subscription_group) == "A类基金份额"
    assert class_label(facts.service_group) == "C类基金份额"
    text, issues = instantiate_slots(
        "将“A类基金份额”替换为“A类、B类基金份额”，将“C类基金份额”替换为“C类、D类基金份额”。", facts)
    assert text == "将“A类基金份额”替换为“A类基金份额”，将“C类基金份额”替换为“C类基金份额”。" and not issues
    # The instantiated replacements equal the references, so the class-substitution rule
    # needs no change for a product that already is A/C (FL-D052).
    from fundlex.recall import class_substitution_noop
    rule = {"class_substitution": True, "operations": [
        {"op": "replace", "old_text": "A类基金份额", "new_text": "A类、B类基金份额",
         "condition": {"kind": "none", "value": None}},
        {"op": "replace", "old_text": "C类基金份额", "new_text": "C类、D类基金份额",
         "condition": {"kind": "none", "value": None}},
    ]}
    assert class_substitution_noop(rule, facts)
    three = Facts(fund_type="open_regular", fund_type_quote="", fund_type_basis="",
                  classes=(ClassFact("A", True, False, True, None, ""),
                           ClassFact("C", False, True, True, None, ""),
                           ClassFact("D", True, False, True, None, "")),
                  floating_management_fee=False, floating_quote="", floating_basis="",
                  notes="", quotes_verified={}, problems=())
    assert not class_substitution_noop(rule, three)
    assert not class_substitution_noop({**rule, "class_substitution": False}, facts)


def test_service_fee_rule_not_applicable_without_service_class():
    """A product with no sales-service-fee class is out of scope, not unresolved.

    The 销售服务费 rules are written with the guide's example C-class group. A contract
    whose classes both charge a subscription fee and no service fee (for example A類/Y類
    divided by sales account) has no C-class clause at all, so ``instantiate_slots`` would
    report ``SERVICE_GROUP_EMPTY`` and the run would stop with ``RECALL_SLOT_UNRESOLVED``.
    The rule must instead be not applicable (FL-D056).
    """
    from fundlex.recall import Facts, ClassFact, slot_group_missing
    facts = Facts(fund_type="open_regular", fund_type_quote="", fund_type_basis="",
                  classes=(ClassFact("A类基金份额", True, False, True, None, ""),
                           ClassFact("Y类基金份额", True, False, True, None, "")),
                  floating_management_fee=False, floating_quote="", floating_basis="",
                  notes="", quotes_verified={}, problems=())
    assert not facts.service_group
    service_rule = {"class_substitution": True, "operations": [
        {"op": "append", "new_text": "销售服务费每日计提。（1）直销机构：对于C类基金份额…",
         "condition": {"kind": "none", "value": None}}]}
    assert slot_group_missing(service_rule, service_rule["operations"], facts)
    # An ordinary rule, and a rule whose text needs no group, are unaffected.
    ordinary = {"class_substitution": False, "operations": [
        {"op": "replace", "old_text": "赎回费用归入基金财产", "new_text": "赎回费用全额计入基金财产",
         "condition": {"kind": "none", "value": None}}]}
    assert not slot_group_missing(ordinary, ordinary["operations"], facts)
    no_group = {"class_substitution": True, "operations": [
        {"op": "replace", "old_text": "A类基金份额", "new_text": "A类、B类基金份额",
         "condition": {"kind": "none", "value": None}}]}
    # A fee-homogeneous product (all classes share one fee profile) has no group for the
    # substitution to select, so the class-substitution rules are out of scope too.
    from fundlex.recall import fee_homogeneous_classes
    assert fee_homogeneous_classes(facts)
    assert slot_group_missing(no_group, no_group["operations"], facts)
    # A rule whose replacement introduces sales-service-fee wording cannot apply.
    rate_rule = {"class_substitution": False, "operations": [
        {"op": "replace", "old_text": "调低申购费率和赎回费率",
         "new_text": "调低申购费率和销售服务费率", "condition": {"kind": "none", "value": None}}]}
    assert slot_group_missing(rate_rule, rate_rule["operations"], facts)
    split = Facts(fund_type="open_regular", fund_type_quote="", fund_type_basis="",
                  classes=(ClassFact("A类基金份额", True, False, True, None, ""),
                           ClassFact("C类基金份额", False, True, True, None, "")),
                  floating_management_fee=False, floating_quote="", floating_basis="",
                  notes="", quotes_verified={}, problems=())
    assert not fee_homogeneous_classes(split)
    assert not slot_group_missing(rate_rule, rate_rule["operations"], split)


def test_delete_target_span_matching_one_sentence_of_reference(tmp_path: Path,
                                                                monkeypatch):
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "legacy")
    """A delete reference spanning two paragraphs still deletes each single-sentence target.

    G17's reference concatenates the 销售服务费用途 sentence with the 募集期间 sentence; the
    two live in different paragraphs. The single-paragraph span can never match the whole
    reference, so the planner must delete the span it does carry (FL-D052).
    """
    from fundlex.recall import Target, plan_target
    first = "C类基金份额的销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、促销活动费、基金份额持有人服务费等。"
    second = "C类基金份额的销售服务费使用范围不包括基金募集期间的上述费用。"
    document = Document()
    document.add_paragraph("七、基金费用与税收")
    document.add_paragraph(first)
    document.add_paragraph(second)
    path = tmp_path / "service_fee.docx"
    document.save(path)
    inspection = inspect_docx(path, full=True)
    facts = _facts()
    rule = {"rule_id": "G17", "title": "销售服务费条款", "section_hint": "基金费用与税收",
            "requirement": f"删除原条款“{first}{second}”。",
            "operations": [{"op": "delete", "old_text": first + second,
                            "condition": {"kind": "none", "value": None}}],
            "old_text_variants": [], "slot_tokens": [], "class_substitution": False,
            "references": [], "source_paragraphs": [62]}
    span = inspection.by_location("p1")
    start = span.text.find(first)
    assert start >= 0 and span.text[start:start + len(first)] == first
    target = Target("T1", "G17", "p1", first, "needs_edit", "", start, start + len(first), "resolved")
    edit = plan_target(rule, target, inspection, facts, gateway=None, compose_budget=[6])
    assert edit is not None
    assert edit.replacement == "" and edit.content_source == "variant_delete"
    assert edit.start == start and edit.end == start + len(first)
    # A span that also keeps text the rule does not delete must not be wiped wholesale:
    # the composer trims only the reference sentence and keeps the rest.
    keeper = first + "本基金的销售服务费率由基金管理人决定。"
    document2 = Document()
    document2.add_paragraph("七、基金费用与税收")
    document2.add_paragraph(keeper)
    path2 = tmp_path / "service_fee_extra.docx"
    document2.save(path2)
    inspection2 = inspect_docx(path2, full=True)
    gateway = ScriptedGateway({"recall_compose": {"replacement": "本基金的销售服务费率由基金管理人决定。",
                                                  "preserved": [], "basis": "scripted trim"}})
    span2 = inspection2.by_location("p1")
    mixed = Target("T2", "G17", "p1", keeper, "needs_edit", "", 0, len(span2.text), "resolved")
    trimmed = plan_target(rule, mixed, inspection2, facts, gateway=gateway, compose_budget=[6])
    assert trimmed is not None and trimmed.content_source in {
        "composed_delete", "tier1_delete", "whole_paragraph_delete"}
    assert "销售服务费率由基金管理人决定" in trimmed.replacement


def test_facts_scope_inventory_planning(contract: Path, tmp_path: Path):
    inspection = inspect_docx(contract, full=True)
    gateway = ScriptedGateway({
        "recall_facts": facts_response,
        "recall_scope": scope_response,
        "recall_inventory": inventory_response,
        "recall_compose": {"replacement": "基金管理人可以适当调低基金申购费率和销售服务费率。",
                           "preserved": [], "basis": "按指引替换变体句"},
    })
    facts = extract_facts(inspection, gateway)
    assert facts.class_count == 3
    assert facts.subscription_group == ("A", "D")
    assert facts.service_group == ("C",)
    assert facts.has_service_fee and not facts.floating_management_fee
    assert all(facts.quotes_verified.values())

    sections = structure_index(inspection)
    titles = {section.title for section in sections}
    assert "二、释义" in titles and "三、基金的基本情况" in titles

    mapping = scope_map(guide_rules(), sections, gateway)
    assert mapping["G01"] and mapping["G22"]

    rules_by_id = {rule["rule_id"]: rule for rule in guide_rules()}
    targets = []
    for window, window_rule_ids in plan_windows(sections, inspection, mapping,
                                                max_chars=gateway.max_context_chars):
        window_rules = [rules_by_id[rule_id] for rule_id in window_rule_ids]
        found, ledger = inventory_window(window, window_rules, facts, gateway, inspection)
        targets.extend(found)
    by_rule: dict[str, list] = {}
    for target in targets:
        by_rule.setdefault(target.rule_id, []).append(target)
    assert by_rule["G05"] and by_rule["G08"] and by_rule["G16"] and by_rule["G22"]
    for target in targets:
        assert target.resolution == "resolved", target.as_dict()

    planned = plan_edits(rules_by_id, targets, inspection, facts, gateway)
    ledger_status = {entry["target_id"]: entry["status"] for entry in planned["ledger"]}
    assert all(status in {"planned", "satisfied", "not_applicable"} for status in ledger_status.values()), planned["ledger"]
    planned_rule_ids: set[str] = set()
    for edit in planned["edits"]:
        planned_rule_ids.update(edit["rule_id"].split("+"))
    assert {"G05", "G06", "G08", "G14", "G15", "G16", "G22"} <= planned_rule_ids

    edits = build_docx_edits(planned["edits"], inspection)
    content_check = content_check_factory(planned["edits"], inspection)
    plan = validate_edit_plan(contract, edits, {}, content_check=content_check)
    redline = tmp_path / "redline.docx"
    final = tmp_path / "final.docx"
    artifacts = generate_artifacts(contract, plan, redline, final)
    assert artifacts.final_sha256 and artifacts.redline_sha256
    final_inspection = inspect_docx(final, full=True)
    final_text = "\n".join(span.text for span in final_inspection.paragraphs)
    assert "调低销售服务费率" in final_text
    assert "并报中国证监会备案" not in final_text
    assert CONVERSION_CLAUSE not in final_text
    assert "本基金A类、D类基金份额在申购时收取申购费用，C类基金份额不收取申购费用。" in final_text
    assert "基金管理人可以适当调低基金申购费率和销售服务费率。" in final_text
    assert "销售服务费每日计提。" in final_text
    assert SERVICE_FEE_CLAUSE not in final_text
    assert "并加上返还基金份额持有人的销售服务费（如有）。" in final_text
    assert CLASS_DEFINITION in final_text  # original definition kept, new text appended
    assert "不收取认购/申购费用；通过直销机构申购的C类基金份额，不收取销售服务费。" in final_text


def test_inventory_quote_not_found_is_flagged(contract: Path):
    inspection = inspect_docx(contract, full=True)
    gateway = ScriptedGateway({"recall_inventory": lambda context: {
        "targets": [{"rule_id": "G01", "location": "p4", "quote": "不存在的引文内容",
                     "state": "needs_edit", "reason": "test"}]}})
    facts = extract_facts(inspection, ScriptedGateway({"recall_facts": facts_response}))
    sections = structure_index(inspection)
    windows = build_windows(sections, inspection, ["S03"], max_chars=60000)
    found, ledger = inventory_window(windows[0], [guide_rules()[0]], facts, gateway, inspection)
    assert found and found[0].resolution == "quote_not_found"
    assert ledger["counts"]["unresolved"] == 1


def test_compose_variant_requires_fixed_text(contract: Path, tmp_path: Path,
                                              monkeypatch):
    monkeypatch.setenv("FUNDLEX_GENERATION_STRATEGY", "legacy")
    inspection = inspect_docx(contract, full=True)
    gateway = ScriptedGateway({"recall_facts": facts_response})
    facts = extract_facts(inspection, gateway)
    rule = guide_rules()[4]  # G14
    from fundlex.recall import Target, plan_target
    target = Target(target_id="T001", rule_id="G14", location="p11",
                    quote="基金管理人可以适当调低基金申购费率、赎回费率和转换费率。",
                    state="needs_edit", reason="变体")
    span = inspection.by_location("p11")
    target.start, target.end = resolve_quote(span, target.quote)
    target.resolution = "resolved"
    gateway.responses["recall_compose"] = {
        "replacement": "基金管理人可以适当调低基金申购费率和销售服务费率。",
        "preserved": [], "basis": "替换"}
    edit = plan_target(rule, target, inspection, facts, gateway, compose_budget=[3])
    assert edit is not None and edit.content_source in {"composed", "tier1_template"}
    gateway.responses["recall_compose"] = {"replacement": "基金管理人可以适当调低基金申购费率。",
                                           "preserved": [], "basis": "缺失固定新文"}
    with pytest.raises(RecallError) as excinfo:
        plan_target(rule, target, inspection, facts, gateway, compose_budget=[3])
    assert excinfo.value.code == "RECALL_COMPOSE_FIXED_TEXT_MISSING"


def test_delete_variant_composes_trimmed_replacement(contract: Path):
    inspection = inspect_docx(contract, full=True)
    facts = extract_facts(inspection, ScriptedGateway({"recall_facts": facts_response}))
    gateway = ScriptedGateway({})
    rule = {
        "rule_id": "G13", "title": "7日赎回费句", "requirement": "删除7日赎回费句",
        "operations": [{"op": "delete",
                        "old_text": "本基金对持续持有期少于7日的投资者收取不低于1.50%的赎回费，并将上述赎回费全额计入基金财产。",
                        "condition": {"kind": "none", "value": None}}],
        "old_text_variants": [], "slot_tokens": [], "class_substitution": False,
        "references": [], "source_paragraphs": [55],
    }
    from fundlex.recall import Target, plan_target
    target = Target(target_id="T001", rule_id="G13", location="p11",
                    quote="基金管理人可以适当调低基金申购费率、赎回费率和转换费率。",
                    state="needs_edit", reason="实际写法与指引参考不同")
    span = inspection.by_location("p11")
    target.start, target.end = resolve_quote(span, target.quote)
    target.resolution = "resolved"
    gateway.responses["recall_compose"] = {"replacement": "", "preserved": [], "basis": "整句删除"}
    edit = plan_target(rule, target, inspection, facts, gateway, compose_budget=[3])
    assert edit is not None and edit.content_source == "composed_delete" and edit.replacement == ""
    gateway.responses["recall_compose"] = {
        "replacement": "基金管理人可以适当调低基金申购费率、赎回费率和转换费率。另加新政策。",
        "preserved": [], "basis": "未删除且新增内容"}
    with pytest.raises(RecallError) as excinfo:
        plan_target(rule, target, inspection, facts, gateway, compose_budget=[3])
    assert excinfo.value.code in {"RECALL_COMPOSE_DELETE_ADDED_TEXT",
                                  "RECALL_COMPOSE_DELETE_NOT_SHRINKING"}
    gateway.responses["recall_compose"] = {"replacement": "基金管理人可以适当调低基金申购费率。",
                                           "preserved": [], "basis": "部分删除"}
    edit = plan_target(rule, target, inspection, facts, gateway, compose_budget=[3])
    assert edit is not None and edit.content_source == "composed_delete"
    assert edit.replacement == "基金管理人可以适当调低基金申购费率。"


# ---------------------------------------------------------------------------
# Deterministic coverage survey and multi-paragraph operations (FL-D047)
# ---------------------------------------------------------------------------

def _facts() -> object:
    from fundlex.recall import ClassFact, Facts
    return Facts(fund_type="open_regular", fund_type_quote="某某灵活配置混合型证券投资基金基金合同",
                 fund_type_basis="标题",
                 classes=(ClassFact(name="A类基金份额", subscription_fee=True, service_fee=False,
                                    current=True, rate="1.2%", quote="A类"),
                          ClassFact(name="C类基金份额", subscription_fee=False, service_fee=True,
                                    current=True, rate="0.4%", quote="C类")),
                 floating_management_fee=False, floating_quote="本基金合同经中国证监会核准。",
                 floating_basis="无浮动安排", notes="", quotes_verified={}, problems=())


def test_survey_finds_summary_counterpart_and_satisfied_targets(tmp_path: Path):
    from fundlex.recall import survey_targets

    path = tmp_path / "contract.docx"
    build_contract(path)
    inspection = inspect_docx(path, full=True)
    hits = survey_targets(inspection, guide_rules(), _facts(), known=[])
    by_rule = {}
    for hit in hits:
        by_rule.setdefault(hit.rule_id, []).append(hit)
    # The liquidation clause the fixture inventory never reports is found by the code.
    locations = {location for hit in by_rule.get("G22", []) for location in hit.locations}
    assert "p17" in locations
    # The promo clause carries one of the rule's own reference variants; the survey
    # reports it as a real target (the model inventory alone can miss it).
    assert [hit.operation for hit in by_rule.get("G14", [])] == ["survey"]


def test_survey_ignores_reference_text_seen_mid_word(tmp_path: Path):
    from fundlex.recall import find_old_text_all

    path = tmp_path / "contract.docx"
    build_contract(path)
    inspection = inspect_docx(path, full=True)
    # "并报中国证监会备案" also appears inside "…调低赎回费率或变更收费方式、或者停止现有
    # 基金份额类别的销售等，调整前基金管理人需及时公告并报中国证监会备案。" — the exact
    # reference exists there, but a mid-word visit ("公告并报中国证监会备案") must not be
    # reported as a separate location.
    occurrences = find_old_text_all(list(inspection.paragraphs), "并报中国证监会备案")
    for parts in occurrences:
        for location, start, end in parts:
            body = display_text(inspection.by_location(location).text)[0]
            assert start == 0 or body[start - 1] in "。；;！？!?…）)\"】》" or body[start - 1] == "告"


SERVICE_FEE_HEAD = "C类基金份额的销售服务费每日计提，按月支付。"
SERVICE_FEE_TAIL = ("由基金管理人向基金托管人发送基金销售服务费划付指令，经基金托管人复核后"
                    "于次月首日起10个工作日内从基金财产中一次性支付给登记机构，由登记机构代付"
                    "给销售机构。若遇法定节假日、公休日等，支付日期顺延。")


def test_composite_delete_spans_two_paragraphs(tmp_path: Path):
    from fundlex.recall import Target, find_old_text_all, plan_composite_target

    path = tmp_path / "composite.docx"
    document = Document()
    # The contract splits the clause over a paragraph break: the head closes the first
    # paragraph after a lead-in, the tail opens the next one.
    document.add_paragraph("3、销售服务费。" + SERVICE_FEE_HEAD)
    document.add_paragraph(SERVICE_FEE_TAIL)
    document.add_paragraph("销售服务费使用范围不包括基金募集期间的上述费用。")
    document.save(path)
    inspection = inspect_docx(path, full=True)

    reference = SERVICE_FEE_HEAD + SERVICE_FEE_TAIL
    rule = {"rule_id": "G15", "title": "销售服务费支付句", "requirement": "删除支付句。",
            "operations": [{"op": "delete", "old_text": reference,
                            "condition": {"kind": "none", "value": None}}]}
    parts = find_old_text_all(list(inspection.paragraphs), reference)
    assert len(parts) == 1 and [item[0] for item in parts[0]] == ["p0", "p1"]

    target = Target(target_id="T001", rule_id="G15", location="p0",
                    quote=reference, state="needs_edit", reason="multi-paragraph")
    target.composite = tuple(f"{location}:{start}:{end}"
                             for location, start, end in parts[0])
    edits = plan_composite_target(rule, target, inspection, _facts(), new_text="")
    assert len(edits) == 2
    merged = {edit.location: edit for edit in edits}
    # The anchor keeps its own lead-in; the follower paragraph is emptied because the
    # reference sentence is all it holds.
    assert merged["p0"].replacement == "3、销售服务费。"
    assert merged["p1"].replacement == ""


def test_composite_replace_separates_unpunctuated_clause(tmp_path: Path):
    """A composite replacement must not run the new policy text into an unterminated link.

    The survey binds a clause that does not end in punctuation and the guide appends the
    direct-sales policy after it; joining them verbatim produced "…份额投资者通过…"
    (FL-D067). The kept prefix here is the paragraph's lead-in before the reference part.
    """
    from fundlex.recall import Target, find_old_text_all, plan_composite_target

    path = tmp_path / "composite_append.docx"
    document = Document()
    lead = "基金份额类别：本基金将基金份额分为 A 类和 C 类两类不同的类别。"
    reference = "在投资者认购、申购基金份额时不收取认购、申购费用，从该类别基金财产中计提销售服务费的，称为 C 类基金份额"
    document.add_paragraph(lead + reference)
    document.add_paragraph("后续无关段落。")
    document.save(path)
    inspection = inspect_docx(path, full=True)

    rule = {"rule_id": "G01", "title": "释义-基金份额类别定义", "requirement": "追加直销条款。",
            "operations": [{"op": "replace", "old_text": reference, "new_text": "（新增）",
                            "condition": {"kind": "none", "value": None}}]}
    parts = find_old_text_all(list(inspection.paragraphs), reference)
    assert len(parts) >= 1
    target = Target(target_id="T001", rule_id="G01", location="p0", quote=reference,
                    state="needs_edit", reason="composite")
    target.composite = tuple(f"{location}:{start}:{end}"
                             for location, start, end in parts[0])
    new_text = "投资者通过直销机构申购本基金A类基金份额时不收取认购/申购费用。"
    edits = plan_composite_target(rule, target, inspection, _facts(), new_text=new_text)
    edit = edits[0]
    # The kept prefix ("…类别。") ends in punctuation, so no separator is added before the
    # new text and the sentence boundary is already correct.
    assert edit.replacement.startswith(lead + new_text)
    assert "。。" not in edit.replacement


def test_find_old_text_all_keeps_each_occurrence(tmp_path: Path):
    from fundlex.recall import find_old_text_all

    path = tmp_path / "duplicates.docx"
    document = Document()
    document.add_paragraph(LIQUIDATION_CLAUSE)
    document.add_paragraph("中间段落。")
    document.add_paragraph(LIQUIDATION_CLAUSE)
    document.save(path)
    inspection = inspect_docx(path, full=True)

    occurrences = find_old_text_all(list(inspection.paragraphs), LIQUIDATION_CLAUSE)
    assert [parts[0][0] for parts in occurrences] == ["p0", "p2"]


def test_survey_quote_offsets_survive_interspersed_spaces(tmp_path: Path):
    """find_old_text_all offsets are compact; the survey quote must be the real text.

    The offset bug shifted the recorded quote by however many spaces preceded the hit
    ("并报中国证监会备案" recorded as "人并报中国证监会备"), which then made
    _splice_fragment miss and pushed the target into compose (FL-D052).
    """
    from fundlex.recall import survey_targets

    path = tmp_path / "spaces.docx"
    document = Document()
    # A full-width/normal space before the hit, one mid-paragraph, one at the boundary.
    document.add_paragraph("并报中国证监会备案。")
    document.add_paragraph("错误偏差达到基金份额净值的 0.25%时，应当通报并报中国证监会备案。")
    document.add_paragraph("本款不涉及任何备案要求，无其他内容。")
    document.add_paragraph("管理人应当公告，并报中国证监会备案。")
    document.save(path)
    inspection = inspect_docx(path, full=True)

    rule = {"rule_id": "G05", "title": "费用调整", "requirement": "删除备案句。",
            "operations": [{"op": "delete", "old_text": "并报中国证监会备案",
                            "condition": {"kind": "none", "value": None}}]}
    hits = survey_targets(inspection, [rule], _facts(), known=[])
    quotes = {}
    for hit in hits:
        if hit.operation != "survey":
            continue
        location, start, end = hit.parts[0]
        span = inspection.by_location(location)
        quotes[location] = hit.quote
        # Every recorded quote is the operation reference itself...
        assert hit.quote.replace(" ", "") == "并报中国证监会备案"
        # ...and its compact slice is exactly that reference.
        compact = "".join(char for char in display_text(span.text)[0] if not char.isspace())
        assert compact[start:end] == "并报中国证监会备案"
    assert {"p0", "p1", "p3"} <= set(quotes)


def test_survey_drops_bare_tail_seen_in_unrelated_clause(tmp_path: Path):
    """A bare clause tail must be corroborated by the rule's own context wording.

    G05 deletes the filing phrase "并报中国证监会备案", which also closes an unrelated
    valuation-error clause; without the fee-adjustment wording ("调低赎回费率") in the
    same paragraph the hit is not this rule's target (FL-D052).
    """
    from fundlex.recall import survey_targets

    path = tmp_path / "tail.docx"
    document = Document()
    document.add_paragraph("错误偏差达到基金份额净值的0.25%时，应当通报基金托管人并报中国证监会备案。")
    document.add_paragraph("调整份额类别时，可以调低赎回费率或变更收费方式，调整前需公告并报中国证监会备案。")
    document.save(path)
    inspection = inspect_docx(path, full=True)

    rule = {"rule_id": "G05", "title": "类别调整句", "section_hint": "基金的基本情况",
            "requirement": "将“调低赎回费率”改为“调低销售服务费率”，并删除“并报中国证监会备案”。",
            "operations": [
                {"op": "replace", "old_text": "调低赎回费率", "new_text": "调低销售服务费率",
                 "condition": {"kind": "has_service_fee", "value": True}},
                {"op": "delete", "old_text": "并报中国证监会备案",
                 "condition": {"kind": "none", "value": None}},
            ]}
    hits = survey_targets(inspection, [rule], _facts(), known=[])
    locations = {location for hit in hits for location in hit.locations}
    assert locations == {"p1"}, locations


def test_plan_target_applies_delete_and_replace_in_one_span(tmp_path: Path):
    """One rule can both delete and replace inside the same span (FL-D052)."""
    from fundlex.recall import Target, plan_target

    path = tmp_path / "mixed.docx"
    document = Document()
    paragraph = ("基金管理人在履行适当程序后可以调整基金份额类别，或者在法律法规和基金合同规定的范围内"
                 "变更现有基金份额类别的申购费率、调低赎回费率或变更收费方式，调整前基金管理人需及时公告"
                 "并报中国证监会备案，无需召开份额持有人大会。")
    document.add_paragraph(paragraph)
    document.save(path)
    inspection = inspect_docx(path, full=True)

    rule = {"rule_id": "G05", "title": "类别调整句", "requirement": "替换费率并删除备案句。",
            "operations": [
                {"op": "replace", "old_text": "调低赎回费率", "new_text": "调低销售服务费率",
                 "condition": {"kind": "has_service_fee", "value": True}},
                {"op": "delete", "old_text": "并报中国证监会备案",
                 "condition": {"kind": "none", "value": None}},
            ], "source_paragraphs": [17]}
    target = Target(target_id="T1", rule_id="G05", location="p0", quote=paragraph,
                    state="needs_edit", reason="")
    target.start, target.end = 0, len(paragraph)
    target.resolution = "resolved"

    class _Gateway:
        max_context_chars = 60000

    edit = plan_target(rule, target, inspection, _facts(), _Gateway(), compose_budget=[1])
    assert edit is not None
    assert "调低销售服务费率" in edit.replacement
    assert "调低赎回费率" not in edit.replacement
    assert "并报中国证监会备案" not in edit.replacement
    assert edit.content_source == "fragment_mixed"


def test_fragment_replace_separates_authorized_text_from_unterminated_lead():
    """A mid-span fragment replace must not glue the new policy sentence to a broken lead.

    Guide G11 replaces a subscriber-fee sentence with the direct-sales policy. When the
    kept lead before the matched fragment does not end in punctuation ("…本基金 A 类基金
    份额的"), splicing it verbatim produced "…份额的通过直销机构申购…" (FL-D067).
    """
    from fundlex.recall import _splice_fragment

    actual = ("本基金 A 类基金份额的申购费用由投资人承担，不列入基金财产。"
              "本基金 C 类基金份额不收取申购费用。")
    old_text = "申购费用由投资人承担，不列入基金财产。"
    new_text = "通过直销机构申购本基金不收取申购费，通过其他销售机构申购本基金的申购费用由投资人承担。"
    replaced = _splice_fragment(actual, old_text, new_text)
    assert replaced is not None
    assert "本基金 A 类基金份额的。通过直销机构申购本基金不收取申购费" in replaced
    assert "份额的通过直销机构" not in replaced

    # A pure deletion keeps the surrounding text verbatim with no inserted punctuation.
    deleted = _splice_fragment(actual, old_text, "")
    assert deleted == "本基金 A 类基金份额的本基金 C 类基金份额不收取申购费用。"


def test_search_keys_add_class_stripped_anchor_for_delete_rules():
    """A delete rule gains a short leading anchor so wording variants are still found.

    The guide writes "C类基金份额的销售服务费每日计提，按月支付…" while contracts write
    "基金销售服务费每日计提，逐日累计至每月月末，按月支付…". The long reference does not
    fuzzy-match either, so the prefix-stripped first clause anchors the clause (FL-D052).
    """
    from fundlex.recall import operation_search_keys
    rule = {"rule_id": "G15", "title": "基金费用与税收-基金销售服务费",
            "requirement": "删除原条款中“C类基金份额的销售服务费每日计提，按月支付。”。",
            "operations": [{"op": "delete",
                            "old_text": "C类基金份额的销售服务费每日计提，按月支付。由基金管理人"
                                        "向基金托管人发送基金销售服务费划付指令。",
                            "condition": {"kind": "none", "value": None}}]}
    keys = operation_search_keys(rule, _facts())
    assert "销售服务费每日计提" in keys
    # A rule that also replaces text keeps no reduced anchor (it is not delete-only).
    replace_rule = {"rule_id": "G05", "title": "基金费用与税收-销售服务费",
                    "requirement": "将“调低赎回费率”改为“调低销售服务费率”。",
                    "operations": [{"op": "replace",
                                    "old_text": "C类基金份额的销售服务费可以适当调低赎回费率",
                                    "new_text": "C类基金份额的销售服务费可以适当调低销售服务费率",
                                    "condition": {"kind": "none", "value": None}}]}
    keys = operation_search_keys(replace_rule, _facts())
    assert keys == ["C类基金份额的销售服务费可以适当调低赎回费率"]


def test_delete_reference_continues_into_next_paragraph(tmp_path: Path):
    """A delete reference split across two paragraphs binds both as one composite target.

    The guide's G17 reference deletes both "…服务费等。" and "销售服务费使用范围不包括…";
    a contract puts the second sentence in the following paragraph. Both must be bound, or
    the second is left behind (FL-D063).
    """
    from fundlex.recall import survey_targets
    document = Document()
    document.add_paragraph("七、基金费用与税收")
    document.add_paragraph("销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、促销活动费、基金份额持有人服务费等。")
    document.add_paragraph("销售服务费使用范围不包括基金募集期间的上述费用。")
    document.add_paragraph("其他无关段落。")
    path = tmp_path / "g17-split.docx"
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G17", "title": "基金销售服务费", "requirement": "删除用途与范围句。",
            "operations": [{"op": "delete",
                            "old_text": "C类基金份额的销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、促销活动费、基金份额持有人服务费等。C类基金份额的销售服务费使用范围不包括基金募集期间的上述费用。",
                            "condition": {"kind": "none", "value": None}}]}
    hits = survey_targets(inspection, [rule], _facts(), known=[])
    survey = [hit for hit in hits if hit.operation == "survey"]
    assert survey, "the split delete reference must produce a survey hit"
    locations = survey[0].locations
    assert len(locations) == 2, locations
    # The composite parts must be usable offsets: measured in each paragraph's compact
    # text and never past its end, or plan_composite_target raises
    # RECALL_TARGET_COMPOSITE_UNRESOLVED and leaves the clause behind (FL-D080).
    from fundlex.recall import compact_offsets
    assert survey[0].parts, "a split delete must carry composite parts"
    for part_location, start, end in survey[0].parts:
        body, _ = compact_offsets(inspection.by_location(part_location).text)
        assert 0 <= start < end <= len(body), (part_location, start, end, len(body))


def test_delete_anchor_extends_to_the_whole_clause(tmp_path: Path):
    """A short reference anchor grows to the sentence it starts for a delete-only rule.

    The guide's G17 lead reduces to "销售服务费主要用于支付销售机构佣金" once the class
    prefix is stripped; binding only that fragment would delete the lead and leave the rest
    of the sentence. The span must extend to the sentence terminator (FL-D063).
    """
    from fundlex.recall import _extend_to_clause_end
    text = "销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费、促销活动费、基金份额持有人服务费等。下一段。"
    end = len("销售服务费主要用于支付销售机构佣金")
    extended = _extend_to_clause_end(text, 0, end, "销售服务费主要用于支付销售机构佣金",
                                     delete_only=True)
    assert text[:extended].endswith("基金份额持有人服务费等。")
    # A replace rule keeps the matched fragment untouched.
    assert _extend_to_clause_end(text, 0, end, "x", delete_only=False) == end
    # A span that already ends at a terminator is left alone.
    already = "短句。后面。"
    assert _extend_to_clause_end(already, 0, 3, "短句。", delete_only=True) == 3


def test_delete_keyword_bag_finds_paraphrased_payment_clause(tmp_path: Path):
    """A paraphrased sales-fee payment sentence is found by the rule's keyword signature.

    The guide's G15 reference says "C类基金份额的销售服务费每日计提，按月支付…" while
    contracts write "基金销售服务费每日计算…" or split the sentence across two paragraphs.
    No single string anchor covers every wording, so the delete-only rule's keyword
    signature (domain term + payment action + period) anchors the clause instead (FL-D062).
    """
    from fundlex.recall import Target, survey_targets
    rewrite = "基金销售服务费每日计算，逐日累计至每月月末，按月支付，经基金管理人与基金托管人核对一致后，由基金托管人于次月前5个工作日内从基金财产中一次性划出。"
    document = Document()
    document.add_paragraph("七、基金费用与税收")
    document.add_paragraph("H为C类基金份额每日应计提的销售服务费")
    document.add_paragraph(rewrite)
    # A fee-purpose sentence in the same section must not satisfy the payment signature
    # (it has no payment-period word).
    document.add_paragraph("销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费。")
    # A托管费 sentence carrying a neighbouring "销售服务费" heading must not be hit either.
    document.add_paragraph(
        "基金托管费每日计算，逐日累计至每月月末，按月支付，经基金管理人与基金托管人双方"
        "核对无误后，基金托管人按照与基金管理人协商一致的方式于次月前5个工作日内从基金财产"
        "中一次性支付给基金托管人。若遇法定节假日、公休日等，支付日期顺延。（三）销售服务费")
    path = tmp_path / "g15-paraphrase.docx"
    document.save(path)
    inspection = inspect_docx(path, full=True)
    hits = survey_targets(inspection, guide_rules(), _facts(), known=[])
    # The fixture guide has no G15; the check is on the derived signature helper.
    from fundlex.recall import _delete_keyword_bag, _keyword_bag_in_clause
    rule = {"rule_id": "G15", "title": "基金费用与税收-基金销售服务费",
            "requirement": "删除销售服务费支付句。",
            "operations": [{"op": "delete",
                            "old_text": "C类基金份额的销售服务费每日计提，按月支付。",
                            "condition": {"kind": "none", "value": None}}]}
    bag = _delete_keyword_bag(rule, _facts())
    assert bag == ("销售服务费", "支付", "按月")
    assert _keyword_bag_in_clause(rewrite, bag)
    assert not _keyword_bag_in_clause(
        "销售服务费主要用于支付销售机构佣金、以及基金管理人的基金行销广告费。", bag)
    assert not _keyword_bag_in_clause(
        "基金托管费每日计算，逐日累计至每月月末，按月支付，经基金管理人与基金托管人双方核对无误后，"
        "基金托管人按照与基金管理人协商一致的方式于次月前5个工作日内从基金财产中一次性支付给基金托管人。"
        "若遇法定节假日、公休日等，支付日期顺延。（三）销售服务费", bag)


def test_whole_paragraph_delete_when_wording_differs(tmp_path: Path):
    """A whole-paragraph delete survives a composer that cannot match the wording.

    The guide's G15 reference says "C类基金份额的销售服务费每日计提…" while contracts also
    write "基金销售服务费每日计提…". The survey binds the whole clause as one target, the
    composer cannot match the text to trim and reports DELETE_NOT_SHRINKING, but the whole
    paragraph is the delete material, so it must be removed rather than left unresolved
    (FL-D052).
    """
    from fundlex.model import ModelError
    from fundlex.recall import Target, plan_target
    sentence = ("基金销售服务费每日计提，逐日累计至每月月末，按月支付，经基金管理人与基金托管人"
                "核对一致后，由基金托管人于次月前5个工作日内从基金财产中一次性划出。")
    document = Document()
    document.add_paragraph("七、基金费用与税收")
    document.add_paragraph(sentence)
    document.add_paragraph("本基金A类基金份额不收取销售服务费。")
    path = tmp_path / "whole-paragraph-delete.docx"
    document.save(path)
    inspection = inspect_docx(path, full=True)

    class _FailingGateway:
        max_context_chars = 60000

        def call(self, **kwargs):
            raise ModelError("MODEL_CALL_BUDGET_EXCEEDED")

    def _composer_refuses(span, target, rule, actual, fixed_text, facts, gateway, *,
                          mode, variants):
        from fundlex.recall import RecallError
        raise RecallError("RECALL_COMPOSE_DELETE_NOT_SHRINKING", {"target": target.target_id})

    rule = {"rule_id": "G15", "title": "销售服务费支付句", "requirement": "删除支付句。",
            "operations": [{"op": "delete",
                            "old_text": "C类基金份额的销售服务费每日计提，按月支付。",
                            "condition": {"kind": "none", "value": None}}]}
    target = Target(target_id="T001", rule_id="G15", location="p1", quote=sentence,
                    state="needs_edit", reason="whole paragraph")
    target.start, target.end, target.resolution = 0, len(sentence), "resolved"

    import fundlex.recall as recall_module
    original = recall_module.compose_variant
    recall_module.compose_variant = _composer_refuses
    try:
        edit = plan_target(rule, target, inspection, _facts(), _FailingGateway(),
                           compose_budget=[1])
    finally:
        recall_module.compose_variant = original
    assert edit is not None
    assert edit.replacement == ""
    assert edit.content_source in {"whole_paragraph_delete", "tier1_delete"}


def test_composite_plan_offsets_after_spaces(tmp_path: Path):
    """plan_composite_target must map compact part offsets through the compact index."""
    from fundlex.recall import Target, find_old_text_all, plan_composite_target

    path = tmp_path / "composite-spaces.docx"
    document = Document()
    # Spaces on both sides of the paragraph break must not shift the deletion geometry.
    document.add_paragraph("C 类基金份额的销售服务费每日计算，逐日累计至每月月末，按月支付，")
    document.add_paragraph("由基金管理人与基金托管人核对一致后按上述方式支付。其余文字应保留。")
    document.save(path)
    inspection = inspect_docx(path, full=True)

    reference = ("C类基金份额的销售服务费每日计算，逐日累计至每月月末，按月支付，"
                 "由基金管理人与基金托管人核对一致后按上述方式支付。")
    parts = find_old_text_all(list(inspection.paragraphs), reference)
    assert len(parts) == 1 and [item[0] for item in parts[0]] == ["p0", "p1"]
    rule = {"rule_id": "G15", "title": "销售服务费支付句", "requirement": "删除支付句。",
            "operations": [{"op": "delete", "old_text": reference,
                            "condition": {"kind": "none", "value": None}}]}
    target = Target(target_id="T001", rule_id="G15", location="p0", quote=reference,
                    state="needs_edit", reason="cross-paragraph")
    target.composite = tuple(f"{location}:{start}:{end}" for location, start, end in parts[0])
    edits = plan_composite_target(rule, target, inspection, _facts(), new_text="")
    merged = {edit.location: edit.replacement for edit in edits}
    # The anchor paragraph is emptied (whitespace may remain); the follower keeps its tail.
    assert merged["p0"].strip() == ""
    assert merged["p1"] == "其余文字应保留。"


@pytest.mark.parametrize("ending,separator", [("", "。"), ("。", ""), ("；", ""), ("，", ""),
                                               ("）", "。"), ("。”", ""), ("”", "。"), ("  ", "。")])
def test_append_at_unpunctuated_paragraph_end_generates_valid_docx(tmp_path, ending, separator):
    from fundlex.recall import Target, plan_target
    source = tmp_path / "source.docx"
    document = Document()
    original = "基金管理人根据约定收取费用" + ending
    document.add_paragraph(original)
    document.save(source)
    inspection = inspect_docx(source, full=True)
    added = "投资者通过直销机构申购时不收取申购费用。"
    rule = {"rule_id": "GTEST", "title": "新增收费约定", "source_paragraphs": [1],
            "operations": [{"op": "append", "new_text": added, "condition": {"kind": "none"}}]}
    target = Target("T1", "GTEST", "p0", original, "needs_edit", "", 0, len(original), "resolved")
    edit = plan_target(rule, target, inspection, _facts(), None, compose_budget=[0])
    assert edit.replacement == separator + added
    edits = build_docx_edits([edit], inspection)
    plan = validate_edit_plan(source, edits, {},
                              content_check=content_check_factory([edit.as_dict()], inspection))
    out = tmp_path / "out"
    out.mkdir()
    final_path = out / "final.docx"
    artifacts = generate_artifacts(source, plan, out / "redline.docx", final_path)
    assert artifacts.final_sha256
    final = inspect_docx(final_path, full=True)
    assert final.by_location("p0").text == original + separator + added


def test_fuzzy_span_clause_snaps_to_sentence_boundary():
    """A fuzzy variant whose span stops mid-sentence must not leave a stub (FL-D079).

    ``find_fuzzy`` matched "…比例进行分" (end=77) of the contract's
    "…比例进行分配。" and the kept suffix "配。" was appended after the replacement,
    yielding "…（如有）。配。". The clause-snapped matcher extends to the real
    sentence stop instead.
    """
    from fundlex.recall import fuzzy_span, fuzzy_span_clause
    actual = ("依据基金财产清算的分配方案，将基金财产清算后的全部剩余资产扣除基金财产清算费用、"
              "交纳所欠税款并清偿基金债务后，按基金份额持有人持有的各类基金份额比例进行分配。")
    reference = ("依据基金财产清算的分配方案，将基金财产清算后的全部剩余资产扣除基金财产清算费用、"
                 "交纳所欠税款并清偿基金债务后，按基金份额持有人持有的基金份额比例进行分配。")
    raw = fuzzy_span(actual, reference, threshold=0.95)
    snapped = fuzzy_span_clause(actual, reference, threshold=0.95)
    assert raw is not None and snapped is not None
    # The raw matcher leaves a "配。" stub; the snapped matcher consumes the whole clause.
    assert actual[raw[1]:] == "配。"
    assert actual[snapped[1]:] == ""
    assert actual[snapped[0]:snapped[1]].endswith("。")


def test_class_label_strips_share_variants_and_drops_placeholders():
    """Class names must render as the canonical "A类基金份额", never doubled (FL-D091).

    Contracts name a class "A类份额" or "A类基金份额" interchangeably, and the fact
    extractor sometimes writes "default" for a single unnamed class. The first produced
    "A类份额类基金份额" and the second "default类基金份额" in a real dev-set run.
    """
    assert class_label(["A类份额", "C类份额"]) == "A类、C类基金份额"
    assert class_label(["A类基金份额", "Y类基金份额"]) == "A类、Y类基金份额"
    assert class_label(["default", "A"]) == "A类基金份额"
    # A single unnamed class still yields grammatical wording rather than a placeholder.
    assert class_label(["default"]) == "基金份额"
    assert "default" not in class_label(["default", "default"])


def test_verify_quote_accepts_ellipsis_and_paragraph_boundary(tmp_path: Path):
    from fundlex.recall import _verify_quote
    path = tmp_path / "quote.docx"
    document = Document()
    document.add_paragraph("本基金销售服务费每日计提，按月支付。")
    document.add_paragraph("基金的运作方式")
    document.add_paragraph("契约型开放式")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    assert _verify_quote(inspection, "销售服务费每日计提……按月支付")
    assert _verify_quote(inspection, "基金的运作方式\n契约型开放式")
    assert not _verify_quote(inspection, "本基金销售服务费按年支付")


def test_composite_part_verification_uses_compact_offsets(tmp_path: Path):
    from fundlex.recall import _component_ok, compact_offsets
    path = tmp_path / "composite.docx"
    document = Document()
    document.add_paragraph("本基金销售服务 费每日计提，按月支付。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    span = inspection.by_location("p0")
    compact, compact_map = compact_offsets(span.text)
    start_clean = compact.index("销售服务费")
    end_clean = start_clean + len("销售服务费")
    start_raw = compact_map[start_clean]
    end_raw = compact_map[end_clean - 1] + 1
    component = {
        "rule_id": "G15", "target_id": "T1", "location": "p0",
        "start": start_raw, "end": end_raw,
        "expected_old_text": span.text[start_raw:end_raw],
        "replacement": span.text[:start_raw] + span.text[end_raw:],
        "content_source": "composite_delete",
        "evidence": {"parts": [f"p0:{start_clean}:{end_clean}"],
                     "part_index": 0, "unit": "G15"},
    }
    assert _component_ok(component, inspection) is True
    tampered = dict(component)
    tampered["replacement"] = component["replacement"][:-2]
    assert _component_ok(tampered, inspection) is False


def test_satisfied_replace_does_not_suppress_delete_survey(tmp_path: Path):
    """G05: a satisfied replace must not hide the still-present delete operation.

    ``satisfied_by_new_text`` is rule-level, so it used to mark the whole
    (rule, location) covered and the survey skipped the deletion ("并报中国证监会
    备案") that was still in the same paragraph.
    """
    from fundlex.recall import survey_targets
    path = tmp_path / "satisfied-and-delete.docx"
    document = Document()
    document.add_paragraph("基金管理人可以适当调低销售服务费率，并报中国证监会备案。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "GX", "title": "类别调整句",
            "requirement": "调低销售服务费率并删除备案句。",
            "operations": [
                {"op": "replace", "old_text": "调低赎回费率",
                 "new_text": "调低销售服务费率", "condition": {"kind": "none", "value": None}},
                {"op": "delete", "old_text": "并报中国证监会备案",
                 "condition": {"kind": "none", "value": None}},
            ]}
    hits = survey_targets(inspection, [rule], _facts())
    assert any(hit.operation == "satisfied" for hit in hits)
    assert any("备案" in hit.quote for hit in hits if hit.operation != "satisfied")

    clean_path = tmp_path / "satisfied-only.docx"
    document = Document()
    document.add_paragraph("基金管理人可以适当调低销售服务费率。")
    document.save(clean_path)
    clean_hits = survey_targets(inspect_docx(clean_path, full=True), [rule], _facts())
    assert any(hit.operation == "satisfied" for hit in clean_hits)
    assert not any("备案" in hit.quote for hit in clean_hits if hit.operation != "satisfied")


def test_satisfied_replace_does_not_suppress_fuzzy_delete_variant(tmp_path: Path):
    """The delete survey also covers a bounded fuzzy variant (FL-D117 residual)."""
    from fundlex.recall import survey_targets
    path = tmp_path / "fuzzy-delete.docx"
    document = Document()
    document.add_paragraph(
        "基金管理人可以适当调低销售服务费率。"
        "基金管理人应当在调整前事先公告并报中国证监会备案。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "GX", "title": "类别调整句",
            "requirement": "调低销售服务费率并删除备案句。",
            "operations": [
                {"op": "replace", "old_text": "调低赎回费率",
                 "new_text": "调低销售服务费率", "condition": {"kind": "none", "value": None}},
                {"op": "delete",
                 "old_text": "基金管理人应当在调整前及时公告并报中国证监会备案",
                 "condition": {"kind": "none", "value": None}},
            ]}
    hits = survey_targets(inspection, [rule], _facts())
    assert any(hit.operation == "satisfied" for hit in hits)
    assert any("事先公告" in hit.quote for hit in hits if hit.operation != "satisfied")


def test_survey_operation_coverage_does_not_hide_other_operations(tmp_path: Path):
    """A target covering only the replace operation must not suppress the delete hit."""
    from fundlex.recall import survey_targets
    path = tmp_path / "operation-coverage.docx"
    document = Document()
    document.add_paragraph("基金管理人可以适当调低赎回费率，并报中国证监会备案。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "GX", "title": "类别调整句",
            "requirement": "调低销售服务费率并删除备案句。",
            "operations": [
                {"op": "replace", "old_text": "调低赎回费率",
                 "new_text": "调低销售服务费率", "condition": {"kind": "none", "value": None}},
                {"op": "delete", "old_text": "并报中国证监会备案",
                 "condition": {"kind": "none", "value": None}},
            ]}
    hits = survey_targets(inspection, [rule], _facts(),
                          known={("GX", "p0"): {"GX:0"}})
    assert not any("调低赎回费率" in hit.quote for hit in hits)
    assert any("备案" in hit.quote for hit in hits)


def test_unknown_operation_state_fails_open_for_specific_hit():
    """An unknown operation identity must not suppress a specific operation hit (FL-D118)."""
    from fundlex.recall import operations_cover_hit
    known = {("GX", "p0"): None}
    assert operations_cover_hit(known, "GX", ["p0"], ("GX:1",)) is False
    assert operations_cover_hit(known, "GX", ["p0"], ()) is True
    assert operations_cover_hit({("GX", "p0"): {"GX:1"}}, "GX", ["p0"], ("GX:1",)) is True
    assert operations_cover_hit({("GX", "p0"): {"GX:0"}}, "GX", ["p0"], ("GX:1",)) is False


def test_g05_rejects_a_target_outside_the_share_class_clause(tmp_path: Path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import RecallError, Target, plan_target

    path = tmp_path / "g05-wrong-site.docx"
    document = Document()
    document.add_paragraph("对于可不经基金份额持有人大会决议通过的事项，由基金管理人和基金托管人同意后变更并公告，并报中国证监会备案。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G05", "title": "类别调整", "requirement": "删除备案句",
            "operations": [{"op": "delete", "old_text": "并报中国证监会备案",
                            "condition": {"kind": "none", "value": None}}]}
    span = inspection.by_location("p0")
    target = Target("T1", "G05", "p0", "并报中国证监会备案", "needs_edit", "",
                    0, len(span.text), "resolved")
    with pytest.raises(RecallError) as excinfo:
        plan_target(rule, target, inspection, _facts(), ScriptedGateway({}),
                    compose_budget=[1])
    assert excinfo.value.code == "TARGET_SOURCE_MISMATCH"


def test_g15_business_notice_understands_chinese_numeral_deadlines():
    from fundlex.recall import g15_contract_deviation
    reasons = g15_contract_deviation(["次月首日起10个工作日内支付"],
                                     "次月初五个工作日内支付")
    assert reasons and "5个工作日" in reasons[0] and "10个工作日" in reasons[0]


def test_delete_residual_copy_gets_a_deterministic_target(tmp_path: Path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, ensure_delete_targets

    path = tmp_path / "delete-copy.docx"
    document = Document()
    document.add_paragraph("销售服务费使用范围不包括基金募集期间的上述费用。")
    document.add_paragraph("若遇法定节假日、公休日等，支付日期顺延。")
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G17", "title": "用途", "requirement": "删除用途条款",
            "operations": [{"op": "delete",
                            "old_text": "C类基金份额的销售服务费使用范围不包括基金募集期间的上述费用。",
                            "condition": {"kind": "none"}}]}
    span = inspection.by_location("p0")
    existing = Target("T1", "G17", "p0", "销售服务费使用范围", "needs_edit", "",
                      0, 5, "resolved")
    created = ensure_delete_targets([existing], {"G17": rule}, inspection)
    # The existing target covers only a 5-character prefix; the full sentence still
    # needs its own target. The generic holiday sentence has no domain anchor.
    assert len(created) == 1
    created = ensure_delete_targets([], {"G17": rule}, inspection)
    assert len(created) == 1
    assert created[0].location == "p0"
    assert created[0].resolution == "resolved"
    assert "使用范围不包括" in created[0].quote


def test_g01_accepts_a_class_definition_without_the_word_fenzlei(tmp_path: Path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import Target, plan_target

    path = tmp_path / "g01-definition.docx"
    document = Document()
    sentence = "本基金根据认购费、申购费、销售服务费收取方式的不同，将基金份额分为不同的类别。"
    document.add_paragraph(sentence)
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G01", "title": "类别定义", "requirement": "追加直销政策",
            "operations": [{"op": "append",
                            "new_text": "通过直销机构申购本基金A类基金份额不收取申购费。",
                            "condition": {"kind": "none", "value": None}}],
            "slot_tokens": ["A类基金份额"], "class_substitution": True}
    span = inspection.by_location("p0")
    facts = _facts()
    facts.fact_status.update({"share_class_structure": "VERIFIED",
                              "subscription_fee_group": "VERIFIED",
                              "service_fee_group": "VERIFIED"})
    target = Target("T1", "G01", "p0", sentence, "needs_edit", "",
                    0, len(span.text), "resolved")
    edit = plan_target(rule, target, inspection, facts, ScriptedGateway({}),
                       compose_budget=[1])
    assert edit is not None
    assert "直销机构" in edit.replacement


def test_g01_rejects_a_product_summary_definition_site(tmp_path: Path):
    from docx import Document
    from fundlex.docx import inspect_docx
    from fundlex.recall import RecallError, Target, plan_target

    path = tmp_path / "g01-wrong-site.docx"
    document = Document()
    sentence = "基金产品资料概要：指基金产品资料概要及其更新，是基金份额持有人了解基金情况的文件。"
    document.add_paragraph(sentence)
    document.save(path)
    inspection = inspect_docx(path, full=True)
    rule = {"rule_id": "G01", "title": "类别定义", "requirement": "追加直销政策",
            "operations": [{"op": "append", "new_text": "通过直销机构申购不收取申购费。",
                            "condition": {"kind": "none", "value": None}}],
            "class_substitution": True}
    span = inspection.by_location("p0")
    facts = _facts()
    facts.fact_status.update({"share_class_structure": "VERIFIED",
                              "subscription_fee_group": "VERIFIED",
                              "service_fee_group": "VERIFIED"})
    target = Target("T1", "G01", "p0", sentence, "needs_edit", "",
                    0, len(span.text), "resolved")
    with pytest.raises(RecallError) as excinfo:
        plan_target(rule, target, inspection, facts, ScriptedGateway({}),
                    compose_budget=[1])
    assert excinfo.value.code == "TARGET_SOURCE_MISMATCH"
