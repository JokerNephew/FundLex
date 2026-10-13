"""Offline gold coordinates must not confuse descendant paragraphs and body IDs."""
from docx import Document
from tools.evaluation.pilot_compare import source_map, check_evidence, ratio
import pytest


def test_independent_gold_paragraph_map_handles_table_without_shifting_body_ids(tmp_path):
    path=tmp_path/'independent.docx'
    doc=Document();doc.add_paragraph('第一段')
    doc.add_table(rows=1,cols=1).cell(0,0).text='表格段'
    doc.add_paragraph('第二段');doc.save(path)
    mapped=source_map(path)
    assert [(m['text'],m['location']) for m in mapped.values()]==[
        ('第一段','p0'),('表格段','tbl0/r0/c0/p0'),('第二段','p1')]
    check_evidence({'paragraph_index':2,'start':0,'end':3,'quote':'第二段'},mapped)
    with pytest.raises(AssertionError,match='coordinate mismatch'):
        check_evidence({'paragraph_index':1,'start':0,'end':3,'quote':'第二段'},mapped)


def test_empty_metric_denominator_is_unknown_not_perfect():
    assert ratio(0,0) is None
    assert ratio(0,3)==0
