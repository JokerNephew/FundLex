"""Recall-first engine binding through the task store, worker and delivery gate.

The scripted processors replace the model calls only; task binding, the guide
snapshot check, artifact publication and the download re-validation are the
real implementations under test.
"""
from __future__ import annotations

import hashlib
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from fundlex import guide_store
from fundlex.api import create_app
from fundlex.docx import generate_unchanged_artifacts, validate_edit_plan
from fundlex.guide import GUIDE_SCHEMA_VERSION
from fundlex.jobs import Worker
from fundlex.model import ModelError
from fundlex.recall_pipeline import RecallPipelineStop
from fundlex.rules import load_ruleset
from fundlex.store import NotFound, RunStore, validate_delivery
from fundlex.tasks import TaskStore

from test_guide import (GUIDE_PARAGRAPHS, ScriptedGateway, build_docx, clean_verification, valid_parsed)
from test_recall_core import build_contract, guide_rules
from test_recall_pipeline import GUIDE_SHA, PipelineGateway, base_responses, guide_document

OWNER = "recall-user"
HEADERS = {"X-FundLex-User": OWNER}


def contract_bytes(label: str = "合同正文示例") -> bytes:
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph(label)
    document.save(stream)
    return stream.getvalue()


def bound_store(tmp_path) -> tuple[RunStore, dict]:
    store = RunStore(tmp_path / "state")
    guide_store.ensure_builtin_guide(store)
    version = guide_store.get_version(store, OWNER, guide_store.BUILTIN_GUIDE_ID)
    return store, version


def submit_task(store, *, data: bytes, guide_version_id: str, guide_rules_sha256: str,
                engine_id: str = "recall_first", owner: str = OWNER):
    tasks = TaskStore(store)
    task = tasks.create(owner, "召回任务", load_ruleset("v1"),
                        [{"client_id": "c1", "name": "contract.docx", "size": len(data)}],
                        engine_id=engine_id, guide_version_id=guide_version_id,
                        guide_rules_sha256=guide_rules_sha256)
    receipts = tasks.uploads(owner, task["id"], [{"client_id": "c1", "name": "contract.docx", "data": data}])
    assert receipts["receipts"][0]["status"] == "accepted"
    return tasks, task["id"], receipts["receipts"][0]["batch_id"], receipts["receipts"][0]["document_id"]


def complete_recall_processor(source, output, ruleset):
    plan = validate_edit_plan(source, [], {})
    validation = generate_unchanged_artifacts(source, plan, output / "redline.docx", output / "final.docx")
    return {"engine_id": "recall_first", "schema_version": "3", "execution_mode": "delivery",
            "qualification_status": "eligible", "delivery_status": "complete",
            "output_validation": {"kind": "recall_validated_unchanged", "checks": list(validation.checks),
                                  "final_sha256": validation.final_sha256,
                                  "redline_sha256": validation.redline_sha256},
            "artifacts": {"final": "final.docx", "redline": "redline.docx"}}


def test_recall_task_records_binding_and_publishes_final_pair(tmp_path):
    app = create_app(tmp_path / "state")
    store = app.state.store
    guide_store.ensure_builtin_guide(store)
    version = guide_store.get_version(store, OWNER, guide_store.BUILTIN_GUIDE_ID)
    data = contract_bytes()
    tasks, task_id, batch_id, document_id = submit_task(
        store, data=data, guide_version_id=version["id"], guide_rules_sha256=version["rules_sha256"])
    assert Worker(store, complete_recall_processor).run_once()
    snapshot = tasks.get(OWNER, task_id)
    assert snapshot["engine_id"] == "recall_first"
    assert snapshot["guide_snapshot"] == {"version_id": guide_store.BUILTIN_GUIDE_ID,
                                          "rules_sha256": version["rules_sha256"]}
    assert snapshot["status"] == "completed"
    with store._connect() as db:
        batch = db.execute("SELECT engine_id,guide_version_id,guide_rules_sha256 FROM batches WHERE id=?",
                           (batch_id,)).fetchone()
    assert (batch["engine_id"], batch["guide_version_id"], batch["guide_rules_sha256"]) == \
        ("recall_first", guide_store.BUILTIN_GUIDE_ID, version["rules_sha256"])
    client = TestClient(app)
    detail = client.get(f"/api/tasks/{task_id}", headers=HEADERS).json()
    assert detail["engine_id"] == "recall_first"
    assert detail["guide_snapshot"]["version_id"] == guide_store.BUILTIN_GUIDE_ID
    document = client.get(f"/api/tasks/{task_id}/documents", headers=HEADERS).json()["items"][0]
    assert document["job_status"] == "succeeded" and document["delivery_status"] == "complete"
    assert {"final", "redline"}.issubset(document["downloads"])
    final = client.get(document["downloads"]["final"], headers=HEADERS)
    assert final.status_code == 200 and final.content == data
    assert client.get(document["downloads"]["redline"], headers=HEADERS).status_code == 200
    report = client.get(document["report_url"], headers=HEADERS).json()
    assert report["engine_id"] == "recall_first" and report["schema_version"] == "3"


def test_incomplete_recall_publishes_working_pair_only(tmp_path):
    app = create_app(tmp_path / "state")
    store = app.state.store
    guide_store.ensure_builtin_guide(store)
    version = guide_store.get_version(store, OWNER, guide_store.BUILTIN_GUIDE_ID)

    def processor(source, output, ruleset):
        plan = validate_edit_plan(source, [], {})
        validation = generate_unchanged_artifacts(source, plan, output / "working_redline.docx",
                                                  output / "working_final.docx")
        return {"engine_id": "recall_first", "schema_version": "3", "execution_mode": "delivery",
                "qualification_status": "eligible", "delivery_status": "incomplete",
                "unresolved": [{"kind": "target", "rule_id": "G01", "reason": "RECALL_TARGET_UNRESOLVED"}],
                "output_validation": {"kind": "recall_working_edits",
                                      "checks": [*validation.checks, "actual_content_verification"],
                                      "final_sha256": validation.final_sha256,
                                      "redline_sha256": validation.redline_sha256},
                "artifacts": {"working_final": "working_final.docx", "working_redline": "working_redline.docx"}}

    data = contract_bytes()
    tasks, task_id, batch_id, document_id = submit_task(
        store, data=data, guide_version_id=version["id"], guide_rules_sha256=version["rules_sha256"])
    assert Worker(store, processor).run_once()
    document = tasks.documents(OWNER, task_id)["items"][0]
    assert document["delivery_status"] == "incomplete"
    assert {"working_final", "working_redline"}.issubset(document["downloads"])
    assert "final" not in document["downloads"] and "redline" not in document["downloads"]
    client = TestClient(app)
    refused = client.get(f"/api/batches/{batch_id}/documents/{document_id}/download/final", headers=HEADERS)
    assert refused.status_code == 404
    # api.py still whitelists legacy kinds; the store layer must serve the working pair.
    working_path, working_sha = store.get_artifact(OWNER, batch_id, document_id, "working_final")
    assert working_path.read_bytes() == data
    assert working_sha == hashlib.sha256(data).hexdigest()
    redline_path, _ = store.get_artifact(OWNER, batch_id, document_id, "working_redline")
    assert redline_path.is_file()
    with pytest.raises(NotFound):
        store.get_artifact(OWNER, batch_id, document_id, "final")


def test_guide_rules_hash_mismatch_fails_document(tmp_path):
    store, version = bound_store(tmp_path)

    def processor(source, output, ruleset):
        raise AssertionError("processor must not run without a matching guide snapshot")

    tasks, task_id, _, _ = submit_task(store, data=contract_bytes(), guide_version_id=version["id"],
                                       guide_rules_sha256="0" * 64)
    assert Worker(store, processor).run_once()
    document = tasks.documents(OWNER, task_id)["items"][0]
    assert document["job_status"] == "failed"
    assert document["error_code"] == "GUIDE_RULESET_HASH_MISMATCH"


def test_missing_guide_version_fails_document(tmp_path):
    store, _ = bound_store(tmp_path)
    tasks, task_id, _, _ = submit_task(store, data=contract_bytes(), guide_version_id="gone-version",
                                       guide_rules_sha256="")
    assert Worker(store, complete_recall_processor).run_once()
    document = tasks.documents(OWNER, task_id)["items"][0]
    assert document["job_status"] == "failed"
    assert document["error_code"] == "GUIDE_VERSION_MISSING"


def test_recall_pipeline_stop_pauses_task(tmp_path):
    store, version = bound_store(tmp_path)

    def processor(source, output, ruleset):
        raise RecallPipelineStop("MODEL_CIRCUIT_OPEN", {"stage": "facts"})

    tasks, task_id, _, _ = submit_task(store, data=contract_bytes(), guide_version_id=version["id"],
                                       guide_rules_sha256=version["rules_sha256"])
    assert Worker(store, processor).run_once()
    snapshot = tasks.get(OWNER, task_id)
    assert snapshot["status"] == "paused"
    assert snapshot["pause_reason"] == "MODEL_CIRCUIT_OPEN"
    document = tasks.documents(OWNER, task_id)["items"][0]
    assert document["job_status"] == "retry_wait"
    assert document["error_code"] == "MODEL_CIRCUIT_OPEN"


def test_validate_delivery_requires_actual_content_verification(tmp_path):
    source = tmp_path / "contract.docx"
    source.write_bytes(contract_bytes())
    plan = validate_edit_plan(source, [], {})
    validation = generate_unchanged_artifacts(source, plan, tmp_path / "redline.docx", tmp_path / "final.docx")
    report = {"engine_id": "recall_first", "schema_version": "3", "execution_mode": "delivery",
              "delivery_status": "complete",
              "output_validation": {"kind": "recall_validated_edits",
                                    "checks": ["package", "revisions", "final_equivalence", "non_target",
                                               "content_verification"],
                                    "final_sha256": validation.final_sha256,
                                    "redline_sha256": validation.redline_sha256}}
    artifacts = {"final": tmp_path / "final.docx", "redline": tmp_path / "redline.docx"}
    with pytest.raises(ValueError, match="MISSING_VERIFIED_ARTIFACTS"):
        validate_delivery(report, artifacts)
    report["output_validation"]["checks"].append("actual_content_verification")
    validate_delivery(report, artifacts)
    # An incomplete recall result must never pass while presenting the final names.
    report["delivery_status"] = "incomplete"
    report["output_validation"]["kind"] = "recall_working_edits"
    report["unresolved"] = [{"kind": "target"}]
    with pytest.raises(ValueError, match="MISSING_VERIFIED_ARTIFACTS"):
        validate_delivery(report, artifacts)
    report["output_validation"]["final_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="ARTIFACT_VERIFICATION_MISMATCH"):
        validate_delivery(report, {"working_final": tmp_path / "final.docx",
                                   "working_redline": tmp_path / "redline.docx"})


def test_worker_runs_real_recall_pipeline_with_bound_guide(tmp_path, monkeypatch):
    """Default-processor dispatch: bound guide -> real pipeline -> published pair."""
    contract = tmp_path / "contract.docx"
    build_contract(contract)
    data = contract.read_bytes()
    store = RunStore(tmp_path / "state")
    guide_store.ensure_builtin_guide(store)
    version = guide_store.create_custom_version(store, OWNER, "guide.docx", b"placeholder")
    from fundlex.guide import certify_guide_document
    document = guide_document()
    document["source_sha256"] = version["source_sha256"]
    document = certify_guide_document(document, source_sha256=version["source_sha256"])
    approved_sha = document["document_sha256"]
    guide_store.finish_guide_job(store, version["job_id"], document=document)
    parsed = guide_store.get_version(store, OWNER, version["id"])
    assert parsed["rules_sha256"] == approved_sha
    tasks, task_id, batch_id, document_id = submit_task(
        store, data=data, guide_version_id=version["id"], guide_rules_sha256=parsed["rules_sha256"])
    gateway = PipelineGateway(base_responses())
    monkeypatch.setattr("fundlex.model.ModelGateway.from_env", classmethod(lambda cls: gateway))
    worker = Worker(store)
    assert worker.run_once() is True
    assert worker.run_once() is False
    item = tasks.documents(OWNER, task_id)["items"][0]
    assert item["job_status"] == "succeeded" and item["delivery_status"] == "complete"
    report = store.get_report(OWNER, batch_id, document_id)
    assert report["engine_id"] == "recall_first" and report["schema_version"] == "3"
    assert report["output_validation"]["kind"] == "recall_validated_edits"
    assert report["guide"]["document_sha256"] == approved_sha
    final_path, _ = store.get_artifact(OWNER, batch_id, document_id, "final")
    final_text = "\n".join(paragraph.text for paragraph in Document(final_path).paragraphs)
    assert "调低销售服务费率" in final_text
    assert gateway.bindings and gateway.bindings[0][1]["guide_document_sha256"] == approved_sha


def test_worker_parses_queued_guide_job(tmp_path, monkeypatch):
    store, _ = bound_store(tmp_path)
    guide_path = build_docx(tmp_path / "guide.docx", GUIDE_PARAGRAPHS)
    version = guide_store.create_custom_version(store, OWNER, "guide.docx", guide_path.read_bytes())
    gateway = ScriptedGateway([("guide_parse", valid_parsed()), ("guide_verify", clean_verification())])
    monkeypatch.setattr("fundlex.model.ModelGateway.from_env", staticmethod(lambda: gateway))
    worker = Worker(store)
    assert worker.run_once() is True
    assert worker.run_once() is False
    parsed = guide_store.get_version(store, OWNER, version["id"])
    assert parsed["status"] == "parsed" and parsed["rules_count"] == 3
    assert parsed["parser_version"] == "guide-parse-1"
    assert parsed["prompt_versions"] == {"guide_parse": "guide_parse_v1", "guide_verify": "guide_verify_v1"}
    document = guide_store.guide_document(store, version["id"])
    assert document["schema_version"] == GUIDE_SCHEMA_VERSION and len(document["groups"]) == 3
    guide_store.set_current_guide(store, OWNER, version["id"])
    assert guide_store.current_guide(store, OWNER)["id"] == version["id"]
    with store._connect() as db:
        job = db.execute("SELECT status,attempts FROM guide_jobs WHERE id=?", (version["job_id"],)).fetchone()
    assert job["status"] == "completed" and job["attempts"] == 1


def test_worker_releases_guide_job_on_model_stop(tmp_path, monkeypatch):
    store, _ = bound_store(tmp_path)
    guide_path = build_docx(tmp_path / "guide.docx", GUIDE_PARAGRAPHS)
    version = guide_store.create_custom_version(store, OWNER, "guide.docx", guide_path.read_bytes())

    def stop():
        raise ModelError("MODEL_CIRCUIT_OPEN")

    monkeypatch.setattr("fundlex.model.ModelGateway.from_env", staticmethod(stop))
    worker = Worker(store)
    assert worker.run_once() is True
    with store._connect() as db:
        job = db.execute("SELECT status,attempts,error_code FROM guide_jobs WHERE id=?",
                         (version["job_id"],)).fetchone()
    assert job["status"] == "queued" and job["attempts"] == 1 and job["error_code"] is None
    assert guide_store.get_version(store, OWNER, version["id"])["status"] == "uploaded"
    assert worker.run_once() is False  # backoff, the attempt is not consumed twice
