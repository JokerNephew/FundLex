"""Independent OOXML fixtures for the conservative DOCX editor."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from fundlex.docx import (
    BARRIER, DOCUMENT, DocxError, generate_artifacts, generate_noop_evaluation_artifacts, inspect_docx,
    validate_edit_plan, Q,
)


def fixture_docx(tmp_path: Path) -> Path:
    doc = Document()
    toc = doc.add_paragraph("目录")
    toc.style = "TOC 1" if "TOC 1" in doc.styles else "Normal"
    doc.add_heading("第一部分 前言", 1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("原有")
    bold = paragraph.add_run("收费条款")
    bold.bold = True
    doc.add_paragraph("完全相同的原有收费条款出现在此非目标段落。")
    doc.add_heading("内容摘要", 1)
    doc.add_paragraph("摘要原有收费条款。")
    doc.add_heading("附录", 1)
    doc.add_paragraph("附录原有收费条款。")
    path = tmp_path / "contract.docx"
    doc.save(path)
    return path


def edit_for(path: Path, location: str, old: str, replacement: str = "已确认的新条款") -> dict:
    inspected = inspect_docx(path, full=True)
    p = inspected.by_location(location)
    start = p.text.index(old)
    return {
        "edit_id": "fixture-edit", "source_sha256": inspected.source_sha256,
        "location": location, "zone": p.zone, "start": start, "end": start + len(old),
        "expected_old_text": old, "expected_old_hash": sha256(old.encode()).hexdigest(),
        "template_id": "T", "replacement": replacement, "verified": True,
        "fingerprint": p.fingerprint, "mapping_version": p.mapping_version,
    }


def test_cross_run_redline_final_and_non_target(tmp_path: Path) -> None:
    source = fixture_docx(tmp_path)
    original_hash = sha256(source.read_bytes()).hexdigest()
    edit = edit_for(source, "p2", "原有收费条款")
    plan = validate_edit_plan(source, [edit], {"T": "已确认的新条款"})
    result = generate_artifacts(source, plan, tmp_path / "redline.docx", tmp_path / "final.docx")
    assert result.checks == ("package", "revisions", "final_equivalence", "non_target", "template_text")
    assert sha256(source.read_bytes()).hexdigest() == original_hash
    final = inspect_docx(tmp_path / "final.docx", full=True)
    assert final.by_location("p2").text == "已确认的新条款"
    assert "原有收费条款" in final.by_location("p3").text
    assert "摘要原有收费条款" in final.by_location("p5").text
    assert final.by_location("p7").zone == "appendix"
    with ZipFile(tmp_path / "redline.docx") as archive:
        xml = etree.fromstring(archive.read(DOCUMENT))
        assert len(xml.findall(".//" + Q("ins"))) == 1
        assert len(xml.findall(".//" + Q("del"))) == 1


def test_template_and_zone_fail_closed(tmp_path: Path) -> None:
    source = fixture_docx(tmp_path)
    edit = edit_for(source, "p2", "原有收费条款")
    edit["replacement"] = "模型创造文本"
    with pytest.raises(DocxError) as error:
        validate_edit_plan(source, [edit], {"T": "固定文本"})
    assert error.value.code == "TEMPLATE_MISMATCH"
    appendix = edit_for(source, "p7", "原有收费条款")
    with pytest.raises(DocxError) as error:
        validate_edit_plan(source, [appendix], {"T": "已确认的新条款"})
    assert error.value.code == "REGION_NOT_EDITABLE"


def test_barrier_prevents_crossing_and_overlap(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("旧文")
    p.add_run().add_break()
    p.add_run("跨界")
    source = tmp_path / "barrier.docx"
    doc.save(source)
    span = inspect_docx(source, full=True).by_location("p0")
    assert BARRIER in span.text
    old = span.text
    edit = edit_for(source, "p0", old)
    with pytest.raises(DocxError) as error:
        validate_edit_plan(source, [edit], {"T": "已确认的新条款"})
    assert error.value.code == "EDIT_CROSSES_BARRIER"


def test_existing_revision_and_duplicate_zip_rejected(tmp_path: Path) -> None:
    source = fixture_docx(tmp_path)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(parts[DOCUMENT])
    p = root.find(".//" + Q("p"))
    assert p is not None
    p.append(OxmlElement("w:ins"))
    parts[DOCUMENT] = etree.tostring(root)
    revision = tmp_path / "revision.docx"
    with ZipFile(revision, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    with pytest.raises(DocxError) as error:
        inspect_docx(revision, full=True)
    assert error.value.code == "EXISTING_REVISION_UNSUPPORTED"
    duplicate = tmp_path / "duplicate.docx"
    with ZipFile(duplicate, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
        archive.writestr(DOCUMENT, parts[DOCUMENT])
    with pytest.raises(DocxError) as error:
        inspect_docx(duplicate)
    assert error.value.code == "ZIP_DUPLICATE_PART"


def test_multiple_edits_in_one_paragraph_and_table_cell(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("旧甲，中段，旧乙。")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格旧词。"
    source = tmp_path / "multi.docx"
    doc.save(source)
    edits = [
        edit_for(source, "p0", "旧甲", "新甲"),
        edit_for(source, "p0", "旧乙", "新乙"),
        edit_for(source, "tbl0/r0/c0/p0", "旧词", "新词"),
    ]
    edits[0]["template_id"] = "A"
    edits[1]["template_id"] = "B"
    edits[2]["template_id"] = "C"
    plan = validate_edit_plan(source, edits, {"A": "新甲", "B": "新乙", "C": "新词"})
    generate_artifacts(source, plan, tmp_path / "redline.docx", tmp_path / "final.docx")
    final = inspect_docx(tmp_path / "final.docx", full=True)
    assert final.by_location("p0").text == "新甲，中段，新乙。"
    assert final.by_location("tbl0/r0/c0/p0").text == "表格新词。"


def test_plan_mutation_is_rejected(tmp_path: Path) -> None:
    source = fixture_docx(tmp_path)
    plan = validate_edit_plan(source, [edit_for(source, "p2", "原有收费条款")], {"T": "已确认的新条款"})
    plan.edits[0]["replacement"] = "篡改内容"
    with pytest.raises(DocxError) as error:
        generate_artifacts(source, plan, tmp_path / "redline.docx", tmp_path / "final.docx")
    assert error.value.code == "TEMPLATE_MISMATCH"
    assert not (tmp_path / "final.docx").exists()


def test_hyperlink_and_field_result_not_editable(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph("前文")
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "链接内容"
    run.append(text)
    hyperlink.append(run)
    p._p.append(hyperlink)
    p.add_run("后文")
    field = doc.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(Q("fldCharType"), "begin")
    separate = OxmlElement("w:fldChar")
    separate.set(Q("fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(Q("fldCharType"), "end")
    for node in (begin, separate):
        field.add_run()._r.append(node)
    field.add_run("字段结果")
    field.add_run()._r.append(end)
    source = tmp_path / "barriers.docx"
    doc.save(source)
    inspected = inspect_docx(source, full=True)
    assert BARRIER in inspected.by_location("p0").text
    assert BARRIER in inspected.by_location("p1").text
    for location, old in (("p0", "链接内容"), ("p1", "字段结果")):
        with pytest.raises(DocxError) as error:
            validate_edit_plan(source, [edit_for(source, location, old)], {"T": "已确认的新条款"})
        assert error.value.code in {"EDIT_CROSSES_BARRIER", "REGION_NOT_EDITABLE"}


def test_merged_table_and_protection_fail_closed(tmp_path: Path) -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "合并单元旧文"
    table.cell(0, 0).merge(table.cell(0, 1))
    source = tmp_path / "merged.docx"
    doc.save(source)
    paragraph = inspect_docx(source, full=True).by_location("tbl0/r0/c0/p0")
    assert paragraph.zone == "unsupported_table"
    with pytest.raises(DocxError) as error:
        validate_edit_plan(source, [edit_for(source, paragraph.location, "旧文")], {"T": "已确认的新条款"})
    assert error.value.code == "REGION_NOT_EDITABLE"
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    settings = etree.fromstring(parts["word/settings.xml"])
    settings.append(OxmlElement("w:documentProtection"))
    parts["word/settings.xml"] = etree.tostring(settings)
    protected = tmp_path / "protected.docx"
    with ZipFile(protected, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    with pytest.raises(DocxError) as error:
        inspect_docx(protected, full=True)
    assert error.value.code == "DOCUMENT_PROTECTED"


def test_pdf_and_traversal_entry_rejected(tmp_path: Path) -> None:
    pdf = tmp_path / "contract.pdf"
    pdf.write_bytes(b"%PDF-1.7")
    with pytest.raises(DocxError) as error:
        inspect_docx(pdf)
    assert error.value.code == "DOCX_ONLY"
    source = fixture_docx(tmp_path)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    bad = tmp_path / "traversal.docx"
    with ZipFile(bad, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
        archive.writestr("../outside", b"x")
    with pytest.raises(DocxError) as error:
        inspect_docx(bad)
    assert error.value.code == "OPC_URI_INVALID"


def test_noop_evaluation_artifacts_are_byte_identical(tmp_path: Path) -> None:
    source = fixture_docx(tmp_path)
    result = generate_noop_evaluation_artifacts(source, tmp_path / "eval-redline.docx", tmp_path / "eval-final.docx")
    assert result.manifest == ()
    assert result.checks == ("package", "zero_edit", "byte_identity")
    assert (tmp_path / "eval-redline.docx").read_bytes() == source.read_bytes()
    assert (tmp_path / "eval-final.docx").read_bytes() == source.read_bytes()
