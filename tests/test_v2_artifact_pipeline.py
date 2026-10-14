"""End-to-end private artifact review and bounded action rollback."""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from fundlex import v2_pipeline
from fundlex.docx import inspect_docx
from fundlex.model import ModelError
from fundlex.rules_v2 import RuleSetV2, load_ruleset_v2
from test_v2_evaluation import EvaluationGateway, build_fixture


class ArtifactGateway(EvaluationGateway):
    def __init__(self, decisions, output):
        super().__init__()
        self.decisions = list(decisions)
        self.review_contexts = []
        self.output = output

    def verify_artifact(self, context, **kwargs):
        # No official download file can exist during any semantic review.
        assert not (self.output / "evaluation_final.docx").exists()
        assert not (self.output / "evaluation_redline.docx").exists()
        assert not (self.output / "final.docx").exists()
        self.review_contexts.append(context)
        action_id = re.search(r"\[ACTION ([^]]+)\]", context).group(1)
        expected_id, decision = self.decisions.pop(0)
        assert action_id == expected_id
        if isinstance(decision, ModelError):
            raise decision
        requirement_ids = re.findall(r"^- (\S+) \[", context, re.MULTILINE)
        edited = context.split("[EDIT_LOCATIONS]\n", 1)[1].split("\n", 1)[0].split(",")[0]
        quote = context.split(f"[AFTER location={edited}]\n", 1)[1].split("\n", 1)[0]
        items = [{"action_id": action_id, "requirement_id": identifier, "result": "satisfied",
                  "evidence_quote": quote, "evidence_source": "after"}
                 for identifier in [*requirement_ids, "PRESERVATION", "COVERAGE"]]
        if decision != "pass":
            # A rejection must be structural (unverifiable evidence), not only a
            # model verdict on PRESERVATION/COVERAGE, which gets a consensus retry.
            for item in items:
                if item["requirement_id"] not in {"PRESERVATION", "COVERAGE"}:
                    item["evidence_quote"] = "该引文并不存在于产物中"
                    item["result"] = "unknown"
        return self._result({"items": items}, "verify_artifact")


def fixture(tmp_path, monkeypatch, ids=("R06-A1", "R07-A1")):
    source = build_fixture(tmp_path / "source.docx")
    package = load_ruleset_v2(path=Path(__file__).resolve().parents[1] / "rules/rules_v2.draft.json")
    actions = tuple(a for a in package.actions if a["action_id"] in ids)
    subset = RuleSetV2(package.version, package.sha256, actions, package.release_policy)
    monkeypatch.setattr(v2_pipeline, "load_ruleset_v2", lambda *args: subset)
    document = Document(source)
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    for action in actions:
        document.add_paragraph(action["locate"]["old_texts"][0])
    document.save(source)
    return source, subset, tmp_path / "out"


def test_pipeline_only_exposes_successfully_reviewed_actual_pair(tmp_path, monkeypatch):
    source, _, output = fixture(tmp_path, monkeypatch)
    gateway = ArtifactGateway([("R06-A1", "pass"), ("R07-A1", "pass")], output)
    report = v2_pipeline.process_v2(source, output, gateway=gateway, evaluation=True)
    assert report["delivery_status"] == "evaluation_modified"
    assert report["edit_plan"]["count"] == 2
    review = report["actual_content_verifications"][0]
    assert review["passed"]
    assert review["source_sha256"] == report["source_sha256"]
    assert review["plan_sha256"] == report["edit_plan"]["sha256"]
    assert review["final_sha256"] == report["output_validation"]["final_sha256"]
    assert review["redline_sha256"] == report["output_validation"]["redline_sha256"]
    assert inspect_docx(Path(report["artifacts"]["evaluation_final"]), full=True).source_sha256 == review["final_sha256"]
    assert not list(output.glob(".v3-review-*"))


def test_first_round_failure_reverts_whole_action_and_rechecks_remaining_pair(tmp_path, monkeypatch):
    source, subset, output = fixture(tmp_path, monkeypatch)
    document = Document(source)
    document.add_paragraph(subset.by_id("R07-A1")["locate"]["old_texts"][0])
    document.save(source)
    gateway = ArtifactGateway([("R06-A1", "pass"), ("R07-A1", "reject"),
                               ("R06-A1", "pass")], output)
    report = v2_pipeline.process_v2(source, output, gateway=gateway, evaluation=True)
    assert report["delivery_status"] == "evaluation_modified_with_exceptions"
    assert report["edit_plan"]["count"] == 1
    assert {r["action_id"] for r in report["content_verifications"]} == {"R06-A1"}
    first, second = report["actual_content_verifications"]
    assert not first["passed"] and second["passed"]
    assert first["plan_sha256"] != second["plan_sha256"] == report["edit_plan"]["sha256"]
    assert first["final_sha256"] != second["final_sha256"] == report["output_validation"]["final_sha256"]
    assert len(gateway.review_contexts) == 3
    final = inspect_docx(Path(report["artifacts"]["evaluation_final"]), full=True)
    original = subset.by_id("R07-A1")["locate"]["old_texts"][0]
    assert sum(span.text == original for span in final.paragraphs) == 2
    assert report["exceptions"][0]["action_id"] == "R07-A1"


def test_second_round_failure_withholds_all_modified_files(tmp_path, monkeypatch):
    source, _, output = fixture(tmp_path, monkeypatch)
    gateway = ArtifactGateway([("R06-A1", "pass"), ("R07-A1", "reject"),
                               ("R06-A1", "reject")], output)
    report = v2_pipeline.process_v2(source, output, gateway=gateway, evaluation=True)
    assert report["delivery_status"] == "incomplete"
    assert report["artifacts"] == {}
    assert report["edit_plan"]["count"] == 0 and not report["edit_plan"]["executable"]
    assert report["content_verifications"] == []
    assert len(report["actual_content_verifications"]) == 2
    assert not list(output.glob("*.docx"))
    assert {entry["action_id"] for entry in report["exceptions"]} == {"R06-A1", "R07-A1"}


def test_all_first_round_actions_rejected_does_not_publish_noop_as_reviewed(tmp_path, monkeypatch):
    source, _, output = fixture(tmp_path, monkeypatch)
    gateway = ArtifactGateway([("R06-A1", "reject"), ("R07-A1", "reject")], output)
    report = v2_pipeline.process_v2(source, output, gateway=gateway, evaluation=True)
    assert report["delivery_status"] == "incomplete"
    assert report["artifacts"] == {}
    assert len(report["actual_content_verifications"]) == 1
    assert not list(output.glob("*.docx"))


def test_second_round_rejection_also_revokes_that_round_passing_actions(tmp_path, monkeypatch):
    source, _, output = fixture(tmp_path, monkeypatch, ids=("R06-A1", "R07-A1", "R13-A1"))
    gateway = ArtifactGateway([("R06-A1", "pass"), ("R07-A1", "reject"), ("R13-A1", "pass"),
                               ("R06-A1", "pass"), ("R13-A1", "reject")], output)
    report = v2_pipeline.process_v2(source, output, gateway=gateway, evaluation=True)
    assert report["delivery_status"] == "incomplete"
    assert report["artifacts"] == {} and report["edit_plan"]["count"] == 0
    assert report["content_verifications"] == []
    assert not list(output.glob("*.docx"))
    reasons = {entry["action_id"]: entry["reason"] for entry in report["exceptions"]}
    assert reasons["R06-A1"] == "V3_ARTIFACT_REVIEW_LIMIT"
    assert reasons["R13-A1"] == "V3_ARTIFACT_SEMANTIC_REJECTED"


def test_actual_review_model_budget_stop_never_exposes_temporary_files(tmp_path, monkeypatch):
    source, _, output = fixture(tmp_path, monkeypatch)
    gateway = ArtifactGateway([("R06-A1", ModelError("MODEL_COST_BUDGET_EXCEEDED"))], output)
    report = v2_pipeline.process_v2(source, output, gateway=gateway, evaluation=True)
    assert report["error_code"] == "MODEL_COST_BUDGET_EXCEEDED"
    assert report["artifacts"] == {}
    assert not list(output.glob("*.docx"))
    assert not list(output.glob(".v3-review-*"))
    assert all(a["status"] == "failed" for a in report["actions"])


def test_zero_edits_no_semantic_review_but_missing_scope_is_exception(tmp_path, monkeypatch):
    source, _, output = fixture(tmp_path, monkeypatch)
    document = Document(source)
    for paragraph in list(document.paragraphs)[-4:]:
        paragraph._element.getparent().remove(paragraph._element)
    document.save(source)
    gateway = ArtifactGateway([], output)
    report = v2_pipeline.process_v2(source, output, gateway=gateway, evaluation=True)
    assert report["delivery_status"] == "incomplete"
    assert report["actual_content_verifications"] == []
    assert {e["action_id"] for e in report["exceptions"]} == {"R06-A1", "R07-A1"}
    assert not gateway.review_contexts


def test_markdown_complete_with_exceptions_describes_available_files(tmp_path):
    path = tmp_path / "report.md"
    v2_pipeline._write_v2_markdown(path, {"execution_mode": "delivery",
                                         "delivery_status": "complete_with_exceptions"})
    text = path.read_text(encoding="utf8")
    assert "已生成正式合同文件" in text
    assert "未生成正式合同文件" not in text
