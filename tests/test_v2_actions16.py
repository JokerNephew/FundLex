"""Fixture tests for the remaining catalogue actions (template and delete modes).

Every replacement text is the published rule literal, so these tests need no
model calls: the engine locates the old clause, instantiates class slots from the
current contract profile, assembles the rule text and verifies it by code.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from fundlex.docx import generate_artifacts, inspect_docx, validate_edit_plan
from fundlex.rules_v2 import load_ruleset_v2
from fundlex.store import engine_build
from fundlex.v2 import plan_v2
from test_v2_engine import ScriptedGateway, build_docx, make_profile


def make_profile_full(classes, subscription, service, *, floating=None, count=None):
    profile = make_profile(classes, subscription, service)
    profile["floating_management_fee"] = floating
    if count is not None:
        profile["share_class_count"] = count
        profile["share_class_count_min"] = count
        profile["share_class_count_verified"] = True
    return profile


def old_texts(action_id):
    action = load_ruleset_v2().by_id(action_id)
    return action["locate"].get("old_texts") or action["locate"]["triggers"]


def plan_one(action_id, path, profile):
    inspection = inspect_docx(path, full=True)
    gateway = ScriptedGateway([])
    outcome = plan_v2(inspection, profile, load_ruleset_v2(), gateway,
                      engine_build_value=engine_build(), actions=[action_id])
    action = next(item for item in outcome["actions"] if item["action_id"] == action_id)
    return outcome, action, inspection


def build_multi_fixture(path: Path) -> Path:
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金的基本情况", level=1)
    document.add_paragraph("基金的类别：混合型证券投资基金")
    document.add_paragraph("份额类别设置")
    document.add_paragraph("本基金将基金份额分为 A 类和 C 类两种不同的类别。在投资者申购 A 类基金份额时收取申购费用，"
                           "不计提销售服务费；在投资者申购 C 类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费。")
    document.add_paragraph("有关基金份额类别的具体设置、费率水平等由基金管理人确定。根据基金销售情况，"
                           "基金管理人在履行适当程序后可以增加新的基金份额类别、或者在基金合同规定的范围内"
                           "变更现有基金份额类别的申购费率、调低赎回费率或变更收费方式。")
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    for text in old_texts("R06-A1") + old_texts("R07-A1") + old_texts("R09-A1")[:1] + old_texts("R13-A1")[:1]:
        document.add_paragraph(text)
    document.add_heading("基金费用与税收", level=1)
    document.add_paragraph("基金销售服务费")
    for text in old_texts("R14-A1") + old_texts("R15-A1") + old_texts("R15-A2"):
        document.add_paragraph(text)
    document.add_heading("基金的收益与分配", level=1)
    document.add_paragraph("收益分配")
    document.add_paragraph(old_texts("R16-A1")[0] + "，各基金份额类别对应的可供分配利润将有所不同。")
    document.add_heading("基金合同的变更、终止与基金财产的清算", level=1)
    document.add_paragraph("基金财产清算剩余资产的分配")
    document.add_paragraph(old_texts("R17-A2")[0])
    document.save(path)
    return path


def build_single_fixture(path: Path) -> Path:
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金的基本情况", level=1)
    document.add_paragraph("基金的类别：混合型证券投资基金")
    document.add_paragraph("份额类别设置")
    document.add_paragraph("本基金不区分份额类别，仅设一种基金份额。")
    document.add_paragraph("有关基金份额类别的具体设置、费率水平等由基金管理人确定。根据基金销售情况，"
                           "基金管理人在履行适当程序后可以增加新的基金份额类别、或者在基金合同规定的范围内"
                           "变更现有基金份额类别的申购费率、调低赎回费率或变更收费方式。")
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    for text in old_texts("R05-A1") + old_texts("R08-A1") + old_texts("R12-A1")[:1]:
        document.add_paragraph(text)
    document.add_heading("基金合同的变更、终止与基金财产的清算", level=1)
    document.add_paragraph("基金财产清算剩余资产的分配")
    document.add_paragraph(old_texts("R17-A1")[0])
    document.save(path)
    return path


def build_r17a3_fixture(path: Path) -> Path:
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金的基本情况", level=1)
    document.add_paragraph("基金的类别：混合型证券投资基金")
    document.add_paragraph("份额类别设置")
    document.add_paragraph("本基金将基金份额分为 A 类和 C 类两种不同的类别。在投资者申购 A 类基金份额时收取申购费用，"
                           "不计提销售服务费；在投资者申购 C 类基金份额时不收取申购费用，而从本类别基金资产中计提销售服务费。")
    document.add_heading("基金份额的申购与赎回", level=1)
    document.add_paragraph("申购和赎回的价格、费用及其用途")
    document.add_heading("基金合同的变更、终止与基金财产的清算", level=1)
    document.add_paragraph("基金财产清算剩余资产的分配")
    document.add_paragraph(old_texts("R17-A3")[0])
    document.save(path)
    return path


MULTI = make_profile_full(["A", "C"], ["A"], ["C"], floating=True)
SINGLE = make_profile_full(["A"], ["A"], [], floating=True, count=1)
MULTI_NONFLOATING = make_profile_full(["A", "C"], ["A"], ["C"], floating=False)


def assert_applied(action_id, outcome, action, expected, *, operation="replace"):
    assert action["status"] == "applied", action.get("failure_reason")
    edit = outcome["edits"][0]
    assert edit["replacement"] == expected
    if operation == "delete":
        assert edit["replacement"] == ""


def test_multi_class_template_actions(tmp_path):
    source = build_multi_fixture(tmp_path / "multi.docx")
    expectations = {
        "R03-A1": "调低销售服务费率",
        "R07-A1": "赎回金额为按实际确认的有效赎回份额乘以当日该类基金份额的基金份额净值并扣除或加上相应的费用，赎回金额单位为元。",
        "R13-A1": "基金管理人可以适当调低基金申购费率和销售服务费率。",
        "R16-A1": "由于本基金A类基金份额和C类基金份额费用收取不同",
        "R17-A2": "依据基金财产清算的分配方案，将基金财产清算后的全部剩余资产扣除基金财产清算费用、交纳所欠税款并清偿基金债务后，"
                  "按基金份额持有人持有的基金份额比例进行分配，并加上返还基金份额持有人的或有管理费（如有）、销售服务费（如有）"
                  "或扣除相应的超额管理费（如有）。",
    }
    for action_id, expected in expectations.items():
        outcome, action, _ = plan_one(action_id, source, MULTI)
        assert_applied(action_id, outcome, action, expected)
    # R06/R09 instantiate the class slots from the current contract.
    r06 = load_ruleset_v2().by_id("R06-A1")["generation"]["template_text"]
    outcome, action, _ = plan_one("R06-A1", source, MULTI)
    assert_applied("R06-A1", outcome, action, r06)
    r09 = load_ruleset_v2().by_id("R09-A1")["generation"]["template_text"]
    outcome, action, _ = plan_one("R09-A1", source, MULTI)
    assert_applied("R09-A1", outcome, action, r09)
    # R14/R15 work on the sales-service-fee section.
    r14 = load_ruleset_v2().by_id("R14-A1")["generation"]["template_text"]
    outcome, action, _ = plan_one("R14-A1", source, MULTI)
    assert_applied("R14-A1", outcome, action, r14)
    for action_id in ("R15-A1", "R15-A2"):
        outcome, action, _ = plan_one(action_id, source, MULTI)
        assert action["status"] == "applied", action.get("failure_reason")
        assert outcome["edits"][0]["replacement"] == ""


def test_single_class_template_actions(tmp_path):
    source = build_single_fixture(tmp_path / "single.docx")
    expectations = {
        "R03-A2": "",
        "R05-A1": "通过直销机构申购本基金不收取申购费，通过直销机构以外的其他销售机构申购本基金的申购费率由基金管理人决定，"
                  "并在招募说明书中列示。",
        "R08-A1": "通过直销机构申购本基金不收取申购费，通过其他销售机构申购本基金的申购费用由投资人承担，不列入基金财产。",
        "R12-A1": "基金管理人可以适当调低基金申购费率。",
        "R17-A1": "依据基金财产清算的分配方案，将基金财产清算后的全部剩余资产扣除基金财产清算费用、交纳所欠税款并清偿基金债务后，"
                  "按基金份额持有人持有的基金份额比例进行分配，并加上返还基金份额持有人的或有管理费（如有）"
                  "或扣除相应的超额管理费（如有）。",
    }
    for action_id, expected in expectations.items():
        outcome, action, _ = plan_one(action_id, source, SINGLE)
        assert_applied(action_id, outcome, action, expected)
    # The single-class fixture must not trigger multi-class actions.
    outcome, action, _ = plan_one("R13-A1", source, SINGLE)
    assert action["status"] == "not_applicable"


def test_r17a3_nonfloating_multi(tmp_path):
    source = build_r17a3_fixture(tmp_path / "r17a3.docx")
    expected = "依据基金财产清算的分配方案，将基金财产清算后的全部剩余资产扣除基金财产清算费用、交纳所欠税款并清偿基金债务后，" \
               "按基金份额持有人持有的基金份额比例进行分配，并加上返还基金份额持有人的销售服务费（如有）。"
    outcome, action, _ = plan_one("R17-A3", source, MULTI_NONFLOATING)
    assert_applied("R17-A3", outcome, action, expected)
    # A floating-management-fee contract must not use the non-floating variant.
    outcome, action, _ = plan_one("R17-A3", source, MULTI)
    assert action["status"] == "not_applicable"


def test_paragraph_scope_mismatch_fails_closed(tmp_path):
    # The paragraph carries extra text after the published old clause; replacing
    # the whole paragraph would drop it, so the engine must refuse.
    document = Document()
    document.add_paragraph("广发示例混合型证券投资基金基金合同")
    document.add_heading("基金费用与税收", level=1)
    document.add_paragraph("基金销售服务费")
    document.add_paragraph(old_texts("R15-A1")[0] + "本段还包含其他必须保留的内容。")
    source = tmp_path / "mismatch.docx"
    document.save(source)
    outcome, action, _ = plan_one("R15-A1", source, MULTI)
    assert action["status"] == "failed"
    assert outcome["edits"] == []
    assert any(target.get("reason") == "V2_SCOPE_TEXT_MISMATCH" for target in action["targets"])


def test_slot_instantiation_and_end_to_end_artifact(tmp_path):
    source = build_multi_fixture(tmp_path / "artifact.docx")
    outcome, action, inspection = plan_one("R06-A1", source, MULTI)
    assert action["status"] == "applied"
    record = outcome["records"][0]
    assert record["content_source"] == "rule_template"
    assert record["slot_instantiation"] == {"A类基金份额": ["A"], "C类基金份额": ["C"]}
    plan = validate_edit_plan(source, outcome["edits"], {}, content_check=outcome["content_check"])
    generate_artifacts(source, plan, tmp_path / "r.docx", tmp_path / "f.docx")
    final = inspect_docx(tmp_path / "f.docx", full=True)
    text = next(span.text for span in final.paragraphs if "在申购时收取申购费用" in span.text)
    assert "通过直销机构申购本基金A类基金份额不收取申购费" in text
