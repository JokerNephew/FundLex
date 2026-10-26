"""The named delivery-e2e workflow is safe by default and exercises delivery."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from tools.evaluation.delivery_e2e import run_delivery_e2e


def _source(path: Path) -> Path:
    document = Document()
    document.add_paragraph("交付流程夹具第一段。")
    document.add_paragraph("交付流程夹具第二段。")
    document.add_paragraph("这一段保持原样，用于验证实际产物保留。")
    document.save(path)
    return path


def test_real_model_requires_explicit_opt_in(tmp_path):
    with pytest.raises(ValueError, match="allow_real_model"):
        run_delivery_e2e(source=_source(tmp_path / "source.docx"),
                         evidence_dir=tmp_path / "evidence", model_mode="real")


def test_fixture_delivery_e2e_runs_complete_user_path_without_model(tmp_path):
    result = run_delivery_e2e(source=_source(tmp_path / "source.docx"),
                              evidence_dir=tmp_path / "evidence")

    assert result["workflow"] == "delivery-e2e"
    assert result["model_mode"] == "fixture"
    assert result["cleanup"].startswith("isolated runtime")
    assert result["report"]["delivery_status"] == "complete"
    assert result["report"]["model"]["calls"] == 0
    assert result["artifact_review"]["mechanical_pass"] is True
    assert set(result["selected_artifacts"]) == {"selected_final", "selected_redline"}
    assert result["review"]["rejected"] and result["review"]["accepted"]
    assert (tmp_path / "evidence" / "result.json").is_file()
