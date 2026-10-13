from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document

from fundlex.model import ModelError, ModelResult
from fundlex.pipeline import process
from fundlex.rules import RuleError, classify_fund_type, load_ruleset, extract_profile, applicability, instantiate_template, evaluate_candidates
from fundlex.docx import inspect_docx


def _contract(path: Path, category: str, operation: str = "契约型开放式") -> None:
    doc = Document()
    doc.add_paragraph("示例证券投资基金基金合同")
    doc.add_paragraph("第四部分 基金的基本情况")
    doc.add_paragraph("二、基金的类别")
    doc.add_paragraph(category)
    doc.add_paragraph("三、基金的运作方式")
    doc.add_paragraph(operation)
    doc.add_paragraph("份额类别设置")
    doc.add_paragraph("本基金可在约定范围内调低赎回费率并报中国证监会备案。")
    doc.add_paragraph("第十二部分 基金的投资")
    doc.add_paragraph("本基金投资于股票和债券，投资于债券的比例不低于基金资产的80%。")
    doc.save(path)


class StubGateway:
    model = "unit-model"
    calls = 0
    base_url = "https://independent-fixture.invalid"
    provider = "independent-fixture"
    request_options = {}
    max_context_chars = 12000
    max_completion_tokens = 1500
    timeout_s = 45
    retries = 1

    def __init__(self, kind: str, fail: bool = False):
        self.kind, self.fail = kind, fail

    def extract_fund_type(self, context: str) -> ModelResult:
        self.calls += 1
        if self.fail:
            raise ModelError("MODEL_SCHEMA_INVALID")
        quote = context.splitlines()[0]
        return ModelResult({"quote": quote, "fund_type": self.kind}, self.model,
                           7, 20, 10, 30, None, 1)


def test_ruleset_gate_is_closed():
    rules = load_ruleset("v1")
    assert not rules.auto_write_enabled
    assert all(not action["auto_write"] for action in rules.actions)
    try:
        load_ruleset("v99")
    except RuleError as exc:
        assert exc.code == "RULESET_NOT_FOUND"
    else:
        raise AssertionError("unknown ruleset accepted")

def test_operation_evidence_preserves_original_spaces(tmp_path: Path):
    source=tmp_path/'spaces.docx'
    _contract(source,'混合型证券投资基金','契 约型 开放式')
    inspection=inspect_docx(source)
    result=classify_fund_type(inspection)
    evidence=result['operation_evidence']
    span=inspection.by_location(evidence['paragraph_id'])
    assert evidence['quote']=='契 约型 开放式'
    assert span.text[evidence['char_start']:evidence['char_end']]==evidence['quote']


def test_ordinary_contract_read_only_artifacts(tmp_path: Path):
    source = tmp_path / "contract.docx"
    _contract(source, "混合型证券投资基金")
    output = tmp_path / "out"
    report = process(source, output, "v1", StubGateway("ordinary"))
    assert report["qualification_status"] == "eligible"
    assert report["release_gate"]["status"] == "closed"
    assert report["edit_plan"]["validated"] and report["edit_plan"]["count"] == 0
    assert report["delivery_status"] != "complete"
    assert report["output_validation"]["not_a_final_contract"]
    assert hashlib.sha256(source.read_bytes()).digest() == hashlib.sha256((output / "evaluation_redline.docx").read_bytes()).digest()
    assert (output / "evaluation_final.docx").read_bytes() == source.read_bytes()
    assert "不完整" in (output / "report.md").read_text(encoding="utf-8")


def test_etf_excluded_from_operation_evidence(tmp_path: Path):
    source = tmp_path / "etf.docx"
    _contract(source, "股票型、指数型证券投资基金", "交易型开放式")
    inspection = inspect_docx(source)
    evidence = classify_fund_type(inspection)
    assert evidence["fund_type"] == "etf"
    report = process(source, tmp_path / "out", "v1", StubGateway("etf"))
    assert report["qualification_status"] == "excluded_by_fund_type"
    assert not (tmp_path / "out" / "evaluation_final.docx").exists()


def test_model_failure_closes_contract(tmp_path: Path):
    source = tmp_path / "contract.docx"
    _contract(source, "混合型证券投资基金")
    report = process(source, tmp_path / "out", "v1", StubGateway("ordinary", fail=True))
    assert report["qualification_status"] == "failed"
    assert report["error_code"] == "MODEL_SCHEMA_INVALID"
    assert report["delivery_status"] == "not_generated"
    assert report["model"]["error_count"] == 1


def test_fee_groups_use_own_clauses_and_conflicts_close(tmp_path: Path):
    path = tmp_path / "profile.docx"
    doc = Document()
    doc.add_paragraph("第四部分 基金的基本情况")
    doc.add_paragraph("B类基金份额收取认购、申购费用，不计提销售服务费。")
    doc.add_paragraph("D类基金份额不收取认购、申购费用，计提销售服务费。")
    doc.save(path)
    profile = extract_profile(inspect_docx(path))
    assert profile["fee_groups_verified"]
    assert profile["fee_groups"] == {"subscription_group": ["B"], "service_group": ["D"]}
    assert applicability({"condition": "has_service_fee"}, profile)[0] == "true"
    assert applicability({"condition": "single_class"}, profile)[0] == "false"
    rules = load_ruleset("v1")
    import pytest
    with pytest.raises(RuleError, match="TEMPLATE_VARIANT_UNAPPROVED"):
        instantiate_template(rules.actions[0], rules, profile)
    doc.add_paragraph("D类基金份额不收取销售服务费。")
    doc.save(path)
    profile = extract_profile(inspect_docx(path))
    assert not profile["fee_groups_verified"]
    assert applicability({"condition": "has_service_fee"}, profile)[0] == "unknown"


def test_subsection_constraint_and_multiple_candidates(tmp_path: Path):
    path = tmp_path / "scope.docx"
    _contract(path, "混合型证券投资基金")
    doc = Document(path)
    doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)
    doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)
    doc.add_paragraph("再次并报中国证监会备案。")
    doc.add_paragraph("五、持有人大会")
    doc.add_paragraph("并报中国证监会备案。")
    doc.save(path)
    actions = evaluate_candidates(inspect_docx(path), load_ruleset("v1"))
    action = next(a for a in actions if a["action_id"] == "R03-A3")
    assert len(action["candidates"]) == 2
    assert all(c["reason"] == "RELEASE_GATE_CLOSED" for c in action["candidates"])


def test_pipeline_released_fixture_executes_fixed_template(tmp_path: Path, monkeypatch):
    from dataclasses import replace
    import fundlex.pipeline as pipeline
    source = tmp_path / "released-fixture.docx"
    _contract(source, "混合型证券投资基金")
    catalogue = load_ruleset("v1")
    action = dict(next(a for a in catalogue.actions if a["action_id"] == "R03-A3"))
    action["auto_write"] = True
    # This is an independent engine fixture, never a published business package.
    fixture = replace(catalogue, actions=(action,))
    from fundlex.rules import release_runtime_context
    tested = release_runtime_context(StubGateway("ordinary"))
    action["evaluation_metrics"] = {"tested_context": tested, "fixture_only": True}
    fixture = replace(fixture, release_policy={"evaluation_context": tested, "basis": "Independent fixed-template engine fixture; not business approval"})
    monkeypatch.setattr(pipeline, "load_ruleset", lambda _: fixture)
    report = pipeline.process(source, tmp_path / "out", "v1", StubGateway("ordinary"))
    assert report["delivery_status"] == "complete"
    assert report["edit_plan"]["count"] == 1
    assert report["actions"][0]["status"] == "applied"
    final = inspect_docx(tmp_path / "out" / "final.docx")
    assert all("并报中国证监会备案" not in span.text for span in final.paragraphs)
    assert report["output_validation"]["kind"] == "validated_edits"
    repeated = pipeline.process(tmp_path / "out" / "final.docx", tmp_path / "repeat", "v1", StubGateway("ordinary"))
    assert repeated["delivery_status"] == "complete"
    assert repeated["edit_plan"]["count"] == 0
    assert repeated["actions"][0]["status"] == "not_matched"
    assert repeated["output_validation"]["kind"] == "validated_unchanged"
    assert (tmp_path / "repeat" / "final.docx").read_bytes() == (tmp_path / "out" / "final.docx").read_bytes()


def test_model_profile_comparison_conflict_fails_closed(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "model-profile.docx"
    doc = Document()
    doc.add_paragraph("第四部分 基金的基本情况")
    quote = "B类基金份额收取申购费用，不计提销售服务费。"
    doc.add_paragraph(quote)
    doc.save(path)
    class ProfileGateway:
        def call(self, **kwargs):
            assert kwargs["task_type"] == "profile"
            return ModelResult({"classes": [{"name": "B", "subscription_fee": {"value": False, "quote": quote}, "sales_service_fee": {"value": False, "quote": quote}}], "floating_management_fee": {"value": None, "quote": ""}}, "fixture", 1, 1, 1, 2, None, 1)
    profile = resolve_profile(inspect_docx(path), ProfileGateway())
    assert profile["source"] == "verified_model_with_independent_comparison"
    assert not profile["fee_groups_verified"]
    assert profile["classes"][0]["subscription_fee"] is None
    assert profile["errors"][0]["code"] == "PROFILE_FEE_POLARITY_UNVERIFIED"


def test_missing_scope_is_failed_and_inline_keyword_cannot_open_scope(tmp_path: Path):
    path = tmp_path / "missing-scope.docx"
    _contract(path, "混合型证券投资基金")
    doc = Document(path)
    doc.paragraphs[6].text = "本基金另有份额类别设置事项。"
    doc.save(path)
    action = next(a for a in evaluate_candidates(inspect_docx(path), load_ruleset("v1")) if a["action_id"] == "R03-A3")
    assert action["status"] == "failed" and not action["scope_verified"]
    assert action["failure_reason"] == "TARGET_SCOPE_UNRESOLVED"
    assert not action["candidates"]


def test_independent_append_fixture_inserts_only_package_literal(tmp_path: Path, monkeypatch):
    from dataclasses import replace
    import fundlex.pipeline as pipeline
    source = tmp_path / "append-fixture.docx"
    _contract(source, "混合型证券投资基金")
    catalogue = load_ruleset("v1")
    action = dict(next(a for a in catalogue.actions if a["action_id"] == "R03-A3"))
    literal = "独立夹具的固定新增条款。"
    action.update(action_id="FIXTURE-APPEND", auto_write=True, operation="append", template_id="FIXTURE-T", new_triggers=[literal])
    fixture = replace(catalogue, actions=(action,), templates={"FIXTURE-T": literal})
    from fundlex.rules import release_runtime_context
    tested = release_runtime_context(StubGateway("ordinary"))
    action["evaluation_metrics"] = {"tested_context": tested, "fixture_only": True}
    fixture = replace(fixture, release_policy={"evaluation_context": tested, "basis": "Independent fixed-template insertion fixture; not business approval"})
    monkeypatch.setattr(pipeline, "load_ruleset", lambda _: fixture)
    report = pipeline.process(source, tmp_path / "out", "v1", StubGateway("ordinary"))
    assert report["delivery_status"] == "complete"
    edit = report["edit_plan"]["edits"][0]
    assert edit["start"] == edit["end"]
    assert edit["expected_old_text"] == ""
    assert edit["replacement"] == literal and edit["template_id"] == "FIXTURE-T"
    final = inspect_docx(tmp_path / "out" / "final.docx")
    assert "并报中国证监会备案" + literal in final.by_location("p7").text
    second = evaluate_candidates(final, fixture)
    assert second[0]["status"] == "satisfied_no_change"
    assert second[0]["candidates"][0]["status"] == "satisfied"
    repeated = pipeline.process(tmp_path / "out" / "final.docx", tmp_path / "repeat", "v1", StubGateway("ordinary"))
    assert repeated["delivery_status"] == "complete" and repeated["edit_plan"]["count"] == 0
    assert repeated["actions"][0]["status"] == "satisfied_no_change"
    assert repeated["output_validation"]["kind"] == "validated_unchanged"


def test_pipeline_rejects_candidate_text_not_in_package(tmp_path: Path, monkeypatch):
    import fundlex.pipeline as pipeline
    source = tmp_path / "tampered.docx"
    _contract(source, "混合型证券投资基金")
    def forged(*args):
        actions = evaluate_candidates(*args)
        target = next(a for a in actions if a["action_id"] == "R03-A3")
        target["candidates"][0].update(status="applied", replacement="模型伪造文本")
        return actions
    monkeypatch.setattr(pipeline, "evaluate_candidates", forged)
    report = pipeline.process(source, tmp_path / "out", "v1", StubGateway("ordinary"))
    assert report["error_code"] == "TEMPLATE_MISMATCH"
    assert report["delivery_status"] == "not_generated"
    assert not (tmp_path / "out" / "final.docx").exists()


def test_cd_mandate_overrides_generic_category_and_generic_cd_list_does_not(tmp_path: Path):
    path = tmp_path / "cd.docx"
    _contract(path, "混合型证券投资基金")
    doc = Document(path)
    doc.paragraphs[-1].text = "本基金投资于同业存单的比例不低于基金资产的80%。"
    doc.save(path)
    kind = classify_fund_type(inspect_docx(path, full=False))
    assert kind["fund_type"] == "interbank_cd"
    assert "同业存单" in kind["supporting_evidence"][0]["quote"]
    report = process(path, tmp_path / "cd-out", "v1", StubGateway("interbank_cd"))
    assert report["qualification_status"] == "excluded_by_fund_type"
    assert not report["actions"] and report["qualification"]["supporting_evidence"]
    doc.paragraphs[-1].text = "本基金投资范围包括债券、同业存单、银行存款。本基金投资于债券的比例不低于基金资产的80%。"
    doc.save(path)
    assert classify_fund_type(inspect_docx(path, full=False))["fund_type"] == "ordinary"


def test_cd_target_without_mandate_and_missing_investment_fail_closed(tmp_path: Path):
    import pytest
    path = tmp_path / "ambiguous-cd.docx"
    _contract(path, "债券型证券投资基金")
    doc = Document(path)
    doc.paragraphs[-1].text = "本基金标的指数为同业存单AAA指数。"
    doc.save(path)
    with pytest.raises(RuleError, match="FUND_TYPE_CD_MANDATE_EVIDENCE_MISSING"):
        classify_fund_type(inspect_docx(path, full=False))
    doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)
    doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)
    doc.save(path)
    with pytest.raises(RuleError, match="FUND_TYPE_INVESTMENT_EVIDENCE_MISSING"):
        classify_fund_type(inspect_docx(path, full=False))


def test_profile_invalid_fields_do_not_block_independent_action(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "bad-profile.docx"
    _contract(path, "混合型证券投资基金")
    doc = Document(path)
    # Put evidence in the permitted profile section, before investment section.
    doc.paragraphs[7].insert_paragraph_before("B类基金份额不收取申购费用。")
    doc.paragraphs[7].insert_paragraph_before("D类基金份额收取销售服务费。")
    doc.save(path)
    class Gateway:
        def call(self, **kwargs):
            return ModelResult({"classes": [{"name": "B", "subscription_fee": {"value": True, "quote": "B类基金份额不收取申购费用。"}, "sales_service_fee": {"value": True, "quote": "D类基金份额收取销售服务费。"}}], "floating_management_fee": {"value": False, "quote": "B类基金份额不收取申购费用。"}}, "fixture", 1, 1, 1, 2, None, 1)
    inspection = inspect_docx(path)
    profile = resolve_profile(inspection, Gateway())
    assert profile["classes"][0]["subscription_fee"] is None
    assert profile["classes"][0]["sales_service_fee"] is None
    assert profile["floating_management_fee"] is None
    assert profile["errors"]
    independent = next(a for a in evaluate_candidates(inspection, load_ruleset("v1"), profile) if a["action_id"] == "R03-A3")
    assert independent["candidates"][0]["target_verified"]
    assert independent["candidates"][0]["reason"] == "RELEASE_GATE_CLOSED"


def test_profile_model_failure_keeps_comparison_but_never_promotes_it(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "profile-failed.docx"
    _contract(path, "混合型证券投资基金")
    class Gateway:
        def call(self, **kwargs):
            raise ModelError("MODEL_SCHEMA_INVALID")
    profile = resolve_profile(inspect_docx(path), Gateway())
    assert not profile["fee_groups_verified"]
    assert profile["source"] == "model_profile_failed" and "comparison" in profile
    assert profile["errors"][0]["code"] == "MODEL_SCHEMA_INVALID"


def test_profile_v3_selects_fee_clauses_after_unrelated_definitions_and_keeps_duplicates(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "priority.docx"
    doc = Document()
    doc.add_paragraph("第一部分 释义")
    for index in range(250):
        doc.add_paragraph(f"投资者信息{index}：" + "法律文件仅供资料说明。" * 20)
    doc.add_paragraph("第四部分 基金的基本情况")
    b = "B类基金份额收取申购费用，不计提销售服务费。"
    d = "D类基金份额不收取申购费用，计提销售服务费。"
    doc.add_paragraph(b)
    doc.add_paragraph(d)
    doc.add_paragraph("基金合同内容摘要")
    doc.add_paragraph("第五部分 基金费用与税收")
    doc.add_paragraph(b)
    doc.add_paragraph(d)
    doc.save(path)
    class Gateway:
        def call(self, **kwargs):
            assert kwargs["prompt_version"] == "profile_v3"
            assert len(kwargs["bounded_context"]) <= 11000
            assert "投资者信息" not in kwargs["bounded_context"]
            assert b in kwargs["bounded_context"] and d in kwargs["bounded_context"]
            return ModelResult({"classes": [{"name": "B", "subscription_fee": {"value": True, "quote": b}, "sales_service_fee": {"value": False, "quote": b}}, {"name": "D", "subscription_fee": {"value": False, "quote": d}, "sales_service_fee": {"value": True, "quote": d}}], "floating_management_fee": {"value": None, "quote": ""}}, "fixture", 1, 1, 1, 2, None, 1)
    profile = resolve_profile(inspect_docx(path), Gateway())
    assert profile["fee_groups_verified"] and not profile["context_truncated"]
    assert profile["context_selection"]["selected_class_fee_clauses"] == 2
    assert profile["fee_groups"] == {"subscription_group": ["B"], "service_group": ["D"]}
    assert len(profile["classes"][0]["evidence"][0]["occurrences"]) == 2


def test_profile_v3_actual_relevant_overflow_stays_unknown(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "overflow.docx"
    doc = Document()
    doc.add_paragraph("第一部分 基金费用与税收")
    b = "B类基金份额收取申购费用，不计提销售服务费。"
    doc.add_paragraph(b)
    for index in range(600):
        doc.add_paragraph(f"B类基金份额申购费用说明{index}：" + "相关费用条款内容。" * 10)
    doc.save(path)
    class Gateway:
        def call(self, **kwargs):
            assert len(kwargs["bounded_context"]) <= 11000
            return ModelResult({"classes": [{"name": "B", "subscription_fee": {"value": True, "quote": b}, "sales_service_fee": {"value": False, "quote": b}}], "floating_management_fee": {"value": None, "quote": ""}}, "fixture", 1, 1, 1, 2, None, 1)
    profile = resolve_profile(inspect_docx(path), Gateway())
    assert profile["context_truncated"] and not profile["fee_groups_verified"]
    assert profile["context_selection"]["selected_unique_clauses"] < profile["context_selection"]["relevant_unique_clauses"]


def test_profile_spaces_direct_asset_negation_and_named_joint_groups(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "joint-fees.docx"
    doc = Document()
    doc.add_paragraph("第四部分 基金的基本情况")
    count_quote = "本基金设置 A 类、C 类及 E 类基金份额。"
    a = "A 类基金份额在申购时收取申购费用，但不从本类别资产中计提销售服务费。"
    ce = "C 类及 E 类基金份额在申购时不收取申购费用，而从本类别基金资产中计提销售服务费。"
    for quote in (count_quote, a, ce):
        doc.add_paragraph(quote)
    doc.save(path)
    class Gateway:
        def call(self, **kwargs):
            return ModelResult({"classes": [{"name": name, "subscription_fee": {"value": name == "A", "quote": a if name == "A" else ce}, "sales_service_fee": {"value": name != "A", "quote": a if name == "A" else ce}} for name in ("A", "C", "E")], "share_class_count": {"value": 3, "quote": count_quote}, "floating_management_fee": {"value": None, "quote": ""}}, "fixture", 1, 1, 1, 2, None, 1)
    profile = resolve_profile(inspect_docx(path), Gateway())
    assert profile["fee_groups_verified"]
    assert profile["share_class_count"] == 3 and profile["share_class_count_verified"]
    assert profile["fee_groups"] == {"subscription_group": ["A"], "service_group": ["C", "E"]}
    evidence = profile["classes"][1]["evidence"][0]
    assert evidence["quote"] == ce
    assert inspect_docx(path).by_location(evidence["paragraph_id"]).text[evidence["char_start"]:evidence["char_end"]] == ce


def test_count_predicates_do_not_depend_on_unrelated_fee_unknowns(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "single-count.docx"
    doc = Document()
    doc.add_paragraph("第四部分 基金的基本情况")
    quote = "本基金不区分份额类别。"
    doc.add_paragraph(quote)
    doc.save(path)
    class Gateway:
        def call(self, **kwargs):
            return ModelResult({"classes": [{"name": "__single__", "subscription_fee": {"value": None, "quote": ""}, "sales_service_fee": {"value": None, "quote": ""}}], "share_class_count": {"value": 1, "quote": quote}, "floating_management_fee": {"value": None, "quote": ""}}, "fixture", 1, 1, 1, 2, None, 1)
    profile = resolve_profile(inspect_docx(path), Gateway())
    assert profile["share_class_count"] == 1 and not profile["fee_groups_verified"]
    assert applicability({"condition": "single_class", "requires_profile": ["share_classes"]}, profile)[0] == "true"
    assert applicability({"condition": "multi_class", "requires_profile": ["share_classes"]}, profile)[0] == "false"
    assert applicability({"condition": "has_service_fee", "requires_profile": ["sales_service_fee_by_class"]}, profile)[0] == "unknown"
    assert applicability({"condition": "single_class", "requires_profile": ["subscription_fee_by_class"]}, profile)[0] == "unknown"


def test_single_class_general_fees_bind_only_explicit_single_declaration(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "single-fees.docx"
    doc = Document()
    doc.add_paragraph("第四部分 基金的基本情况")
    declaration = "本基金仅有一种基金份额。"
    fees = "本基金申购时收取申购费用，不收取销售服务费。"
    doc.add_paragraph(declaration)
    doc.add_paragraph(fees)
    doc.save(path)
    class Gateway:
        def call(self, **kwargs):
            return ModelResult({"classes": [{"name": "__single__", "subscription_fee": {"value": True, "quote": fees}, "sales_service_fee": {"value": False, "quote": fees}}], "share_class_count": {"value": 1, "quote": declaration}, "floating_management_fee": {"value": None, "quote": ""}}, "fixture", 1, 1, 1, 2, None, 1)
    profile = resolve_profile(inspect_docx(path), Gateway())
    assert profile["fee_groups_verified"]
    assert applicability({"condition": "no_service_fee", "requires_profile": ["sales_service_fee_by_class"]}, profile)[0] == "true"
    doc.paragraphs[1].text = "本基金可以在未来设置不同基金份额类别。"
    doc.save(path)
    profile = resolve_profile(inspect_docx(path), Gateway())
    assert not profile["share_class_count_verified"] and not profile["fee_groups_verified"]


def test_one_observed_named_class_does_not_prove_single_and_services_are_three_valued():
    unknown_count = {"share_classes": ["A"], "classes": [{"name": "A", "subscription_fee": None, "sales_service_fee": True}], "fee_groups_verified": False, "share_class_count": None, "share_class_count_min": 1, "share_class_count_verified": False, "class_universe_verified": False, "floating_management_fee": True}
    assert applicability({"condition": "single_class", "requires_profile": ["share_classes"]}, unknown_count)[0] == "unknown"
    assert applicability({"condition": "has_service_fee", "requires_profile": ["sales_service_fee_by_class"]}, unknown_count)[0] == "true"
    assert applicability({"condition": "no_service_fee", "requires_profile": ["sales_service_fee_by_class"]}, unknown_count)[0] == "false"
    unknown_count.update(share_class_count=2, share_class_count_verified=True)
    assert applicability({"condition": "floating_multi", "requires_profile": ["share_classes", "floating_management_fee"]}, unknown_count)[0] == "true"


def test_joint_names_with_different_class_predicates_are_not_shared_evidence(tmp_path: Path):
    from fundlex.rules import resolve_profile
    path = tmp_path / "mixed-attribution.docx"
    doc = Document()
    doc.add_paragraph("第四部分 基金的基本情况")
    quote = "本基金设立 A 类及 C 类基金份额。"
    mixed = "A 类基金份额收取申购费用，C 类基金份额不收取申购费用。"
    doc.add_paragraph(quote)
    doc.add_paragraph(mixed)
    doc.save(path)
    class Gateway:
        def call(self, **kwargs):
            return ModelResult({"classes": [{"name": name, "subscription_fee": {"value": True, "quote": mixed}, "sales_service_fee": {"value": None, "quote": ""}} for name in ("A", "C")], "floating_management_fee": {"value": None, "quote": ""}}, "fixture", 1, 1, 1, 2, None, 1)
    profile = resolve_profile(inspect_docx(path), Gateway())
    by_class={c['name']:c for c in profile['classes']}
    assert by_class['A']['subscription_fee'] is True
    assert by_class['C']['subscription_fee'] is None
    assert not profile["fee_groups_verified"]
    assert applicability({"condition": "multi_class", "requires_profile": ["share_classes"]}, profile)[0] == "true"


def test_unsupported_target_tables_fail_with_locations_non_target_tables_ignore(tmp_path: Path):
    path = tmp_path / "target-table.docx"
    doc = Document()
    doc.add_paragraph("第六部分 基金份额的申购与赎回")
    doc.add_paragraph("六、申购和赎回的价格、费用及其用途")
    doc.add_paragraph("本基金按照合同处理申购。")
    old = next(a for a in load_ruleset("v1").actions if a["action_id"] == "R10-A1")["old_texts"][0]
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = old
    doc.save(path)
    inspection = inspect_docx(path)
    span = next(s for s in inspection.paragraphs if s.zone == "unsupported_table")
    assert span.logical_zone == "body" and not span.editable
    action = next(a for a in evaluate_candidates(inspection, load_ruleset("v1")) if a["action_id"] == "R10-A1")
    assert action["status"] == "failed" and not action["scope_verified"]
    assert action["candidates"][0]["reason"] == "UNSUPPORTED_TARGET_STRUCTURE"
    assert action["candidates"][0]["location"].startswith("tbl")
    assert "unsupported_table_structure" in action["candidates"][0]["structure_reasons"]
    # The same phrase in an ordinary table in another section is not a target.
    table._tbl.getparent().remove(table._tbl)
    doc.add_paragraph("第七部分 其他约定")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = old
    doc.save(path)
    action = next(a for a in evaluate_candidates(inspect_docx(path), load_ruleset("v1")) if a["action_id"] == "R10-A1")
    assert action["status"] == "not_matched" and action["scope_verified"]
    assert not action["candidates"]


def test_nested_target_table_and_appendix_table_have_distinct_logical_regions(tmp_path: Path):
    path = tmp_path / "nested.docx"
    doc = Document()
    doc.add_paragraph("第六部分 基金份额的申购与赎回")
    doc.add_paragraph("六、申购和赎回的价格、费用及其用途")
    old = next(a for a in load_ruleset("v1").actions if a["action_id"] == "R10-A1")["old_texts"][0]
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = old
    doc.save(path)
    action = next(a for a in evaluate_candidates(inspect_docx(path), load_ruleset("v1")) if a["action_id"] == "R10-A1")
    assert action["status"] == "failed"
    assert any(c["location"].startswith("readonly:") and old in c["quote"] for c in action["candidates"])
    doc.tables[0]._tbl.addprevious(doc.add_paragraph("附录")._p)
    doc.save(path)
    action = next(a for a in evaluate_candidates(inspect_docx(path), load_ruleset("v1")) if a["action_id"] == "R10-A1")
    assert action["status"] == "not_matched" and action["scope_verified"]
    assert not action["candidates"]


def test_release_context_rejects_missing_engine_model_prompt_or_transport_drift(tmp_path: Path, monkeypatch):
    from dataclasses import replace
    from copy import deepcopy
    import pytest
    import fundlex.rules as rules
    source = tmp_path / "release-context.docx"
    _contract(source, "混合型证券投资基金")
    catalogue = load_ruleset("v1")
    action = dict(next(a for a in catalogue.actions if a["action_id"] == "R03-A3"))
    action["auto_write"] = True
    gateway = StubGateway("ordinary")
    tested = rules.release_runtime_context(gateway)
    action["evaluation_metrics"] = {"tested_context": deepcopy(tested), "fixture_only": True}
    fixture = replace(catalogue, actions=(action,), release_policy={"evaluation_context": deepcopy(tested)})
    rules.verify_release_context(fixture, gateway, gateway.model)
    gateway.api_key = "independent-fixture-credential"
    assert rules.release_runtime_context(gateway) == tested
    for field, value in (("engine_build", "f" * 64), ("model_transport_sha256", "e" * 64), ("actual_model", "changed-model"), ("prompt_versions", {})):
        changed = deepcopy(tested)
        changed[field] = value
        with pytest.raises(RuleError, match="RELEASE_CONTEXT_UNVERIFIED"):
            rules.verify_release_context(replace(fixture, release_policy={"evaluation_context": changed}), gateway, gateway.model)
    with pytest.raises(RuleError, match="RELEASE_CONTEXT_UNVERIFIED"):
        rules.verify_release_context(replace(fixture, release_policy={}), gateway, gateway.model)
    gateway.timeout_s = 90
    with pytest.raises(RuleError, match="RELEASE_CONTEXT_UNVERIFIED"):
        rules.verify_release_context(fixture, gateway, gateway.model)


def test_output_validation_failure_reverts_applied_statuses(tmp_path: Path, monkeypatch):
    from dataclasses import replace
    import fundlex.pipeline as pipeline
    from fundlex.rules import release_runtime_context
    from fundlex.docx import DocxError
    source = tmp_path / "output-failure.docx"
    _contract(source, "混合型证券投资基金")
    catalogue = load_ruleset("v1")
    action = dict(next(a for a in catalogue.actions if a["action_id"] == "R03-A3"))
    action["auto_write"] = True
    gateway = StubGateway("ordinary")
    tested = release_runtime_context(gateway)
    action["evaluation_metrics"] = {"tested_context": tested, "fixture_only": True}
    fixture = replace(catalogue, actions=(action,), release_policy={"evaluation_context": tested})
    monkeypatch.setattr(pipeline, "load_ruleset", lambda _: fixture)
    def failed(*args):
        raise DocxError("FINAL_REVISION_MISMATCH")
    monkeypatch.setattr(pipeline, "generate_artifacts", failed)
    report = pipeline.process(source, tmp_path / "out", "v1", gateway)
    assert report["delivery_status"] == "not_generated" and not report["edit_plan"]["executable"]
    assert report["actions"][0]["status"] == "failed"
    assert report["actions"][0]["candidates"][0]["status"] == "failed"
    assert report["actions"][0]["candidates"][0]["reason"] == "FINAL_REVISION_MISMATCH"


def test_non_contract_mirrors_reject_before_model_and_record_title_evidence(tmp_path: Path):
    for kind in ("招募说明书", "托管协议", "公告", "决议"):
        source = tmp_path / (kind + ".docx")
        _contract(source, "混合型证券投资基金")
        doc = Document(source)
        doc.paragraphs[0].text = "示例证券投资基金" + kind
        doc.save(source)
        gateway = StubGateway("ordinary")
        report = process(source, tmp_path / kind, "v1", gateway)
        assert report["qualification_status"] == "failed"
        assert report["error_code"] == "NON_CONTRACT_DOCUMENT" and gateway.calls == 0
        assert report["document_identity"]["evidence"][0]["paragraph_id"] == "p0"
        assert report["document_identity"]["evidence"][0]["quote"].endswith(kind)


def test_cover_fund_name_does_not_classify_type_and_unknown_identity_closes(tmp_path: Path):
    source = tmp_path / "misleading-name.docx"
    _contract(source, "混合型证券投资基金")
    doc = Document(source)
    doc.paragraphs[0].text = "示例交易型开放式指数ETF证券投资基金基金合同"
    doc.save(source)
    gateway = StubGateway("ordinary")
    report = process(source, tmp_path / "ordinary", "v1", gateway)
    assert report["qualification_status"] == "eligible"
    assert report["qualification"]["fund_type"] == "ordinary"
    assert report["document_identity"]["status"] == "verified"
    doc.paragraphs[0].text = "本文件依据基金合同提供资料。"
    doc.save(source)
    gateway = StubGateway("ordinary")
    report = process(source, tmp_path / "unknown", "v1", gateway)
    assert report["error_code"] == "DOCUMENT_IDENTITY_EVIDENCE_MISSING"
    assert gateway.calls == 0


def test_title_conflict_and_title_without_category_never_qualify(tmp_path: Path):
    source = tmp_path / "conflicting-title.docx"
    _contract(source, "混合型证券投资基金")
    doc = Document(source)
    doc.paragraphs[1].insert_paragraph_before("示例证券投资基金招募说明书")
    doc.save(source)
    gateway = StubGateway("ordinary")
    report = process(source, tmp_path / "conflict", "v1", gateway)
    assert report["error_code"] == "DOCUMENT_TYPE_CONFLICT" and gateway.calls == 0
    doc = Document()
    doc.add_paragraph("示例证券投资基金基金合同")
    doc.add_paragraph("本文无可靠基金类型条款。")
    doc.save(source)
    gateway = StubGateway("ordinary")
    report = process(source, tmp_path / "no-category", "v1", gateway)
    assert report["error_code"] == "FUND_TYPE_EVIDENCE_MISSING" and gateway.calls == 0


def test_identity_without_cover_requires_operative_preamble_not_generic_mentions(tmp_path: Path):
    source = tmp_path / "preamble.docx"
    _contract(source, "混合型证券投资基金")
    doc = Document(source)
    doc.paragraphs[0].text = "第一部分 前言"
    doc.paragraphs[1].insert_paragraph_before("基金管理人、基金托管人与基金份额持有人依照法律法规订立本基金合同。")
    doc.paragraphs[2].insert_paragraph_before("本基金合同约定基金管理人、基金托管人与基金份额持有人的权利义务。")
    doc.save(source)
    report = process(source, tmp_path / "preamble", "v1", StubGateway("ordinary"))
    assert report["qualification_status"] == "eligible"
    assert report["document_identity"]["basis"] == "operative_contract_preamble"
    assert len(report["document_identity"]["evidence"]) == 2


def test_pipeline_missing_release_context_fails_before_write(tmp_path: Path, monkeypatch):
    from dataclasses import replace
    import fundlex.pipeline as pipeline
    source = tmp_path / "unbound.docx"
    _contract(source, "混合型证券投资基金")
    catalogue = load_ruleset("v1")
    action = dict(next(a for a in catalogue.actions if a["action_id"] == "R03-A3"))
    action["auto_write"] = True
    monkeypatch.setattr(pipeline, "load_ruleset", lambda _: replace(catalogue, actions=(action,)))
    report = pipeline.process(source, tmp_path / "out", "v1", StubGateway("ordinary"))
    assert report["error_code"] == "RELEASE_CONTEXT_UNVERIFIED"
    assert report["release_gate"]["status"] == "closed" and not report["edit_plan"]["executable"]
    assert not (tmp_path / "out" / "final.docx").exists()


def test_whitespace_trigger_maps_original_half_open_source_and_punctuation(tmp_path: Path):
    from fundlex.rules import NORMALIZER_VERSION
    path = tmp_path / "spacing.docx"
    rules = load_ruleset("v1")
    action = next(a for a in rules.actions if a["action_id"] == "R11-A1")
    raw = action["old_texts"][0].replace("少于7日", "少于7 日").removesuffix("。")
    doc = Document()
    doc.add_paragraph("第六部分 基金份额的申购与赎回")
    doc.add_paragraph("六、申购和赎回的价格、费用及其用途")
    host = "本基金的申购费率和赎回费率由基金管理人确定。"
    original = host + "  " + raw + "   "
    doc.add_paragraph(original)
    doc.save(path)
    result = next(a for a in evaluate_candidates(inspect_docx(path), rules) if a["action_id"] == "R11-A1")
    assert len(result["candidates"]) == 1
    candidate = result["candidates"][0]
    assert candidate["target_verified"] and candidate["reason"] == "RELEASE_GATE_CLOSED"
    assert "少于7 日" in candidate["quote"]
    assert candidate["char_start"] == len(host)+2
    assert original[candidate["char_start"]:candidate["char_end"]] == candidate["quote"]
    assert candidate["normalizer_version"] == NORMALIZER_VERSION


def test_normalizer_does_not_join_word_break_or_field_barrier(tmp_path: Path):
    from docx.oxml import OxmlElement
    from fundlex.docx import Q
    from fundlex.rules import _normalized_matches
    rules = load_ruleset("v1")
    action = next(a for a in rules.actions if a["action_id"] == "R11-A1")
    old = action["old_texts"][0]
    split = old.index("7日") + 1
    for kind in ("br", "field"):
        path = tmp_path / (kind + ".docx")
        doc = Document()
        doc.add_paragraph("第六部分 基金份额的申购与赎回")
        doc.add_paragraph("六、申购和赎回的价格、费用及其用途")
        paragraph = doc.add_paragraph(old[:split])
        if kind == "br":
            paragraph.add_run().add_break()
        else:
            field = OxmlElement("w:fldSimple")
            field.set(Q("instr"), "PAGE")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = " "
            run.append(text)
            field.append(run)
            paragraph._p.append(field)
        paragraph.add_run(old[split:])
        doc.save(path)
        inspection = inspect_docx(path)
        assert not _normalized_matches(inspection.by_location("p2").text, old)
        result = next(a for a in evaluate_candidates(inspection, rules) if a["action_id"] == "R11-A1")
        assert not result["candidates"]
