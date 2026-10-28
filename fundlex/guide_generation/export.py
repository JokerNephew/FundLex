"""Candidate guide DOCX export and evidence package rendering."""
from __future__ import annotations

import html
import json
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .models import GuideItem, canonical_json, sha256_json, utc_now


def export_guide_docx(items: list[GuideItem], path: str | Path, *,
                      title: str = "新法规候选修改指引",
                      regulation_title: str = "") -> dict[str, Any]:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    intro = document.add_paragraph()
    intro.add_run(
        f"依据：{regulation_title or '本次上传的监管规定'}。本文件为 AI 生成的候选修改指引，"
        "不是最终正式业务指引，须经人工复核、修订和批准。")
    intro.runs[0].italic = True
    for index, item in enumerate(items, start=1):
        document.add_paragraph()
        group = document.add_paragraph()
        label = f"{item.section} - {item.title}：" if item.section else f"{item.title}："
        group_run = group.add_run(label)
        group_run.bold = True
        requirement = document.add_paragraph()
        requirement.add_run("法规要求：").bold = True
        requirement.add_run(item.business_requirement or item.instruction)
        if item.applicability or item.conditions or item.exceptions:
            applicability = document.add_paragraph()
            applicability.add_run("适用条件：").bold = True
            parts = [item.applicability] if item.applicability else []
            parts.extend(f"条件：{value}" for value in item.conditions)
            parts.extend(f"例外：{value}" for value in item.exceptions)
            applicability.add_run("；".join(part for part in parts if part))
        if item.old_text_patterns:
            old = document.add_paragraph()
            old.add_run("原条款，参考如下描述：").bold = True
            for pattern in item.old_text_patterns:
                paragraph = document.add_paragraph()
                paragraph.add_run(pattern)
        new = document.add_paragraph()
        new.add_run("新条款：").bold = True
        if item.new_text.strip():
            document.add_paragraph(item.new_text.strip())
        else:
            document.add_paragraph(item.instruction or "待人工补充新条款文案。")
        source = document.add_paragraph()
        source.add_run("指引说明：").bold = True
        source.add_run(item.instruction or item.title)
    document.add_paragraph()
    note = document.add_paragraph()
    note.add_run(
        "说明：候选指引中的新文案来源已在 Evidence Package 中逐条记录（法规原文 / "
        "模型起草 / 历史合同措辞改编）。人工修订后系统会保存 AI→Human Diff；人工版本才是业务最终版本。")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    import hashlib
    data = path.read_bytes()
    return {"path": str(path), "sha256": hashlib.sha256(data).hexdigest(), "size": len(data),
            "item_count": len(items)}


def _json_dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2, default=str)


def build_evidence_package(*, run: dict[str, Any], regulation: dict[str, Any],
                           requirements: list[dict[str, Any]], impacts: list[dict[str, Any]],
                           local_candidates: list[dict[str, Any]], families: list[dict[str, Any]],
                           guide_items: list[dict[str, Any]], findings: list[dict[str, Any]],
                           coverage: list[dict[str, Any]], applicability: list[dict[str, Any]],
                           residual_rounds: list[dict[str, Any]], model_calls: list[dict[str, Any]],
                           model_identity: dict[str, Any], engine_compatibility: dict[str, Any] | None = None,
                           human_versions: list[dict[str, Any]] | None = None,
                           human_changes: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    impact_map = {item["impact_id"]: item for item in impacts}
    coverage_by_ref: dict[str, list[dict[str, Any]]] = {}
    for row in coverage:
        coverage_by_ref.setdefault(row.get("ref_id", ""), []).append(row)
    applicability_by_item: dict[str, list[dict[str, Any]]] = {}
    for row in applicability:
        applicability_by_item.setdefault(row.get("guide_item_id", ""), []).append(row)
    findings_by_item: dict[str, list[dict[str, Any]]] = {}
    for finding in findings:
        findings_by_item.setdefault(finding.get("guide_item_id", ""), []).append(finding)
    local_by_id = {item["local_id"]: item for item in local_candidates}
    chain: list[dict[str, Any]] = []
    for item in guide_items:
        item_id = item["guide_item_id"]
        linked_impacts = [impact_map[value] for value in item.get("impact_ids", []) if value in impact_map]
        local_ids = []
        for family_id in item.get("clause_family_ids", []):
            family = next((value for value in families if value.get("family_id") == family_id), None)
            if family:
                local_ids.extend(family.get("local_ids") or [])
        chain.append({
            "guide_item_id": item_id,
            "guide_item": item,
            "requirements": [value for value in requirements if value["requirement_id"] in item.get("requirement_ids", [])],
            "impacts": linked_impacts,
            "local_candidates": [local_by_id[value] for value in local_ids if value in local_by_id],
            "clause_family": next((value for value in families if value.get("family_id") in item.get("clause_family_ids", [])), None),
            "quality_findings": findings_by_item.get(item_id, []),
            "coverage": coverage_by_ref.get(item_id, []),
            "semantic_applicability": applicability_by_item.get(item_id, []),
            "new_text_origin": item.get("new_text_origin", ""),
            "residual_history": [round_item for round_item in residual_rounds
                                 if any(impact.get("impact_id") in (round_item.get("new_impact_ids") or [])
                                        for impact in linked_impacts)],
        })
    return {
        "schema_version": "guide-generation-evidence-1",
        "generated_at": utc_now(),
        "run": run,
        "regulation": regulation,
        "requirements": requirements,
        "impacts": impacts,
        "local_guide_candidates": local_candidates,
        "clause_families": families,
        "guide_items": guide_items,
        "quality_findings": findings,
        "coverage": coverage,
        "semantic_applicability": applicability,
        "residual_rounds": residual_rounds,
        "model_identity": model_identity,
        "model_calls": model_calls,
        "engine_compatibility": engine_compatibility or {"status": "not_run"},
        "human_versions": human_versions or [],
        "human_changes": human_changes or [],
        "traceability_chain": chain,
        "qualification": {
            "guide_is_candidate_only": True,
            "engine_compatibility_is_separate": True,
            "independent_model_review": model_identity.get("independent_model_review_status"),
        },
    }


def _e(value: Any) -> str:
    return html.escape("" if value is None else str(value))


def render_evidence_html(package: dict[str, Any], *, title: str = "候选修改指引 Evidence Report") -> str:
    run = package.get("run") or {}
    stats = run.get("stats") or {}
    requirements = package.get("requirements") or []
    impacts = package.get("impacts") or []
    items = package.get("guide_items") or []
    findings = package.get("quality_findings") or []
    coverage = package.get("coverage") or []
    rounds = package.get("residual_rounds") or []
    model_identity = package.get("model_identity") or {}
    chain = package.get("traceability_chain") or []
    requirement_map = {item["requirement_id"]: item for item in requirements}
    impact_map = {item["impact_id"]: item for item in impacts}
    coverage_by_ref: dict[str, list[dict[str, Any]]] = {}
    for row in coverage:
        coverage_by_ref.setdefault(row.get("ref_id", ""), []).append(row)

    def card(label: str, value: Any, hint: str = "") -> str:
        return (f'<div class="card"><div class="card-label">{_e(label)}</div>'
                f'<div class="card-value">{_e(value)}</div>'
                f'<div class="card-hint">{_e(hint)}</div></div>')

    requirement_counts = {}
    for item in requirements:
        requirement_counts[item.get("disposition", "")] = requirement_counts.get(item.get("disposition", ""), 0) + 1
    impact_counts = {}
    for item in impacts:
        impact_counts[item.get("impact_status", "")] = impact_counts.get(item.get("impact_status", ""), 0) + 1
    critical = [item for item in findings if item.get("severity") == "critical"]
    coverage_summary: dict[str, int] = {}
    for row in coverage:
        if row.get("level") == "requirement":
            coverage_summary[row.get("status", "")] = coverage_summary.get(row.get("status", ""), 0) + 1
    cards = "".join([
        card("法规 Requirements", len(requirements), "所有要求均有 disposition"),
        card("需要改合同", requirement_counts.get("modify_contract", 0), "modify_contract"),
        card("Impact Candidates", len(impacts), f"definite {impact_counts.get('definite', 0)} / possible {impact_counts.get('possible', 0)}"),
        card("Guide Items", len(items), "候选，不是正式指引"),
        card("Quality Findings", len(findings), f"critical {len(critical)}"),
        card("Requirement Coverage", f"{coverage_summary.get('covered', 0)} covered / {coverage_summary.get('gap', 0)} gap", "独立覆盖门"),
        card("Residual Rounds", len(rounds), f"status {run.get('residual_status', '')}"),
        card("Independent Review", model_identity.get("independent_model_review_status", ""), "Author/Reviewer 模型身份"),
    ])
    item_rows = []
    for chain_item in chain:
        item = chain_item["guide_item"]
        requirements_html = "".join(
            f'<div class="req"><b>{_e(req.get("requirement_id"))}</b> [{_e(req.get("article"))}] '
            f'{_e(req.get("requirement"))}<br><span class="quote">{_e(req.get("source_quote"))}</span></div>'
            for req in chain_item.get("requirements") or [])
        impact_html = "".join(
            f'<div class="impact"><b>{_e(impact.get("impact_id"))}</b> '
            f'{_e(impact.get("impact_status"))} / {_e(impact.get("contract_version_id"))}<br>'
            f'{_e((impact.get("source_text") or "")[:600])}<br>'
            f'<span class="muted">{_e(impact.get("impact_reason"))}</span></div>'
            for impact in chain_item.get("impacts") or [])
        finding_html = "".join(
            f'<li class="{_e(f.get("severity"))}"><b>{_e(f.get("severity"))}</b> '
            f'[{_e(f.get("category"))}] {_e(f.get("description"))}'
            f'<br><span class="quote">{_e(f.get("evidence_quote"))}</span></li>'
            for f in chain_item.get("quality_findings") or [])
        applicability_html = "".join(
            f'<li>{_e(row.get("clause_id"))}: <b>{_e(row.get("verdict"))}</b> — {_e(row.get("reason"))}</li>'
            for row in chain_item.get("semantic_applicability") or [])
        item_rows.append(f"""
        <details class="guide-item">
          <summary><b>{_e(item.get('guide_item_id'))}</b> {_e(item.get('title'))}
            <span class="badge {_e(item.get('quality_status'))}">{_e(item.get('quality_status'))}</span>
            <span class="badge">{_e(item.get('new_text_origin'))}</span></summary>
          <div class="section"><h4>法规来源 / Requirement</h4>{requirements_html or '<p class="muted">无</p>'}</div>
          <div class="section"><h4>真实合同证据 / ImpactCandidate</h4>{impact_html or '<p class="muted">无</p>'}</div>
          <div class="section"><h4>为什么形成这条 Guide</h4>
            <p><b>业务目标：</b>{_e(item.get('business_requirement'))}</p>
            <p><b>指引：</b>{_e(item.get('instruction'))}</p>
            <p><b>适用：</b>{_e(item.get('applicability'))}</p>
            <p><b>条件：</b>{_e('；'.join(item.get('conditions') or []))}</p>
            <p><b>例外：</b>{_e('；'.join(item.get('exceptions') or []))}</p>
            <p><b>原条款模式：</b>{_e('；'.join(item.get('old_text_patterns') or []))}</p>
            <p><b>新文案：</b>{_e(item.get('new_text'))}</p>
            <p><b>来源类型：</b>{_e(item.get('new_text_origin'))}</p>
          </div>
          <div class="section"><h4>Reviewer Findings</h4><ul>{finding_html or '<li>无</li>'}</ul></div>
          <div class="section"><h4>Semantic Applicability</h4><ul>{applicability_html or '<li>未测试</li>'}</ul></div>
        </details>""")
    residual_rows = "".join(
        f"<tr><td>{_e(round_item.get('round_no'))}</td><td>{_e(round_item.get('status'))}</td>"
        f"<td>{_e(round_item.get('new_patterns'))}</td><td>{_e(round_item.get('new_impact_count', ''))}</td></tr>"
        for round_item in rounds)
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{_e(title)}</title>
<style>
body{{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI","Noto Sans CJK SC",sans-serif;margin:0;background:#f6f7f9;color:#1f2937}}
header{{background:#12263f;color:white;padding:28px 36px}}
header h1{{margin:0 0 8px;font-size:24px}} header p{{margin:0;color:#c9d6e3}}
main{{max-width:1240px;margin:0 auto;padding:24px 32px 80px}}
.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:12px;margin:20px 0}}
.card{{background:white;border:1px solid #e5e7eb;border-radius:10px;padding:16px}}
.card-label{{font-size:12px;color:#6b7280;text-transform:uppercase;letter-spacing:.04em}}
.card-value{{font-size:28px;font-weight:700;margin:6px 0}}
.card-hint{{font-size:12px;color:#6b7280}}
section.panel{{background:white;border:1px solid #e5e7eb;border-radius:10px;padding:18px 20px;margin:16px 0}}
h2{{font-size:18px;margin:0 0 12px}} h3{{font-size:15px;margin:18px 0 8px}}
details.guide-item{{background:white;border:1px solid #e5e7eb;border-radius:10px;padding:12px 16px;margin:10px 0}}
details.guide-item summary{{cursor:pointer;font-size:15px}}
.badge{{display:inline-block;font-size:11px;padding:2px 7px;border-radius:999px;background:#eef2ff;color:#3730a3;margin-left:6px}}
.badge.fail{{background:#fee2e2;color:#991b1b}} .badge.pass{{background:#dcfce7;color:#166534}} .badge.pass_with_findings{{background:#fef3c7;color:#92400e}}
.section{{margin:12px 0}} .section h4{{margin:0 0 6px;font-size:13px;color:#374151;text-transform:uppercase;letter-spacing:.04em}}
.req,.impact{{border-left:3px solid #c7d2fe;background:#f8faff;padding:8px 10px;margin:6px 0;border-radius:4px;font-size:13px}}
.quote{{color:#4b5563;font-family:ui-monospace,SFMono-Regular,Menlo,monospace;font-size:12px}}
.muted{{color:#6b7280}} .critical{{color:#991b1b}} .major{{color:#92400e}}
table{{border-collapse:collapse;width:100%;font-size:13px}} td,th{{border:1px solid #e5e7eb;padding:7px 9px;text-align:left}}
.warn{{background:#fffbeb;border:1px solid #fde68a;padding:12px 14px;border-radius:8px;color:#92400e}}
</style></head><body>
<header><h1>{_e(title)}</h1><p>Run {_e(run.get('id'))} · {_e(run.get('title'))} · 本报告是候选指引证据，不是正式业务指引</p></header>
<main>
<div class="warn">候选指引只代表在当前多路发现、完整 corpus sweep、独立复核和 residual mining 策略下未发现新的 uncovered impact；不构成数学意义上的 100% 无遗漏证明。Engine Compatibility 单独报告，不影响 Guide Quality。</div>
<div class="cards">{cards}</div>
<section class="panel"><h2>运行与可追溯身份</h2>
<p><b>Regulation SHA-256：</b><span class="quote">{_e(run.get('source_sha256'))}</span></p>
<p><b>Corpus Snapshot：</b>{_e(run.get('corpus_snapshot_id'))} · {_e(stats.get('corpus', {}).get('contracts', ''))} contracts</p>
<p><b>Code commit：</b>{_e(run.get('commit_sha'))} · <b>Engine build：</b>{_e(run.get('code_build'))}</p>
<p><b>Author model：</b>{_e((model_identity.get('guide_author') or {}).get('model'))} · <b>Reviewer model：</b>{_e((model_identity.get('guide_reviewer') or {}).get('model'))} · <b>Coverage auditor：</b>{_e((model_identity.get('coverage_auditor') or {}).get('model'))}</p>
</section>
<section class="panel"><h2>Residual Mining</h2>
<table><thead><tr><th>Round</th><th>Status</th><th>New patterns</th><th>New impacts</th></tr></thead><tbody>{residual_rows or '<tr><td colspan="4">未运行</td></tr>'}</tbody></table>
</section>
<section class="panel"><h2>Guide Items（点击展开证据链）</h2>{''.join(item_rows) or '<p class="muted">没有候选指引。</p>'}</section>
<section class="panel"><h2>全部 Quality Findings</h2>
<table><thead><tr><th>Severity</th><th>Category</th><th>GuideItem</th><th>Description</th></tr></thead><tbody>
{''.join(f"<tr><td class='{_e(f.get('severity'))}'>{_e(f.get('severity'))}</td><td>{_e(f.get('category'))}</td><td>{_e(f.get('guide_item_id'))}</td><td>{_e(f.get('description'))}</td></tr>" for f in findings) or '<tr><td colspan="4">无</td></tr>'}
</tbody></table></section>
<section class="panel"><h2>Requirement × Contract Coverage 摘要</h2>
<table><thead><tr><th>Status</th><th>Count</th></tr></thead><tbody>
{''.join(f"<tr><td>{_e(k)}</td><td>{_e(v)}</td></tr>" for k, v in sorted(_count_coverage(coverage).items())) or '<tr><td colspan="2">无</td></tr>'}
</tbody></table></section>
</main></body></html>"""


def _count_coverage(rows: list[dict[str, Any]]) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in rows:
        if row.get("level") == "contract_requirement":
            key = row.get("status", "")
            result[key] = result.get(key, 0) + 1
    return result
